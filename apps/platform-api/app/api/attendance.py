"""attendance API - endpoints for attendance management and scoring.

Endpoints:
- POST /attendance/sync: Trigger incremental sync from signin system
- GET /attendance/event-groups: List event groups
- GET /attendance/event-groups/{id}: Get event group detail with sessions
- GET /attendance/records: List attendance records with filters
- POST /attendance/records/{id}/adjudications: Create adjudication
- POST /attendance/event-groups/{id}/recalculate: Recalculate scores
"""
from __future__ import annotations

import hmac
import re
from io import BytesIO
from datetime import UTC, datetime
from typing import Literal
from fastapi import APIRouter, Depends, Header, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field

from app.api.auth import require_permission
from app.core.settings import get_settings
from app.db import execute, fetch_all, fetch_one, transaction
from app.services.audit import write_audit
from app.services.attendance_scoring import (
    get_active_rule,
    normalize_activity_type,
    recalculate_event_group,
    score_is_applicable,
    upsert_score_record,
)
from app.services.attendance_sync import sync_from_signin
from app.services.iam import accessible_org_ids


router = APIRouter(prefix="/api/v1/attendance", tags=["attendance"])

CURRENT_CLASS_NAME_SQL = (
    "COALESCE((SELECT ou.name FROM member_org_relations mor "
    "JOIN org_units ou ON ou.id=mor.org_unit_id "
    "WHERE mor.member_id=m.id AND mor.relation_type IN ('STUDY_CLASS','SPECIAL_COHORT') "
    "ORDER BY mor.id DESC LIMIT 1), m.class_name, '')"
)


class AdjudicationPayload(BaseModel):
    adjudication_type: str = Field(pattern="^(EARLY_LEAVE|CANCEL_EARLY_LEAVE|MANUAL_CHECKIN|INVALIDATE_CHECKIN|LEAVE|CANCEL_LEAVE|MEMBER_RELINK)$")
    occurred_at: str | None = None
    reason: str = Field(min_length=4, max_length=500)
    member_id: int | None = None


def _attendance_sync_health() -> dict:
    rows = fetch_all(
        "SELECT id, status, started_at, finished_at, received_sessions, "
        "received_records, error_count, error_summary "
        "FROM attendance_sync_runs WHERE source_key=? "
        "ORDER BY id DESC LIMIT 20",
        ("signin",),
    )
    if not rows:
        return {
            "state": "NO_RUNS",
            "alert_threshold": 3,
            "consecutive_failure_count": 0,
            "last_run": None,
        }

    latest = rows[0]
    consecutive_failures = 0
    for row in rows:
        status = str(row["status"]).upper()
        if status == "RUNNING":
            continue
        if status == "SUCCESS":
            break
        if status in {"ERROR", "PARTIAL"}:
            consecutive_failures += 1

    latest_status = str(latest["status"]).upper()
    if consecutive_failures >= 3:
        state = "CRITICAL"
    elif consecutive_failures:
        state = "WARNING"
    elif latest_status == "RUNNING":
        state = "RUNNING"
    else:
        state = "HEALTHY"
    return {
        "state": state,
        "alert_threshold": 3,
        "consecutive_failure_count": consecutive_failures,
        "last_run": {
            "status": latest_status,
            "started_at": latest["started_at"],
            "finished_at": latest["finished_at"],
            "received_sessions": latest["received_sessions"] or 0,
            "received_records": latest["received_records"] or 0,
            "error_count": latest["error_count"] or 0,
            "has_error_summary": bool(latest["error_summary"]),
        },
    }


def _write_failure_alert_if_threshold_reached() -> None:
    health = _attendance_sync_health()
    if health["consecutive_failure_count"] != health["alert_threshold"]:
        return
    with transaction() as connection:
        write_audit(
            connection,
            actor_user_id=None,
            action="attendance.sync.failure_alert",
            resource_type="attendance_sync",
            resource_id="signin",
            after={
                "state": health["state"],
                "consecutive_failure_count": health[
                    "consecutive_failure_count"
                ],
                "alert_threshold": health["alert_threshold"],
            },
        )


def _member_activity_month_clause(month: str | None) -> tuple[str, tuple[str, ...]]:
    """Limit member-quality review to members appearing in the selected event month."""
    if not month:
        return "", ()
    return (
        " AND EXISTS (SELECT 1 FROM attendance_records ar "
        "JOIN attendance_sessions mas ON mas.id=ar.attendance_session_id "
        "JOIN attendance_event_groups meg ON meg.id=mas.event_group_id "
        "WHERE ar.member_id=m.id AND substr(meg.event_date, 1, 7)=?)",
        (month,),
    )


def _execute_count(statement: str, params: tuple = ()) -> int:
    row = fetch_one(statement, params) if params else fetch_one(statement)
    return int((row or {"count": 0})["count"] or 0)


def _attendance_reconciliation_summary(month: str | None = None) -> dict:
    """Return only review workload counts; never expose attendance snapshots."""
    member_month_clause, member_month_params = _member_activity_month_clause(month)
    unmatched_month_clause = " AND substr(eg.event_date, 1, 7)=?" if month else ""
    queries = {
        "unmatched_attendance_records": (
            "SELECT COUNT(*) AS count FROM attendance_records ar "
            "JOIN attendance_sessions s ON s.id=ar.attendance_session_id "
            "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
            "WHERE (ar.attendance_status='UNMATCHED' OR ar.member_id IS NULL)"
            + unmatched_month_clause,
            (month,) if month else (),
        ),
        "active_members_missing_phone_hash": (
            "SELECT COUNT(*) AS count FROM members "
            "WHERE status='ACTIVE' AND (phone_hash IS NULL OR phone_hash='')"
            + (member_month_clause.replace("m.id", "members.id") if month else ""),
            member_month_params,
        ),
        "active_members_missing_primary_region": (
            "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='PRIMARY_REGION')"
            + member_month_clause,
            member_month_params,
        ),
        "active_members_missing_study_class": (
            "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='STUDY_CLASS')"
            + member_month_clause,
            member_month_params,
        ),
        "active_members_missing_study_group": (
            "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='STUDY_GROUP') "
            f"AND NOT ({CURRENT_CLASS_NAME_SQL} IN ('先锋班','神仙班') "
            "OR COALESCE(m.notes,'') LIKE '%目前不读书%')"
            + member_month_clause,
            member_month_params,
        ),
        "active_members_expected_no_study_group": (
            "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='STUDY_GROUP') "
            f"AND ({CURRENT_CLASS_NAME_SQL} IN ('先锋班','神仙班') "
            "OR COALESCE(m.notes,'') LIKE '%目前不读书%')"
            + member_month_clause,
            member_month_params,
        ),
    }
    return {
        "scope": "AGGREGATE_ONLY",
        "write_enabled": False,
        "items": [
            {
                "key": key,
                "count": _execute_count(statement, params),
            }
            for key, (statement, params) in queries.items()
        ],
    }


def _org_is_allowed(
    primary_org_id: str,
    study_org_unit_id: str | None,
    allowed: set[str] | None,
) -> bool:
    return allowed is None or bool(
        {primary_org_id, study_org_unit_id}.intersection(allowed)
    )


def _verify_signin_service_key(x_api_key: str | None) -> None:
    expected = get_settings().signin_service_api_key
    if not expected:
        raise HTTPException(503, "签到服务密钥未配置")
    if not x_api_key:
        raise HTTPException(401, "缺少 X-API-Key 请求头")
    if not hmac.compare_digest(x_api_key, expected):
        raise HTTPException(403, "API Key 无效")


def _run_sync(cursor: str | None, actor_user_id: int | None, action: str) -> dict:
    try:
        result = sync_from_signin(cursor)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, f"同步失败: {exc}") from exc
    with transaction() as connection:
        write_audit(
            connection,
            actor_user_id=actor_user_id,
            action=action,
            resource_type="attendance_sync_run",
            resource_id=str(result["run_id"]),
            after={
                "status": result["status"],
                "received_sessions": result["received_sessions"],
                "received_records": result["received_records"],
                "errors": result["errors"],
            },
        )
    return {"success": True, "data": result}


@router.post("/sync")
def sync(
    cursor: str | None = None,
    user: dict = Depends(require_permission("attendance:sync")),
) -> dict:
    """Trigger incremental sync from signin system."""
    return _run_sync(cursor, user["id"], "attendance.sync.trigger")


@router.post("/sync/scheduled")
def scheduled_sync(
    cursor: str | None = None,
    x_api_key: str | None = Header(default=None),
) -> dict:
    """Run the weekday timer pull using the signin service credential."""
    _verify_signin_service_key(x_api_key)
    try:
        result = _run_sync(cursor, None, "attendance.sync.scheduled")
    except HTTPException:
        _write_failure_alert_if_threshold_reached()
        raise
    if result["data"]["status"] == "PARTIAL":
        _write_failure_alert_if_threshold_reached()
    return result


@router.get("/sync/status")
def sync_status(
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """Return privacy-safe health details for the signin synchronization."""
    return {"success": True, "data": _attendance_sync_health()}


@router.get("/reconciliation-summary")
def reconciliation_summary(
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """Return privacy-safe aggregate workload for manual data review."""
    month = month if isinstance(month, str) else None
    return {"success": True, "data": _attendance_reconciliation_summary(month)}


@router.get("/reconciliation-queue")
def reconciliation_queue(
    issue: Literal[
        "unmatched_attendance_records",
        "active_members_missing_phone_hash",
        "active_members_missing_primary_region",
        "active_members_missing_study_class",
        "active_members_missing_study_group",
    ] = "unmatched_attendance_records",
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
    user: dict = Depends(require_permission("members:detail_view")),
) -> dict:
    """Return a paged read-only queue for authorized manual review."""
    month = month if isinstance(month, str) else None
    member_month_clause, member_month_params = _member_activity_month_clause(month)
    unmatched_month_clause = " AND substr(eg.event_date, 1, 7)=?" if month else ""
    queries = {
        "unmatched_attendance_records": {
            "count": "SELECT COUNT(*) AS count FROM attendance_records r "
            "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
            "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
            "WHERE (r.attendance_status='UNMATCHED' OR r.member_id IS NULL)"
            + unmatched_month_clause,
            "count_params": (month,) if month else (),
            "rows": "SELECT r.id, r.member_code_snapshot, r.name_snapshot, "
            "r.attendance_status, r.checked_at, s.session_name, "
            "eg.title, eg.event_date, eg.org_unit_id, eg.study_org_unit_id "
            "FROM attendance_records r JOIN attendance_sessions s "
            "ON s.id=r.attendance_session_id JOIN attendance_event_groups eg "
            "ON eg.id=s.event_group_id WHERE (r.attendance_status='UNMATCHED' "
            "OR r.member_id IS NULL)"
            + unmatched_month_clause
            + " ORDER BY eg.event_date DESC, r.id DESC",
            "row_params": (month,) if month else (),
        },
        "active_members_missing_phone_hash": {
            "count": "SELECT COUNT(*) AS count FROM members "
            "WHERE status='ACTIVE' AND (phone_hash IS NULL OR phone_hash='')"
            + (member_month_clause.replace("m.id", "members.id") if month else ""),
            "count_params": member_month_params,
            "rows": "SELECT id, member_code AS member_code_snapshot, name AS name_snapshot, "
            "status AS attendance_status, NULL AS checked_at, NULL AS session_name, "
            "NULL AS title, NULL AS event_date, org_unit_id, development_org_unit_id AS study_org_unit_id "
            "FROM members WHERE status='ACTIVE' AND (phone_hash IS NULL OR phone_hash='') "
            + (member_month_clause.replace("m.id", "members.id") if month else "")
            + " ORDER BY id DESC",
            "row_params": member_month_params,
        },
        "active_members_missing_primary_region": {
            "count": "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id "
            "AND r.relation_type='PRIMARY_REGION')" + member_month_clause,
            "count_params": member_month_params,
            "rows": "SELECT m.id, m.member_code AS member_code_snapshot, m.name AS name_snapshot, "
            "m.status AS attendance_status, NULL AS checked_at, NULL AS session_name, NULL AS title, "
            "NULL AS event_date, m.org_unit_id, m.development_org_unit_id AS study_org_unit_id "
            "FROM members m WHERE m.status='ACTIVE' AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='PRIMARY_REGION')" + member_month_clause
            + " ORDER BY m.id DESC",
            "row_params": member_month_params,
        },
        "active_members_missing_study_class": {
            "count": "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id "
            "AND r.relation_type='STUDY_CLASS')" + member_month_clause,
            "count_params": member_month_params,
            "rows": "SELECT m.id, m.member_code AS member_code_snapshot, m.name AS name_snapshot, "
            "m.status AS attendance_status, NULL AS checked_at, NULL AS session_name, NULL AS title, "
            "NULL AS event_date, m.org_unit_id, m.development_org_unit_id AS study_org_unit_id "
            "FROM members m WHERE m.status='ACTIVE' AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='STUDY_CLASS')" + member_month_clause
            + " ORDER BY m.id DESC",
            "row_params": member_month_params,
        },
        "active_members_missing_study_group": {
            "count": "SELECT COUNT(*) AS count FROM members m WHERE m.status='ACTIVE' "
            "AND NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id "
            "AND r.relation_type='STUDY_GROUP')" + member_month_clause,
            "count_params": member_month_params,
            "rows": "SELECT m.id, m.member_code AS member_code_snapshot, m.name AS name_snapshot, "
            "m.status AS attendance_status, NULL AS checked_at, NULL AS session_name, NULL AS title, "
            "NULL AS event_date, m.org_unit_id, m.development_org_unit_id AS study_org_unit_id "
            "FROM members m WHERE m.status='ACTIVE' AND NOT EXISTS (SELECT 1 FROM member_org_relations r "
            "WHERE r.member_id=m.id AND r.relation_type='STUDY_GROUP') "
            f"AND NOT ({CURRENT_CLASS_NAME_SQL} IN ('先锋班','神仙班') "
            "OR COALESCE(m.notes,'') LIKE '%目前不读书%')" + member_month_clause
            + " ORDER BY m.id DESC",
            "row_params": member_month_params,
        },
    }[issue]
    total = _execute_count(queries["count"], queries["count_params"])
    row_params = (*queries["row_params"], limit, offset)
    rows = fetch_all(f"{queries['rows']} LIMIT ? OFFSET ?", row_params)
    return {
        "success": True,
        "data": {
            "scope": "MANUAL_REVIEW_READ_ONLY",
            "issue": issue,
            "write_enabled": False,
            "total": total,
            "limit": limit,
            "offset": offset,
            "rows": rows,
        },
    }


@router.get("/reconciliation-breakdown")
def reconciliation_breakdown(
    issue: Literal[
        "unmatched_attendance_records",
        "active_members_missing_phone_hash",
        "active_members_missing_primary_region",
        "active_members_missing_study_class",
        "active_members_missing_study_group",
    ] = "unmatched_attendance_records",
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """Return aggregate issue distribution by primary operating organization."""
    month = month if isinstance(month, str) else None
    if issue == "unmatched_attendance_records":
        statement = (
            "SELECT eg.org_unit_id, o.name AS org_name, COUNT(*) AS count "
            "FROM attendance_records r JOIN attendance_sessions s ON s.id=r.attendance_session_id "
            "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
            "JOIN org_units o ON o.id=eg.org_unit_id "
            "WHERE (r.attendance_status='UNMATCHED' OR r.member_id IS NULL) "
            + ("AND substr(eg.event_date, 1, 7)=? " if month else "")
            + "GROUP BY eg.org_unit_id, o.name ORDER BY count DESC, eg.org_unit_id"
        )
        rows = fetch_all(statement, (month,)) if month else fetch_all(statement)
    else:
        member_month_clause, member_month_params = _member_activity_month_clause(month)
        conditions = {
            "active_members_missing_phone_hash": "m.phone_hash IS NULL OR m.phone_hash=''",
            "active_members_missing_primary_region": "NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id AND r.relation_type='PRIMARY_REGION')",
            "active_members_missing_study_class": "NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id AND r.relation_type='STUDY_CLASS')",
            "active_members_missing_study_group": "NOT EXISTS (SELECT 1 FROM member_org_relations r WHERE r.member_id=m.id AND r.relation_type='STUDY_GROUP') AND NOT ("
            + CURRENT_CLASS_NAME_SQL
            + " IN ('先锋班','神仙班') OR COALESCE(m.notes,'') LIKE '%目前不读书%')",
        }
        statement = (
            "SELECT m.org_unit_id, o.name AS org_name, COUNT(*) AS count "
            "FROM members m JOIN org_units o ON o.id=m.org_unit_id "
            "WHERE m.status='ACTIVE' AND (" + conditions[issue] + ") "
            + member_month_clause
            + "GROUP BY m.org_unit_id, o.name ORDER BY count DESC, m.org_unit_id"
        )
        rows = (
            fetch_all(statement, member_month_params)
            if member_month_params
            else fetch_all(statement)
        )
    return {
        "success": True,
        "data": {"scope": "AGGREGATE_ONLY", "issue": issue, "rows": rows},
    }


@router.get("/event-groups")
def list_event_groups(
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    org_unit_id: str | None = None,
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """List attendance event groups."""
    params: list = []
    sql = (
        "SELECT eg.id, eg.source_key, eg.title, eg.event_date, "
        "eg.activity_type, eg.status, "
        "eg.org_unit_id, o.name AS org_name, o.unit_type AS org_unit_type, "
        "eg.study_org_unit_id, "
        "class_org.name AS class_name, "
        "(SELECT COUNT(*) FROM attendance_sessions s "
        "WHERE s.event_group_id=eg.id) AS session_count, "
        "(SELECT COUNT(*) FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "WHERE s.event_group_id=eg.id) AS record_count, "
        "(SELECT COUNT(*) FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "WHERE s.event_group_id=eg.id "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS present_count "
        ", (SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members cm ON cm.id=mor.member_id "
        "WHERE mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id AND cm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS class_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) "
        "FROM attendance_records r JOIN attendance_sessions s "
        "ON s.id=r.attendance_session_id "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE s.event_group_id=eg.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) "
        "AS class_present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members rm ON rm.id=mor.member_id "
        "WHERE mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id AND rm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS region_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) "
        "FROM attendance_records r JOIN attendance_sessions s "
        "ON s.id=r.attendance_session_id "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE s.event_group_id=eg.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) "
        "AS region_present_count "
        "FROM attendance_event_groups eg "
        "JOIN org_units o ON o.id=eg.org_unit_id "
        "LEFT JOIN org_units class_org ON class_org.id=eg.study_org_unit_id "
        "WHERE 1=1"
    )
    if month:
        sql += " AND substr(eg.event_date, 1, 7)=?"
        params.append(month)
    if org_unit_id:
        sql += " AND eg.org_unit_id=?"
        params.append(org_unit_id)
    allowed = accessible_org_ids(user["id"])
    if allowed is not None:
        if not allowed:
            return {"success": True, "data": []}
        placeholders = ",".join("?" for _ in allowed)
        sql += (
            f" AND (eg.org_unit_id IN ({placeholders}) "
            f"OR eg.study_org_unit_id IN ({placeholders}))"
        )
        params.extend(sorted(allowed))
        params.extend(sorted(allowed))
    sql += " ORDER BY eg.event_date DESC"
    rows = fetch_all(sql, tuple(params))
    return {"success": True, "data": rows}


def _session_display_title(title: str | None, session_name: str | None, session_code: str) -> str:
    """Build a stable, human-readable title for one attendance session.

    The source system groups several sessions under one event title.  Keep the
    source title intact for traceability, but remove a trailing session suffix
    before appending the canonical session label so each list row is distinct.
    """
    source_title = (title or "未命名活动").strip()
    label_map = {
        "MORNING": "上午",
        "AM": "上午",
        "AFTERNOON": "下午",
        "PM": "下午",
        "EVENING": "空巴",
        "KONPA": "空巴",
        "KONPAI": "空巴",
    }
    label = (session_name or "").strip() or label_map.get(session_code.upper(), session_code)
    base = re.sub(
        r"\s*[-－—·]\s*(?:上午|下午|晚上?空巴|空巴|AM|PM|MORNING|AFTERNOON|EVENING|KONPA)\s*$",
        "",
        source_title,
        flags=re.IGNORECASE,
    ).strip()
    if not label or label in base:
        return source_title
    return f"{base} · {label}"


def _as_comparable_datetime(value: object | None) -> datetime | None:
    """Parse a database/source timestamp without changing the stored source fact."""
    if value in (None, ""):
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None
    if parsed.tzinfo is not None:
        return parsed.astimezone(UTC).replace(tzinfo=None)
    return parsed


def _apply_checkin_time_display_policy(row: dict) -> dict:
    """Fail closed for a check-in time before the explicit check-in opening.

    The source snapshot remains untouched.  An invalid source time is not shown
    to operators or exported as a genuine check-in time; it is marked for data
    review instead.  We only compare with an explicit check-in opening, because
    a scheduled start alone does not rule out an approved early arrival.
    """
    display = dict(row)
    checked_at = _as_comparable_datetime(row.get("checked_at"))
    checkin_start = _as_comparable_datetime(row.get("checkin_start_at"))
    if checked_at and checkin_start and checked_at < checkin_start:
        display["checked_at"] = None
        display["checked_at_review_status"] = "TIME_BEFORE_CHECKIN_START"
    else:
        display["checked_at_review_status"] = None
    return display


def _attendance_flag_label(value: object | None, score_present: bool) -> str:
    if not score_present:
        return "—"
    return "是" if bool(value) else "否"


@router.get("/activity-sessions")
def list_activity_sessions(
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    org_unit_id: str | None = None,
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """List one row per attendance session with session-scoped statistics."""
    params: list = []
    sql = (
        "SELECT eg.id, eg.source_key, eg.external_group_id, eg.title, eg.event_date, "
        "eg.activity_type, eg.status, eg.org_unit_id, o.name AS org_name, "
        "o.unit_type AS org_unit_type, eg.study_org_unit_id, class_org.name AS class_name, "
        "s.id AS session_id, s.session_code, s.session_name, s.session_order, "
        "s.scheduled_start_at, s.scheduled_end_at, s.status AS session_status, "
        "1 AS session_count, "
        "(SELECT COUNT(*) FROM attendance_records r WHERE r.attendance_session_id=s.id) AS record_count, "
        "(SELECT COUNT(*) FROM attendance_records r WHERE r.attendance_session_id=s.id "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) FROM member_org_relations mor "
        "JOIN members cm ON cm.id=mor.member_id WHERE mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id AND cm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) AS class_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) FROM attendance_records r "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='STUDY_CLASS' AND mor.org_unit_id=eg.study_org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE r.attendance_session_id=s.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS class_present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) FROM member_org_relations mor "
        "JOIN members rm ON rm.id=mor.member_id WHERE mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id AND rm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) AS region_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) FROM attendance_records r "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='PRIMARY_REGION' AND mor.org_unit_id=eg.org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE r.attendance_session_id=s.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS region_present_count "
        "FROM attendance_event_groups eg "
        "JOIN attendance_sessions s ON s.event_group_id=eg.id "
        "JOIN org_units o ON o.id=eg.org_unit_id "
        "LEFT JOIN org_units class_org ON class_org.id=eg.study_org_unit_id "
        "WHERE 1=1"
    )
    if month:
        sql += " AND substr(eg.event_date, 1, 7)=?"
        params.append(month)
    if org_unit_id:
        sql += " AND eg.org_unit_id=?"
        params.append(org_unit_id)
    allowed = accessible_org_ids(user["id"])
    if allowed is not None:
        if not allowed:
            return {"success": True, "data": []}
        placeholders = ",".join("?" for _ in allowed)
        sql += (
            f" AND (eg.org_unit_id IN ({placeholders}) "
            f"OR eg.study_org_unit_id IN ({placeholders}))"
        )
        params.extend(sorted(allowed))
        params.extend(sorted(allowed))
    sql += " ORDER BY eg.event_date DESC, eg.id DESC, s.session_order, s.id"
    rows = fetch_all(sql, tuple(params))
    for row in rows:
        row["display_title"] = _session_display_title(
            row.get("title"), row.get("session_name"), str(row["session_code"])
        )
    return {"success": True, "data": rows}


@router.get("/event-groups/{group_id}")
def event_group_detail(
    group_id: int,
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """Get event group detail with sessions and score summary."""
    group = fetch_one(
        "SELECT eg.*, o.name AS org_name, o.unit_type AS org_unit_type, "
        "class_org.name AS class_name, "
        "(SELECT COUNT(*) FROM attendance_sessions s "
        "WHERE s.event_group_id=eg.id) AS session_count, "
        "(SELECT COUNT(*) FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "WHERE s.event_group_id=eg.id) AS record_count, "
        "(SELECT COUNT(*) FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "WHERE s.event_group_id=eg.id "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS present_count "
        ", (SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members cm ON cm.id=mor.member_id "
        "WHERE mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id AND cm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS class_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) "
        "FROM attendance_records r JOIN attendance_sessions s "
        "ON s.id=r.attendance_session_id "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE s.event_group_id=eg.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) "
        "AS class_present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members rm ON rm.id=mor.member_id "
        "WHERE mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id AND rm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS region_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) "
        "FROM attendance_records r JOIN attendance_sessions s "
        "ON s.id=r.attendance_session_id "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE s.event_group_id=eg.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) "
        "AS region_present_count "
        "FROM attendance_event_groups eg "
        "JOIN org_units o ON o.id=eg.org_unit_id "
        "LEFT JOIN org_units class_org ON class_org.id=eg.study_org_unit_id "
        "WHERE eg.id=?",
        (group_id,),
    )
    if not group:
        raise HTTPException(404, "活动组不存在")
    allowed = accessible_org_ids(user["id"])
    if not _org_is_allowed(
        group["org_unit_id"], group.get("study_org_unit_id"), allowed
    ):
        raise HTTPException(403, "不在组织授权范围内")
    sessions = fetch_all(
        "SELECT s.*, "
        "(SELECT COUNT(*) FROM attendance_records r WHERE r.attendance_session_id=s.id) AS record_count, "
        "(SELECT COUNT(*) FROM attendance_records r WHERE r.attendance_session_id=s.id "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS present_count, "
        "(SELECT COUNT(DISTINCT r.member_id) FROM attendance_records r "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE r.attendance_session_id=s.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) AS class_present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members cm ON cm.id=mor.member_id "
        "WHERE mor.relation_type='STUDY_CLASS' "
        "AND mor.org_unit_id=eg.study_org_unit_id AND cm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS class_member_count, "
        "(SELECT COUNT(DISTINCT r.member_id) FROM attendance_records r "
        "JOIN member_org_relations mor ON mor.member_id=r.member_id "
        "AND mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date) "
        "WHERE r.attendance_session_id=s.id AND r.member_id IS NOT NULL "
        "AND r.participant_type='MEMBER' "
        "AND r.attendance_status IN ('PRESENT','MANUAL_PRESENT')) "
        "AS region_present_count, "
        "(SELECT COUNT(DISTINCT mor.member_id) "
        "FROM member_org_relations mor JOIN members rm ON rm.id=mor.member_id "
        "WHERE mor.relation_type='PRIMARY_REGION' "
        "AND mor.org_unit_id=eg.org_unit_id AND rm.status='ACTIVE' "
        "AND (mor.valid_from IS NULL OR mor.valid_from<=eg.event_date) "
        "AND (mor.valid_until IS NULL OR mor.valid_until>=eg.event_date)) "
        "AS region_member_count, "
        "(SELECT SUM(sr.final_points) FROM attendance_score_records sr "
        "JOIN attendance_records r ON r.id=sr.attendance_record_id "
        "WHERE r.attendance_session_id=s.id) AS total_points "
        "FROM attendance_sessions s "
        "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
        "WHERE s.event_group_id=? "
        "ORDER BY s.session_order",
        (group_id,),
    )
    class_breakdown: list[dict] = []
    if not group.get("study_org_unit_id"):
        class_breakdown = fetch_all(
            "SELECT class_org.id AS class_org_unit_id, class_org.name AS class_name, "
            "COUNT(DISTINCT cm.id) AS class_member_count, "
            "COUNT(DISTINCT CASE WHEN r.attendance_status IN ('PRESENT','MANUAL_PRESENT') "
            "THEN r.member_id END) AS class_present_count "
            "FROM org_units class_org "
            "JOIN member_org_relations class_rel "
            "ON class_rel.org_unit_id=class_org.id "
            "AND class_rel.relation_type='STUDY_CLASS' "
            "AND (class_rel.valid_from IS NULL OR class_rel.valid_from<=?) "
            "AND (class_rel.valid_until IS NULL OR class_rel.valid_until>=?) "
            "JOIN members cm ON cm.id=class_rel.member_id AND cm.status='ACTIVE' "
            "JOIN member_org_relations region_rel "
            "ON region_rel.member_id=cm.id "
            "AND region_rel.relation_type='PRIMARY_REGION' "
            "AND region_rel.org_unit_id=? "
            "AND (region_rel.valid_from IS NULL OR region_rel.valid_from<=?) "
            "AND (region_rel.valid_until IS NULL OR region_rel.valid_until>=?) "
            "LEFT JOIN attendance_records r "
            "ON r.member_id=cm.id AND r.participant_type='MEMBER' "
            "AND r.attendance_session_id IN "
            "(SELECT s.id FROM attendance_sessions s WHERE s.event_group_id=?) "
            "WHERE class_org.unit_type='CLASS' AND class_org.is_active=1 "
            "GROUP BY class_org.id, class_org.name "
            "ORDER BY class_org.name, class_org.id",
            (
                group["event_date"],
                group["event_date"],
                group["org_unit_id"],
                group["event_date"],
                group["event_date"],
                group_id,
            ),
        )
    return {
        "success": True,
        "data": {
            "group": group,
            "sessions": sessions,
            "class_breakdown": class_breakdown,
        },
    }


@router.get("/event-groups/{group_id}/records.xlsx")
def download_event_group_records(
    group_id: int,
    session_id: int | None = None,
    user: dict = Depends(require_permission("members:read")),
) -> StreamingResponse:
    """Download a privacy-minimized Excel workbook of read-only sign-in details."""
    group = fetch_one(
        "SELECT id, title, event_date, org_unit_id, study_org_unit_id "
        "FROM attendance_event_groups WHERE id=?",
        (group_id,),
    )
    if not group:
        raise HTTPException(404, "活动组不存在")
    allowed = accessible_org_ids(user["id"])
    if not _org_is_allowed(
        group["org_unit_id"], group.get("study_org_unit_id"), allowed
    ):
        raise HTTPException(403, "不在组织授权范围内")
    params: list = [group_id]
    session_clause = ""
    if session_id is not None:
        session = fetch_one(
            "SELECT id FROM attendance_sessions WHERE id=? AND event_group_id=?",
            (session_id, group_id),
        )
        if not session:
            raise HTTPException(404, "活动场次不存在")
        session_clause = " AND s.id=?"
        params.append(session_id)

    sql = (
        "SELECT r.name_snapshot, r.member_code_snapshot, r.participant_type, "
        "r.attendance_status, r.checked_at, s.session_code, s.session_name, "
        "s.checkin_start_at, eg.title, eg.event_date, sr.final_points, sr.is_late, "
        "sr.is_early_leave, sr.id AS score_record_id "
        "FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
        "LEFT JOIN attendance_score_records sr ON sr.attendance_record_id=r.id "
        "WHERE eg.id=?" + session_clause
    )
    if allowed is not None:
        if not allowed:
            rows = []
        else:
            placeholders = ",".join("?" for _ in allowed)
            sql += (
                f" AND (eg.org_unit_id IN ({placeholders}) "
                f"OR eg.study_org_unit_id IN ({placeholders}))"
            )
            params.extend(sorted(allowed))
            params.extend(sorted(allowed))
            sql += " ORDER BY eg.event_date DESC, s.session_order, r.name_snapshot"
            rows = fetch_all(sql, tuple(params))
    else:
        sql += " ORDER BY eg.event_date DESC, s.session_order, r.name_snapshot"
        rows = fetch_all(sql, tuple(params))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "签到明细"
    headers = [
        "活动",
        "活动日期",
        "场次",
        "姓名",
        "学员编号",
        "参与类型",
        "签到状态",
        "签到时间",
        "迟到",
        "早退",
        "积分",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    participant_labels = {"MEMBER": "学员", "GUEST": "嘉宾", "OBSERVER": "旁听"}
    status_labels = {
        "PRESENT": "已签到",
        "MANUAL_PRESENT": "人工确认",
        "ABSENT": "未签到",
        "LEAVE": "请假",
        "INVALIDATED": "已作废",
        "UNMATCHED": "待匹配",
    }
    for source_row in rows:
        row = _apply_checkin_time_display_policy(source_row)
        score_present = row.get("score_record_id") is not None
        sheet.append(
            [
                _session_display_title(
                    row.get("title"), row.get("session_name"), str(row["session_code"])
                ),
                row.get("event_date") or "",
                row.get("session_name") or row.get("session_code") or "",
                row.get("name_snapshot") or "",
                row.get("member_code_snapshot") or "",
                participant_labels.get(str(row.get("participant_type") or ""), row.get("participant_type") or ""),
                status_labels.get(str(row.get("attendance_status") or ""), row.get("attendance_status") or ""),
                "待核对" if row.get("checked_at_review_status") else row.get("checked_at") or "",
                _attendance_flag_label(row.get("is_late"), score_present),
                _attendance_flag_label(row.get("is_early_leave"), score_present),
                row.get("final_points") if score_present else "—",
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 36)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = f"-{session_id}" if session_id is not None else "-all"
    filename = f"attendance-details-{group_id}{suffix}.xlsx"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/records")
def list_records(
    event_group_id: int | None = None,
    session_id: int | None = None,
    member_id: int | None = None,
    session_code: str | None = None,
    attendance_status: str | None = None,
    is_late: bool | None = None,
    user: dict = Depends(require_permission("members:read")),
) -> dict:
    """List attendance records with filters."""
    params: list = []
    sql = (
        "SELECT r.id, r.attendance_session_id, r.member_id, r.member_code_snapshot, "
        "r.name_snapshot, r.participant_type, r.score_eligible, r.attendance_status, "
        "r.checked_at, r.checkin_source, s.session_code, s.session_name, s.checkin_start_at, "
        "eg.title, eg.event_date, eg.org_unit_id, eg.study_org_unit_id, "
        "sr.final_points, sr.is_late, sr.is_early_leave, sr.id AS score_record_id "
        "FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
        "LEFT JOIN attendance_score_records sr ON sr.attendance_record_id=r.id "
        "WHERE 1=1"
    )
    if event_group_id:
        sql += " AND s.event_group_id=?"
        params.append(event_group_id)
    if session_id:
        sql += " AND r.attendance_session_id=?"
        params.append(session_id)
    if member_id:
        sql += " AND r.member_id=?"
        params.append(member_id)
    if session_code:
        sql += " AND s.session_code=?"
        params.append(session_code)
    if attendance_status:
        sql += " AND r.attendance_status=?"
        params.append(attendance_status)
    if is_late is not None:
        sql += " AND sr.is_late=?"
        params.append(1 if is_late else 0)

    allowed = accessible_org_ids(user["id"])
    if allowed is not None:
        if not allowed:
            return {"success": True, "data": []}
        placeholders = ",".join("?" for _ in allowed)
        sql += (
            f" AND (eg.org_unit_id IN ({placeholders}) "
            f"OR eg.study_org_unit_id IN ({placeholders}))"
        )
        params.extend(sorted(allowed))
        params.extend(sorted(allowed))

    sql += " ORDER BY eg.event_date DESC, s.session_order, r.name_snapshot"
    rows = [
        _apply_checkin_time_display_policy(row)
        for row in fetch_all(sql, tuple(params))
    ]
    return {"success": True, "data": rows}


@router.post("/records/{record_id}/adjudications")
def create_adjudication(
    record_id: int,
    payload: AdjudicationPayload,
    user: dict = Depends(require_permission("attendance:adjudicate")),
) -> dict:
    """Create an adjudication for an attendance record."""
    record = fetch_one(
        "SELECT r.id, r.attendance_session_id, r.member_id, r.attendance_status, "
        "r.checked_at, eg.org_unit_id, eg.study_org_unit_id "
        "FROM attendance_records r "
        "JOIN attendance_sessions s ON s.id=r.attendance_session_id "
        "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
        "WHERE r.id=?",
        (record_id,),
    )
    if not record:
        raise HTTPException(404, "出勤记录不存在")
    allowed = accessible_org_ids(user["id"])
    if not _org_is_allowed(
        record["org_unit_id"], record.get("study_org_unit_id"), allowed
    ):
        raise HTTPException(403, "不在组织授权范围内")

    if payload.adjudication_type == "MEMBER_RELINK":
        if payload.member_id is None:
            raise HTTPException(400, "MEMBER_RELINK 必须指定 member_id")
        member = fetch_one(
            "SELECT id, org_unit_id FROM members WHERE id=?", (payload.member_id,)
        )
        if not member:
            raise HTTPException(400, "目标学长不存在")
        if allowed is not None and member["org_unit_id"] not in allowed:
            raise HTTPException(403, "目标学长不在组织授权范围内")

    now = datetime.now(UTC).isoformat()
    cancel_targets = {
        "CANCEL_EARLY_LEAVE": "EARLY_LEAVE",
        "CANCEL_LEAVE": "LEAVE",
    }
    status_updates = {
        "MANUAL_CHECKIN": "MANUAL_PRESENT",
        "INVALIDATE_CHECKIN": "INVALIDATED",
        "LEAVE": "LEAVE",
        "CANCEL_LEAVE": "PRESENT" if record["checked_at"] else "ABSENT",
    }

    with transaction() as connection:
        target_type = cancel_targets.get(payload.adjudication_type)
        if target_type:
            cancelled = execute(
                connection,
                "UPDATE attendance_adjudications SET superseded_at=? "
                "WHERE attendance_record_id=? AND adjudication_type=? "
                "AND superseded_at IS NULL",
                (now, record_id, target_type),
            )
            if cancelled.rowcount == 0:
                raise HTTPException(400, f"没有可撤销的 {target_type} 裁定")
        else:
            execute(
                connection,
                "UPDATE attendance_adjudications SET superseded_at=? "
                "WHERE attendance_record_id=? AND adjudication_type=? "
                "AND superseded_at IS NULL",
                (now, record_id, payload.adjudication_type),
            )
        cursor = execute(
            connection,
            "INSERT INTO attendance_adjudications"
            "(attendance_record_id, adjudication_type, occurred_at, reason, actor_user_id, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (record_id, payload.adjudication_type, payload.occurred_at or now,
             payload.reason, user["id"], now),
        )
        adj_id = cursor.lastrowid
        if payload.adjudication_type in status_updates:
            execute(
                connection,
                "UPDATE attendance_records SET attendance_status=?, updated_at=? WHERE id=?",
                (status_updates[payload.adjudication_type], now, record_id),
            )
        elif payload.adjudication_type == "MEMBER_RELINK":
            execute(
                connection,
                "UPDATE attendance_records SET member_id=?, "
                "attendance_status=CASE WHEN attendance_status='UNMATCHED' "
                "THEN CASE WHEN checked_at IS NULL THEN 'ABSENT' ELSE 'PRESENT' END "
                "ELSE attendance_status END, score_eligible=1, updated_at=? WHERE id=?",
                (payload.member_id, now, record_id),
            )
        write_audit(
            connection,
            actor_user_id=user["id"],
            action="attendance.adjudication.create",
            resource_type="attendance_record",
            resource_id=str(record_id),
            org_unit_id=record["org_unit_id"],
            purpose=payload.reason,
            before={
                "member_id": record["member_id"],
                "attendance_status": record["attendance_status"],
            },
            after={
                "adjudication_id": adj_id,
                "type": payload.adjudication_type,
                "member_id": payload.member_id or record["member_id"],
                "attendance_status": status_updates.get(
                    payload.adjudication_type, record["attendance_status"]
                ),
            },
        )

    # Recalculate score for this record
    session = fetch_one(
        "SELECT s.session_code, s.scheduled_start_at, eg.activity_type "
        "FROM attendance_sessions s "
        "JOIN attendance_event_groups eg ON eg.id=s.event_group_id "
        "WHERE s.id=?",
        (record["attendance_session_id"],),
    )
    rec = fetch_one("SELECT * FROM attendance_records WHERE id=?", (record_id,))
    activity_type = (
        normalize_activity_type(session["activity_type"]) if session else None
    )
    if (
        session
        and rec
        and score_is_applicable(
            rec["member_id"], bool(rec["score_eligible"]), activity_type
        )
        and get_active_rule(session["session_code"], activity_type)
    ):
        upsert_score_record(
            attendance_record_id=record_id,
            member_id=rec["member_id"],
            session_code=session["session_code"],
            attendance_status=rec["attendance_status"],
            checked_at=rec["checked_at"],
            scheduled_start_at=session["scheduled_start_at"],
            score_eligible=bool(rec["score_eligible"]),
            activity_type=activity_type,
        )

    return {"success": True, "data": {"id": adj_id}}


@router.post("/event-groups/{group_id}/recalculate")
def recalculate(
    group_id: int,
    user: dict = Depends(require_permission("attendance:adjudicate")),
) -> dict:
    """Recalculate all scores for an event group."""
    group = fetch_one(
        "SELECT org_unit_id, study_org_unit_id "
        "FROM attendance_event_groups WHERE id=?",
        (group_id,),
    )
    if not group:
        raise HTTPException(404, "活动组不存在")
    allowed = accessible_org_ids(user["id"])
    if not _org_is_allowed(
        group["org_unit_id"], group.get("study_org_unit_id"), allowed
    ):
        raise HTTPException(403, "不在组织授权范围内")
    result = recalculate_event_group(group_id)
    with transaction() as connection:
        write_audit(
            connection,
            actor_user_id=user["id"],
            action="attendance.score.recalculate",
            resource_type="attendance_event_group",
            resource_id=str(group_id),
            org_unit_id=group["org_unit_id"],
            after=result,
        )
    return {"success": True, "data": result}
