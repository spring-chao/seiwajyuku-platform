"""One privacy-minimized workbook shared by onsite and all-session downloads."""
from __future__ import annotations
from io import BytesIO
from datetime import datetime,timedelta,timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from app.services import signin_engine as engine
from app.services.attendance_export_context import _hydrate, participation_statistics, apply_session_requirements

DETAIL_HEADERS = ['场次','实际到场姓名','原报名姓名','联系方式（脱敏）','公司','分中心','班级','小组','班会小组参会率','签到身份','原班级','桌号','签到状态','签到时间','统计范围说明']


def _display_time(value):
    if not value:return ''
    try:
        parsed=datetime.fromisoformat(str(value).replace('Z','+00:00'))
        if parsed.tzinfo:
            return parsed.astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
    except ValueError:
        pass
    return str(value)


def _session_name(event):
    return event.get('session_name') or {'MORNING':'上午','AFTERNOON':'下午','KONPA':'晚上空巴'}.get(str(event.get('session_code','')).upper(), '整场')


def _read(operation,payload,user,allowed):
    return engine.engine_request('/ops/v1/manage/'+operation,{'payload':payload,'actor':{'id':user['id'],'permissions':['attendance:view' if operation=='stats' else 'attendance:export']},'allowed_org_unit_ids':allowed})


def live_dataset(event_ids, *, user, expected=None):
    if 'attendance:export' not in user.get('permissions',[]):
        raise engine.SigninEngineError('需要活动签到导出权限，请在员工管理中授权活动签到运营能力',403)
    engine.enabled()
    allowed = engine.allowed_scope(user['id'],'attendance:export')
    if not event_ids or len(event_ids)>30 or len(set(event_ids))!=len(event_ids):
        raise engine.SigninEngineError('场次编号为空、重复或超过30场，请缩小范围',400)
    bundles=[]
    for event_id in event_ids:
        data = _read('export',{'event_id':event_id},user,allowed)
        event = data.get('event') or engine.event_detail(event_id,user=user,permission='attendance:export')
        engine.validate_organizations(event,allowed,required=True)
        if str(event.get('event_id')) != str(event_id):
            raise engine.SigninEngineError('导出场次发生变化，请刷新重试',409)
        if bundles:
            first=bundles[0]['event']
            for field in ('event_group_id','class_org_unit_id','org_unit_id','event_date'):
                if str(event.get(field) or '') != str(first.get(field) or ''):
                    raise engine.SigninEngineError('导出场次不属于同一活动、班级和日期',409)
        if expected and (str(event.get('event_date',''))[:10] != str(expected['event_date'])[:10] or
                expected.get('study_org_unit_id') and (event.get('class_org_unit_id') or event.get('org_unit_id')) != expected['study_org_unit_id']):
            raise engine.SigninEngineError('实时场次与所选活动的日期或班级不一致，请先核对关联',409)
        rows = data.get('rows')
        if not isinstance(rows,list):
            raise engine.SigninEngineError('实时名单不完整，本次未生成导出文件',502)
        if sum(len(b['rows']) for b in bundles)+len(rows)>15000:
            raise engine.SigninEngineError('导出超过15000条，请按场次导出',400)
        _hydrate(rows,event_id)
        stats = _read('stats',{'event_id':event_id},user,allowed)
        additions = [row for row in rows if row.get('_staff_makeup_added')]
        for row in additions:
            home = row.get('attendance_role') == 'HOME_CLASS_MEMBER' or (row.get('current_class_org_unit_id') or row.get('class_org_unit_id')) == (event.get('class_org_unit_id') or event.get('org_unit_id'))
            key = 'home_class_checked' if home else 'cross_class_checked'
            for field in (key, 'checked', 'onsite_total'):
                if isinstance(stats.get(field), (int, float)):
                    stats[field] += 1
        if additions and stats.get('home_class_total'):
            stats['rate'] = stats.get('home_class_checked', 0) * 100 / stats['home_class_total']
        bundles.append({'event':event,'rows':engine.redact(rows),'stats':engine.redact(stats)})
    shared = participation_statistics(event_id=event_ids[0],user=user,permission='attendance:export') if str(bundles[0]['event'].get('activity_type','')).lower()=='class_meeting' else {}
    for bundle in bundles:
        apply_session_requirements(bundle['stats'],shared,str(bundle['event']['event_id']))
    from app.db import transaction
    from app.services.learning_requirement_policy import excluded_members
    ids = [r['current_member_id'] for b in bundles for r in b['rows'] if r.get('current_member_id')]
    with transaction() as connection:
        excluded = excluded_members(connection,ids,metric='class_meeting') if shared else {}
    rates={g['group_org_unit_id']:g['participation_rate'] for g in shared.get('group_statistics',[])}
    for bundle in bundles:
        for row in bundle['rows']:
            row['statistics_exclusion_reason']=excluded.get(row.get('current_member_id'),'')
            row['class_group_participation_rate']=None if row['statistics_exclusion_reason'] else rates.get(row.get('current_group_org_unit_id'))
    return {'sessions':bundles,**shared,'data_source':'LIVE_SIGNIN'}


def _append_sheet(book,name,headers,rows,percent_columns=()):
    sheet=book.create_sheet(name);sheet.append(headers)
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in sheet[1]:cell.font=Font(bold=True)
    for col in percent_columns:
        for row in sheet.iter_rows(min_row=2,min_col=col,max_col=col):
            if isinstance(row[0].value,(int,float)):row[0].number_format='0.0%'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=min(50,max(15,max(len(str(c.value or '')) for c in column)+2))
    for row in sheet:
        for cell in row:cell.alignment=Alignment(vertical='center',wrap_text=True)
    return sheet


def workbook_bytes(data):
    book=Workbook();book.remove(book.active)
    details=[];session_rows=[]
    for bundle in data['sessions']:
        event,stats=bundle['event'],bundle['stats'];label=_session_name(event)
        for row in bundle['rows']:
            group_name=row.get('current_group_name') if row.get('current_membership_resolved') else row.get('ops_study_group_name') or row.get('group_name')
            if group_name in {row.get('class_name'),row.get('home_class_name')}:group_name=None
            details.append([label,row.get('actual_attendee_name') or row.get('registered_display_name') or row.get('name') or '',row.get('registered_display_name') or row.get('registered_name') or row.get('name') or '',row.get('phone_masked') or '',row.get('company') or '',row.get('center') or '',row.get('class_name') or '',group_name or '未分组',row.get('class_group_participation_rate') if row.get('class_group_participation_rate') is not None else '—',row.get('attendance_role_label') or row.get('attendance_role') or '',row.get('home_class_name') or '',row.get('dinner_table_num') or '',row.get('sign_status') or row.get('attendance_status_label') or ('已签到' if row.get('checked') or row.get('sign_time') or row.get('checked_at') else '未签到'),_display_time(row.get('sign_time') or row.get('checked_at')),row.get('statistics_exclusion_reason') or ''])
        session_rows.append([label,stats.get('home_class_total',stats.get('total')),stats.get('home_class_checked',stats.get('checked')),stats.get('cross_class_checked'),stats.get('guest_checked'),stats.get('onsite_total',stats.get('checked')),stats['rate']/100 if isinstance(stats.get('rate'),(int,float)) else None,stats.get('requirement_roster_count'),stats.get('requirement_present_count'),stats.get('requirement_participation_rate')])
    _append_sheet(book,'签到名单',DETAIL_HEADERS,details,(9,))
    _append_sheet(book,'场次汇总',['场次','本班报名（含自愿参加）','本班已签到（报名口径）','外班实际到场','来宾实际到场','现场真实人数','报名签到率','应参人数（排除精进组）','应参到场（含本班师资）','班级参会率（含本班师资，排除精进组）'],session_rows,(7,10))
    total=data.get('class_statistics') or {};groups=data.get('group_totals') or {}
    summary = _append_sheet(book,'签到汇总',['统计项','数值'],[
        ['统计来源','实时签到；报名明细逐场保留，班会整体仅上午下午合并去重'],
        ['班会应参人数（排除精进组）',total.get('roster_count')],['在册学长去重到场',total.get('member_present_count')],['额外正式本班师资到场',total.get('staff_present_count')],['整体去重到场',total.get('present_count')],['总参与率（含本班师资）',total.get('participation_rate')],['总小组参与率',groups.get('participation_rate')],['统计口径',data.get('group_statistics_basis','按所选活动场次')],['待核对',data.get('group_statistics_message','')]])
    if not total:
        summary.delete_rows(2,summary.max_row)
        for values in [['统计来源','实时签到，按当前场次报名和签到统计'],['场次数',len(data['sessions'])]]:
            summary.append(values)
        if len(data['sessions'])==1:
            stats=data['sessions'][0]['stats']
            for values in [['报名人数',stats.get('total',stats.get('home_class_total'))],['已签到人数',stats.get('checked',stats.get('home_class_checked'))],['总参与率',stats['rate']/100 if isinstance(stats.get('rate'),(int,float)) else None]]:
                summary.append(values)
    for row in summary.iter_rows(min_row=2):
        if '率' in str(row[0].value) and isinstance(row[1].value,(int,float)):
            row[1].number_format='0.0%'
    rows=[[g['group_name'],g['roster_count'],g['present_count'],g['participation_rate']] for g in data.get('group_statistics',[])]
    if total.get('staff_present_count'):rows.append(['本班师资（额外去重到场）',None,total['staff_present_count'],None])
    if total:rows.append([(total.get('class_name') or '班级')+'总参与率',total.get('roster_count'),total.get('present_count'),total.get('participation_rate')])
    if total or rows:_append_sheet(book,'小组参会统计',['小组','本组应参人数','本次班会去重到场人数','小组参与率／班级总参与率'],rows,(4,))
    output=BytesIO();book.save(output);return output.getvalue()


def live_workbook(event_ids, *, user, expected=None):
    content = workbook_bytes(live_dataset(event_ids,user=user,expected=expected))
    from app.db import transaction
    from app.services.audit import write_audit
    with transaction() as connection:
        write_audit(connection,actor_user_id=user['id'],action='attendance.workbook_export',resource_type='signin_event',resource_id=event_ids[0],after={'session_count':len(event_ids),'data_source':'LIVE_SIGNIN','sensitive_contact':False})
    return content
