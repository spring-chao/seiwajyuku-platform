"""Confirmed opening balances: source preview, audited review, atomic post.

An as-of balance is not a course event or an annual/monthly learning score.
The separate append-only ledger preserves that distinction and works at the
original production schema boundary, without repricing historic credits.
"""
from __future__ import annotations

import hashlib
import io
import json
import sqlite3
import threading
import zipfile
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.core.settings import get_settings
from app.core.build_info import get_build_info
from app.db import connect, execute, transaction
from app.migrations import MIGRATION_ROOT, _split_mysql
from app.services.audit import write_audit
from app.services.iam import accessible_org_ids, user_context
from app.services.learning_credits import LearningCreditError, _db_timestamp, _write_allowed
from pathlib import Path

MIGRATION = "0069_learning_credit_opening_balances.sql"
MIGRATION_HASHES = {
    "mysql": "56227861767ab78f2cd0ff136a6e947943b35e87674376017e4b1594f26c7224",
    "sqlite": "f8751617eccbfa95e4b40308eadcd8c09a423c1da00e91793a59c9fb28b33d56",
}
TABLES = ("learning_credit_opening_imports", "learning_credit_opening_rows", "learning_credit_opening_balances")
HEADERS = ("学长编号", "姓名", "班级/分中心", "累计学分")
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
_lock = threading.Lock()


def _permission(actor: int, *permissions: str) -> dict:
    user = user_context(actor) or {}
    if not user or any(p not in user.get("permissions", []) for p in permissions):
        raise PermissionError("当前账号没有这项学分操作权限，请由已授权人员处理")
    return user


OPENING_PERMISSION = "plans:credit_opening_manage"


def _manage(actor: int) -> dict:
    user = _permission(actor, "members:read")
    if OPENING_PERMISSION not in user["permissions"] and "plans:historical_credit_import_manage" not in user["permissions"]:
        raise PermissionError("当前账号没有期初学分管理权限")
    return user


def _scope(actor: int):
    user = _manage(actor)
    capability = OPENING_PERMISSION if OPENING_PERMISSION in user["permissions"] else "plans:historical_credit_import_manage"
    return accessible_org_ids(actor, capability)


def _tables(connection) -> set[str]:
    if isinstance(connection, sqlite3.Connection):
        rows = execute(connection, "SELECT name AS table_name FROM sqlite_master WHERE type='table'").fetchall()
    else:
        rows = execute(connection, "SELECT table_name AS table_name FROM information_schema.tables WHERE table_schema=DATABASE()").fetchall()
    return {r["table_name"] for r in rows} & set(TABLES)


def storage_available(connection) -> bool:
    if _tables(connection) != set(TABLES):
        return False
    expected = (
        {"id", "content_fingerprint", "file_sha256", "original_filename", "cutoff_date", "status", "created_by", "approved_by", "posted_by", "created_at", "approved_at", "posted_at"},
        {"id", "import_id", "member_id", "org_unit_id", "member_name", "class_label", "excel_row", "points"},
        {"member_id", "import_id", "org_unit_id", "points", "cutoff_date", "posted_by", "posted_at"},
    )
    return all({c[0] for c in execute(connection, "SELECT * FROM " + table + " WHERE 1=0").description} == fields for table, fields in zip(TABLES, expected))


def guard_regular_post(connection, item: dict, occurrence: tuple) -> None:
    """Serialize with opening posts and refuse overlapping learning facts."""
    if not storage_available(connection):
        return
    execute(connection, "SELECT id FROM members WHERE id=?" + ("" if isinstance(connection, sqlite3.Connection) else " FOR UPDATE"), (item["member_id"],)).fetchone()
    row = execute(connection, "SELECT cutoff_date FROM learning_credit_opening_balances WHERE member_id=?", (item["member_id"],)).fetchone()
    if not row or item.get("source_type") == "REVERSAL":
        return
    cutoff = date.fromisoformat(str(row["cutoff_date"])[:10])
    occurred_at, precision, year, month = occurrence
    if precision == "EXACT_DATE":
        parsed = datetime.fromisoformat(str(occurred_at).replace("Z", "+00:00"))
        from datetime import UTC
        effective = (parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else parsed).astimezone(ZoneInfo("Asia/Shanghai")).date()
    else:
        effective = date(int(year), int(month or 1), 1)
    if effective <= cutoff:
        raise LearningCreditError("这条学分的期间已包含在期初余额中，不能重复入账；请核对统计截止日")


def opening_summary(connection, *, member_id: int | None = None, allowed: set[str] | None = None) -> dict:
    empty = {"entry_count": 0, "total_points": "0.00", "cutoff_date": None}
    if not storage_available(connection) or allowed == set() or not execute(connection, "SELECT version FROM schema_migrations WHERE version=?", (MIGRATION,)).fetchone():
        return empty
    conditions, params = [], []
    if member_id is not None:
        conditions.append("member_id=?")
        params.append(member_id)
    if allowed is not None:
        conditions.append("org_unit_id IN (" + ",".join("?" for _ in allowed) + ")")
        params.extend(sorted(allowed))
    from app.services import credit_year_reconciliation as reconciliation
    ledger = 'learning_credit_opening_balances'
    if reconciliation.available(connection):
        ledger = "(SELECT member_id,org_unit_id,points,cutoff_date FROM learning_credit_opening_balances UNION ALL SELECT a.member_id,a.org_unit_id,a.points,a.cutoff_date FROM learning_credit_opening_adjustments a JOIN learning_credit_opening_imports i ON i.id=a.import_id AND i.status='POSTED' JOIN learning_credit_year_imports y ON y.id=a.year_import_id AND y.status='POSTED') confirmed_opening"
    row = execute(connection, "SELECT COUNT(DISTINCT member_id) AS n,COALESCE(SUM(points),0) AS points,"
                  "MAX(cutoff_date) AS cutoff_date FROM " + ledger +
                  (" WHERE " + " AND ".join(conditions) if conditions else ""), tuple(params)).fetchone()
    return {"entry_count": int(row["n"]), "total_points": format(Decimal(str(row["points"])), ".2f"),
            "cutoff_date": str(row["cutoff_date"])[:10] if row["cutoff_date"] else None}


@contextmanager
def _operation(connection):
    sqlite = isinstance(connection, sqlite3.Connection)
    if not sqlite:
        execute(connection, "SET TRANSACTION ISOLATION LEVEL READ COMMITTED")
    acquired = _lock.acquire(blocking=False) if sqlite else int(execute(
        connection, "SELECT GET_LOCK('seiwajyuku:credit-opening',0) AS acquired").fetchone()["acquired"] or 0) == 1
    if not acquired:
        raise LearningCreditError("其他学分操作正在处理，请稍后刷新状态")
    try:
        yield
    finally:
        if sqlite:
            _lock.release()
        else:
            execute(connection, "SELECT RELEASE_LOCK('seiwajyuku:credit-opening')")


def _roster(connection, actor: int) -> dict[str, dict]:
    allowed = _scope(actor)
    if allowed == set():
        return {}
    rows = execute(connection, "SELECT m.id,m.member_code,m.name,m.org_unit_id,o.name AS org_name "
                   "FROM members m JOIN org_units o ON o.id=m.org_unit_id "
                   "WHERE m.status='ACTIVE' AND o.is_active=1 ORDER BY m.id").fetchall()
    relations = execute(connection, "SELECT r.member_id,r.org_unit_id,o.name FROM member_org_relations r "
                        "JOIN org_units o ON o.id=r.org_unit_id WHERE r.relation_type='STUDY_CLASS' "
                        "AND o.is_active=1 AND (r.valid_from IS NULL OR r.valid_from<=?) "
                        "AND (r.valid_until IS NULL OR r.valid_until>=?)", (datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat(),) * 2).fetchall()
    classes: dict[int, list] = {}
    for r in relations:
        classes.setdefault(int(r["member_id"]), []).append(dict(r))
    result = {}
    for r in rows:
        r = dict(r)
        current = sorted(classes.get(int(r["id"]), []), key=lambda c: c["org_unit_id"])
        # One balance belongs to one accountable organization. Ambiguous class
        # membership is resolved in the roster, never by fuzzy name matching.
        r["ambiguous_class"] = len(current) > 1
        r["class_label"] = " / ".join(c["name"] for c in current) if current else r["org_name"]
        r["org_unit_id"] = current[0]["org_unit_id"] if len(current) == 1 else r["org_unit_id"]
        if allowed is not None and r["org_unit_id"] not in allowed:
            continue
        result[str(r["member_code"])] = r
    return result


def template_bytes(actor: int) -> bytes:
    _manage(actor)
    if OPENING_PERMISSION not in _manage(actor)["permissions"]:
        _permission(actor, "exports:normal")
    connection = connect()
    try:
        roster = _roster(connection, actor)
        if len(roster) > MAX_ROWS:
            raise LearningCreditError("当前范围超过5000人，请缩小组织范围后下载")
        template = load_workbook(Path(__file__).resolve().parents[1] / "assets/credit-opening-template.xlsx")
        sheet = template["期初学分"]
        for r in roster.values():
            # Literal text prevents spreadsheet formula injection in roster labels.
            sheet.append([str(r["member_code"]), r["name"], r["class_label"], None])
            for cell in sheet[sheet.max_row][:3]:
                cell.data_type = "s"
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            row[0].number_format = "@"
            row[3].number_format = "0.00"
        output = io.BytesIO()
        template.save(output)
        write_audit(connection, actor_user_id=actor, action="learning.credit_opening.template",
                    resource_type="credit_opening", after={"roster_count": len(roster)}, purpose="期初余额模板（不含联系方式）")
        connection.commit()
        return output.getvalue()
    finally:
        connection.close()


def _parse(content: bytes, cutoff_date: str) -> tuple[date, list[dict], int]:
    try:
        cutoff = date.fromisoformat(cutoff_date)
    except ValueError as exc:
        raise LearningCreditError("请选择有效的统计截止日期") from exc
    if cutoff.year < 2000 or cutoff > datetime.now(ZoneInfo("Asia/Shanghai")).date():
        raise LearningCreditError("截止日期必须在2000年以后，且不能晚于今天")
    if not content or len(content) > MAX_BYTES:
        raise LearningCreditError("只支持5MB以内的Excel文件")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            files = archive.infolist()
            if sum(f.file_size for f in files) > 50 * 1024 * 1024 or len(files) > 1000:
                raise LearningCreditError("Excel内容过大，请使用下载的模板分批上传")
            if any("externallinks" in f.filename.lower() or "vbaproject" in f.filename.lower() for f in files):
                raise LearningCreditError("不接受宏或外部链接，请把分值粘贴为纯数字")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except LearningCreditError:
        raise
    except Exception as exc:
        raise LearningCreditError("无法读取Excel，请使用下载的.xlsx模板") from exc
    try:
        if "期初学分" not in book.sheetnames:
            raise LearningCreditError("找不到“期初学分”工作表，请使用下载的模板")
        sheet = book["期初学分"]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column != 4:
            raise LearningCreditError("模板只接受4列、最多5000位学长，请勿添加列")
        if tuple(c.value for c in next(sheet.iter_rows(max_row=1))) != HEADERS:
            raise LearningCreditError("表头已改变，请保留模板的4列表头")
        rows, skipped = [], 0
        for i, cells in enumerate(sheet.iter_rows(min_row=2, max_col=4), start=2):
            values = [c.value for c in cells]
            if all(v is None or v == "" for v in values):
                continue
            if values[3] is None or values[3] == "":
                skipped += 1
                continue
            errors = []
            if any(c.data_type == "f" for c in cells):
                errors.append("含公式，请复制后粘贴为数值")
            if any(isinstance(v, bool) for v in values):
                errors.append("不能填写布尔值")
            code = str(values[0] or "").strip()
            if not isinstance(values[0], str):
                errors.append("学长编号必须为文本，请从模板保留原编号")
            try:
                points = Decimal(str(values[3]))
                if not points.is_finite() or points < 0 or points > Decimal("99999999.99") or points != points.quantize(Decimal("0.01")):
                    raise InvalidOperation
                points_text = format(points, ".2f")
            except (InvalidOperation, ValueError):
                errors.append("累计学分需为非负数字，最多两位小数")
                points_text = None
            rows.append({"excel_row": i, "member_code": code, "member_name": str(values[1] or "").strip(),
                         "class_label": str(values[2] or "").strip(), "points": points_text, "errors": errors})
        if not rows:
            raise LearningCreditError("没有填写累计学分。空白表示暂不导入，已确认0分请填0")
        return cutoff, rows, skipped
    finally:
        book.close()


def _fingerprint(cutoff: date, rows: list[dict]) -> str:
    payload = {"cutoff_date": cutoff.isoformat(), "rows": sorted(
        [{k: r[k] for k in ("member_id", "org_unit_id", "member_name", "class_label", "points")} for r in rows], key=lambda r: r["member_id"])}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _validate(connection, actor: int, cutoff: date, rows: list[dict], *, check_balances=True):
    roster, seen = _roster(connection, actor), set()
    ids = sorted({int(roster[r["member_code"]]["id"]) for r in rows if r["member_code"] in roster})
    existing, overlap = set(), set()
    ready = storage_available(connection)
    period = {c[0] for c in execute(connection, "SELECT * FROM learning_credit_entries WHERE 1=0").description}
    from datetime import UTC, timedelta
    upper = datetime.combine(cutoff + timedelta(days=1), datetime.min.time(), ZoneInfo("Asia/Shanghai")).astimezone(UTC).replace(tzinfo=None).isoformat(sep=" ")
    exact = "datetime(occurred_at)<datetime(?)" if isinstance(connection, sqlite3.Connection) else "occurred_at<?"
    predicate = "((occurred_precision='EXACT_DATE' AND " + exact + ") OR (occurred_precision<>'EXACT_DATE' AND (occurred_year<? OR (occurred_year=? AND (occurred_month IS NULL OR occurred_month<=?)))))" if "occurred_year" in period else exact
    period_params = (upper, cutoff.year, cutoff.year, cutoff.month) if "occurred_year" in period else (upper,)
    for start in range(0, len(ids), 200):
        chunk = ids[start:start + 200]
        placeholders = ",".join("?" for _ in chunk)
        if check_balances and ready:
            existing.update(int(r["member_id"]) for r in execute(connection, "SELECT member_id FROM learning_credit_opening_balances WHERE member_id IN (" + placeholders + ")", tuple(chunk)).fetchall())
        overlap.update(int(r["member_id"]) for r in execute(connection, "SELECT DISTINCT member_id FROM learning_credit_entries WHERE member_id IN (" + placeholders + ") AND status IN ('POSTED','REVERSED') AND " + predicate, (*chunk, *period_params)).fetchall())
    for row in rows:
        member = roster.get(row["member_code"])
        if row["member_code"] in seen:
            row["errors"].append("同一学长出现多行，请只保留一行累计余额")
        seen.add(row["member_code"])
        if not member:
            row["errors"].append("编号不在当前有效名册或授权范围内，请重新下载模板")
            continue
        row.update(member_id=int(member["id"]), org_unit_id=member["org_unit_id"])
        if member["name"] != row["member_name"] or member["class_label"] != row["class_label"]:
            row["errors"].append("姓名或班级与名册不一致，请重新下载模板，保留编号/姓名/班级")
        if member["ambiguous_class"]:
            row["errors"].append("名册存在多个有效班级，请先在学员管理核对班级关系")
        if row["member_id"] in existing:
            row["errors"].append("这位学长已录入期初余额，不可再次累计；请核对原导入记录")
        # A cumulative amount must never be added on top of overlapping facts.
        if row["member_id"] in overlap:
            row["errors"].append("截止日期以前已有正式学分，累计余额会重复计算，请先核对已有账本")


def preview(content: bytes, cutoff_date: str, actor: int) -> dict:
    _manage(actor)
    cutoff, rows, skipped = _parse(content, cutoff_date)
    connection = connect()
    try:
        _validate(connection, actor, cutoff, rows)
        errors = sum(bool(r["errors"]) for r in rows)
        return {"cutoff_date": cutoff.isoformat(), "rows": rows, "error_count": errors,
                "row_count": len(rows), "skipped_count": skipped,
                "total_points": format(sum((Decimal(r["points"]) for r in rows if r["points"] and not r["errors"]), Decimal(0)), ".2f"),
                "fingerprint": _fingerprint(cutoff, rows) if not errors else None,
                "storage_available": storage_available(connection)}
    finally:
        connection.close()


def _ready(connection):
    _write_allowed()
    if not get_settings().learning_credit_opening_balance_enabled:
        raise LearningCreditError("期初余额导入暂未启用，请联系系统管理员完成设置")
    if not storage_available(connection):
        raise LearningCreditError("期初余额存储尚未准备完成，请由系统管理员在本页完成设置")
    if not execute(connection, "SELECT version FROM schema_migrations WHERE version=?", (MIGRATION,)).fetchone():
        raise LearningCreditError("期初余额存储执行记录不完整，请先核验存储")


def _batch(connection, actor: int, import_id: int) -> tuple[dict, list[dict]]:
    row = execute(connection, "SELECT * FROM learning_credit_opening_imports WHERE id=?", (import_id,)).fetchone()
    if not row:
        raise LearningCreditError("导入记录不存在或不在当前授权范围内")
    rows = [dict(r) for r in execute(connection, "SELECT r.*,m.member_code FROM learning_credit_opening_rows r JOIN members m ON m.id=r.member_id WHERE import_id=? ORDER BY r.excel_row", (import_id,)).fetchall()]
    allowed = _scope(actor)
    if allowed is not None and any(r["org_unit_id"] not in allowed for r in rows):
        raise PermissionError("该导入包含当前账号范围外的学长，请由覆盖全部范围的人员处理")
    return dict(row), rows


def register(content: bytes, cutoff_date: str, filename: str, expected_fingerprint: str, actor: int) -> dict:
    _permission(actor, OPENING_PERMISSION)
    _manage(actor)
    cutoff, rows, _ = _parse(content, cutoff_date)
    with transaction() as connection, _operation(connection):
        _ready(connection)
        # Re-upload of an already posted identical source is a receipt, not a
        # second insert. Validate identities before finding that receipt.
        _validate(connection, actor, cutoff, rows, check_balances=False)
        if any(r["errors"] for r in rows):
            raise LearningCreditError("文件或名册已变化，请重新检查并修改错误行")
        fingerprint = _fingerprint(cutoff, rows)
        if fingerprint != expected_fingerprint:
            raise LearningCreditError("文件或名册已变化，请重新检查后提交")
        existing = execute(connection, "SELECT id FROM learning_credit_opening_imports WHERE content_fingerprint=?", (fingerprint,)).fetchone()
        if existing:
            batch, _ = _batch(connection, actor, int(existing["id"]))
            if batch["status"] == "CANCELLED":
                raise LearningCreditError("这份来源已取消，请核对并修正分值后重新上传，原记录继续保留")
            return {"id": batch["id"], "status": batch["status"], "idempotent": True}
        _validate(connection, actor, cutoff, rows)
        if any(r["errors"] for r in rows):
            raise LearningCreditError("已有期初余额或名册变化，请刷新检查结果")
        now = _db_timestamp(connection)
        cursor = execute(connection, "INSERT INTO learning_credit_opening_imports(content_fingerprint,file_sha256,original_filename,cutoff_date,created_by,created_at) VALUES (?,?,?,?,?,?)",
                         (fingerprint, hashlib.sha256(content).hexdigest(), filename[:255], cutoff.isoformat(), actor, now))
        import_id = int(cursor.lastrowid)
        for r in rows:
            execute(connection, "INSERT INTO learning_credit_opening_rows(import_id,member_id,org_unit_id,member_name,class_label,excel_row,points) VALUES (?,?,?,?,?,?,?)",
                    (import_id, r["member_id"], r["org_unit_id"], r["member_name"], r["class_label"], r["excel_row"], r["points"]))
        write_audit(connection, actor_user_id=actor, action="learning.credit_opening.submit", resource_type="credit_opening_import", resource_id=str(import_id),
                    after={"content_fingerprint": fingerprint, "file_sha256": hashlib.sha256(content).hexdigest(), "count": len(rows), "cutoff_date": cutoff.isoformat()})
        return {"id": import_id, "status": "PENDING_APPROVAL", "idempotent": False}


def _recheck(connection, actor, batch, rows):
    cutoff = date.fromisoformat(str(batch["cutoff_date"])[:10])
    for r in rows:
        r["errors"] = []
        r["points"] = format(Decimal(str(r["points"])), ".2f")
    _validate(connection, actor, cutoff, rows)
    if any(r["errors"] for r in rows) or _fingerprint(cutoff, rows) != batch["content_fingerprint"]:
        raise LearningCreditError("名册、余额或已有学分发生变化，不能入账，请重新核对来源")


def action(import_id: int, actor: int, operation: str, expected_fingerprint: str) -> dict:
    _manage(actor)
    if operation not in {"approve", "post", "cancel"}:
        raise LearningCreditError("不支持这项导入操作")
    _permission(actor, OPENING_PERMISSION)
    with transaction() as connection, _operation(connection):
        _ready(connection)
        batch, rows = _batch(connection, actor, import_id)
        from app.services import credit_year_reconciliation as reconciliation
        if reconciliation.managed_opening(connection,import_id):
            raise LearningCreditError('这份补差来源由年度导入管理，请从历史年度入口复核和入账')
        if batch["content_fingerprint"] != expected_fingerprint:
            raise LearningCreditError("导入版本已变化，请刷新后核对")
        if operation == "cancel":
            if batch["status"] == "POSTED":
                raise LearningCreditError("已入账的余额不能取消，需凭核对后的来源追加更正记录")
            if batch["status"] == "CANCELLED":
                return {"id": import_id, "status": "CANCELLED", "idempotent": True}
            execute(connection, "UPDATE learning_credit_opening_imports SET status='CANCELLED' WHERE id=?", (import_id,))
            status = "CANCELLED"
        elif operation == "approve":
            if batch["status"] == "CANCELLED":
                raise LearningCreditError("该来源已取消，不能审批")
            if batch["status"] in {"APPROVED", "POSTED"}:
                return {"id": import_id, "status": batch["status"], "idempotent": True}
            _recheck(connection, actor, batch, rows)
            execute(connection, "UPDATE learning_credit_opening_imports SET status='APPROVED',approved_by=?,approved_at=? WHERE id=? AND status='PENDING_APPROVAL'", (actor, _db_timestamp(connection), import_id))
            status = "APPROVED"
        else:
            if batch["status"] == "POSTED":
                return {"id": import_id, "status": "POSTED", "idempotent": True}
            if batch["status"] != "APPROVED" or not batch["approved_by"]:
                raise LearningCreditError("文件尚未确认，请先完成复核")
            # Lock member rows in a deterministic order. Regular ledger writers
            # acquire the same lock before checking the cutoff boundary.
            for r in sorted(rows, key=lambda r: r["member_id"]):
                execute(connection, "SELECT id FROM members WHERE id=?" + ("" if isinstance(connection, sqlite3.Connection) else " FOR UPDATE"), (r["member_id"],)).fetchone()
            _recheck(connection, actor, batch, rows)
            now = _db_timestamp(connection)
            for r in rows:
                execute(connection, "INSERT INTO learning_credit_opening_balances(member_id,import_id,org_unit_id,points,cutoff_date,posted_by,posted_at) VALUES (?,?,?,?,?,?,?)", (r["member_id"], import_id, r["org_unit_id"], r["points"], batch["cutoff_date"], actor, now))
            execute(connection, "UPDATE learning_credit_opening_imports SET status='POSTED',posted_by=?,posted_at=? WHERE id=? AND status='APPROVED'", (actor, now, import_id))
            status = "POSTED"
        write_audit(connection, actor_user_id=actor, action="learning.credit_opening." + operation, resource_type="credit_opening_import", resource_id=str(import_id),
                    after={"content_fingerprint": expected_fingerprint, "status": status, "count": len(rows), "total_points": format(sum((Decimal(r["points"]) for r in rows), Decimal(0)), ".2f")})
        return {"id": import_id, "status": status, "idempotent": False}


def workbench(actor: int) -> dict:
    user = _manage(actor)
    connection = connect()
    try:
        dialect = "sqlite" if isinstance(connection, sqlite3.Connection) else "mysql"
        tables = _tables(connection)
        marker = execute(connection, "SELECT version FROM schema_migrations WHERE version=?", (MIGRATION,)).fetchone()
        reserved = execute(connection, "SELECT result FROM audit_logs WHERE action='production.credit_opening.setup' ORDER BY id DESC LIMIT 1").fetchone()
        available = storage_available(connection) and bool(marker)
        imports = []
        if available:
            ids = execute(connection, "SELECT id FROM learning_credit_opening_imports ORDER BY id DESC LIMIT 100").fetchall()
            for i in ids:
                try:
                    batch, rows = _batch(connection, actor, i["id"])
                except PermissionError:
                    continue
                imports.append({**batch, "rows": rows, "row_count": len(rows), "total_points": format(sum((Decimal(str(r["points"])) for r in rows), Decimal(0)), ".2f")})
        permissions = set(user.get("permissions", []))
        return {"actor_user_id": actor, "storage_available": available, "enabled": get_settings().learning_credit_opening_balance_enabled,
                "setup_allowed": "system_admin" in user.get("roles", []) and "plans:production_rule_reconciliation_apply" in permissions and get_settings().credit_opening_setup_enabled and not tables and not marker and not reserved,
                "setup_incomplete": not available and bool(tables or marker or reserved),
                "release_commit": get_build_info()["commit_sha"], "migration_sha256": MIGRATION_HASHES[dialect],
                "can_approve": OPENING_PERMISSION in permissions, "can_post": OPENING_PERMISSION in permissions,
                "imports": imports, "summary": opening_summary(connection, allowed=_scope(actor))}
    finally:
        connection.close()


def setup(actor: int, expected_release_commit: str, expected_migration_sha256: str) -> dict:
    user = _manage(actor)
    _permission(actor, "plans:production_rule_reconciliation_apply")
    if "system_admin" not in user.get("roles", []):
        raise PermissionError("仅系统管理员可以准备期初余额存储")
    _write_allowed()
    settings = get_settings()
    if not settings.credit_opening_setup_enabled or settings.run_bootstrap_on_startup:
        raise LearningCreditError("当前发布未开放这项固定存储准备操作")
    commit = get_build_info()["commit_sha"]
    if len(commit) != 40 or commit != expected_release_commit:
        raise LearningCreditError("服务版本发生变化，请刷新后重新核对")
    connection = connect()
    try:
        with _operation(connection):
            dialect = "sqlite" if isinstance(connection, sqlite3.Connection) else "mysql"
            content = (MIGRATION_ROOT / dialect / MIGRATION).read_bytes()
            digest = hashlib.sha256(content).hexdigest()
            if digest != MIGRATION_HASHES[dialect] or digest != expected_migration_sha256:
                raise LearningCreditError("存储脚本校验失败，已停止操作")
            marker = execute(connection, "SELECT version FROM schema_migrations WHERE version=?", (MIGRATION,)).fetchone()
            if marker and storage_available(connection):
                return {"status": "READY", "idempotent": True}
            if marker or _tables(connection) or execute(connection, "SELECT id FROM audit_logs WHERE action='production.credit_opening.setup' LIMIT 1").fetchone():
                raise LearningCreditError("发现已执行或中断记录，请核验实际存储；系统不会自动重试迁移")
            write_audit(connection, actor_user_id=actor, action="production.credit_opening.setup", resource_type="schema_migration", resource_id=MIGRATION,
                        result="STARTED", after={"release_commit": commit, "migration_sha256": digest, "scope": list(TABLES), "permission_policy": {"capability": OPENING_PERMISSION, "roles": ["employee_learning_management", "ops_center_learning"], "individual_assignments_changed": 0}}, request_id=commit)
            # MySQL DDL implicitly commits: reserve BEFORE the first statement.
            connection.commit()
            for statement in _split_mysql(content.decode()):
                execute(connection, statement)
            if not storage_available(connection):
                raise LearningCreditError("存储准备未完整完成，请核验执行结果")
            execute(connection, "INSERT INTO schema_migrations(version,applied_at) VALUES (?,?)", (MIGRATION, _db_timestamp(connection)))
            write_audit(connection, actor_user_id=actor, action="production.credit_opening.setup.complete", resource_type="schema_migration", resource_id=MIGRATION,
                        after={"release_commit": commit, "migration_sha256": digest, "existing_ledger_delta": 0, "permission_policy": {"capability": OPENING_PERMISSION, "roles": ["employee_learning_management", "ops_center_learning"], "individual_assignments_changed": 0}}, request_id=commit)
            connection.commit()
            return {"status": "READY", "idempotent": False}
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
