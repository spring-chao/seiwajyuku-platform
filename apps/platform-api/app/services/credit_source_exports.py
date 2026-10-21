"""Audited, scoped credit exports with typed values and source row references."""
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from app.db import connect,execute
from app.services import learning_credits as credits
from app.services import credit_opening_balances as opening
from app.services import credit_year_allocations as annual
from app.services import credit_source_imports as sources
from app.services.audit import write_audit
from app.services.credit_monthly_targets import monthly_target

LABELS={'CLASS_MEETING_SCORE':'班会签到','DAILY_READING':'读书打卡','EXCELLENT_SHARE':'优秀分享','GROUP_MEETING_ATTENDANCE':'小组学习会','COURSE_COMPLETION':'课程完成'}


def entry_sources(c,entry):
    if not sources.available(c):return []
    original=entry
    if entry.get('reversal_of_entry_id'):
        row=execute(c,'SELECT * FROM learning_credit_entries WHERE id=? AND member_id=?',(entry['reversal_of_entry_id'],entry['member_id'])).fetchone()
        if not row:return []
        original=dict(row)
    try:snapshot=json.loads(original.get('rule_snapshot_json') or '{}')
    except (ValueError,TypeError):return []
    if not isinstance(snapshot,dict):return []
    facts=[v for v in snapshot.get('fact_ids',[]) if type(v) is int and v>0][:500]
    attendance=[v.get('attendance_record_id') for v in snapshot.get('score_records',[]) if isinstance(v,dict) and type(v.get('attendance_record_id')) is int and v['attendance_record_id']>0][:500]
    if not facts and not attendance:return []
    parts=[];params=[original['member_id'],original.get('class_org_unit_id')]
    if facts:parts.append('r.fact_id IN ('+','.join('?' for _ in facts)+')');params.extend(facts)
    if attendance:parts.append('r.attendance_record_id IN ('+','.join('?' for _ in attendance)+')');params.extend(attendance)
    rows=execute(c,'SELECT DISTINCT i.id AS import_id,i.original_filename,i.file_sha256,i.source_kind,r.excel_row,r.occurred_on,r.source_reference FROM learning_credit_source_rows r JOIN learning_credit_source_imports i ON i.id=r.import_id WHERE r.member_id=? AND i.class_org_unit_id=? AND ('+' OR '.join(parts)+') ORDER BY i.id,r.excel_row,r.source_reference',tuple(params)).fetchall()
    return [{**dict(r),'original_filename':_safe_text(r['original_filename']),'occurred_on':str(r['occurred_on'])[:10]} for r in rows]


def _safe_text(value):
    return re.sub(r'(?<!\d)(1[3-9]\d{2})\d{4}(\d{4})(?!\d)',r'\1****\2',str(value or ''))


def _literal(sheet,row,col,value):
    cell=sheet.cell(row,col,value)
    if isinstance(value,str):cell.data_type='s'
    return cell


def _sheet(book,name,headers,rows):
    s=book.create_sheet(name);s.sheet_view.showGridLines=False
    for i,label in enumerate(headers,1):
        cell=_literal(s,1,i,label);cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11);cell.fill=PatternFill('solid',fgColor='275747');cell.alignment=Alignment(horizontal='center',vertical='center')
    s.row_dimensions[1].height=28
    for row_index,values in enumerate(rows,2):
        for col,value in enumerate(values,1):
            cell=_literal(s,row_index,col,value);cell.font=Font(name='Arial',size=11,color='263238');cell.alignment=Alignment(vertical='center')
            if isinstance(value,(Decimal,float,int)) and not isinstance(value,bool):cell.number_format='#,##0.00' if isinstance(value,(Decimal,float)) else '#,##0'
            if isinstance(value,(date,datetime)):cell.number_format='yyyy-mm-dd'
        s.row_dimensions[row_index].height=22
    for col,label in enumerate(headers,1):
        from openpyxl.utils import get_column_letter
        longest=max([len(str(label))]+[len(str(r[col-1] or '')) for r in rows[:200]])
        s.column_dimensions[get_column_letter(col)].width=min(60,max(16,longest*1.6+2))
    s.freeze_panes='C2';s.auto_filter.ref=s.dimensions
    return s


def export_credits(actor,class_id=None,occurred_from='2026-10-01',occurred_to=None,member_id=None):
    if member_id is not None and (type(member_id) is not int or member_id <= 0):raise credits.LearningCreditError('学员编号无效')
    user=credits.user_context(actor) or {}
    if not {'exports:normal','members:read','plans:credit_settlement_preview'}<=set(user.get('permissions',[])):raise PermissionError('当前账号没有学分导出权限')
    start=date.fromisoformat(occurred_from);end=date.fromisoformat(occurred_to) if occurred_to else datetime.now(sources.activity.SHANGHAI).date()
    if start>end:raise credits.LearningCreditError('导出开始日期不能晚于结束日期')
    export_permissions=('exports:normal','members:read','plans:credit_settlement_preview')
    if class_id and not sources._has_scope(actor,class_id,export_permissions):raise PermissionError('导出班级不在授权范围')
    scope={class_id} if class_id else sources.effective_scope(actor,export_permissions)
    c=connect()
    try:
        # An all-scope export includes confirmed balances held by a centre
        # before class assignment. A class filter still limits that class.
        classes=[dict(r) for r in execute(c,"SELECT id,name FROM org_units ORDER BY name").fetchall() if scope is None or r['id'] in scope]
        class_ids={r['id'] for r in classes}
        if not class_ids:raise credits.LearningCreditError('当前范围没有可导出的班级')
        placeholders=','.join('?' for _ in class_ids);params=tuple(sorted(class_ids))
        roster=execute(c,"SELECT DISTINCT m.id,m.member_code,m.name,r.org_unit_id,o.name AS class_name FROM members m JOIN member_org_relations r ON r.member_id=m.id JOIN org_units o ON o.id=r.org_unit_id WHERE m.status='ACTIVE' AND r.relation_type='STUDY_CLASS' AND r.org_unit_id IN ("+placeholders+") AND (r.valid_from IS NULL OR r.valid_from<=?) AND (r.valid_until IS NULL OR r.valid_until>=?)"+(' AND m.id=?' if member_id else '')+" ORDER BY m.id,r.org_unit_id",(*params,end.isoformat(),end.isoformat(),*((member_id,) if member_id else ()))).fetchall()
        members={}
        for r in roster:
            member=members.setdefault(int(r['id']),{'member_code':r['member_code'],'name':r['name'],'classes':[],'class_ids':set()});member['classes'].append(r['class_name']);member['class_ids'].add(r['org_unit_id'])
        if len(members)>5000:raise credits.LearningCreditError('导出超过5000位学长，请选择班级缩小范围')
        ledger_scope='e.class_org_unit_id IN ('+placeholders+')'
        ledger_params=params
        if not class_id:
            ledger_scope+=' OR (e.class_org_unit_id IS NULL AND m.org_unit_id IN ('+placeholders+'))'
            ledger_params=(*params,*params)
        if member_id:
            ledger_scope='('+ledger_scope+') AND e.member_id=?'
            ledger_params=(*ledger_params,member_id)
        ledger=[dict(r) for r in execute(c,"SELECT e.*,m.member_code,m.name AS member_name,o.name AS class_name FROM learning_credit_entries e JOIN members m ON m.id=e.member_id LEFT JOIN org_units o ON o.id=e.class_org_unit_id WHERE e.status IN ('POSTED','REVERSED') AND ("+ledger_scope+") ORDER BY e.member_id,e.occurred_at,e.id",ledger_params).fetchall()]
        if len(ledger)>50000:raise credits.LearningCreditError('学分明细超过50000条，请缩小班级范围')
        # Include retained balances/ledger for former members as well as the
        # current roster, so summary totals do not silently lose old credits.
        retained=[]
        if opening.storage_available(c):retained=execute(c,'SELECT b.member_id AS id,m.member_code,m.name,o.name AS class_name FROM learning_credit_opening_balances b JOIN members m ON m.id=b.member_id JOIN org_units o ON o.id=b.org_unit_id WHERE b.org_unit_id IN ('+placeholders+')'+(' AND b.member_id=?' if member_id else ''),(*params,*((member_id,) if member_id else ()))).fetchall()
        for r in retained:
            members.setdefault(int(r['id']),{'member_code':r['member_code'],'name':r['name'],'classes':[r['class_name']+'（余额归属）']})
        for r in ledger:
            members.setdefault(int(r['member_id']),{'member_code':r['member_code'],'name':r['member_name'],'classes':[str(r['class_name'] or '')+'（账本归属）']})
        if len(members)>5000:raise credits.LearningCreditError('导出超过5000位学长，请选择班级缩小范围')
        if member_id and member_id not in members:raise PermissionError('该学员不在当前学分导出范围内')
        groups={}
        for r in execute(c,"SELECT DISTINCT r.member_id,o.name,o.parent_id FROM member_org_relations r JOIN org_units o ON o.id=r.org_unit_id WHERE r.relation_type='STUDY_GROUP' AND (r.valid_from IS NULL OR r.valid_from<=?) AND (r.valid_until IS NULL OR r.valid_until>=?) ORDER BY o.id",(end.isoformat(),end.isoformat())).fetchall():
            if r['member_id'] in members and r['parent_id'] in members[r['member_id']].get('class_ids',set()):groups.setdefault(r['member_id'],[]).append(r['name'])
        month_start=end.replace(day=1)
        target_classes={cid for m in members.values() for cid in m.get('class_ids',set())}
        targets={cid:monthly_target(c,cid,end) for cid in target_classes}
        totals={};added={};monthly={};details=[]
        for e in ledger:
            member=int(e['member_id']);point=Decimal(str(e['points']));totals[member]=totals.get(member,Decimal(0))+point
            # A partial year/month never masquerades as a confirmed October date.
            if e.get('occurred_precision')!='EXACT_DATE' or not e.get('occurred_at'):continue
            day=date.fromisoformat(str(e['occurred_at'])[:10])
            if month_start<=day<=end:monthly[member]=monthly.get(member,Decimal(0))+point
            if not start<=day<=end:continue
            added[member]=added.get(member,Decimal(0))+point
            trace=entry_sources(c,e)
            descriptions='；'.join(f"{r['original_filename']} · 第{r['excel_row']}行 · {r['source_reference']}" for r in trace)
            details.append([str(e['id']),str(e['member_code']),_safe_text(e['member_name']),e['class_name'],' / '.join(groups.get(member,[])),day,point,LABELS.get(e['credit_type'],e['credit_type']),e['source_type'],str(e['source_id']),descriptions or '业务记录来源（编号见前列）',e['rule_version'],e['status'],str(e['reversal_of_entry_id']) if e.get('reversal_of_entry_id') else ''])
        summaries=[]
        for member,m in members.items():
            base=Decimal(opening.opening_summary(c,member_id=member,allowed=class_ids)['total_points'])
            ids=m.get('class_ids',set())
            target=targets[next(iter(ids))] if len(ids)==1 else {'points':None,'reason':'没有唯一有效的班级学习关系','plan':''}
            full=target['points'];earned=monthly.get(member,Decimal(0))
            summaries.append([str(m['member_code']),_safe_text(m['name']),' / '.join(m['classes']),' / '.join(groups.get(member,[])),base,added.get(member,Decimal(0)),base+totals.get(member,Decimal(0)),earned,full if full is not None else '待配置',earned/full if full is not None and full>0 else None,target['reason'],target['plan']])
        history=[]
        if annual.available(c):
            for r in execute(c,"SELECT a.member_id,m.member_code,m.name,o.name AS class_name,a.year_index,a.points,i.original_filename,i.cutoff_date FROM learning_credit_year_allocations a JOIN learning_credit_year_imports i ON i.id=a.import_id AND i.status='POSTED' JOIN class_learning_bindings b ON b.id=a.binding_id JOIN org_units o ON o.id=b.class_org_unit_id JOIN members m ON m.id=a.member_id WHERE b.class_org_unit_id IN ("+placeholders+")"+(' AND a.member_id=?' if member_id else '')+" ORDER BY m.id,a.year_index",(*params,*((member_id,) if member_id else ()))).fetchall():
                history.append([str(r['member_code']),_safe_text(r['name']),r['class_name'],int(r['year_index']),Decimal(str(r['points'])),date.fromisoformat(str(r['cutoff_date'])[:10]),_safe_text(r['original_filename'])])
        book=Workbook();book.remove(book.active)
        summary_sheet=_sheet(book,'学分汇总',['学长编号','姓名','班级','小组','已确认期初及补差学分',f'{start}至{end}新增学分','累计学分（账本原值）',f'{month_start}至{end}当月获得学分','当月满分','学分率','满分构成/待配置原因','学习计划及规则'],summaries)
        for row in summary_sheet.iter_rows(min_row=2):row[9].number_format='0.0%'
        summary_sheet['I1'].comment=Comment(f'月份：{end:%Y-%m}。满分按整月已发布计划、已排期学习周期和绑定规则计算。学分率=该月1日至截止日实际账本学分/整月满分；不含期初余额。多班归属或配置不全时留空；比例不封顶。','苏州盛和塾')
        _sheet(book,'新增学分明细',['账本编号','学长编号','姓名','班级','小组','业务日期','学分','计分项目','来源类型','业务来源编号','原文件与行号','规则版本','账本状态','冲销原账本编号'],details)
        _sheet(book,'已导入历史年度',['学长编号','姓名','班级','学习年度','原年度学分','统计截止日','原文件'],history)
        output=io.BytesIO();book.save(output)
        write_audit(c,actor_user_id=actor,action='learning.credit.export',resource_type='credit_export',org_unit_id=class_id,after={'selected_member_id':member_id,'member_count':len(summaries),'detail_count':len(details),'historical_year_count':len(history),'occurred_from':start.isoformat(),'occurred_to':end.isoformat(),'contact_fields':False});c.commit()
        return output.getvalue()
    finally:c.close()


def export_options(actor):
    user=credits.user_context(actor) or {}
    if not {'exports:normal','members:read','plans:credit_settlement_preview'}<=set(user.get('permissions',[])):raise PermissionError('当前账号没有学分导出权限')
    c=connect()
    try:
        scope=sources.effective_scope(actor,('exports:normal','members:read','plans:credit_settlement_preview'))
        return {'classes':[dict(r) for r in execute(c,"SELECT id,name FROM org_units WHERE unit_type IN ('CLASS','SPECIAL_COHORT') ORDER BY name").fetchall() if scope is None or r['id'] in scope]}
    finally:c.close()
