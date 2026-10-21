"""Scoped Excel source receipts feeding the existing reviewed credit ledger."""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from app.db import connect, execute, transaction
from app.core.settings import get_settings
from app.core.build_info import get_build_info
from app.migrations import MIGRATION_ROOT, _split_mysql
from app.services import learning_credits as credits
from app.services import learning_activity_credits as activity
from app.services import class_meeting_credits as class_credits
from app.services import credit_settlement_batches as batches
from app.services import credit_opening_balances as opening
from app.services.attendance_scoring import calculate_score, upsert_score_record
from app.services.audit import write_audit
from app.services.learning_org_catalog import class_catalog

MIGRATION = '0072_learning_credit_source_imports.sql'
TABLES = {'learning_credit_source_imports','learning_credit_source_rows','learning_credit_source_batches'}
HASHES = {'mysql': '65a05f28adeb1d317579be449ab68ca874b4663a6298dd9211394b7de149a926', 'sqlite': 'af4bc023b8a77dbe8a81a4c050dbc6836b3e46fb956e10df6312889c676a948d'}
START = date(2026,10,1)
KINDS = {'CLASS_MEETING','DAILY_READING','EXCELLENT_SHARE'}
SLOTS = [('MORNING','上午'),('AFTERNOON','下午'),('KONPA','空巴')]
STATES = {'正常到场':('MANUAL_PRESENT',False,False),'迟到':('MANUAL_PRESENT',True,False),
          '早退':('MANUAL_PRESENT',False,True),'迟到且早退':('MANUAL_PRESENT',True,True),
          '请假':('LEAVE',False,False),'缺席':('ABSENT',False,False)}
SHEETS = {'CLASS_MEETING':'班会线下签到','DAILY_READING':'读书打卡','EXCELLENT_SHARE':'优秀分享'}
HEADERS = {'CLASS_MEETING':['学长编号','姓名','上午状态','下午状态','空巴状态'],
           'DAILY_READING':['学长编号','姓名','打卡日期','已完成录音'],
           'EXCELLENT_SHARE':['学长编号','姓名','分享日期','分享编号','分享标题','已评定优秀']}
READING_REPORT_HEADERS = (
    ['姓名','账户','小组','日期','内容标题','录音'],
    ['姓名','账户','小组','日期','内容标题','完成率','打卡状态','读书','分享','回应','录音','准时'],
)
RECORDING_COMPLETE = {'按时完成','提前完成','过期完成'}


def _json(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,default=str,separators=(',',':'))


def _hash(value):
    return hashlib.sha256(_json(value).encode()).hexdigest()


def present_tables(c):
    query="SELECT name AS n FROM sqlite_master WHERE type='table'" if isinstance(c,sqlite3.Connection) else 'SELECT TABLE_NAME AS n FROM information_schema.tables WHERE TABLE_SCHEMA=DATABASE()'
    return {r['n'] for r in execute(c,query).fetchall()} & TABLES


def available(c):
    if present_tables(c)!=TABLES or not execute(c,'SELECT version FROM schema_migrations WHERE version=?',(MIGRATION,)).fetchone():
        return False
    fields = {'learning_credit_source_imports':{'id','fingerprint','file_sha256','original_filename','source_kind','class_org_unit_id','binding_id','cycle_id','event_group_id','options_json','source_snapshot_json','source_note','created_by','created_at'},
              'learning_credit_source_rows':{'id','import_id','member_id','excel_row','occurred_on','source_reference','record_json','fact_id','attendance_record_id'},
              'learning_credit_source_batches':{'import_id','batch_id'}}
    return all(keys=={col[0] for col in execute(c,'SELECT * FROM '+table+' WHERE 1=0').description} for table,keys in fields.items())


def _principal(actor,kind=None,write=False):
    user=credits.user_context(actor) or {}
    required={'members:read','plans:credit_settlement_manage','plans:credit_settlement_preview'}
    if write and kind=='CLASS_MEETING':required.update({'attendance:adjudicate','plans:period_write'})
    elif write:required.add('plans:credit_activity_fact_manage')
    if not required<=set(user.get('permissions',[])):raise PermissionError('当前账号没有这项来源导入权限')
    if write:credits._write_allowed()
    return user


def _class(c,actor,class_id):
    if not _has_scope(actor,class_id):raise PermissionError('班级不在当前学分授权范围')
    row=execute(c,"SELECT b.*,o.name AS class_name,p.status AS plan_status,p.plan_key,p.version_label FROM class_learning_bindings b JOIN org_units o ON o.id=b.class_org_unit_id JOIN learning_plan_versions p ON p.id=b.plan_version_id WHERE b.class_org_unit_id=? AND b.status='ACTIVE' AND o.is_active=1 ORDER BY b.started_at DESC,b.id DESC LIMIT 1",(class_id,)).fetchone()
    if not row:raise credits.LearningCreditError('班级没有可用的学习计划，请先在班级学习计划中准备')
    return dict(row)


def _has_scope(actor,class_id,permissions=('members:read','plans:credit_settlement_manage','plans:credit_settlement_preview')):
    if not credits._scope_allows(actor,class_id):return False
    for permission in permissions:
        allowed=credits.accessible_org_ids(actor,permission)
        if allowed is not None and class_id not in allowed:return False
    return True


def effective_scope(actor,permissions=('members:read','plans:credit_settlement_manage','plans:credit_settlement_preview')):
    """Compute this request's permission intersection once, never cache IAM globally."""
    scopes=[credits.accessible_org_ids(actor),*(credits.accessible_org_ids(actor,p) for p in permissions)]
    restricted=[set(scope) for scope in scopes if scope is not None]
    return set.intersection(*restricted) if restricted else None


def _write_scope(actor,class_id,kind):
    required=('attendance:adjudicate','plans:period_write') if kind=='CLASS_MEETING' else ('plans:credit_activity_fact_manage',)
    if not _has_scope(actor,class_id,required):raise PermissionError('班级不在这项来源操作的授权范围')


def _members(c,class_id):
    return {r['member_code']:dict(r) for r in execute(c,"SELECT DISTINCT m.id,m.member_code,m.name,m.phone_last4,m.phone_masked FROM members m JOIN member_org_relations r ON r.member_id=m.id AND r.org_unit_id=? AND r.relation_type='STUDY_CLASS' WHERE m.status='ACTIVE' ORDER BY m.id",(class_id,)).fetchall()}


def _date(value):
    if isinstance(value,datetime):value=value.date()
    if isinstance(value,date):result=value
    else:
        try:result=date.fromisoformat(str(value).strip())
        except ValueError as exc:raise credits.LearningCreditError('日期需为完整的YYYY-MM-DD') from exc
    if result<START or result>datetime.now(activity.SHANGHAI).date():raise credits.LearningCreditError('新增学分来源日期需在2026-10-01至今天之间')
    return result


def _options(kind,value):
    try:
        if len(value)>5000:raise ValueError
        result=json.loads(value)
        if not isinstance(result,dict):raise ValueError
        if kind!='CLASS_MEETING':
            if result:raise ValueError
            return {}
        if set(result)!={'cycle_id','actual_start'} or type(result['cycle_id']) is not int or result['cycle_id']<=0:raise ValueError
        actual=datetime.fromisoformat(result['actual_start'])
        if actual.tzinfo is None:raise ValueError
        actual=actual.astimezone(activity.SHANGHAI)
        _date(actual.date())
        if actual>datetime.now(activity.SHANGHAI):raise ValueError
        return {'cycle_id':result['cycle_id'],'actual_start':actual.isoformat()}
    except (TypeError,ValueError) as exc:raise credits.LearningCreditError('请填写本次学习日及有效的实际班会开始时间') from exc


def _parse(content,kind):
    if kind not in KINDS:raise credits.LearningCreditError('请选择班会、读书打卡或优秀分享')
    if not content or len(content)>opening.MAX_BYTES:raise credits.LearningCreditError('只支持5MB以内的.xlsx文件')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if len(z.infolist())>1000 or sum(f.file_size for f in z.infolist())>50*1024*1024 or any('externallinks' in f.lower() or 'vbaproject' in f.lower() for f in z.namelist()):raise ValueError
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
    except Exception as exc:raise credits.LearningCreditError('无法读取文件，请上传原始.xlsx文件，保留普通数值与文字') from exc
    try:
        if kind=='DAILY_READING' and '导出数据' in book.sheetnames:
            sheet=book['导出数据'];headers=[v.value for v in next(sheet.iter_rows(max_row=1))]
            if headers not in READING_REPORT_HEADERS or sheet.max_row>5001 or sheet.max_column!=len(headers):raise credits.LearningCreditError('读书会原表需保留姓名、账户、小组、日期、内容标题、录音列，单次最多5000行')
            rows=[]
            for i,cells in enumerate(sheet.iter_rows(min_row=2),2):
                values=[cell.value for cell in cells]
                if all(value in (None,'') for value in values):continue
                raw=dict(zip(headers,values));account=str(raw.get('账户') or '').strip();status=str(raw.get('录音') or '').strip()
                row={'excel_row':i,'member_code':'','member_name':str(raw.get('姓名') or '').strip(),'original_report':True,'errors':[],
                     'raw':{**raw,'打卡日期':raw.get('日期'),'已完成录音':'是' if status in RECORDING_COMPLETE else '否'}}
                if any(cell.data_type=='f' for cell in cells):row['errors'].append('含公式，请粘贴为数值或文字')
                if not re.fullmatch(r'1[3-9]\d\*{4}\d{4}',account):row['errors'].append('账户需为原表中的脱敏手机号，不能填写完整号码')
                if status not in RECORDING_COMPLETE|{'未提交'}:row['errors'].append('录音状态只接受按时完成、提前完成、过期完成或未提交')
                if any(len(str(v or ''))>255 for v in values):row['errors'].append('单元格文字不能超过255字')
                rows.append(row)
            if not rows:raise credits.LearningCreditError('原表没有读书打卡记录')
            return rows
        if SHEETS[kind] not in book.sheetnames:raise credits.LearningCreditError('找不到'+SHEETS[kind]+'工作表')
        sheet=book[SHEETS[kind]];headers=HEADERS[kind]
        if sheet.max_row>5001 or sheet.max_column!=len(headers) or [v.value for v in next(sheet.iter_rows(max_row=1))]!=headers:raise credits.LearningCreditError('请保留模板表头，单次最多5000行')
        rows=[]
        for i,cells in enumerate(sheet.iter_rows(min_row=2),2):
            values=[v.value for v in cells]
            if all(v in (None,'') for v in values[2:]):continue
            row={'excel_row':i,'member_code':str(values[0] or '').strip(),'member_name':str(values[1] or '').strip(),'errors':[],'raw':dict(zip(headers,values))}
            if any(cell.data_type=='f' for cell in cells):row['errors'].append('含公式，请粘贴为数值或文字')
            if not isinstance(values[0],str):row['errors'].append('学长编号需保留为文本')
            if any(len(str(v or ''))>255 for v in values):row['errors'].append('单元格文字不能超过255字')
            rows.append(row)
        if not rows:raise credits.LearningCreditError('没有填写有效记录，空白表示暂不导入')
        return rows
    finally:book.close()


def _cycle(c,binding,options):
    row=execute(c,'SELECT * FROM class_learning_cycles WHERE id=? AND binding_id=? AND class_org_unit_id=?',(options['cycle_id'],binding['id'],binding['class_org_unit_id'])).fetchone()
    if not row or row['cycle_status']=='UPCOMING':raise credits.LearningCreditError('请选择本班已开放或已举办的学习日')
    actual=datetime.fromisoformat(options['actual_start'])
    opened=activity._parse_storage_timestamp(row['opened_at'])
    if opened and actual<=opened:raise credits.LearningCreditError('实际班会时间需晚于所选周期开始时间')
    if row['actual_class_meeting_at'] and activity._local_date_from_timestamp(row['actual_class_meeting_at'])!=actual.date():raise credits.LearningCreditError('本期已登记其他班会日期，请核对本次学习日')
    if not row['actual_class_meeting_at'] and row['cycle_status']!='OPEN':
        closed=activity._parse_storage_timestamp(row['closed_at'])
        if not closed or actual>=closed:raise credits.LearningCreditError('实际班会日期不属于所选周期，无法补登')
    groups=execute(c,"SELECT id FROM attendance_event_groups WHERE study_org_unit_id=? AND event_date=? AND activity_type='CLASS_MEETING' AND status='ACTIVE' ORDER BY id",(binding['class_org_unit_id'],actual.date().isoformat())).fetchall()
    source_id=row['source_event_group_id']
    if source_id:
        if source_id not in {r['id'] for r in groups}:raise credits.LearningCreditError('学习日签到来源日期或状态不同，请核对')
    elif len(groups)>1:raise credits.LearningCreditError('同日有多个班会来源，请先在班级学习计划绑定正确活动')
    elif groups:source_id=groups[0]['id']
    return dict(row),int(source_id) if source_id else None


def _cutoff(c,member,day):
    if opening.storage_available(c):
        row=execute(c,'SELECT cutoff_date FROM learning_credit_opening_balances WHERE member_id=?',(member,)).fetchone()
        if row and day<=date.fromisoformat(str(row['cutoff_date'])[:10]):raise credits.LearningCreditError('这一天已包含在期初余额内，不能再次导入加分')


def _fact_id(kind,member,day,reference):
    return 'EXCEL:'+_hash([kind,member,day,reference])


def _checked(c,actor,content,kind,class_id,options_text,note):
    _principal(actor,kind)
    note=str(note).strip()
    if not 4<=len(note)<=1000:raise credits.LearningCreditError('请填写至少4字的来源确认说明')
    binding=_class(c,actor,class_id);options=_options(kind,options_text);rows=_parse(content,kind)
    roster=_members(c,class_id);seen=set();facts=[];cycle=None;event_id=None
    roster_by_account={}
    if kind=='DAILY_READING':
        for member in roster.values():
            last4=member['phone_last4'] or str(member['phone_masked'] or '')[-4:]
            roster_by_account.setdefault((''.join(member['name'].split()),last4),[]).append(member)
    if kind=='CLASS_MEETING':cycle,event_id=_cycle(c,binding,options)
    for r in rows:
        if r.get('original_report'):
            account=str(r['raw']['账户']).strip();matches=[member for member in roster_by_account.get((''.join(r['member_name'].split()),account[-4:]),[])
                if not member['phone_masked'] or member['phone_masked']==account]
            m=matches[0] if len(matches)==1 else None
            if m:r['member_code']=m['member_code']
        else:m=roster.get(r['member_code'])
        if not m or ''.join(m['name'].split())!=''.join(r['member_name'].split()):
            r['errors'].append('姓名与账户尾号未唯一匹配本班名册，请核对后台联系方式' if r.get('original_report') else '编号或姓名未匹配本班名册，请使用预填模板')
            continue
        r['member_id']=int(m['id']);r['member_name']=m['name']
        try:
            day=_date(datetime.fromisoformat(options['actual_start']).date() if kind=='CLASS_MEETING' else r['raw'][HEADERS[kind][2]])
            r['occurred_on']=day.isoformat();_cutoff(c,m['id'],day)
            reason=activity._member_relation_reason(c,m['id'],class_id,day)
            if reason:raise credits.LearningCreditError('学长在这一天没有唯一有效的本班学习关系')
            if kind=='CLASS_MEETING':
                for _,label in SLOTS:
                    if r['raw'][label+'状态'] not in STATES:raise credits.LearningCreditError('上午、下午、空巴均需明确填写到场、迟到、早退、请假或缺席状态')
                key=(m['id'],day.isoformat())
            else:
                yes=r['raw'][HEADERS[kind][-1]]
                if yes not in ('是','否'):raise credits.LearningCreditError('完成或优秀确认需明确填写“是”或“否”')
                if kind=='EXCELLENT_SHARE' and yes!='是':raise credits.LearningCreditError('本入口只接受已经评定优秀的分享')
                reference=str(r['raw'].get('分享编号') or '').strip() if kind=='EXCELLENT_SHARE' else day.isoformat()
                if not reference or len(reference)>128 or re.search(r'(?<!\d)1[3-9]\d{9}(?!\d)',reference):raise credits.LearningCreditError('请填写稳定的分享编号，不能使用手机号；重复上传保留同一编号')
                r['source_reference']=reference;r['source_id']=_fact_id(kind,m['id'],day.isoformat(),reference)
                r['participation_status']='CONFIRMED' if yes=='是' else 'CANCELLED'
                existing=execute(c,'SELECT * FROM learning_credit_activity_facts WHERE source_type=? AND source_id=?',('CREDIT_EXCEL_'+kind,r['source_id'])).fetchone()
                if existing and (existing['class_org_unit_id']!=class_id or existing['participation_status']!=r['participation_status'] or str(existing['title'] or '')!=str(r['raw'].get('分享标题') or '')):raise credits.LearningCreditError('已有同一来源但内容不同，请核对原记录，不能覆盖')
                if existing and r.get('original_report') and json.loads(existing['metadata_json'] or '{}').get('source_format')!='READING_REPORT':raise credits.LearningCreditError('这一天已有其他格式的读书事实，请核对原来源，不能重复改写')
                # Existing IDs sort before new IDs, matching the durable insert order.
                fact={'id':int(existing['id']) if existing else 10**15+r['excel_row'],'activity_type':kind,'member_id':m['id'],'member_name':m['name'],'class_org_unit_id':class_id,'class_name':binding['class_name'],'binding_id':binding['id'],'occurred_on':day.isoformat(),'_occurred_date':day,'participation_status':r['participation_status'],'source_type':'CREDIT_EXCEL_'+kind,'source_id':r['source_id'],'title':r['raw'].get('分享标题'),'metadata':{'source_format':'READING_REPORT'} if r.get('original_report') else {}}
                r['fact_id']=fact['id'] if existing else None
                key=(m['id'],day.isoformat(),reference)
            if key in seen:raise credits.LearningCreditError('同一学长的同一来源在表内重复，请只保留一行')
            if not r['errors']:seen.add(key)
            if kind!='CLASS_MEETING' and not r['errors']:facts.append((r,fact))
        except (ValueError,TypeError) as exc:r['errors'].append(str(exc))
    if kind=='CLASS_MEETING':_class_preview(c,binding,cycle,event_id,options,rows)
    elif facts:
        valid_facts=[f for _,f in facts]
        proposals=activity._daily_entries(c,valid_facts) if kind=='DAILY_READING' else activity._excellent_entries(c,valid_facts)
        by_fact={fid:p for p in proposals for fid in p['fact_ids']}
        for r,f in facts:
            p=by_fact[f['id']];r.update(proposed_points=format(Decimal(str(p['points'])),'.2f'),credit_status=p['status'],credit_reasons=p['reasons'],rule_snapshot=p['rule_snapshot'])
            if p['status']=='BLOCKED':r['errors'].extend(_reason(x) for x in p['reasons'])
    material={'kind':kind,'class_id':class_id,'binding_id':binding['id'],'options':options,'note':note,'file_sha256':hashlib.sha256(content).hexdigest(),
              'rows':[{k:r.get(k) for k in ('excel_row','member_id','member_code','member_name','raw','occurred_on','source_reference','source_id','participation_status')} for r in rows if kind!='DAILY_READING' or not r['errors']]}
    if kind=='DAILY_READING':
        # Keep row numbers and reasons for traceability, never persist rejected
        # raw cells (which may contain an unmasked account or formula).
        material['ignored_rows']=[{'excel_row':r['excel_row'],'reasons':r['errors']} for r in rows if r['errors']]
    if len(_json(material).encode())>8*1024*1024:raise credits.LearningCreditError('来源记录过大，请分批导入')
    return binding,options,rows,_hash(material),note,event_id,material


def _reason(value):
    return {'BUSINESS_CALENDAR_MISSING':'该日期工作日日历尚未发布，请由管理员补齐日历后重检','BUSINESS_CALENDAR_INVALID':'工作日日历配置无效，请修复后重检','LEARNING_ROUND_MISSING':'这一天没有有效学习计划','MEMBER_CLASS_RELATION_MISSING':'缺少当天班级学习关系'}.get(value,value)


def _class_preview(c,binding,cycle,event_id,options,rows):
    starts=credits._db_timestamp(c,datetime.fromisoformat(options['actual_start']))
    sessions=class_credits._meeting_context(c,event_id)['sessions'] if event_id else [{'id':-i,'session_code':code,'session_name':label,'scheduled_start_at':starts if i==1 else None,'status':'ACTIVE'} for i,(code,label) in enumerate(SLOTS,1)]
    if {s['session_code'] for s in sessions}!={x[0] for x in SLOTS}:raise credits.LearningCreditError('本入口使用上午、下午、空巴三场次，请在活动管理核对场次')
    # Resolve the same bound generic policy used by the formal class-meeting projector.
    from app.services.credit_rule_mapping import bound_generic_rule
    rule=bound_generic_rule(c,binding['credit_rule_version_id'],'CLASS_MEETING_SCORE')
    meeting={'id':event_id or -int(cycle['id']),'study_org_unit_id':binding['class_org_unit_id'],'class_org_unit_id':binding['class_org_unit_id'],'org_unit_id':binding['class_org_unit_id'],'class_name':binding['class_name'],'title':'班会线下签到','activity_type':'CLASS_MEETING','event_date':datetime.fromisoformat(options['actual_start']).date().isoformat(),'status':'ACTIVE','binding_id':binding['id'],'learning_cycle_id':cycle['id'],'sessions':sessions,'credit_rule':rule,'plan_status':binding['plan_status'],'plan_key':binding['plan_key'],'version_label':binding['version_label']}
    for r in rows:
        if r['errors']:continue
        meeting['credit_rule_version_id']=binding['credit_rule_version_id']
        details=[];adjudications={}
        for slot in sessions:
            label=dict(SLOTS)[slot['session_code']];status,late,early=STATES[r['raw'][label+'状态']]
            old=execute(c,'SELECT * FROM attendance_records WHERE attendance_session_id=? AND member_id=?',(slot['id'],r['member_id'])).fetchall() if event_id else []
            if len(old)>1:r['errors'].append('这一场次有重复签到事实，请先在签到管理核对');continue
            accepted={status,'ABSENT'}|({'PRESENT'} if status=='MANUAL_PRESENT' else set())
            if old and (old[0]['attendance_status'] not in accepted or old[0]['participant_type']!='MEMBER' or not old[0]['score_eligible']):r['errors'].append('已有码签到或请假记录与原表不同，请通过签到裁定核对');continue
            saved_score=execute(c,'SELECT * FROM attendance_score_records WHERE attendance_record_id=?',(old[0]['id'],)).fetchone() if old else None
            if old and old[0]['attendance_status']=='ABSENT' and status!='ABSENT' and execute(c,"SELECT id FROM learning_credit_entries WHERE member_id=? AND class_org_unit_id=? AND credit_type='CLASS_MEETING_SCORE' AND substr(occurred_at,1,10)=? AND status IN ('POSTED','REVERSED') LIMIT 1",(r['member_id'],binding['class_org_unit_id'],r['occurred_on'])).fetchone():
                r['errors'].append('当天班会已经入账，请通过签到裁定与账本冲销处理补签，不能覆盖已入账事实');continue
            if old and old[0]['attendance_status'] in {'PRESENT','MANUAL_PRESENT'} and saved_score and (bool(saved_score['is_late'])!=late or bool(saved_score['is_early_leave'])!=early):r['errors'].append('原表迟到或早退状态与已有签到不同，请通过签到裁定核对');continue
            if old and old[0]['attendance_status']=='PRESENT' and not saved_score:
                actual_score=calculate_score(attendance_record_id=int(old[0]['id']),member_id=r['member_id'],session_code=slot['session_code'],attendance_status='PRESENT',checked_at=old[0]['checked_at'],scheduled_start_at=slot['scheduled_start_at'],score_eligible=True,_connection=c)
                if bool(actual_score['is_late'])!=late or bool(actual_score['is_early_leave'])!=early:
                    r['errors'].append('原表状态与实际扫码时间不同，请通过签到裁定核对');continue
            record_id=int(old[0]['id']) if old else -(r['excel_row']*10+abs(slot['id']))
            try:
                score=calculate_score(attendance_record_id=record_id,member_id=r['member_id'],session_code=slot['session_code'],attendance_status=status,checked_at=None,scheduled_start_at=slot['scheduled_start_at'],score_eligible=True,_connection=c,_late_override=late,_early_override=early)
                if saved_score and old[0]['attendance_status'] in {status,'PRESENT'}:
                    score=dict(saved_score)
            except ValueError as exc:r['errors'].append(str(exc));continue
            details.append({'attendance_record_id':record_id,'attendance_session_id':slot['id'],'attendance_member_id':r['member_id'],'member_name':r['member_name'],'score_record_id':record_id,'score_member_id':r['member_id'],'score_rule_id':score['rule_id'],'score_rule_version':score['rule_version'],'session_code':slot['session_code'],'session_name':label,'attendance_status':status,'attendance_updated_at':None,'score_calculated_at':None,'latest_adjudication_at':None,**score})
            details[-1]['score_eligible']=True
            adjudications[record_id]=['EARLY_LEAVE'] if early else []
        if r['errors']:continue
        p=class_credits._member_proposal(c,meeting,details,adjudications,class_credits._common_meeting_reasons(meeting))
        r.update(source_score_points=format(Decimal(str(p['points'])),'.2f'),proposed_points=format(Decimal(str(p['points'] if p['status']=='READY' else 0)),'.2f'),credit_status=p['status'],credit_reasons=p['reasons'],rule_snapshot=p['rule_snapshot'])
        if p['status']=='BLOCKED':r['errors'].extend(p['reasons'])


def preview(content,kind,class_id,options,note,actor):
    c=connect()
    try:
        binding,options,rows,fingerprint,note,_,_=_checked(c,actor,content,kind,class_id,options,note)
        errors=sum(bool(r['errors']) for r in rows)
        accepted=len(rows)-errors
        return {'rows':rows,'row_count':len(rows),'error_count':errors,'accepted_count':accepted,'ignored_count':errors if kind=='DAILY_READING' else 0,'fingerprint':fingerprint if accepted and (not errors or kind=='DAILY_READING') else None,'proposed_points':format(sum((Decimal(r.get('proposed_points','0')) for r in rows if not r['errors'] and r.get('credit_status')=='READY'),Decimal(0)),'.2f'),'class_name':binding['class_name'],'storage_available':available(c)}
    finally:c.close()


def _receipt(c,actor,import_id,*,allowed_scope=...,summary=False):
    columns='id,fingerprint,file_sha256,original_filename,source_kind,class_org_unit_id,binding_id,cycle_id,event_group_id,options_json,source_note,created_by,created_at' if summary else '*'
    row=execute(c,'SELECT '+columns+' FROM learning_credit_source_imports WHERE id=?',(import_id,)).fetchone()
    if not row:raise credits.LearningCreditError('导入来源不存在')
    permitted=_has_scope(actor,row['class_org_unit_id']) if allowed_scope is ... else allowed_scope is None or row['class_org_unit_id'] in allowed_scope
    if not permitted:raise PermissionError('导入来源不在当前授权范围')
    item=dict(row);item['class_name']=execute(c,'SELECT name FROM org_units WHERE id=?',(row['class_org_unit_id'],)).fetchone()['name']
    rows=[] if summary else [dict(r) for r in execute(c,'SELECT r.*,m.member_code,m.name AS member_name FROM learning_credit_source_rows r JOIN members m ON m.id=r.member_id WHERE r.import_id=? ORDER BY r.excel_row,r.id',(import_id,)).fetchall()]
    if not summary:_validate_receipt(c,item,rows)
    if not summary:item['ignored_rows']=json.loads(item['source_snapshot_json']).get('ignored_rows',[])
    stats=execute(c,'SELECT COUNT(DISTINCT excel_row) AS n,COUNT(fact_id) AS fact_count FROM learning_credit_source_rows WHERE import_id=?',(import_id,)).fetchone() if summary else None
    item['rows']=rows;item['row_count']=int(stats['n']) if summary else len({r['excel_row'] for r in rows})
    item['source_only']=item['source_kind']=='DAILY_READING' and (int(stats['fact_count'])==0 if summary else not any(r['fact_id'] is not None for r in rows))
    source_stats=execute(c,'SELECT COUNT(DISTINCT member_id) AS members,MIN(occurred_on) AS first_day,MAX(occurred_on) AS last_day FROM learning_credit_source_rows WHERE import_id=?',(import_id,)).fetchone()
    item['member_count']=int(source_stats['members'])
    item['period_start']=str(source_stats['first_day'])[:10] if source_stats['first_day'] else None
    item['period_end']=str(source_stats['last_day'])[:10] if source_stats['last_day'] else None

    item['batches']=[dict(r) for r in execute(c,'SELECT b.id,b.batch_no,b.status,b.period_start,b.period_end,b.blocked_count,b.proposed_entry_count,b.proposed_points,b.posted_entry_count,b.posted_points,b.created_by,b.source_fingerprint,b.rule_fingerprint FROM learning_credit_source_batches l JOIN learning_credit_settlement_batches b ON b.id=l.batch_id WHERE l.import_id=? ORDER BY b.id',(import_id,)).fetchall()]
    for batch in item['batches']:
        if summary:
            count=execute(c,"SELECT COUNT(DISTINCT member_id) AS n FROM learning_credit_settlement_batch_items WHERE batch_id=? AND status<>'BLOCKED'",(batch['id'],)).fetchone()
            batch['member_count']=int(count['n'])
            continue
        batch_items=execute(c,"SELECT * FROM learning_credit_settlement_batch_items WHERE batch_id=? AND status<>'BLOCKED' ORDER BY id",(batch['id'],)).fetchall()
        batch['confirmation_fingerprint']=batches._approval_fingerprint({**batch,'id':batch['id']},batch_items)
        batch['member_count']=len({r['member_id'] for r in batch_items})
    item['settlement_fingerprint']='' if summary else _hash([(b['id'],b['confirmation_fingerprint']) for b in item['batches']])
    item['settlement_summary']={'proposed_points':format(sum((Decimal(str(b['proposed_points'])) for b in item['batches']),Decimal(0)),'.2f'),
        'posted_points':format(sum((Decimal(str(b['posted_points'])) for b in item['batches']),Decimal(0)),'.2f'),
        'entry_count':sum(int(b['proposed_entry_count']) for b in item['batches'])}
    if not summary:
        people={r['member_id']:{'member_code':r['member_code'],'member_name':r['member_name'],
            'credit_count':0,'proposed_points':Decimal(0),'posted_points':Decimal(0),'blocked_count':0} for r in rows}
        entries=execute(c,"SELECT i.*,m.name AS result_member_name,m.member_code AS result_member_code FROM learning_credit_settlement_batch_items i JOIN members m ON m.id=i.member_id JOIN learning_credit_source_batches l ON l.batch_id=i.batch_id WHERE l.import_id=? ORDER BY i.id",(import_id,)).fetchall()
        for entry in entries:
            person=people.setdefault(entry['member_id'],{'member_code':entry['result_member_code'],'member_name':entry['result_member_name'],
                'credit_count':0,'proposed_points':Decimal(0),'posted_points':Decimal(0),'blocked_count':0})
            if entry['status']=='BLOCKED':person['blocked_count']+=1;continue
            person['credit_count']+=1
            person['proposed_points']+=Decimal(str(entry['points']))
            if entry['status']=='POSTED':person['posted_points']+=Decimal(str(entry['points']))
        item['source_member_count']=item['member_count']
        item['member_count']=len(people)
        item['member_results']=[{**r,'proposed_points':format(r['proposed_points'],'.2f'),
            'posted_points':format(r['posted_points'],'.2f')} for r in people.values()]
    for batch in item['batches']:
        if summary:break
        stored=execute(c,'SELECT source_type,source_snapshot_json FROM learning_credit_settlement_batches WHERE id=?',(batch['id'],)).fetchone()
        saved=json.loads(stored['source_snapshot_json'])
        source_kind=item['source_kind']
        if source_kind=='CLASS_MEETING':
            matches=stored['source_type']=='CLASS_MEETING' and saved.get('event_group_id')==item['event_group_id']
        else:
            ids={int(r['fact_id']) for r in rows if r['fact_id'] is not None}
            selected=saved.get('fact_ids_filter') or []
            matches=stored['source_type']==source_kind and saved.get('class_org_unit_id')==item['class_org_unit_id'] and bool(selected) and set(selected)<=ids
        if not matches:raise credits.LearningCreditError('结算批次与原文件来源不同，请核对关联')
    item['created_by_current_actor']=int(row['created_by'])==actor
    for b in item['batches']:b['created_by_current_actor']=int(b['created_by'])==actor
    item.pop('source_snapshot_json',None)
    return item


def _validate_receipt(c,source,rows):
    snapshot=json.loads(source['source_snapshot_json'])
    if _hash(snapshot)!=source['fingerprint'] or snapshot['kind']!=source['source_kind'] or snapshot['class_id']!=source['class_org_unit_id'] or snapshot['binding_id']!=source['binding_id'] or snapshot['file_sha256']!=source['file_sha256'] or snapshot['note']!=source['source_note'] or _json(snapshot['options'])!=source['options_json']:
        raise credits.LearningCreditError('原文件来源快照校验失败，请核对导入记录')
    frozen={int(r['excel_row']):r for r in snapshot['rows']}
    expected_count=len(frozen)*(3 if source['source_kind']=='CLASS_MEETING' else 1)
    if len(rows)!=expected_count:raise credits.LearningCreditError('原表行记录不完整，请核对')
    seen=set()
    for r in rows:
        key=(r['excel_row'],r['source_reference'])
        saved=frozen.get(int(r['excel_row']))
        if not saved or key in seen or int(saved['member_id'])!=int(r['member_id']) or saved['occurred_on']!=str(r['occurred_on'])[:10]:raise credits.LearningCreditError('来源行的学长、日期或编号已变化')
        seen.add(key)
        record=json.loads(r['record_json'])
        if source['source_kind']=='CLASS_MEETING':
            label=dict(SLOTS).get(r['source_reference'])
            if not label or record['state']!=saved['raw'][label+'状态'] or (record['attendance_status'],bool(record['late']),bool(record['early']))!=STATES[record['state']]:raise credits.LearningCreditError('线下签到来源状态校验失败')
            ar=execute(c,'SELECT a.member_id,s.event_group_id,s.session_code FROM attendance_records a JOIN attendance_sessions s ON s.id=a.attendance_session_id WHERE a.id=?',(r['attendance_record_id'],)).fetchone()
            if not ar or ar['member_id']!=r['member_id'] or ar['event_group_id']!=source['event_group_id'] or ar['session_code']!=r['source_reference']:raise credits.LearningCreditError('签到事实身份或场次已变化')
        else:
            if _json(saved['raw'])!=_json(record) or saved['source_reference']!=r['source_reference']:raise credits.LearningCreditError('原始打卡或分享来源已变化')
            if r['fact_id'] is None:
                if source['source_kind']=='DAILY_READING' and saved['participation_status']=='CANCELLED' and record.get('已完成录音')=='否':continue
                raise credits.LearningCreditError('来源缺少对应的确认事实，请核对')
            fact=execute(c,'SELECT * FROM learning_credit_activity_facts WHERE id=?',(r['fact_id'],)).fetchone()
            if not fact or fact['member_id']!=r['member_id'] or fact['class_org_unit_id']!=source['class_org_unit_id'] or fact['source_type']!='CREDIT_EXCEL_'+source['source_kind'] or fact['source_id']!=saved['source_id'] or fact['participation_status']!=saved['participation_status'] or str(fact['occurred_on'])[:10]!=saved['occurred_on'] or str(fact['title'] or '')!=str(saved['raw'].get('分享标题') or ''):raise credits.LearningCreditError('打卡或分享事实已变化，请重新核对')


def register(content,kind,class_id,options_text,note,filename,fingerprint,actor):
    _principal(actor,kind,write=True)
    _write_scope(actor,class_id,kind)
    with transaction() as c,opening._operation(c,wait_seconds=10):
        if not available(c):raise credits.LearningCreditError('来源导入存储待管理员准备')
        prior=execute(c,'SELECT id,file_sha256,source_kind,class_org_unit_id,options_json,source_note FROM learning_credit_source_imports WHERE fingerprint=?',(fingerprint,)).fetchone()
        digest=hashlib.sha256(content).hexdigest()
        if prior:
            if prior['file_sha256']!=digest or prior['source_kind']!=kind or prior['class_org_unit_id']!=class_id or prior['options_json']!=_json(_options(kind,options_text)) or prior['source_note']!=note.strip():raise credits.LearningCreditError('来源已变化，请重新检查')
            c.commit()
            return {'id':int(prior['id']),'idempotent':True}
        binding,options,rows,current,note,event_id,source_snapshot=_checked(c,actor,content,kind,class_id,options_text,note)
        if current!=fingerprint or (kind!='DAILY_READING' and any(r['errors'] for r in rows)):raise credits.LearningCreditError('来源、名册、规则或账本已变化，请重新检查')
        ignored_count=sum(bool(r['errors']) for r in rows)
        if kind=='DAILY_READING':rows=[r for r in rows if not r['errors']]
        if not rows:raise credits.LearningCreditError('没有可保存的有效记录，请核对表格后重新检查')
        now=credits._db_timestamp(c)
        if kind=='CLASS_MEETING' and not event_id:
            event_id=int(execute(c,"INSERT INTO attendance_event_groups(source_key,external_group_id,org_unit_id,study_org_unit_id,title,activity_type,event_date,status,source_updated_at,created_at,updated_at) VALUES ('CREDIT_EXCEL',?,?,?,?,'CLASS_MEETING',?,'ACTIVE',?,?,?)",('cycle:'+str(options['cycle_id']),class_id,class_id,binding['class_name']+'线下班会',rows[0]['occurred_on'],now,now,now)).lastrowid)
            for i,(code,label) in enumerate(SLOTS):execute(c,"INSERT INTO attendance_sessions(event_group_id,external_session_id,session_code,session_name,session_order,scheduled_start_at,status,source_revision,source_updated_at,finalized_at,created_at,updated_at) VALUES (?,?,?,?,?,?,'ACTIVE',1,?,?,?,?)",(event_id,'OFFLINE:'+code,code,label,i,credits._db_timestamp(c,datetime.fromisoformat(options['actual_start'])) if i==0 else None,now,now,now,now))
        import_id=int(execute(c,'INSERT INTO learning_credit_source_imports(fingerprint,file_sha256,original_filename,source_kind,class_org_unit_id,binding_id,cycle_id,event_group_id,options_json,source_snapshot_json,source_note,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',(fingerprint,digest,Path(filename).name[:255],kind,class_id,binding['id'],options.get('cycle_id'),event_id,_json(options),_json(source_snapshot),note,actor,now)).lastrowid)
        if kind=='CLASS_MEETING':_store_attendance(c,actor,import_id,event_id,rows,note,now)
        else:
            for r in rows:
                existing=execute(c,'SELECT id FROM learning_credit_activity_facts WHERE source_type=? AND source_id=?',('CREDIT_EXCEL_'+kind,r['source_id'])).fetchone()
                if kind=='DAILY_READING' and r['participation_status']=='CANCELLED':fact_id=None
                elif existing:fact_id=int(existing['id'])
                else:
                    fact=activity.record_learning_activity_fact(actor_user_id=actor,activity_type=kind,member_id=r['member_id'],class_org_unit_id=class_id,occurred_on=r['occurred_on'],source_type='CREDIT_EXCEL_'+kind,source_id=r['source_id'],participation_status=r['participation_status'],title=r['raw'].get('分享标题'),metadata={'source_import_id':import_id,'file_sha256':digest,'excel_row':r['excel_row'],'reference':r['source_reference'],**({'source_format':'READING_REPORT'} if r.get('original_report') else {})},binding_id=binding['id'],_connection=c)
                    fact_id=int(fact['id'])
                execute(c,'INSERT INTO learning_credit_source_rows(import_id,member_id,excel_row,occurred_on,source_reference,record_json,fact_id) VALUES (?,?,?,?,?,?,?)',(import_id,r['member_id'],r['excel_row'],r['occurred_on'],r['source_reference'],_json(r['raw']),fact_id))
        write_audit(c,actor_user_id=actor,action='learning.credit_source.import',resource_type='credit_source_import',resource_id=str(import_id),org_unit_id=class_id,after={'kind':kind,'count':len(rows),'ignored_count':ignored_count,'file_sha256':digest,'fingerprint':fingerprint,'ledger_write':False})
        # Publish the receipt before releasing the shared operation lock so a
        # waiting retry observes this committed idempotency key.
        c.commit()
        return {'id':import_id,'idempotent':False}


def _store_attendance(c,actor,import_id,event_id,rows,note,now):
    sessions={r['session_code']:dict(r) for r in execute(c,'SELECT * FROM attendance_sessions WHERE event_group_id=?',(event_id,)).fetchall()}
    for r in rows:
        for code,label in SLOTS:
            status,late,early=STATES[r['raw'][label+'状态']];slot=sessions[code]
            previous=execute(c,'SELECT * FROM attendance_records WHERE attendance_session_id=? AND member_id=?',(slot['id'],r['member_id'])).fetchone()
            if previous:
                rid=int(previous['id'])
                # Import complements an absent roster. Existing scan/manual facts
                # are retained; discrepancies require the ordinary adjudication UI.
                if previous['attendance_status']=='ABSENT' and status!='ABSENT':
                    execute(c,"UPDATE attendance_records SET attendance_status=?,checked_at=NULL,checkin_source='EXCEL_OFFLINE',source_updated_at=?,updated_at=? WHERE id=?",(status,now,now,rid))
                elif previous['checkin_source']=='EXCEL_OFFLINE':
                    old=offline_flags(c,rid)
                    if old is not None and old!=(late,early):raise credits.LearningCreditError('已存在不同的线下签到状态，请核对原来源')
            else:
                rid=int(execute(c,"INSERT INTO attendance_records(attendance_session_id,external_record_id,member_id,member_code_snapshot,name_snapshot,participant_type,score_eligible,attendance_status,checkin_source,source_updated_at,received_at,created_at,updated_at) VALUES (?,?,?,?,?,'MEMBER',1,?,'EXCEL_OFFLINE',?,?,?,?)",(slot['id'],'EXCEL:'+str(r['member_id']),r['member_id'],r['member_code'],r['member_name'],status,now,now,now,now)).lastrowid)
            record={'state':r['raw'][label+'状态'],'session_code':code,'late':late,'early':early,'attendance_status':status}
            execute(c,'INSERT INTO learning_credit_source_rows(import_id,member_id,excel_row,occurred_on,source_reference,record_json,attendance_record_id) VALUES (?,?,?,?,?,?,?)',(import_id,r['member_id'],r['excel_row'],r['occurred_on'],code,_json(record),rid))
            is_offline=not previous or previous['checkin_source']=='EXCEL_OFFLINE' or previous['attendance_status']=='ABSENT'
            previous_score=execute(c,'SELECT id FROM attendance_score_records WHERE attendance_record_id=?',(rid,)).fetchone() if previous else None
            if is_offline and (not previous or previous['attendance_status']!=status or not previous_score):
                if not previous or previous['attendance_status']!=status:
                    kind='MANUAL_CHECKIN' if status=='MANUAL_PRESENT' else 'LEAVE' if status=='LEAVE' else None
                    if kind:execute(c,'INSERT INTO attendance_adjudications(attendance_record_id,adjudication_type,reason,actor_user_id,created_at) VALUES (?,?,?,?,?)',(rid,kind,'线下签到表确认：'+note[:450],actor,now))
                    if early:execute(c,"INSERT INTO attendance_adjudications(attendance_record_id,adjudication_type,reason,actor_user_id,created_at) VALUES (?,'EARLY_LEAVE',?,?,?)",(rid,'线下签到表确认早退：'+note[:450],actor,now))
                upsert_score_record(attendance_record_id=rid,member_id=r['member_id'],session_code=code,attendance_status=status,checked_at=None,scheduled_start_at=slot['scheduled_start_at'],score_eligible=True,source_updated_at=now,_connection=c)
            elif previous and previous['attendance_status']=='PRESENT' and not previous_score:
                upsert_score_record(attendance_record_id=rid,member_id=r['member_id'],session_code=code,attendance_status='PRESENT',checked_at=previous['checked_at'],scheduled_start_at=slot['scheduled_start_at'],score_eligible=True,source_updated_at=previous['source_updated_at'],_connection=c)


def offline_flags(c,record_id):
    if not available(c):return None
    rows=execute(c,"SELECT r.record_json FROM learning_credit_source_rows r JOIN learning_credit_source_imports i ON i.id=r.import_id JOIN attendance_records a ON a.id=r.attendance_record_id WHERE r.attendance_record_id=? AND i.source_kind='CLASS_MEETING' AND a.checkin_source='EXCEL_OFFLINE' ORDER BY i.id DESC,r.id DESC LIMIT 1",(record_id,)).fetchall()
    values={(bool(v['late']),bool(v['early'])) for v in (json.loads(r['record_json']) for r in rows)}
    if len(values)>1:raise credits.LearningCreditError('线下签到来源状态有冲突，请核对')
    return next(iter(values)) if values else None


def create_batches(import_id,actor):
    _principal(actor)
    c=connect()
    try:
        if not available(c):raise credits.LearningCreditError('来源存储尚未准备')
        source=_receipt(c,actor,import_id)
    finally:c.close()
    kind=source['source_kind']
    if any(b['status'] in {'POSTING','PARTIAL_FAILED'} for b in source['batches']):
        raise credits.LearningCreditError('本次来源正在入账或有部分失败，请先核对结果并继续入账')
    if source['batches'] and all(b['status']=='POSTED' or not int(b['proposed_entry_count']) for b in source['batches']):
        return source
    if kind=='CLASS_MEETING':
        _principal(actor,kind,write=True)
        _write_scope(actor,source['class_org_unit_id'],kind)
        options=json.loads(source['options_json'])
        # Freeze the source/cycle association for the preview and ledger.
        # Confirmation and progression belong to successful formal posting.
        with transaction() as c:
            from app.services.learning_cycles import _lock_class_for_update
            _lock_class_for_update(c,source['class_org_unit_id'])
            binding=_class(c,actor,source['class_org_unit_id'])
            cycle,event_id=_cycle(c,binding,options)
            if event_id!=source['event_group_id']:raise credits.LearningCreditError('学习日签到来源已变化，请核对')
            if cycle['source_event_group_id'] is None:
                execute(c,'UPDATE class_learning_cycles SET source_event_group_id=?,updated_at=? WHERE id=?',(source['event_group_id'],credits._db_timestamp(c),source['cycle_id']))
                write_audit(c,actor_user_id=actor,action='learning.credit_source.link_cycle',resource_type='class_learning_cycle',resource_id=str(source['cycle_id']),org_unit_id=source['class_org_unit_id'],after={'source_event_group_id':source['event_group_id'],'progress_changed':False})
        plans=[{'event_group_id':source['event_group_id']}]
    else:
        by_period={}
        for r in source['rows']:
            if r['fact_id'] is None:continue
            period=str(r['occurred_on'])[:10] if kind=='DAILY_READING' else str(r['occurred_on'])[:7]
            by_period.setdefault(period,[]).append(r)
        plans=[]
        for period,rows in by_period.items():
            # Keep every person's month together so monthly caps are explained
            # before approval. Reading's day key is also kept together.
            if len(rows)>500:raise credits.LearningCreditError('同一天读书或同月分享超过500行，请缩小班级/文件范围')
            dates=sorted(str(r['occurred_on'])[:10] for r in rows)
            plans.append({'occurred_from':dates[0],'occurred_to':dates[-1],'fact_ids':sorted({int(r['fact_id']) for r in rows})})
    new_ids=set()
    for plan in plans:
        if kind=='CLASS_MEETING':batch=batches.dry_run_class_meeting_batch(actor_user_id=actor,**plan)
        else:batch=batches.dry_run_activity_import_batch(actor_user_id=actor,activity_type=kind,class_org_unit_id=source['class_org_unit_id'],**plan)
        new_ids.add(batch['id'])
        with transaction() as c:
            if not execute(c,'SELECT batch_id FROM learning_credit_source_batches WHERE import_id=? AND batch_id=?',(import_id,batch['id'])).fetchone():execute(c,'INSERT INTO learning_credit_source_batches(import_id,batch_id) VALUES (?,?)',(import_id,batch['id']))
    with transaction() as c:
        for old in source['batches']:
            if old['id'] not in new_ids and old['status'] in {'DRY_RUN','PENDING_APPROVAL','APPROVED'}:
                changed=execute(c,"UPDATE learning_credit_settlement_batches SET status='CANCELLED',updated_at=? WHERE id=? AND status=?",(credits._db_timestamp(c),old['id'],old['status']))
                if changed.rowcount!=1:raise credits.LearningCreditError('原结算状态已变化，请刷新查看实际结果')
                execute(c,'DELETE FROM learning_credit_source_batches WHERE import_id=? AND batch_id=?',(import_id,old['id']))
                write_audit(c,actor_user_id=actor,action='learning.credit_source.replace_preview',resource_type='credit_source_import',resource_id=str(import_id),org_unit_id=source['class_org_unit_id'],after={'previous_batch_id':old['id'],'new_batch_ids':sorted(new_ids),'ledger_write':False})
    return get_import(import_id,actor)


def get_import(import_id,actor):
    _principal(actor);c=connect()
    try:return _receipt(c,actor,import_id)
    finally:c.close()


def process(import_id,actor,operation,expected_fingerprint=None,resume_partial_failure=False):
    source=get_import(import_id,actor)
    if operation not in ('submit','approve','post','confirm'):raise credits.LearningCreditError('不支持的结算操作')
    capability={'submit':'plans:credit_settlement_manage','approve':'plans:credit_settlement_approve','post':'plans:credit_settlement_post','confirm':'plans:credit_settlement_post'}[operation]
    if not _has_scope(actor,source['class_org_unit_id'],(capability,)):raise PermissionError('这份来源不在当前复核或入账授权范围')
    if not source['batches']:raise credits.LearningCreditError('请先生成结算预览')
    if operation=='confirm':
        if not expected_fingerprint or expected_fingerprint!=source['settlement_fingerprint']:
            raise credits.LearningCreditError('结算预览已变化，请重新查看后确认入账')
        _principal(actor,source['source_kind'],write=True)
        _write_scope(actor,source['class_org_unit_id'],source['source_kind'])
    for b in source['batches']:
        status=b['status']
        if not int(b['proposed_entry_count']):continue
        if operation=='confirm':
            if status not in {'DRY_RUN','PENDING_APPROVAL','APPROVED','POSTED','PARTIAL_FAILED'}:
                raise credits.LearningCreditError('入账未全部完成，请核对结算结果；部分失败请到日常活动结算处理')
            result=batches._post_batch(actor_user_id=actor,batch_id=b['id'],source_type=source['source_kind'],resume_partial_failure=resume_partial_failure,
                source_confirmation={'import_id':import_id,'fingerprint':b['confirmation_fingerprint']})
            if result['status']!='POSTED':raise credits.LearningCreditError('本次入账未全部完成，请查看结果；已入账条目会保留，未入账条目请核对后处理')
        elif operation=='submit' and status=='DRY_RUN':batches.submit_settlement_batch_for_approval(actor_user_id=actor,batch_id=b['id'])
        elif operation=='approve' and status=='PENDING_APPROVAL':batches.approve_settlement_batch(actor_user_id=actor,batch_id=b['id'])
        elif operation=='post' and status=='APPROVED':batches.post_settlement_batch(actor_user_id=actor,batch_id=b['id'])
        elif status not in {'POSTED','PENDING_APPROVAL','APPROVED'}:raise credits.LearningCreditError('批次状态需要核对，请到日常活动结算查看具体原因')
    return get_import(import_id,actor)


def workbench(actor):
    user=_principal(actor);c=connect()
    try:
        scope=effective_scope(actor)
        bound={r['class_org_unit_id'] for r in execute(c,"SELECT DISTINCT class_org_unit_id FROM class_learning_bindings WHERE status='ACTIVE'").fetchall()}
        classes=[{'id':r['id'],'name':r['name'],'center_name':r['parent_name']} for r in class_catalog(c,scope) if r['id'] in bound]
        ids={r['id'] for r in classes}
        cycles=[dict(r) for r in execute(c,"SELECT c.id,c.class_org_unit_id,c.learning_cycle_index,c.cycle_status,c.class_meeting_status,c.actual_class_meeting_at,c.planned_class_meeting_at FROM class_learning_cycles c JOIN class_learning_bindings b ON b.id=c.binding_id WHERE b.status='ACTIVE' AND c.cycle_status IN ('OPEN','CLOSED') ORDER BY c.id DESC").fetchall() if r['class_org_unit_id'] in ids]
        ready=available(c);imports=[]
        if ready:
            for r in execute(c,'SELECT id,class_org_unit_id FROM learning_credit_source_imports ORDER BY id DESC LIMIT 100').fetchall():
                if scope is None or r['class_org_unit_id'] in scope:imports.append(_receipt(c,actor,r['id'],allowed_scope=scope,summary=True))
        tables=present_tables(c);reserved=execute(c,"SELECT id FROM audit_logs WHERE action='production.credit_source.setup' LIMIT 1").fetchone();marker=execute(c,'SELECT version FROM schema_migrations WHERE version=?',(MIGRATION,)).fetchone()
        return {'classes':classes,'cycles':cycles,'imports':imports,'storage_available':ready,'release_commit':get_build_info()['commit_sha'],'migration_sha256':HASHES['sqlite' if isinstance(c,sqlite3.Connection) else 'mysql'],
                'setup_allowed':'system_admin' in user.get('roles',[]) and 'plans:production_rule_reconciliation_apply' in user['permissions'] and get_settings().credit_opening_setup_enabled and not tables and not reserved and not marker,
                'setup_incomplete':not ready and bool(tables or reserved or marker),'can_import_activity':'plans:credit_activity_fact_manage' in user['permissions'],'can_import_attendance':{'attendance:adjudicate','plans:period_write'}<=set(user['permissions']), 'reading_calendar':reading_calendar(c)}
    finally:c.close()


def reading_calendar(c):
    row=execute(c,"SELECT v.version_label,COUNT(d.id) AS day_count FROM learning_business_calendar_versions v LEFT JOIN learning_business_calendar_days d ON d.calendar_version_id=v.id WHERE v.calendar_key='CHINA_MAINLAND' AND v.calendar_year=2026 AND v.status='PUBLISHED' GROUP BY v.id,v.version_label ORDER BY v.id DESC LIMIT 1").fetchone()
    return {'status':'PUBLISHED' if row else 'MISSING','version_label':row['version_label'] if row else None,'day_count':int(row['day_count']) if row else 0}


def prepare_reading_calendar(actor):
    user=_principal(actor)
    if 'plans:business_calendar_manage' not in user['permissions']:raise PermissionError('当前账号没有工作日日历维护权限')
    credits._write_allowed()
    path=Path(__file__).resolve().parents[1]/'assets/credit-reading-calendar-2026.json'
    payload=json.loads(path.read_text(encoding='utf-8'))
    with transaction() as c,opening._operation(c):
        result=activity.save_business_calendar(actor_user_id=actor,calendar_year=2026,version_label=payload['version_label'],days=payload['days'],status='PUBLISHED',_connection=c,_only_if_missing=True)
        return {'reading_calendar':reading_calendar(c),'idempotent':not result['persisted']}


def template(kind,class_id,actor):
    _principal(actor,kind);c=connect()
    try:
        binding=_class(c,actor,class_id);members=_members(c,class_id)
        if len(members)>5000:raise credits.LearningCreditError('名册超过5000位，请缩小范围')
        if kind not in KINDS:raise credits.LearningCreditError('模板类型无效')
        book=load_workbook(Path(__file__).resolve().parents[1]/'assets'/('credit-source-'+kind.lower()+'.xlsx'));sheet=book[SHEETS[kind]]
        for index,member in enumerate(members.values(),2):
            for col,value in enumerate([str(member['member_code']),member['name']],1):
                cell=sheet.cell(index,col,value);cell.data_type='s'
        stream=io.BytesIO();book.save(stream)
        write_audit(c,actor_user_id=actor,action='learning.credit_source.template',resource_type='credit_source',org_unit_id=class_id,after={'kind':kind,'count':len(members),'contact_fields':False});c.commit()
        return stream.getvalue()
    finally:c.close()


def setup(actor,commit,digest):
    user=_principal(actor)
    if 'system_admin' not in user.get('roles',[]) or 'plans:production_rule_reconciliation_apply' not in user['permissions']:raise PermissionError('仅有存储准备权限的管理员可以执行')
    credits._write_allowed();settings=get_settings()
    if not settings.credit_opening_setup_enabled or settings.run_bootstrap_on_startup or commit!=get_build_info()['commit_sha'] or not re.fullmatch('[a-f0-9]{40}',commit):raise credits.LearningCreditError('发布版本或设置已变化，请刷新核对')
    c=connect()
    try:
        with opening._operation(c):
            dialect='sqlite' if isinstance(c,sqlite3.Connection) else 'mysql';content=(MIGRATION_ROOT/dialect/MIGRATION).read_bytes()
            if hashlib.sha256(content).hexdigest()!=digest or digest!=HASHES[dialect]:raise credits.LearningCreditError('固定迁移校验失败')
            if available(c):return {'status':'READY','idempotent':True}
            if present_tables(c) or execute(c,'SELECT version FROM schema_migrations WHERE version=?',(MIGRATION,)).fetchone() or execute(c,"SELECT id FROM audit_logs WHERE action='production.credit_source.setup' LIMIT 1").fetchone():raise credits.LearningCreditError('存在中断或已执行记录，请核对结构，不自动重试')
            if not batches._batch_tables_available(c):raise credits.LearningCreditError('请先准备正式结算批次存储')
            write_audit(c,actor_user_id=actor,action='production.credit_source.setup',resource_type='schema_migration',resource_id=MIGRATION,result='STARTED',after={'commit':commit,'sha256':digest,'tables':sorted(TABLES),'ledger_write':False});c.commit()
            for sql in _split_mysql(content.decode()):execute(c,sql)
            execute(c,'INSERT INTO schema_migrations(version,applied_at) VALUES (?,?)',(MIGRATION,credits._db_timestamp(c)))
            write_audit(c,actor_user_id=actor,action='production.credit_source.setup.complete',resource_type='schema_migration',resource_id=MIGRATION,after={'commit':commit,'sha256':digest,'ledger_write':False});c.commit()
            return {'status':'READY','idempotent':False}
    finally:c.close()
