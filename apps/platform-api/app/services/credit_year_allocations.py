"""Audited annual classification of existing balances, without changing points."""
import hashlib
import ast
import io
import json
import posixpath
import re
import sqlite3
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from app.db import connect, execute, transaction
from app.core.settings import get_settings
from app.core.build_info import get_build_info
from app.migrations import MIGRATION_ROOT, _split_mysql
from app.services import credit_opening_balances as opening
from app.services.audit import write_audit
from app.services.learning_credits import LearningCreditError, _write_allowed, _db_timestamp

MIGRATION = '0070_learning_credit_year_allocations.sql'
HASHES = {'mysql': 'e334147567494353772e010aa0352e020dfa2dc54d158729a170cffdeb46700b',
          'sqlite': 'f1807c500c01891fed94b509c804e0aecc03005ffef65449c73a899ac5c29200'}
TABLES = {'learning_credit_year_imports', 'learning_credit_year_rows',
          'learning_credit_year_allocations', 'learning_credit_class_progress'}
NS = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
ACTION = 'learning.credit_year'


def present_tables(c):
    sql = "SELECT name AS n FROM sqlite_master WHERE type='table'" if isinstance(c, sqlite3.Connection) else 'SELECT TABLE_NAME AS n FROM information_schema.tables WHERE TABLE_SCHEMA=DATABASE()'
    return {r['n'] for r in execute(c, sql).fetchall()} & TABLES


def available(c):
    if present_tables(c) != TABLES or not execute(c, 'SELECT version FROM schema_migrations WHERE version=?', (MIGRATION,)).fetchone():
        return False
    expected = {
        'learning_credit_year_imports': {'id','content_fingerprint','file_sha256','original_filename','binding_id','year_index','cutoff_date','source_note','status','created_by','approved_by','posted_by','created_at','approved_at','posted_at'},
        'learning_credit_year_rows': {'id','import_id','member_id','opening_import_id','excel_row','points'},
        'learning_credit_year_allocations': {'member_id','opening_import_id','binding_id','year_index','points','import_id'},
        'learning_credit_class_progress': {'binding_id','completed_days','cutoff_date','import_id'},
    }
    return all(fields <= {col[0] for col in execute(c,'SELECT * FROM '+table+' WHERE 1=0').description} for table,fields in expected.items())


def normalized_name(value):
    return ''.join(unicodedata.normalize('NFKC', str(value or '')).split())


def _xml(content):
    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
        raise LearningCreditError('不接受含外部定义的Excel，请另存为普通.xlsx文件')
    return ET.fromstring(content)


def _cells(content, sheet_name):
    # Read source values without parsing irrelevant, sometimes invalid WPS fills.
    # No extraction, formula execution, external links or invented cached scores.
    if not content or len(content) > opening.MAX_BYTES:
        raise LearningCreditError('只支持5MB以内的.xlsx文件')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if len(z.infolist()) > 1000 or sum(f.file_size for f in z.infolist()) > 50 * 1024 * 1024:
                raise LearningCreditError('Excel内容过大，请分批处理')
            if any('externallinks' in n.lower() or 'vbaproject' in n.lower() for n in z.namelist()):
                raise LearningCreditError('不接受宏或外部链接')
            strings = []
            if 'xl/sharedStrings.xml' in z.namelist():
                strings = [''.join(t.text or '' for t in si.findall('.//x:t', NS)) for si in _xml(z.read('xl/sharedStrings.xml')).findall('x:si', NS)]
            rels = {r.attrib['Id']: r.attrib['Target'] for r in _xml(z.read('xl/_rels/workbook.xml.rels')) if r.attrib.get('TargetMode') != 'External'}
            sheets = _xml(z.read('xl/workbook.xml')).findall('x:sheets/x:sheet', NS)
            found = [s for s in sheets if s.attrib['name'] == sheet_name]
            if len(found) != 1:
                raise LearningCreditError(f'找不到“{sheet_name}”工作表')
            target = rels[found[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            if not path.startswith('xl/worksheets/') or not path.endswith('.xml'):
                raise LearningCreditError('工作表来源路径不合法')
            root = _xml(z.read(path))
            cells = {}
            for cell in root.findall('x:sheetData/x:row/x:c', NS):
                coord = cell.attrib['r']
                if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,3}', coord):
                    raise LearningCreditError('工作表行列超出支持范围')
                value = cell.findtext('x:v', None, NS)
                kind = cell.attrib.get('t', 'n')
                if kind == 's' and value is not None:
                    value = strings[int(value)]
                elif kind == 'inlineStr':
                    value = ''.join(t.text or '' for t in cell.findall('.//x:t', NS))
                cells[coord] = {'value': value, 'formula': cell.findtext('x:f', None, NS), 'kind': kind}
            return cells
    except LearningCreditError:
        raise
    except (KeyError, ValueError, IndexError, ET.ParseError, zipfile.BadZipFile) as exc:
        raise LearningCreditError('无法读取Excel原始值，请检查.xlsx文件') from exc


def _number(cell, label):
    try:
        value = Decimal(str(cell.get('value')))
        if cell.get('kind') in {'e', 'b'} or not value.is_finite() or value < 0 or value > Decimal('99999999.99') or value != value.quantize(Decimal('.01')):
            raise InvalidOperation
        return value
    except (InvalidOperation, ValueError):
        raise LearningCreditError(f'{label}缺少有效分值，不能用空白代替0分')


def _source_value(cells, coordinate, row, seen=None, sum_context=False):
    from openpyxl.utils import column_index_from_string, get_column_letter
    seen = set() if seen is None else seen
    if coordinate in seen or len(seen) > 260:
        raise LearningCreditError(f'{coordinate}存在循环公式')
    cell = cells.get(coordinate, {})
    if sum_context and cell.get('kind') in {'s','inlineStr','str'} and not cell.get('formula'):
        return Decimal(0)  # Excel SUM(range) ignores text, including attendance marks.
    if cell.get('value') is None:
        if cell.get('formula'):
            raise LearningCreditError(f'{coordinate}公式缺少缓存，请在Excel计算后保存')
        return Decimal(0)
    try:
        cached = Decimal(str(cell['value']))
        if cell.get('kind') in {'b','e'} or not cached.is_finite():
            raise InvalidOperation
    except InvalidOperation as exc:
        raise LearningCreditError(f'{coordinate}明细不是数字，请核对') from exc
    if not cell.get('formula'):
        return cached
    formula = cell['formula'].lstrip('=').replace('$','').replace(' ','').upper()
    seen = seen | {coordinate}
    ref = re.fullmatch(r'([A-Z]+)' + str(row), formula)
    total = re.fullmatch(r'SUM\(([A-Z]+)' + str(row) + r':([A-Z]+)' + str(row) + r'\)', formula)
    if ref:
        result = _source_value(cells, formula, row, seen)
    elif total:
        first,last = map(column_index_from_string,total.groups())
        if first < 5 or last > 256 or first > last:
            raise LearningCreditError(f'{coordinate}总分范围不合法')
        result = sum((_source_value(cells,f'{get_column_letter(col)}{row}',row,seen,True) for col in range(first,last+1)),Decimal(0))
    else:
        if len(formula)>128 or not re.fullmatch(r'[0-9.+*/()\-]+',formula):
            raise LearningCreditError(f'{coordinate}公式不支持，请核对后粘贴为数值')
        def arithmetic(node):
            if isinstance(node,ast.Constant) and type(node.value) in (int,float) and abs(node.value)<=100000000:
                return Decimal(str(node.value))
            if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
                value = arithmetic(node.operand)
                return -value if isinstance(node.op,ast.USub) else value
            if isinstance(node,ast.BinOp):
                a,b = arithmetic(node.left),arithmetic(node.right)
                if isinstance(node.op,ast.Add): return a+b
                if isinstance(node.op,ast.Sub): return a-b
                if isinstance(node.op,ast.Mult): return a*b
                if isinstance(node.op,ast.Div) and b: return a/b
            raise LearningCreditError(f'{coordinate}算术公式不支持')
        try:
            tree = ast.parse(formula,mode='eval')
            if len(list(ast.walk(tree)))>64: raise ValueError
            result = arithmetic(tree.body)
        except (ValueError,SyntaxError,InvalidOperation) as exc:
            raise LearningCreditError(f'{coordinate}公式无法核对') from exc
    if result != cached:
        raise LearningCreditError(f'{coordinate}缓存与原公式不一致，请重新计算保存')
    return result


def parse_source(content, year):
    if type(year) is not int or year not in (1, 2, 3):
        raise LearningCreditError('请选择第一、第二或第三学年')
    sheet = f'第{year}学年学分统计'
    cells = _cells(content, sheet)
    if cells.get('B2', {}).get('value') != '姓名' or cells.get('D3', {}).get('value') != '总分值':
        raise LearningCreditError('年度统计表需要保留姓名和总分值表头')
    records, seen = [], set()
    for index in range(5, opening.MAX_ROWS + 5):
        serial = cells.get(f'A{index}', {}).get('value')
        name = cells.get(f'B{index}', {}).get('value')
        if normalized_name(name) in {'合计','总计','平均','平均分'}:
            continue
        if not serial or not str(serial).isdigit() or not name:
            continue
        key = normalized_name(name)
        if key in seen:
            raise LearningCreditError(f'第{index}行姓名重复，请先核对同名学长')
        seen.add(key)
        score_cell = cells.get(f'D{index}', {})
        score = _number(score_cell, f'{sheet}!D{index}')
        if _source_value(cells,f'D{index}',index) != score:
            raise LearningCreditError(f'D{index}原总分不一致')
        records.append({'excel_row': index, 'source_name': str(name), 'name_key': key,
                        'points': format(score, '.2f'), 'errors': []})
    if not records:
        raise LearningCreditError('没有找到有效学长年度分值')
    return cells.get('A1', {}).get('value', ''), records


def _binding(c, actor, binding_id):
    opening._permission(actor, opening.OPENING_PERMISSION, 'members:read')
    allowed = opening._scope(actor)
    row = execute(c, 'SELECT b.*,o.name AS class_name FROM class_learning_bindings b JOIN org_units o ON o.id=b.class_org_unit_id WHERE b.id=? AND b.status=\'ACTIVE\' AND o.is_active=1', (binding_id,)).fetchone()
    if not row or (allowed is not None and row['class_org_unit_id'] not in allowed):
        raise PermissionError('班级计划不存在或不在当前授权范围')
    current = execute(c, "SELECT id FROM class_learning_bindings WHERE class_org_unit_id=? AND status='ACTIVE' ORDER BY started_at DESC,id DESC LIMIT 1", (row['class_org_unit_id'],)).fetchone()
    if not current or current['id'] != binding_id:
        raise LearningCreditError('班级当前学习计划已变化，请选择最新计划')
    return dict(row)


def _checked(c, actor, content, binding_id, year, cutoff_text, note):
    cutoff, note = date.fromisoformat(cutoff_text), str(note).strip()
    if cutoff > datetime.now(ZoneInfo('Asia/Shanghai')).date() or cutoff.year < 2000 or len(note) < 8 or len(note) > 1000:
        raise LearningCreditError('请填写有效截止日期及至少8个字的年度确认说明')
    binding = _binding(c, actor, binding_id)
    title, rows = parse_source(content, year)
    if normalized_name(binding['class_name']) not in normalized_name(title):
        raise LearningCreditError('原表班级与选择的班级不一致')
    roster = opening._roster(c, actor)
    names = {}
    ready = available(c)
    opening_ready = opening.storage_available(c)
    for member in roster.values():
        if member['org_unit_id'] == binding['class_org_unit_id'] and not member['ambiguous_class']:
            names.setdefault(normalized_name(member['name']), []).append(member)
    for row in rows:
        members = names.get(row['name_key'], [])
        if len(members) != 1:
            row['errors'].append('姓名未唯一匹配本班有效名册，请核对学员管理中的班级和姓名')
            continue
        m = members[0]
        row.update(member_id=int(m['id']), member_code=m['member_code'], member_name=m['name'])
        balance = execute(c, 'SELECT * FROM learning_credit_opening_balances WHERE member_id=? AND org_unit_id=?', (m['id'], binding['class_org_unit_id'])).fetchone() if opening_ready else None
        if not balance:
            row['errors'].append('这位学长尚无已入账期初余额，请先按累计分模板录入')
            continue
        row['opening_import_id'] = int(balance['import_id'])
        if str(balance['cutoff_date'])[:10] != cutoff.isoformat():
            row['errors'].append('统计截止日期与已入账余额不一致，请核对原来源')
        allocations = execute(c, 'SELECT year_index,binding_id,points FROM learning_credit_year_allocations WHERE member_id=? AND opening_import_id=?', (m['id'], balance['import_id'])).fetchall() if ready else []
        if any(int(a['year_index']) == year for a in allocations):
            row['errors'].append('这一年度已归属，不可重复导入')
        if Decimal(row['points']) + sum((Decimal(str(a['points'])) for a in allocations), Decimal(0)) > Decimal(str(balance['points'])):
            row['errors'].append('年度合计超过已有累计余额，请核对分值，不能重复增加余额')
        if any(int(a['binding_id']) != binding_id for a in allocations):
            row['errors'].append('同一余额已有其他班级学习计划归属，请核对转班记录')
    if ready:
        progress = execute(c, 'SELECT * FROM learning_credit_class_progress WHERE binding_id=?', (binding_id,)).fetchone()
        if progress and (int(progress['completed_days']) != year * 12 or str(progress['cutoff_date'])[:10] != cutoff.isoformat()):
            raise LearningCreditError('该班级已有不同的历史学习进度，请先核对，不能覆盖')
    return binding, rows, cutoff, note


def _fingerprint(binding, rows, year, cutoff, note):
    values = {'binding_id': binding['id'], 'class_id': binding['class_org_unit_id'], 'year_index': year,
              'cutoff_date': cutoff.isoformat(), 'source_note': note, 'completed_days': year * 12,
              'rows': sorted([{k: r[k] for k in ('member_id','opening_import_id','points')} for r in rows], key=lambda r: r['member_id'])}
    return hashlib.sha256(json.dumps(values, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def preview(content, binding_id, year, cutoff, note, actor):
    c = connect()
    try:
        binding, rows, cutoff, note = _checked(c, actor, content, binding_id, year, cutoff, note)
        errors = sum(bool(r['errors']) for r in rows)
        return {'rows': rows, 'row_count': len(rows), 'error_count': errors,
                'total_points': format(sum((Decimal(r['points']) for r in rows), Decimal(0)), '.2f'),
                'fingerprint': _fingerprint(binding, rows, year, cutoff, note) if not errors else None,
                'year_index': year, 'current_year_after_import': year + 1, 'completed_days': year * 12,
                'class_name': binding['class_name'], 'cutoff_date': cutoff.isoformat(), 'source_note': note,
                'storage_available': available(c), 'adds_points': False}
    finally:
        c.close()


def _ready(c):
    _write_allowed()
    if not get_settings().learning_credit_opening_balance_enabled or not available(c):
        raise LearningCreditError('年度归属存储尚未准备完成或期初学分管理未启用')


def register(content, binding_id, year, cutoff, note, filename, fingerprint, actor):
    with transaction() as c, opening._operation(c):
        _ready(c)
        binding, rows, cutoff, note = _checked(c, actor, content, binding_id, year, cutoff, note)
        existing = execute(c, 'SELECT id,status FROM learning_credit_year_imports WHERE content_fingerprint=?', (fingerprint,)).fetchone()
        if existing:
            if existing['status'] == 'CANCELLED':
                raise LearningCreditError('这一来源已取消，请核对后重新提交')
            # Identity and immutable source are both checked on repeated submit.
            old, saved = _batch(c, actor, existing['id'])
            if old['file_sha256'] != hashlib.sha256(content).hexdigest() or old['binding_id'] != binding_id or old['year_index'] != year or str(old['cutoff_date'])[:10] != cutoff.isoformat() or old['source_note'] != note or any('opening_import_id' not in r for r in rows) or _fingerprint(binding,rows,year,cutoff,note) != fingerprint:
                raise LearningCreditError('来源文件已变化，请重新检查')
            return {'id': existing['id'], 'status': existing['status'], 'idempotent': True}
        if any(r['errors'] for r in rows) or _fingerprint(binding, rows, year, cutoff, note) != fingerprint:
            raise LearningCreditError('文件、余额或名册已变化，请重新检查')
        now = _db_timestamp(c)
        row = execute(c, 'INSERT INTO learning_credit_year_imports(content_fingerprint,file_sha256,original_filename,binding_id,year_index,cutoff_date,source_note,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?)',
                      (fingerprint, hashlib.sha256(content).hexdigest(), filename[:255], binding_id, year, cutoff.isoformat(), note, actor, now))
        import_id = int(row.lastrowid)
        for r in rows:
            execute(c, 'INSERT INTO learning_credit_year_rows(import_id,member_id,opening_import_id,excel_row,points) VALUES (?,?,?,?,?)', (import_id, r['member_id'], r['opening_import_id'], r['excel_row'], r['points']))
        write_audit(c, actor_user_id=actor, action=ACTION+'.submit', resource_type='credit_year_import', resource_id=str(import_id), org_unit_id=binding['class_org_unit_id'], after={'file_sha256': hashlib.sha256(content).hexdigest(), 'content_fingerprint': fingerprint, 'year_index': year, 'cutoff_date': cutoff.isoformat(), 'source_note': note, 'count': len(rows), 'adds_points': False})
        return {'id': import_id, 'status': 'PENDING_APPROVAL', 'idempotent': False}


def _batch(c, actor, import_id):
    batch = execute(c, 'SELECT * FROM learning_credit_year_imports WHERE id=?', (import_id,)).fetchone()
    if not batch:
        raise LearningCreditError('年度来源不存在')
    batch = dict(batch)
    binding = _binding(c, actor, int(batch['binding_id']))
    rows = [dict(r) for r in execute(c, 'SELECT r.*,m.member_code,m.name AS member_name FROM learning_credit_year_rows r JOIN members m ON m.id=r.member_id WHERE r.import_id=? ORDER BY r.excel_row', (import_id,)).fetchall()]
    batch.update(class_name=binding['class_name'], row_count=len(rows), total_points=format(sum((Decimal(str(r['points'])) for r in rows), Decimal(0)), '.2f'))
    return batch, rows


def _recheck(c, actor, batch, rows):
    binding = _binding(c, actor, batch['binding_id'])
    roster = {int(m['id']): m for m in opening._roster(c, actor).values()}
    for row in rows:
        member = roster.get(int(row['member_id']))
        if not member or member['ambiguous_class'] or member['org_unit_id'] != binding['class_org_unit_id']:
            raise LearningCreditError('学长名册或班级已变化，请重新核对')
        b = execute(c, 'SELECT * FROM learning_credit_opening_balances WHERE member_id=?', (row['member_id'],)).fetchone()
        if not b or b['import_id'] != row['opening_import_id'] or str(b['cutoff_date'])[:10] != str(batch['cutoff_date'])[:10]:
            raise LearningCreditError('期初余额来源已变化，请重新核对')
        allocated = execute(c, 'SELECT year_index,points,binding_id FROM learning_credit_year_allocations WHERE member_id=? AND opening_import_id=?', (row['member_id'], row['opening_import_id'])).fetchall()
        if any(a['year_index'] == batch['year_index'] or a['binding_id'] != batch['binding_id'] for a in allocated) or sum((Decimal(str(a['points'])) for a in allocated), Decimal(0)) + Decimal(str(row['points'])) > Decimal(str(b['points'])):
            raise LearningCreditError('年度归属已变化或超过余额，请重新核对')
        row['points'] = format(Decimal(str(row['points'])), '.2f')
    if _fingerprint(binding, rows, batch['year_index'], date.fromisoformat(str(batch['cutoff_date'])[:10]), batch['source_note']) != batch['content_fingerprint']:
        raise LearningCreditError('年度分值发生变化，请重新检查')
    progress = execute(c, 'SELECT * FROM learning_credit_class_progress WHERE binding_id=?', (batch['binding_id'],)).fetchone()
    if progress and (progress['completed_days'] != batch['year_index']*12 or str(progress['cutoff_date'])[:10] != str(batch['cutoff_date'])[:10]):
        raise LearningCreditError('班级历史进度已变化，请重新核对')


def action(import_id, actor, operation, fingerprint):
    opening._permission(actor, opening.OPENING_PERMISSION, 'members:read')
    if operation not in ('approve', 'post', 'cancel'):
        raise LearningCreditError('不支持这项年度操作')
    with transaction() as c, opening._operation(c):
        _ready(c)
        batch, rows = _batch(c, actor, import_id)
        if batch['content_fingerprint'] != fingerprint:
            raise LearningCreditError('来源版本已变化，请刷新检查')
        status = batch['status']
        if status == 'POSTED':
            if operation == 'cancel':
                raise LearningCreditError('已归属年度不能取消，请凭来源做前向更正')
            return {'id': import_id, 'status': status, 'idempotent': True}
        if status == 'CANCELLED':
            if operation != 'cancel':
                raise LearningCreditError('这一来源已取消')
            return {'id': import_id, 'status': status, 'idempotent': True}
        if operation == 'cancel':
            status = 'CANCELLED'
            execute(c, 'UPDATE learning_credit_year_imports SET status=? WHERE id=?', (status, import_id))
        else:
            for r in sorted(rows, key=lambda r: r['member_id']):
                execute(c, 'SELECT id FROM members WHERE id=?' + ('' if isinstance(c, sqlite3.Connection) else ' FOR UPDATE'), (r['member_id'],)).fetchone()
            _recheck(c, actor, batch, rows)
            if operation == 'approve':
                if status == 'APPROVED':
                    return {'id': import_id, 'status': status, 'idempotent': True}
                status = 'APPROVED'
                execute(c, 'UPDATE learning_credit_year_imports SET status=?,approved_by=?,approved_at=? WHERE id=?', (status, actor, _db_timestamp(c), import_id))
            else:
                if status != 'APPROVED' or not batch['approved_by']:
                    raise LearningCreditError('请先核对年度来源并完成复核')
                for r in rows:
                    execute(c, 'INSERT INTO learning_credit_year_allocations(member_id,opening_import_id,binding_id,year_index,points,import_id) VALUES (?,?,?,?,?,?)', (r['member_id'],r['opening_import_id'],batch['binding_id'],batch['year_index'],r['points'],import_id))
                if not execute(c, 'SELECT binding_id FROM learning_credit_class_progress WHERE binding_id=?', (batch['binding_id'],)).fetchone():
                    execute(c, 'INSERT INTO learning_credit_class_progress(binding_id,completed_days,cutoff_date,import_id) VALUES (?,?,?,?)', (batch['binding_id'],batch['year_index']*12,batch['cutoff_date'],import_id))
                status = 'POSTED'
                execute(c, 'UPDATE learning_credit_year_imports SET status=?,posted_by=?,posted_at=? WHERE id=?', (status,actor,_db_timestamp(c),import_id))
        write_audit(c, actor_user_id=actor, action=ACTION+'.'+operation, resource_type='credit_year_import',resource_id=str(import_id),after={'content_fingerprint':fingerprint,'status':status,'count':len(rows),'adds_points':False})
        return {'id':import_id,'status':status,'idempotent':False}


def workbench(actor):
    opening._permission(actor, opening.OPENING_PERMISSION, 'members:read')
    c = connect()
    try:
        allowed = opening._scope(actor)
        bindings = [dict(r) for r in execute(c, "SELECT b.id,b.class_org_unit_id,o.name AS class_name FROM class_learning_bindings b JOIN org_units o ON o.id=b.class_org_unit_id WHERE b.status='ACTIVE' AND o.is_active=1 AND b.id=(SELECT b2.id FROM class_learning_bindings b2 WHERE b2.class_org_unit_id=b.class_org_unit_id AND b2.status='ACTIVE' ORDER BY b2.started_at DESC,b2.id DESC LIMIT 1) ORDER BY o.name").fetchall() if allowed is None or r['class_org_unit_id'] in allowed]
        tables = present_tables(c)
        ready = available(c)
        reserved = execute(c, "SELECT id FROM audit_logs WHERE action='production.credit_year.setup' LIMIT 1").fetchone()
        marker = execute(c, 'SELECT version FROM schema_migrations WHERE version=?', (MIGRATION,)).fetchone()
        imports = []
        if ready:
            for item in execute(c,'SELECT id FROM learning_credit_year_imports ORDER BY id DESC LIMIT 100').fetchall():
                try:
                    batch, rows = _batch(c,actor,item['id'])
                    imports.append({**batch,'rows':rows})
                except PermissionError:
                    continue
        user = opening._manage(actor)
        dialect = 'sqlite' if isinstance(c,sqlite3.Connection) else 'mysql'
        return {'storage_available':ready,'bindings':bindings,'imports':imports,'release_commit':get_build_info()['commit_sha'],'migration_sha256':HASHES[dialect],
                'setup_allowed':'system_admin' in user.get('roles',[]) and 'plans:production_rule_reconciliation_apply' in user['permissions'] and get_settings().credit_opening_setup_enabled and not tables and not marker and not reserved,
                'setup_incomplete':not ready and bool(tables or marker or reserved)}
    finally:
        c.close()


def setup(actor, commit, digest):
    user = opening._permission(actor, opening.OPENING_PERMISSION,'plans:production_rule_reconciliation_apply')
    if 'system_admin' not in user.get('roles',[]):
        raise PermissionError('仅管理员可以准备年度归属存储')
    _write_allowed()
    settings = get_settings()
    if not settings.credit_opening_setup_enabled or settings.run_bootstrap_on_startup or commit != get_build_info()['commit_sha'] or not re.fullmatch('[a-f0-9]{40}',commit):
        raise LearningCreditError('发布设置或版本已变化，请刷新核对')
    c = connect()
    try:
        with opening._operation(c):
            dialect = 'sqlite' if isinstance(c,sqlite3.Connection) else 'mysql'
            content = (MIGRATION_ROOT/dialect/MIGRATION).read_bytes()
            if hashlib.sha256(content).hexdigest() != digest or digest != HASHES[dialect]:
                raise LearningCreditError('固定迁移校验失败')
            if available(c):
                return {'status':'READY','idempotent':True}
            if present_tables(c) or execute(c,'SELECT version FROM schema_migrations WHERE version=?',(MIGRATION,)).fetchone() or execute(c,"SELECT id FROM audit_logs WHERE action='production.credit_year.setup' LIMIT 1").fetchone():
                raise LearningCreditError('存在已执行或中断记录，请核验实际存储，不自动重试')
            if not opening.storage_available(c):
                raise LearningCreditError('请先完成期初余额存储准备')
            write_audit(c,actor_user_id=actor,action='production.credit_year.setup',resource_type='schema_migration',resource_id=MIGRATION,result='STARTED',after={'commit':commit,'sha256':digest,'tables':sorted(TABLES),'adds_points':False})
            c.commit()
            for sql in _split_mysql(content.decode('utf-8')):
                execute(c,sql)
            execute(c,'INSERT INTO schema_migrations(version,applied_at) VALUES (?,?)',(MIGRATION,_db_timestamp(c)))
            write_audit(c,actor_user_id=actor,action='production.credit_year.setup.complete',resource_type='schema_migration',resource_id=MIGRATION,after={'commit':commit,'sha256':digest,'adds_points':False})
            c.commit()
            return {'status':'READY','idempotent':False}
    finally:
        c.close()
