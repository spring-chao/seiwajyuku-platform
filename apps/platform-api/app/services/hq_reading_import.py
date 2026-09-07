"""C7.1.2 HQ daily-reading import and identity reconciliation.

The headquarters export is an upstream observation, not a score input.  This
module keeps the three boundaries explicit:

``Excel -> source identity/group reconciliation -> source fact -> C7 DRY-RUN``

It never writes ``learning_credit_entries`` and refuses to operate when the
formal settlement flag is enabled.  Source observations are deliberately
kept even when they are incomplete or blocked, so a later headquarters
export can move ``NOT_COMPLETE`` to ``COMPLETE`` without inventing a fact.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.core.settings import get_settings
from app.db import connect, execute, transaction
from app.services.audit import write_audit
from app.services.iam import accessible_org_ids, user_context
from app.services.learning_activity_credits import (
    DAILY_READING,
    _calendar_day,
    _date_text,
    _member_relation_reason,
    _resolve_binding_for_fact,
    dry_run_daily_reading,
    record_learning_activity_fact,
)
from app.services.learning_credits import (
    LearningCreditError,
    _ledger_count,
    _write_allowed,
)


HQ_SOURCE_TYPE = "HQ_READING_EXPORT"
HQ_SOURCE_SYSTEM = "HQ"
MANUAL_EXCELLENT_SHARE_SOURCE_TYPE = "EXCELLENT_SHARE_MANUAL_VERIFIED"

CLASS_REGULAR = "CLASS_REGULAR"
CLASS_ADVANCED = "CLASS_ADVANCED"
NO_GROUP = "NO_GROUP"
UNKNOWN = "UNKNOWN"

MATCH_AUTO = "AUTO_MATCHED"
MATCH_CANDIDATE = "CANDIDATE"
MATCH_CONFIRMED = "CONFIRMED_BINDING"

COMPLETE_RECORDING_VALUES = {"提前完成", "按时完成", "过期完成"}
NOT_COMPLETE_RECORDING_VALUES = {"未提交", "-", "—", "–", "未完成"}

_HEADER_ALIASES: dict[str, set[str]] = {
    "occurred_on": {"日期", "业务日期", "记录日期", "打卡日期", "完成日期", "businessdate"},
    "name": {"姓名", "学员姓名", "学长姓名", "成员姓名", "name"},
    "group": {"小组", "所属小组", "组别", "学习小组", "group", "groupname"},
    "recording": {"录音", "录音状态", "录音完成", "录音情况", "recording", "recordingstatus"},
    "masked_account": {
        "账号", "脱敏账号", "会员账号", "学号", "编号", "手机号", "手机号脱敏",
        "maskedaccount", "account", "membercode",
    },
    "is_staff": {"工作人员", "人员类型", "身份", "是否工作人员", "staff", "staffflag"},
    "learning_qualification": {
        "学习资格", "是否参与学习", "学习身份", "是否学习", "learningqualification",
    },
}
_REQUIRED_FIELDS = ("occurred_on", "name", "group", "recording")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _compact(value: Any) -> str:
    return re.sub(r"[\s\u3000]+", "", _text(value))


def _header_key(value: Any) -> str:
    return re.sub(r"[\s\u3000_\-/:：()（）【】\[\]]+", "", _text(value)).lower()


def _safe_filename(value: str | None) -> str:
    name = PurePath(value or "hq-reading-export.xlsx").name.strip()
    if not name:
        name = "hq-reading-export.xlsx"
    return name[:255]


def _parse_date(value: Any, label: str = "总部业务日期") -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value)
        except (TypeError, ValueError, OverflowError) as exc:
            raise LearningCreditError(f"{label}不是有效日期") from exc
        if isinstance(parsed, datetime):
            return parsed.date()
        if isinstance(parsed, date):
            return parsed
    text = _text(value)
    for pattern in (
        r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?$",
        r"^(\d{4})年(\d{1,2})月(\d{1,2})日?$",
    ):
        match = re.fullmatch(pattern, text)
        if match:
            try:
                return date(
                    int(match.group(1)), int(match.group(2)), int(match.group(3))
                )
            except ValueError as exc:
                raise LearningCreditError(f"{label}不是有效日期") from exc
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return date.fromisoformat(text)
        except ValueError as exc:
            raise LearningCreditError(f"{label}不是有效日期") from exc
    raise LearningCreditError(f"{label}必须包含完整的YYYY-MM-DD日期")


def _masked_account(value: Any) -> str | None:
    text = re.sub(r"[\s\u3000]+", "", _text(value))
    if not text:
        return None
    if len(text) > 128:
        raise LearningCreditError("脱敏账号不能超过128个字符")
    digits_only = re.sub(r"[-()（）]", "", text)
    if re.fullmatch(r"\d{7,}", digits_only):
        raise LearningCreditError("总部账号必须是脱敏账号或非手机号标识，不能上传完整手机号")
    return text


def _recording_status(value: Any) -> str:
    text = _text(value)
    if not text or text in NOT_COMPLETE_RECORDING_VALUES:
        return "NOT_COMPLETE"
    if text in COMPLETE_RECORDING_VALUES:
        return "COMPLETE"
    return "UNKNOWN"


def _staff_flag(value: Any) -> bool:
    text = _compact(value).lower()
    return text in {"是", "工作人员", "staff", "yes", "true", "1"}


def _qualification(value: Any) -> str:
    text = _compact(value).lower()
    if text in {"是", "学习", "学员", "学长", "learning", "yes", "true", "1"}:
        return "LEARNING"
    if text in {"否", "非学习", "不学习", "notlearning", "no", "false", "0"}:
        return "NOT_LEARNING"
    return "UNKNOWN"


def _source_identity_key(
    *, target_class_org_unit_id: str, name: str, masked_account: str | None, group: str | None
) -> str:
    # An account-backed identity remains stable if the learner changes groups.
    # Without an account, name+group is the minimum safe key; name alone is
    # intentionally never enough.
    material = (
        f"{HQ_SOURCE_SYSTEM}|class|{target_class_org_unit_id}|account|{_compact(masked_account)}"
        if masked_account
        else f"{HQ_SOURCE_SYSTEM}|class|{target_class_org_unit_id}|name|{_compact(name)}|group|{_compact(group) or '<NO_GROUP>'}"
    )
    return "HQID:" + hashlib.sha256(material.encode("utf-8")).hexdigest()


def _source_fact_id(*, source_identity_key: str, occurred_on: str) -> str:
    material = f"{HQ_SOURCE_TYPE}|{source_identity_key}|{occurred_on}|{DAILY_READING}"
    return "HQDR:" + hashlib.sha256(material.encode("utf-8")).hexdigest()


def _content_hash(row: dict[str, Any]) -> str:
    material = {
        "occurred_on": row.get("occurred_on"),
        "name": row.get("name"),
        "masked_account": row.get("masked_account"),
        "group": row.get("group"),
        "recording_raw": row.get("recording_raw"),
        "recording_status": row.get("recording_status"),
        "is_staff": row.get("is_staff"),
        "learning_qualification": row.get("learning_qualification"),
    }
    return hashlib.sha256(_json(material).encode("utf-8")).hexdigest()


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _find_header(values: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    aliases = {
        field: {_header_key(alias) for alias in names}
        for field, names in _HEADER_ALIASES.items()
    }
    for row_index, row in enumerate(values[:50]):
        mapping: dict[str, int] = {}
        for index, value in enumerate(row):
            key = _header_key(value)
            if not key:
                continue
            for field, candidates in aliases.items():
                if field not in mapping and key in candidates:
                    mapping[field] = index
        if all(field in mapping for field in _REQUIRED_FIELDS):
            return row_index, mapping
    return None


def parse_hq_reading_workbook(
    content: bytes, *, target_class_org_unit_id: str
) -> dict[str, Any]:
    """Parse only per-day detail sheets; aggregate-only workbooks are rejected."""

    if not content:
        raise LearningCreditError("总部每日读书工作簿不能为空")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several format-specific errors
        raise LearningCreditError("总部每日读书工作簿无法读取") from exc

    rows: list[dict[str, Any]] = []
    sheets_with_detail = 0
    detail_sheet_names: list[str] = []
    skipped_sheets: list[str] = []
    try:
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            header = _find_header(values)
            if header is None:
                skipped_sheets.append(sheet.title)
                continue
            sheets_with_detail += 1
            detail_sheet_names.append(sheet.title)
            header_index, columns = header
            for index, value_row in enumerate(values[header_index + 1 :], start=header_index + 2):
                if not any(_text(value) for value in value_row):
                    continue
                source = {
                    field: value_row[column] if column < len(value_row) else None
                    for field, column in columns.items()
                }
                name = _text(source.get("name"))
                if not name:
                    raise LearningCreditError(f"工作表“{sheet.title}”第{index}行姓名不能为空")
                try:
                    occurred = _parse_date(source.get("occurred_on"))
                except LearningCreditError as exc:
                    raise LearningCreditError(f"工作表“{sheet.title}”第{index}行：{exc}") from exc
                masked_account = _masked_account(source.get("masked_account"))
                group = _text(source.get("group")) or None
                normalized = {
                    "source_sheet_name": sheet.title,
                    "source_row_number": index,
                    "source_row_key": f"{sheet.title}:{index}"[:255],
                    "occurred_on": occurred.isoformat(),
                    "name": name,
                    "masked_account": masked_account,
                    "group": group,
                    "recording_raw": _text(source.get("recording")) or None,
                    "recording_status": _recording_status(source.get("recording")),
                    "is_staff": _staff_flag(source.get("is_staff")),
                    "learning_qualification": _qualification(
                        source.get("learning_qualification")
                    ),
                }
                normalized["source_identity_key"] = _source_identity_key(
                    target_class_org_unit_id=target_class_org_unit_id,
                    name=name,
                    masked_account=masked_account,
                    group=group,
                )
                normalized["source_id"] = _source_fact_id(
                    source_identity_key=normalized["source_identity_key"],
                    occurred_on=normalized["occurred_on"],
                )
                normalized["content_hash"] = _content_hash(normalized)
                rows.append(normalized)
    finally:
        workbook.close()

    if not sheets_with_detail:
        raise LearningCreditError(
            "工作簿缺少每日明细必需列：日期、姓名、小组、录音；不能使用完成率等汇总列替代"
        )
    if not rows:
        raise LearningCreditError("总部每日读书工作簿没有可导入的每日明细记录")

    dates = [row["occurred_on"] for row in rows]
    group_counts = Counter(row["group"] or "<NO_GROUP>" for row in rows)
    return {
        "rows": rows,
        "sheet_names": detail_sheet_names,
        "skipped_sheets": skipped_sheets,
        "row_count": len(rows),
        "unique_person_count": len({row["source_identity_key"] for row in rows}),
        "date_from": min(dates),
        "date_to": max(dates),
        "group_distribution": dict(sorted(group_counts.items())),
        "recording_distribution": dict(
            sorted(Counter(row["recording_status"] for row in rows).items())
        ),
    }


def _require_import_access(actor_user_id: int) -> None:
    user = user_context(actor_user_id) or {}
    if "plans:hq_reading_import_manage" not in user.get("permissions", []):
        raise PermissionError("无权导入总部每日读书并进行身份核验")
    if get_settings().learning_credit_settlement_enabled:
        raise LearningCreditError("C7.1.2 只允许在正式学分结算关闭时运行")
    _write_allowed()


def _require_preview_access(actor_user_id: int) -> None:
    user = user_context(actor_user_id) or {}
    permissions = set(user.get("permissions", []))
    if not {
        "plans:hq_reading_import_manage",
        "plans:credit_settlement_preview",
    }.intersection(permissions):
        raise PermissionError("无权预览总部每日读书导入结果")
    if get_settings().learning_credit_settlement_enabled:
        raise LearningCreditError("C7.1.2 只允许在正式学分结算关闭时运行")


def _scope_check(actor_user_id: int, class_id: str) -> None:
    allowed = accessible_org_ids(actor_user_id)
    if allowed is not None and class_id not in allowed:
        raise PermissionError("总部每日读书导入不在当前组织授权范围内")


def _class_row(connection, class_id: str) -> dict[str, Any]:
    row = execute(
        connection,
        "SELECT id, name, unit_type, is_active FROM org_units "
        "WHERE id=? AND unit_type IN ('CLASS', 'SPECIAL_COHORT') LIMIT 1",
        (class_id,),
    ).fetchone()
    if not row or not row["is_active"]:
        raise LearningCreditError("目标班级不存在、不是班级或已停用")
    return dict(row)


def _group_type(source_group_name: str | None, class_name: str) -> str:
    value = _compact(source_group_name)
    if not value:
        return NO_GROUP
    if value == "精进组" or value == f"{_compact(class_name)}精进组":
        return CLASS_ADVANCED
    if re.fullmatch(r"\d+组", value):
        return CLASS_REGULAR
    if value.startswith(_compact(class_name)) and re.fullmatch(
        r"\d+组", value[len(_compact(class_name)) :]
    ):
        return CLASS_REGULAR
    return UNKNOWN


def _group_mapping(
    connection, *, class_row: dict[str, Any], source_group_name: str | None
) -> tuple[str | None, str, str | None]:
    group_type = _group_type(source_group_name, str(class_row["name"]))
    source = _compact(source_group_name)
    if not source:
        return None, group_type, None
    class_name = _compact(class_row["name"])
    names = {source, f"{class_name}{source}"}
    if source.startswith(class_name):
        names.add(source[len(class_name) :])
    groups = execute(
        connection,
        "SELECT id, name FROM org_units "
        "WHERE parent_id=? AND unit_type='GROUP' AND is_active=1",
        (class_row["id"],),
    ).fetchall()
    matches = [row for row in groups if _compact(row["name"]) in names]
    if len(matches) != 1:
        return None, group_type, "GROUP_MAPPING_MISSING"
    return str(matches[0]["id"]), group_type, None


def _member_candidates(
    connection, *, class_id: str, mapped_group_id: str | None, name: str
) -> list[dict[str, Any]]:
    params: list[Any] = [class_id, name]
    group_clause = ""
    if mapped_group_id:
        group_clause = (
            " AND EXISTS (SELECT 1 FROM member_org_relations gr "
            "WHERE gr.member_id=m.id AND gr.org_unit_id=? "
            "AND gr.relation_type='STUDY_GROUP')"
        )
        params.append(mapped_group_id)
    rows = execute(
        connection,
        "SELECT DISTINCT m.id, m.member_code, m.name, m.phone_masked, m.phone_last4 "
        "FROM members m JOIN member_org_relations cr ON cr.member_id=m.id "
        "WHERE cr.org_unit_id=? AND cr.relation_type='STUDY_CLASS' "
        "AND m.name=?" + group_clause + " ORDER BY m.id",
        tuple(params),
    ).fetchall()
    return [dict(row) for row in rows]


def _account_matches(source: str | None, member: dict[str, Any]) -> bool:
    if not source:
        return False
    source_key = _compact(source).lower()
    for field in ("phone_masked", "member_code"):
        if _compact(member.get(field)).lower() == source_key:
            return True
    tail = "".join(re.findall(r"\d", source_key))[-4:]
    return bool("*" in source_key and tail and tail == _compact(member.get("phone_last4")))


def _identity_conflicts(existing: dict[str, Any], person: dict[str, Any]) -> list[str]:
    comparisons = (
        ("姓名", existing.get("latest_seen_name"), person.get("name")),
        ("脱敏账号", existing.get("latest_seen_masked_account"), person.get("masked_account")),
        ("小组", existing.get("latest_seen_group"), person.get("group")),
    )
    conflicts: list[str] = []
    for label, old, new in comparisons:
        old_key = _compact(old)
        new_key = _compact(new)
        if old_key and new_key and old_key != new_key:
            conflicts.append(label)
    return conflicts


def _match_person(
    connection,
    *,
    class_row: dict[str, Any],
    person: dict[str, Any],
    mapped_group_id: str | None,
    group_reason: str | None,
) -> dict[str, Any]:
    key = person["source_identity_key"]
    if group_reason:
        return {
            "member_id": None,
            "status": "GROUP_MAPPING_MISSING",
            "method": None,
            "reason": group_reason,
            "candidates": [],
        }
    existing = execute(
        connection,
        "SELECT * FROM hq_reading_source_identities "
        "WHERE source_system=? AND source_identity_key=? LIMIT 1",
        (HQ_SOURCE_SYSTEM, key),
    ).fetchone()
    if existing:
        existing_dict = dict(existing)
        conflicts = _identity_conflicts(existing_dict, person)
        if conflicts or str(existing_dict.get("status") or "").upper() != "CONFIRMED":
            return {
                "member_id": None,
                "status": "SOURCE_IDENTITY_CONFLICT",
                "method": None,
                "reason": "SOURCE_IDENTITY_CONFLICT",
                "conflicts": conflicts,
                "candidates": [],
            }
        candidates = _member_candidates(
            connection,
            class_id=str(class_row["id"]),
            mapped_group_id=mapped_group_id,
            name=person["name"],
        )
        if not any(int(candidate["id"]) == int(existing_dict["member_id"]) for candidate in candidates):
            # The source identity is already permanently bound.  If the
            # platform member is still in the target class but has moved to a
            # different group, retain the member id so row eligibility can
            # report GROUP_MISMATCH.  Never reinterpret this as permission to
            # mutate the platform organization relation or switch identity.
            class_candidates = _member_candidates(
                connection,
                class_id=str(class_row["id"]),
                mapped_group_id=None,
                name=person["name"],
            )
            if any(
                int(candidate["id"]) == int(existing_dict["member_id"])
                for candidate in class_candidates
            ):
                return {
                    "member_id": int(existing_dict["member_id"]),
                    "status": MATCH_CONFIRMED,
                    "method": "CONFIRMED_BINDING",
                    "reason": None,
                    "candidates": class_candidates,
                }
            return {
                "member_id": None,
                "status": "SOURCE_IDENTITY_CONFLICT",
                "method": None,
                "reason": "SOURCE_IDENTITY_CONFLICT",
                "candidates": [],
            }
        return {
            "member_id": int(existing_dict["member_id"]),
            "status": MATCH_CONFIRMED,
            "method": "CONFIRMED_BINDING",
            "reason": None,
            "candidates": candidates,
        }

    candidates = _member_candidates(
        connection,
        class_id=str(class_row["id"]),
        mapped_group_id=mapped_group_id,
        name=person["name"],
    )
    if person.get("masked_account"):
        account_candidates = [
            candidate for candidate in candidates
            if _account_matches(person["masked_account"], candidate)
        ]
        if len(account_candidates) == 1:
            return {
                "member_id": int(account_candidates[0]["id"]),
                "status": MATCH_AUTO,
                "method": "CLASS_GROUP_NAME_MASKED_ACCOUNT",
                "reason": None,
                "candidates": candidates,
            }
        if len(account_candidates) > 1:
            candidates = account_candidates
        return {
            "member_id": None,
            "status": "MEMBER_MAPPING_REQUIRED",
            "method": None,
            "reason": "MEMBER_MAPPING_REQUIRED",
            "candidates": candidates,
        }
    if len(candidates) == 1:
        return {
            "member_id": int(candidates[0]["id"]),
            "status": MATCH_CANDIDATE,
            "method": "CLASS_GROUP_NAME_UNIQUE_CANDIDATE",
            "reason": "IDENTITY_CONFIRMATION_REQUIRED",
            "candidates": candidates,
        }
    return {
        "member_id": None,
        "status": "MEMBER_MAPPING_REQUIRED",
        "method": None,
        "reason": "MEMBER_MAPPING_REQUIRED",
        "candidates": candidates,
    }


def _group_relations_at(
    connection, *, member_id: int, class_id: str, occurred_on: date
) -> list[str]:
    rows = execute(
        connection,
        "SELECT r.org_unit_id FROM member_org_relations r "
        "JOIN org_units g ON g.id=r.org_unit_id "
        "WHERE r.member_id=? AND r.relation_type='STUDY_GROUP' "
        "AND g.parent_id=? AND g.unit_type='GROUP' "
        "AND (r.valid_from IS NULL OR r.valid_from<=?) "
        "AND (r.valid_until IS NULL OR r.valid_until>=?)",
        (member_id, class_id, occurred_on.isoformat(), occurred_on.isoformat()),
    ).fetchall()
    return [str(row["org_unit_id"]) for row in rows]


def _row_eligibility(
    connection, *, row: dict[str, Any], member_id: int | None, match: dict[str, Any]
) -> dict[str, Any]:
    group_type = row["group_type"]
    match_status = match["status"]
    reason = match.get("reason")
    is_staff = bool(row.get("is_staff", row.get("source_is_staff")))
    learning_qualification = row.get(
        "learning_qualification", row.get("learning_qualification_status", "UNKNOWN")
    )
    personal = False
    denominator = False
    numerator = False
    try:
        occurred = date.fromisoformat(row["occurred_on"])
    except (TypeError, ValueError):
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "INVALID_OCCURRED_DATE",
        }

    if match_status not in {MATCH_AUTO, MATCH_CANDIDATE, MATCH_CONFIRMED}:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": reason or match_status,
        }
    if match_status != MATCH_CONFIRMED:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "IDENTITY_CONFIRMATION_REQUIRED",
        }
    if member_id is None:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "MEMBER_MAPPING_REQUIRED",
        }
    relation_reason = _member_relation_reason(
        connection, member_id, row["target_class_org_unit_id"], occurred
    )
    if relation_reason:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": relation_reason,
        }

    mapped_group_id = row.get("mapped_group_org_unit_id")
    platform_groups = _group_relations_at(
        connection,
        member_id=member_id,
        class_id=row["target_class_org_unit_id"],
        occurred_on=occurred,
    )
    if (mapped_group_id and mapped_group_id not in platform_groups) or (
        not mapped_group_id and platform_groups
    ):
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "GROUP_MISMATCH",
        }
    if len(platform_groups) > 1:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "GROUP_RELATION_AMBIGUOUS",
        }
    if is_staff and learning_qualification == "UNKNOWN":
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "REVIEW_REQUIRED",
        }
    if is_staff and learning_qualification == "NOT_LEARNING":
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "STAFF_NOT_LEARNING",
        }
    if group_type == UNKNOWN:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "GROUP_CLASSIFICATION_REQUIRED",
        }
    if group_type == NO_GROUP:
        return {
            "personal": False,
            "denominator": False,
            "numerator": False,
            "reason": "NO_GROUP",
        }
    if row["recording_status"] == "UNKNOWN":
        return {
            "personal": False,
            "denominator": group_type == CLASS_REGULAR,
            "numerator": False,
            "reason": "RECORDING_STATUS_UNRECOGNIZED",
        }
    if group_type == CLASS_REGULAR:
        denominator = True
        numerator = row["recording_status"] == "COMPLETE"
    if row["recording_status"] == "COMPLETE":
        personal = True
        reason = (
            "ADVANCED_GROUP_PERSONAL_ONLY"
            if group_type == CLASS_ADVANCED
            else "READY"
        )
    else:
        reason = "NOT_COMPLETE"
    return {
        "personal": personal,
        "denominator": denominator,
        "numerator": numerator,
        "reason": reason,
    }


def _person_view(
    *, person: dict[str, Any], match: dict[str, Any], mapped_group_id: str | None, group_type: str
) -> dict[str, Any]:
    candidates = []
    for candidate in match.get("candidates", []):
        candidates.append(
            {
                "member_id": int(candidate["id"]),
                "member_name": candidate.get("name"),
                "member_code": candidate.get("member_code"),
                "phone_masked": candidate.get("phone_masked"),
            }
        )
    return {
        "source_identity_key": person["source_identity_key"],
        "source_name": person["name"],
        "source_masked_account": person.get("masked_account"),
        "source_group_name": person.get("group"),
        "mapped_group_org_unit_id": mapped_group_id,
        "group_type": group_type,
        "member_id": match.get("member_id"),
        "match_status": match.get("status"),
        "match_method": match.get("method"),
        "reason": match.get("reason"),
        "candidates": candidates,
    }


def _collapse_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["source_id"]].append(row)
    result: list[dict[str, Any]] = []
    for source_id, candidates in sorted(grouped.items()):
        first = dict(candidates[0])
        hashes = {str(candidate["content_hash"]) for candidate in candidates}
        first["seen_count_in_file"] = len(candidates)
        first["source_row_keys"] = [candidate["source_row_key"] for candidate in candidates]
        if len(hashes) > 1:
            first["source_conflict"] = True
            first["source_conflict_reason"] = "SOURCE_FACT_CONFLICT"
        else:
            first["source_conflict"] = False
        result.append(first)
    return result


def _observation_payload(
    *,
    batch_id: int,
    first_batch_id: int,
    row: dict[str, Any],
    target_class_id: str,
    mapped_group_id: str | None,
    group_type: str,
    match: dict[str, Any],
    eligibility: dict[str, Any],
) -> dict[str, Any]:
    match_status = match["status"]
    reason = eligibility["reason"]
    if row.get("source_conflict"):
        match_status = "SOURCE_IDENTITY_CONFLICT"
        reason = row.get("source_conflict_reason", "SOURCE_FACT_CONFLICT")
    return {
        "batch_id": batch_id,
        "first_batch_id": first_batch_id,
        "source_type": HQ_SOURCE_TYPE,
        "source_id": row["source_id"],
        "source_identity_key": row["source_identity_key"],
        "source_row_key": row["source_row_key"],
        "source_row_number": int(row["source_row_number"]),
        "target_class_org_unit_id": target_class_id,
        "member_id": None if row.get("source_conflict") else match.get("member_id"),
        "source_name": row["name"],
        "source_masked_account": row.get("masked_account"),
        "source_group_name": row.get("group"),
        "mapped_group_org_unit_id": mapped_group_id,
        "group_type": group_type,
        "source_is_staff": int(bool(row.get("is_staff"))),
        "learning_qualification_status": row.get("learning_qualification", "UNKNOWN"),
        "recording_raw": row.get("recording_raw"),
        "recording_status": row["recording_status"],
        "occurred_on": row["occurred_on"],
        "match_status": match_status,
        "match_method": match.get("method"),
        "personal_credit_eligible": int(bool(eligibility["personal"]))
        if not row.get("source_conflict")
        else 0,
        "class_rate_denominator_eligible": int(bool(eligibility["denominator"]))
        if not row.get("source_conflict")
        else 0,
        "class_rate_numerator_eligible": int(bool(eligibility["numerator"]))
        if not row.get("source_conflict")
        else 0,
        "eligibility_reason": reason,
        "content_hash": row["content_hash"],
        "metadata_json": _json(
            {
                "source_sheet_name": row.get("source_sheet_name"),
                "source_row_keys": row.get("source_row_keys", [row["source_row_key"]]),
                "seen_count_in_file": row.get("seen_count_in_file", 1),
            }
        ),
    }


def _upsert_observation(
    connection, payload: dict[str, Any], *, actor_user_id: int | None = None
) -> dict[str, Any]:
    existing = execute(
        connection,
        "SELECT * FROM hq_reading_import_observations "
        "WHERE source_type=? AND source_id=? LIMIT 1",
        (HQ_SOURCE_TYPE, payload["source_id"]),
    ).fetchone()
    now = datetime.now(UTC).isoformat()
    if not existing:
        columns = [
            "batch_id", "first_batch_id", "source_type", "source_id", "source_identity_key",
            "source_row_key", "source_row_number", "target_class_org_unit_id", "member_id",
            "source_name", "source_masked_account", "source_group_name", "mapped_group_org_unit_id",
            "group_type", "source_is_staff", "learning_qualification_status", "recording_raw",
            "recording_status", "occurred_on", "match_status", "match_method",
            "personal_credit_eligible", "class_rate_denominator_eligible",
            "class_rate_numerator_eligible", "eligibility_reason", "content_hash", "metadata_json",
            "seen_count", "created_at", "updated_at",
        ]
        values = [payload.get(column) for column in columns[:-3]] + [1, now, now]
        placeholders = ", ".join("?" for _ in columns)
        execute(
            connection,
            f"INSERT INTO hq_reading_import_observations({', '.join(columns)}) "
            f"VALUES ({placeholders})",
            tuple(values),
        )
        created = dict(
            execute(
                connection,
                "SELECT * FROM hq_reading_import_observations WHERE source_type=? AND source_id=?",
                (HQ_SOURCE_TYPE, payload["source_id"]),
            ).fetchone()
        )
        if actor_user_id is not None and payload.get("match_status") in {
            "SOURCE_IDENTITY_CONFLICT",
            "GROUP_MAPPING_MISSING",
        }:
            write_audit(
                connection,
                actor_user_id=actor_user_id,
                action="learning.hq_reading_import.observation.block",
                resource_type="hq_reading_import_observation",
                resource_id=str(created["id"]),
                org_unit_id=payload["target_class_org_unit_id"],
                purpose="记录总部每日读书来源冲突或小组映射缺失",
                after={
                    "observation_id": created["id"],
                    "source_id": payload["source_id"],
                    "source_identity_key": payload["source_identity_key"],
                    "match_status": payload["match_status"],
                    "eligibility_reason": payload.get("eligibility_reason"),
                },
            )
        return created

    current = dict(existing)
    seen_count = int(current.get("seen_count") or 0) + 1
    conflict_statuses = {
        "SOURCE_IDENTITY_CONFLICT",
        "GROUP_MAPPING_MISSING",
        "GROUP_MISMATCH",
    }
    source_conflict = str(payload.get("match_status") or "") in conflict_statuses
    member_changed = False
    # A source-id collision or later anomaly is never allowed to clear or
    # switch the permanently bound platform member id.
    if current.get("member_id") is not None:
        if payload.get("member_id") is None:
            member_changed = True
            payload = {
                **payload,
                "member_id": current["member_id"],
                "personal_credit_eligible": 0,
                "class_rate_denominator_eligible": 0,
                "class_rate_numerator_eligible": 0,
                "eligibility_reason": payload.get("eligibility_reason")
                or "SOURCE_IDENTITY_CONFLICT",
            }
        elif int(current["member_id"]) != int(payload["member_id"]):
            member_changed = True
            payload = {
                **payload,
                "member_id": current["member_id"],
                "match_status": "SOURCE_IDENTITY_CONFLICT",
                "personal_credit_eligible": 0,
                "class_rate_denominator_eligible": 0,
                "class_rate_numerator_eligible": 0,
                "eligibility_reason": "SOURCE_IDENTITY_CONFLICT",
            }
    mutable = [
        "batch_id", "source_row_key", "source_row_number", "member_id", "source_name",
        "source_masked_account", "source_group_name", "mapped_group_org_unit_id", "group_type",
        "source_is_staff", "learning_qualification_status", "recording_raw", "recording_status",
        "occurred_on", "match_status", "match_method", "personal_credit_eligible",
        "class_rate_denominator_eligible", "class_rate_numerator_eligible", "eligibility_reason",
        "content_hash", "metadata_json",
    ]
    assignments = ", ".join(f"{column}=?" for column in mutable)
    execute(
        connection,
        f"UPDATE hq_reading_import_observations SET {assignments}, seen_count=?, updated_at=? "
        "WHERE id=?",
        tuple(payload.get(column) for column in mutable) + (seen_count, now, current["id"]),
    )
    if actor_user_id is not None and (
        member_changed
        or source_conflict
        or str(current.get("content_hash")) != str(payload.get("content_hash"))
    ):
        write_audit(
            connection,
            actor_user_id=actor_user_id,
            action="learning.hq_reading_import.observation.update",
            resource_type="hq_reading_import_observation",
            resource_id=str(current["id"]),
            org_unit_id=payload["target_class_org_unit_id"],
            purpose="记录总部每日读书来源观察更新或冲突，不自动切换学员身份",
            before={
                "observation_id": current["id"],
                "source_id": current["source_id"],
                "content_hash": current.get("content_hash"),
                "member_id": current.get("member_id"),
                "match_status": current.get("match_status"),
            },
            after={
                "observation_id": current["id"],
                "source_id": payload["source_id"],
                "content_hash": payload.get("content_hash"),
                "member_id": payload.get("member_id"),
                "match_status": payload.get("match_status"),
                "eligibility_reason": payload.get("eligibility_reason"),
            },
        )
    return dict(
        execute(
            connection,
            "SELECT * FROM hq_reading_import_observations WHERE id=?",
            (current["id"],),
        ).fetchone()
    )


def _fact_metadata(observation: dict[str, Any], *, batch_id: int, binding_reason: str | None) -> dict[str, Any]:
    # This deliberately contains no points/score/credit_points field.
    return {
        "hq_import_batch_id": batch_id,
        "source_row_key": observation["source_row_key"],
        "source_row_number": observation["source_row_number"],
        "source_identity_key": observation["source_identity_key"],
        "source_group_name": observation.get("source_group_name"),
        "mapped_group_org_unit_id": observation.get("mapped_group_org_unit_id"),
        "group_type": observation["group_type"],
        "recording_status": observation["recording_status"],
        "personal_credit_eligible": bool(observation["personal_credit_eligible"]),
        "class_rate_denominator_eligible": bool(
            observation["class_rate_denominator_eligible"]
        ),
        "class_rate_numerator_eligible": bool(observation["class_rate_numerator_eligible"]),
        "eligibility_reason": observation.get("eligibility_reason"),
        "binding_resolution": binding_reason or "RESOLVED",
    }


def _sync_source_fact(connection, *, observation: dict[str, Any], batch_id: int, actor_user_id: int) -> str:
    existing = execute(
        connection,
        "SELECT * FROM learning_credit_activity_facts "
        "WHERE source_type=? AND source_id=? LIMIT 1",
        (HQ_SOURCE_TYPE, observation["source_id"]),
    ).fetchone()
    if not observation["personal_credit_eligible"]:
        if existing and str(existing["participation_status"]).upper() in {"RECORDED", "CONFIRMED"}:
            metadata = json.loads(existing["metadata_json"] or "{}")
            metadata.update(
                {
                    "hq_import_batch_id": batch_id,
                    "eligibility_reason": observation.get("eligibility_reason"),
                }
            )
            execute(
                connection,
                "UPDATE learning_credit_activity_facts SET participation_status='CANCELLED', "
                "metadata_json=?, updated_at=? WHERE id=?",
                (_json(metadata), datetime.now(UTC).isoformat(), existing["id"]),
            )
            return "CANCELLED"
        return "NOT_PRODUCED"

    occurred = date.fromisoformat(str(observation["occurred_on"]))
    existing_binding_id = (
        int(existing["binding_id"])
        if existing and existing["binding_id"]
        else None
    )
    if existing_binding_id is not None:
        # Once a source fact has been attached to a learning round, later
        # overlapping imports must retain that frozen binding.  A new current
        # binding must not rewrite historical C7 context.
        binding_id = existing_binding_id
        binding_reason = None
    else:
        binding, binding_reason = _resolve_binding_for_fact(
            connection,
            {
                "class_org_unit_id": observation["target_class_org_unit_id"],
                "occurred_on": observation["occurred_on"],
                "_occurred_date": occurred,
                "binding_id": None,
            },
        )
        binding_id = int(binding["id"]) if binding and not binding_reason else None
    metadata_json = _json(
        _fact_metadata(observation, batch_id=batch_id, binding_reason=binding_reason)
    )
    now = datetime.now(UTC).isoformat()
    if existing:
        current = dict(existing)
        if (
            int(current["member_id"]) != int(observation["member_id"])
            or str(current["class_org_unit_id"]) != str(observation["target_class_org_unit_id"])
            or str(current["occurred_on"]) != str(observation["occurred_on"])
        ):
            return "SOURCE_IDENTITY_CONFLICT"
        execute(
            connection,
            "UPDATE learning_credit_activity_facts SET binding_id=?, participation_status='CONFIRMED', "
            "metadata_json=?, updated_at=? WHERE id=?",
            (binding_id, metadata_json, now, current["id"]),
        )
        return "UPDATED"
    cursor = execute(
        connection,
        "INSERT INTO learning_credit_activity_facts "
        "(activity_type, member_id, class_org_unit_id, binding_id, occurred_on, "
        "participation_status, source_type, source_id, title, metadata_json, created_by, created_at, updated_at) "
        "VALUES (?, ?, ?, ?, ?, 'CONFIRMED', ?, ?, ?, ?, ?, ?, ?)",
        (
            DAILY_READING,
            int(observation["member_id"]),
            observation["target_class_org_unit_id"],
            binding_id,
            observation["occurred_on"],
            HQ_SOURCE_TYPE,
            observation["source_id"],
            "每日读书（总部导入事实）",
            metadata_json,
            actor_user_id,
            now,
            now,
        ),
    )
    write_audit(
        connection,
        actor_user_id=actor_user_id,
        action="learning.hq_reading_fact.create",
        resource_type="learning_credit_activity_fact",
        resource_id=str(cursor.lastrowid),
        org_unit_id=observation["target_class_org_unit_id"],
        purpose="由总部每日读书导入生成事实，不直接赋分",
        after={
            "id": cursor.lastrowid,
            "source_type": HQ_SOURCE_TYPE,
            "source_id": observation["source_id"],
            "member_id": observation["member_id"],
            "occurred_on": observation["occurred_on"],
            "binding_id": binding_id,
            "binding_resolution": binding_reason or "RESOLVED",
        },
    )
    return "CREATED"


def _person_groups(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        item = result.setdefault(
            row["source_identity_key"],
            {
                "source_identity_key": row["source_identity_key"],
                "name": row["name"],
                "masked_account": row.get("masked_account"),
                "group": row.get("group"),
                "rows": [],
            },
        )
        item["rows"].append(row)
        # An account-backed key can span a group change.  That is a conflict,
        # not permission to silently move the identity.
        if _compact(item.get("name")) != _compact(row.get("name")):
            item["person_conflict"] = True
        if _compact(item.get("group")) != _compact(row.get("group")):
            item["person_conflict"] = True
        if _compact(item.get("masked_account")) != _compact(row.get("masked_account")):
            item["person_conflict"] = True
    return result


def _batch_summary(
    *, batch: dict[str, Any], observations: list[dict[str, Any]], identities: list[dict[str, Any]]
) -> dict[str, Any]:
    return {
        "batch_id": int(batch["id"]),
        "source_type": HQ_SOURCE_TYPE,
        "target_class_org_unit_id": batch["target_class_org_unit_id"],
        "original_filename": batch["original_filename"],
        "file_sha256": batch["file_sha256"],
        "row_count": int(batch["row_count"]),
        "unique_person_count": int(batch["unique_person_count"]),
        "date_from": batch.get("date_from"),
        "date_to": batch.get("date_to"),
        "status": batch["status"],
        "observation_count": len(observations),
        "identity_status_counts": dict(
            sorted(Counter(item.get("match_status") for item in identities).items())
        ),
        "group_type_counts": dict(
            sorted(Counter(item.get("group_type") for item in observations).items())
        ),
        "recording_status_counts": dict(
            sorted(Counter(item.get("recording_status") for item in observations).items())
        ),
        "ready_personal_fact_count": sum(
            int(item.get("personal_credit_eligible") or 0) for item in observations
        ),
        "class_rate_denominator_count": sum(
            int(item.get("class_rate_denominator_eligible") or 0) for item in observations
        ),
        "class_rate_numerator_count": sum(
            int(item.get("class_rate_numerator_eligible") or 0) for item in observations
        ),
    }


def import_hq_reading_workbook(
    *,
    actor_user_id: int,
    target_class_org_unit_id: str,
    original_filename: str,
    content: bytes,
) -> dict[str, Any]:
    """Persist an HQ import batch and source observations, never ledger rows."""

    _require_import_access(actor_user_id)
    if not target_class_org_unit_id:
        raise LearningCreditError("导入前必须选择目标班级")
    _scope_check(actor_user_id, target_class_org_unit_id)
    parsed = parse_hq_reading_workbook(
        content, target_class_org_unit_id=target_class_org_unit_id
    )
    file_sha256 = hashlib.sha256(content).hexdigest()
    filename = _safe_filename(original_filename)
    collapsed = _collapse_rows(parsed["rows"])
    with transaction() as connection:
        class_row = _class_row(connection, target_class_org_unit_id)
        duplicate = execute(
            connection,
            "SELECT * FROM hq_reading_import_batches WHERE source_type=? "
            "AND target_class_org_unit_id=? AND file_sha256=? LIMIT 1",
            (HQ_SOURCE_TYPE, target_class_org_unit_id, file_sha256),
        ).fetchone()
        if duplicate:
            existing_batch = dict(duplicate)
            observations = [
                dict(row)
                for row in execute(
                    connection,
                    "SELECT * FROM hq_reading_import_observations WHERE batch_id=? ORDER BY occurred_on, id",
                    (existing_batch["id"],),
                ).fetchall()
            ]
            identities = []
            for key in sorted({item["source_identity_key"] for item in observations}):
                row = next(item for item in observations if item["source_identity_key"] == key)
                identities.append(
                    {
                        "source_identity_key": key,
                        "source_name": row["source_name"],
                        "source_group_name": row.get("source_group_name"),
                        "match_status": row["match_status"],
                        "member_id": row.get("member_id"),
                    }
                )
            return {
                "duplicate": True,
                "message": "相同总部文件已导入，未重复创建批次或事实",
                "summary": _batch_summary(
                    batch=existing_batch, observations=observations, identities=identities
                ),
                "identities": identities,
            }

        now = datetime.now(UTC).isoformat()
        cursor = execute(
            connection,
            "INSERT INTO hq_reading_import_batches "
            "(source_type, target_class_org_unit_id, original_filename, file_sha256, imported_at, "
            "imported_by, row_count, unique_person_count, date_from, date_to, status, summary_json, created_at, updated_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'MATCHING', ?, ?, ?)",
            (
                HQ_SOURCE_TYPE,
                target_class_org_unit_id,
                filename,
                file_sha256,
                now,
                actor_user_id,
                parsed["row_count"],
                parsed["unique_person_count"],
                parsed["date_from"],
                parsed["date_to"],
                "{}",
                now,
                now,
            ),
        )
        batch_id = int(cursor.lastrowid)
        people = _person_groups(parsed["rows"])
        person_matches: dict[str, dict[str, Any]] = {}
        person_views: list[dict[str, Any]] = []
        person_mapping: dict[str, tuple[str | None, str, str | None]] = {}
        for key in sorted(people):
            person = people[key]
            first = person["rows"][0]
            mapped_group_id, group_type, group_reason = _group_mapping(
                connection,
                class_row=class_row,
                source_group_name=person.get("group"),
            )
            person_mapping[key] = (mapped_group_id, group_type, group_reason)
            if person.get("person_conflict"):
                match = {
                    "member_id": None,
                    "status": "SOURCE_IDENTITY_CONFLICT",
                    "method": None,
                    "reason": "SOURCE_IDENTITY_CONFLICT",
                    "candidates": [],
                }
            else:
                match = _match_person(
                    connection,
                    class_row=class_row,
                    person=person,
                    mapped_group_id=mapped_group_id,
                    group_reason=group_reason,
                )
            person_matches[key] = match
            person_views.append(
                _person_view(
                    person=person,
                    match=match,
                    mapped_group_id=mapped_group_id,
                    group_type=group_type,
                )
            )

        observations: list[dict[str, Any]] = []
        fact_actions: Counter[str] = Counter()
        for row in collapsed:
            mapped_group_id, group_type, group_reason = person_mapping[row["source_identity_key"]]
            match = person_matches[row["source_identity_key"]]
            row_with_scope = {
                **row,
                "target_class_org_unit_id": target_class_org_unit_id,
                "mapped_group_org_unit_id": mapped_group_id,
                "group_type": group_type,
            }
            eligibility = _row_eligibility(
                connection,
                row=row_with_scope,
                member_id=match.get("member_id"),
                match=match,
            )
            if group_reason:
                eligibility = {
                    "personal": False,
                    "denominator": False,
                    "numerator": False,
                    "reason": group_reason,
                }
            payload = _observation_payload(
                batch_id=batch_id,
                first_batch_id=batch_id,
                row=row_with_scope,
                target_class_id=target_class_org_unit_id,
                mapped_group_id=mapped_group_id,
                group_type=group_type,
                match=match,
                eligibility=eligibility,
            )
            observation = _upsert_observation(
                connection, payload, actor_user_id=actor_user_id
            )
            observations.append(observation)
            action = _sync_source_fact(
                connection,
                observation=observation,
                batch_id=batch_id,
                actor_user_id=actor_user_id,
            )
            fact_actions[action] += 1

        unresolved = sum(
            int(item["match_status"] not in {MATCH_CONFIRMED}) for item in observations
        )
        status = "CONFIRMED" if unresolved == 0 else "MATCHING"
        summary = _batch_summary(
            batch={
                "id": batch_id,
                "target_class_org_unit_id": target_class_org_unit_id,
                "original_filename": filename,
                "file_sha256": file_sha256,
                "row_count": parsed["row_count"],
                "unique_person_count": parsed["unique_person_count"],
                "date_from": parsed["date_from"],
                "date_to": parsed["date_to"],
                "status": status,
            },
            observations=observations,
            identities=person_views,
        )
        summary["fact_actions"] = dict(sorted(fact_actions.items()))
        execute(
            connection,
            "UPDATE hq_reading_import_batches SET status=?, summary_json=?, updated_at=? WHERE id=?",
            (status, _json(summary), now, batch_id),
        )
        write_audit(
            connection,
            actor_user_id=actor_user_id,
            action="learning.hq_reading_import.create",
            resource_type="hq_reading_import_batch",
            resource_id=str(batch_id),
            org_unit_id=target_class_org_unit_id,
            purpose="导入总部每日读书原始观察并执行身份/小组预匹配",
            after={
                "batch_id": batch_id,
                "source_type": HQ_SOURCE_TYPE,
                "target_class_org_unit_id": target_class_org_unit_id,
                "file_sha256": file_sha256,
                "row_count": parsed["row_count"],
                "unique_person_count": parsed["unique_person_count"],
                "status": status,
                "fact_actions": dict(sorted(fact_actions.items())),
            },
        )
        return {
            "duplicate": False,
            "batch": summary,
            "identities": person_views,
            "fact_actions": dict(sorted(fact_actions.items())),
            "settlement_enabled": False,
            "formal_settlement_allowed": False,
        }


def _load_batch(connection, batch_id: int) -> dict[str, Any]:
    row = execute(
        connection,
        "SELECT * FROM hq_reading_import_batches WHERE id=? LIMIT 1",
        (batch_id,),
    ).fetchone()
    if not row:
        raise LearningCreditError("总部每日读书导入批次不存在")
    return dict(row)


def _validate_confirm_member(
    connection, *, class_id: str, observation: dict[str, Any], member_id: int
) -> dict[str, Any]:
    candidates = _member_candidates(
        connection,
        class_id=class_id,
        mapped_group_id=observation.get("mapped_group_org_unit_id"),
        name=observation["source_name"],
    )
    member = next((item for item in candidates if int(item["id"]) == int(member_id)), None)
    if not member:
        raise LearningCreditError(
            f"学员{member_id}不在目标班级/小组及姓名候选范围内，不能建立总部身份绑定"
        )
    return member


def confirm_hq_reading_identities(
    *, actor_user_id: int, batch_id: int, confirmations: list[dict[str, Any]]
) -> dict[str, Any]:
    """Confirm selected identities and produce/update only source facts."""

    _require_import_access(actor_user_id)
    if not isinstance(confirmations, list) or not confirmations:
        raise LearningCreditError("至少需要一条总部来源身份确认")
    with transaction() as connection:
        batch = _load_batch(connection, int(batch_id))
        _scope_check(actor_user_id, batch["target_class_org_unit_id"])
        seen_keys: set[str] = set()
        confirmed = 0
        fact_actions: Counter[str] = Counter()
        for item in confirmations:
            key = _text(item.get("source_identity_key"))
            if not key or key in seen_keys:
                raise LearningCreditError("身份确认列表存在空键或重复键")
            seen_keys.add(key)
            try:
                member_id = int(item.get("member_id"))
            except (TypeError, ValueError) as exc:
                raise LearningCreditError("身份确认必须指定有效member_id") from exc
            observations = [
                dict(row)
                for row in execute(
                    connection,
                    "SELECT * FROM hq_reading_import_observations "
                    "WHERE batch_id=? AND source_identity_key=? ORDER BY occurred_on, id",
                    (batch_id, key),
                ).fetchall()
            ]
            if not observations:
                raise LearningCreditError(f"批次中不存在来源身份{key}")
            first = observations[0]
            if any(
                row["match_status"]
                in {"SOURCE_IDENTITY_CONFLICT", "GROUP_MAPPING_MISSING"}
                for row in observations
            ):
                raise LearningCreditError(
                    "该来源身份存在冲突或小组映射缺失，不能直接确认；请先处理阻塞项"
                )
            if first["group_type"] == UNKNOWN and not first.get("mapped_group_org_unit_id"):
                raise LearningCreditError("未知小组不能建立身份绑定，请先完成小组映射")
            member = _validate_confirm_member(
                connection,
                class_id=batch["target_class_org_unit_id"],
                observation=first,
                member_id=member_id,
            )
            existing = execute(
                connection,
                "SELECT * FROM hq_reading_source_identities "
                "WHERE source_system=? AND source_identity_key=? LIMIT 1",
                (HQ_SOURCE_SYSTEM, key),
            ).fetchone()
            now = datetime.now(UTC).isoformat()
            if existing and int(existing["member_id"]) != member_id:
                raise LearningCreditError("SOURCE_IDENTITY_CONFLICT:该总部来源身份已有不同学员绑定")
            if existing:
                conflicts = _identity_conflicts(
                    dict(existing),
                    {
                        "name": first["source_name"],
                        "masked_account": first.get("source_masked_account"),
                        "group": first.get("source_group_name"),
                    },
                )
                if conflicts:
                    raise LearningCreditError("SOURCE_IDENTITY_CONFLICT:来源身份字段发生变化")
                execute(
                    connection,
                    "UPDATE hq_reading_source_identities SET latest_seen_name=?, "
                    "latest_seen_masked_account=?, latest_seen_group=?, updated_at=? WHERE id=?",
                    (
                        first["source_name"],
                        first.get("source_masked_account"),
                        first.get("source_group_name"),
                        now,
                        existing["id"],
                    ),
                )
            else:
                execute(
                    connection,
                    "INSERT INTO hq_reading_source_identities "
                    "(source_system, source_identity_key, target_class_org_unit_id, member_id, "
                    "first_confirmed_at, first_confirmed_by, latest_seen_name, latest_seen_masked_account, "
                    "latest_seen_group, status, created_at, updated_at) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'CONFIRMED', ?, ?)",
                    (
                        HQ_SOURCE_SYSTEM,
                        key,
                        batch["target_class_org_unit_id"],
                        member_id,
                        now,
                        actor_user_id,
                        first["source_name"],
                        first.get("source_masked_account"),
                        first.get("source_group_name"),
                        now,
                        now,
                    ),
                )
            for observation in observations:
                row = dict(observation)
                row["member_id"] = member_id
                row["match_status"] = MATCH_CONFIRMED
                row["match_method"] = "MANUAL_CONFIRMED"
                eligibility = _row_eligibility(
                    connection,
                    row=row,
                    member_id=member_id,
                    match={"status": MATCH_CONFIRMED, "reason": None},
                )
                row["personal_credit_eligible"] = int(eligibility["personal"])
                row["class_rate_denominator_eligible"] = int(eligibility["denominator"])
                row["class_rate_numerator_eligible"] = int(eligibility["numerator"])
                row["eligibility_reason"] = eligibility["reason"]
                execute(
                    connection,
                    "UPDATE hq_reading_import_observations SET member_id=?, match_status=?, "
                    "match_method=?, personal_credit_eligible=?, class_rate_denominator_eligible=?, "
                    "class_rate_numerator_eligible=?, eligibility_reason=?, updated_at=? WHERE id=?",
                    (
                        member_id,
                        row["match_status"],
                        row["match_method"],
                        row["personal_credit_eligible"],
                        row["class_rate_denominator_eligible"],
                        row["class_rate_numerator_eligible"],
                        row["eligibility_reason"],
                        now,
                        row["id"],
                    ),
                )
                refreshed = dict(
                    execute(
                        connection,
                        "SELECT * FROM hq_reading_import_observations WHERE id=?",
                        (row["id"],),
                    ).fetchone()
                )
                fact_actions[
                    _sync_source_fact(
                        connection,
                        observation=refreshed,
                        batch_id=batch_id,
                        actor_user_id=actor_user_id,
                    )
                ] += 1
            confirmed += 1
            write_audit(
                connection,
                actor_user_id=actor_user_id,
                action="learning.hq_reading_identity.confirm",
                resource_type="hq_reading_source_identity",
                resource_id=key,
                org_unit_id=batch["target_class_org_unit_id"],
                purpose="确认总部来源身份与平台学员绑定",
                after={
                    "source_identity_key": key,
                    "member_id": member_id,
                    "batch_id": batch_id,
                    "observation_count": len(observations),
                },
            )
        remaining = execute(
            connection,
            "SELECT COUNT(*) AS count FROM hq_reading_import_observations "
            "WHERE batch_id=? AND match_status<>?",
            (batch_id, MATCH_CONFIRMED),
        ).fetchone()
        status = "CONFIRMED" if int(remaining["count"] or 0) == 0 else "MATCHING"
        execute(
            connection,
            "UPDATE hq_reading_import_batches SET status=?, updated_at=? WHERE id=?",
            (status, datetime.now(UTC).isoformat(), batch_id),
        )
        return {
            "batch_id": int(batch_id),
            "status": status,
            "confirmed_identity_count": confirmed,
            "fact_actions": dict(sorted(fact_actions.items())),
            "settlement_enabled": False,
            "formal_settlement_allowed": False,
        }


def _class_rates(observations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_date: dict[str, dict[str, set[str]]] = defaultdict(
        lambda: {"denominator": set(), "numerator": set()}
    )
    for row in observations:
        key = str(row["member_id"]) if row.get("member_id") is not None else str(row["source_identity_key"])
        if row.get("class_rate_denominator_eligible"):
            by_date[row["occurred_on"]]["denominator"].add(key)
        if row.get("class_rate_numerator_eligible"):
            by_date[row["occurred_on"]]["numerator"].add(key)
    result = []
    for occurred_on in sorted(by_date):
        denominator = len(by_date[occurred_on]["denominator"])
        numerator = len(by_date[occurred_on]["numerator"])
        result.append(
            {
                "occurred_on": occurred_on,
                "denominator_person_count": denominator,
                "numerator_person_count": numerator,
                "rate": round(numerator / denominator, 6) if denominator else None,
                "denominator_basis": "unique_person_id",
                "numerator_basis": "unique_person_id",
            }
        )
    return result


def _fact_preview_map(preview: dict[str, Any]) -> dict[int, dict[str, Any]]:
    result: dict[int, dict[str, Any]] = {}
    for entry in preview.get("entries", []):
        for fact_id in entry.get("fact_ids", []):
            result[int(fact_id)] = entry
    return result


def dry_run_hq_reading_import(
    *, actor_user_id: int, batch_id: int, limit: int = 500
) -> dict[str, Any]:
    """Return an operator-readable chain; the operation is strictly read-only."""

    _require_preview_access(actor_user_id)
    if limit < 1 or limit > 5000:
        raise LearningCreditError("总部每日读书DRY-RUN数量必须在1到5000之间")
    connection = connect()
    try:
        batch = _load_batch(connection, int(batch_id))
        _scope_check(actor_user_id, batch["target_class_org_unit_id"])
        before_ledger = _ledger_count(connection)
        observations = [
            dict(row)
            for row in execute(
                connection,
                "SELECT o.*, m.name AS member_name "
                "FROM hq_reading_import_observations o "
                "LEFT JOIN members m ON m.id=o.member_id "
                "WHERE o.batch_id=? ORDER BY o.occurred_on, o.id LIMIT ?",
                (batch_id, limit),
            ).fetchall()
        ]
        daily = dry_run_daily_reading(
            actor_user_id=actor_user_id,
            class_org_unit_id=batch["target_class_org_unit_id"],
            occurred_from=batch.get("date_from"),
            occurred_to=batch.get("date_to"),
            limit=500,
        )
        entry_by_fact = _fact_preview_map(daily)
        details: list[dict[str, Any]] = []
        for observation in observations:
            fact = execute(
                connection,
                "SELECT id FROM learning_credit_activity_facts "
                "WHERE source_type=? AND source_id=? LIMIT 1",
                (HQ_SOURCE_TYPE, observation["source_id"]),
            ).fetchone()
            fact_id = int(fact["id"]) if fact else None
            projected = entry_by_fact.get(fact_id) if fact_id else None
            if projected:
                projected_status = projected.get("status")
                projected_points = float(projected.get("points") or 0)
                projected_reasons = projected.get("reasons", [])
                binding = {
                    "binding_id": projected.get("binding_id"),
                    "learning_round": projected.get("learning_round"),
                    "plan_key": projected.get("plan_key"),
                    "plan_version": projected.get("plan_version"),
                    "rule_version_id": projected.get("rule_version_id"),
                    "rule_version": projected.get("rule_version"),
                    "rule_key": projected.get("rule_key"),
                    "calendar": projected.get("calendar"),
                }
            else:
                projected_status = "BLOCKED"
                projected_points = 0.0
                projected_reasons = [observation.get("eligibility_reason") or "FACT_NOT_PRODUCED"]
                binding = {
                    "binding_id": None,
                    "learning_round": None,
                    "plan_key": None,
                    "plan_version": None,
                    "rule_version_id": None,
                    "rule_version": None,
                    "rule_key": DAILY_READING,
                    "calendar": None,
                }
            details.append(
                {
                    "observation_id": int(observation["id"]),
                    "source_row_key": observation["source_row_key"],
                    "source_row_number": int(observation["source_row_number"]),
                    "source_name": observation["source_name"],
                    "source_masked_account": observation.get("source_masked_account"),
                    "source_group_name": observation.get("source_group_name"),
                    "mapped_group_org_unit_id": observation.get("mapped_group_org_unit_id"),
                    "group_type": observation["group_type"],
                    "member_id": int(observation["member_id"]) if observation.get("member_id") else None,
                    "member_name": observation.get("member_name"),
                    "match_status": observation["match_status"],
                    "match_method": observation.get("match_method"),
                    "recording_raw": observation.get("recording_raw"),
                    "recording_status": observation["recording_status"],
                    "occurred_on": observation["occurred_on"],
                    "personal_credit_eligible": bool(observation["personal_credit_eligible"]),
                    "class_rate_denominator_eligible": bool(
                        observation["class_rate_denominator_eligible"]
                    ),
                    "class_rate_numerator_eligible": bool(
                        observation["class_rate_numerator_eligible"]
                    ),
                    "eligibility_reason": observation.get("eligibility_reason"),
                    "fact_id": fact_id,
                    "binding": binding,
                    "projected_status": projected_status,
                    "projected_points": projected_points,
                    "projected_reasons": projected_reasons,
                }
            )
        after_ledger = _ledger_count(connection)
        summary = {
            "total_observation_count": len(observations),
            "unique_person_count": len(
                {item.get("member_id") or item["source_identity_key"] for item in observations}
            ),
            "matched_person_count": len(
                {
                    item.get("member_id")
                    for item in observations
                    if item.get("member_id") is not None
                }
            ),
            "recording_complete_count": sum(
                item["recording_status"] == "COMPLETE" for item in observations
            ),
            "recording_not_complete_count": sum(
                item["recording_status"] == "NOT_COMPLETE" for item in observations
            ),
            "group_type_counts": dict(
                sorted(Counter(item["group_type"] for item in observations).items())
            ),
            "ready_count": sum(item["projected_status"] == "READY" for item in details),
            "blocked_count": sum(item["projected_status"] == "BLOCKED" for item in details),
            "no_credit_count": sum(
                item["projected_status"] == "NO_CREDIT" for item in details
            ),
            "proposed_points": sum(item["projected_points"] for item in details),
        }
        return {
            "mode": "DRY_RUN",
            "persisted": False,
            "settlement_enabled": False,
            "formal_settlement_allowed": False,
            "batch": {
                "id": int(batch["id"]),
                "target_class_org_unit_id": batch["target_class_org_unit_id"],
                "original_filename": batch["original_filename"],
                "file_sha256": batch["file_sha256"],
                "date_from": batch.get("date_from"),
                "date_to": batch.get("date_to"),
            },
            "summary": summary,
            "class_rates": _class_rates(observations),
            "details": details,
            "daily_reading_preview": daily,
            "write_proof": {
                "ledger_entries_before": before_ledger,
                "ledger_entries_after": after_ledger,
                "ledger_entries_delta": after_ledger - before_ledger,
                "formal_settlement_called": False,
            },
        }
    finally:
        connection.close()


def record_manual_verified_excellent_share(
    *,
    actor_user_id: int,
    member_id: int,
    class_org_unit_id: str,
    occurred_on: str,
    note: str | None = None,
    evidence: str | None = None,
) -> dict[str, Any]:
    """Record a verified excellent-share fact without accepting any score input."""

    _require_import_access(actor_user_id)
    normalized_date = _date_text(occurred_on, "优秀分享发生日期")
    note_text = _text(note)
    evidence_text = _text(evidence)
    if len(note_text) > 1000 or len(evidence_text) > 1000:
        raise LearningCreditError("优秀分享说明或依据不能超过1000个字符")
    _scope_check(actor_user_id, class_org_unit_id)
    connection = connect()
    try:
        relation_reason = _member_relation_reason(
            connection,
            int(member_id),
            class_org_unit_id,
            date.fromisoformat(normalized_date),
        )
    finally:
        connection.close()
    if relation_reason:
        raise LearningCreditError(f"手工核验的学员与班级关系无效：{relation_reason}")
    material = f"{member_id}|{class_org_unit_id}|{normalized_date}|{note_text}|{evidence_text}"
    source_id = "MANUAL_ES:" + hashlib.sha256(material.encode("utf-8")).hexdigest()
    return record_learning_activity_fact(
        actor_user_id=actor_user_id,
        activity_type="EXCELLENT_SHARE",
        member_id=member_id,
        class_org_unit_id=class_org_unit_id,
        occurred_on=normalized_date,
        source_type=MANUAL_EXCELLENT_SHARE_SOURCE_TYPE,
        source_id=source_id,
        participation_status="CONFIRMED",
        metadata={
            "manual_verified": True,
            "note": note_text or None,
            "evidence": evidence_text or None,
        },
        _allow_controlled_source=True,
    )
