"""Expandable layouts taken from the supplied Wujiang September workbook.

The bundled workbook contains styles and course headings only, never source people or results.
"""
from collections import defaultdict
from copy import copy
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as letter

from app.services.credit_source_exports import _literal, _safe_text

ASSET = Path(__file__).resolve().parents[1] / 'assets' / 'monthly-core-template.xlsx'


class Layout:
    def __init__(self):
        self.reference = load_workbook(ASSET)
        self.book = Workbook(); self.book.remove(self.book.active)
        self.styles={}
        edge=Side(style='thin',color='000000')
        self.border=Border(left=edge,right=edge,top=edge,bottom=edge)
        self.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        self.yellow=PatternFill('solid',fgColor='FFFF00')

    def sheet(self, name, source=None):
        ws = self.book.create_sheet(name)
        ws._template_name = source or name
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def cell(self, ws, r, c, value=None, ref=(4,14), percent=False):
        key=(ws._template_name,ref)
        if key not in self.styles:
            source = self.reference[ws._template_name].cell(*ref)
            self.styles[key]={attr:copy(getattr(source,attr)) for attr in ('font','fill','protection','number_format')}
        cell = _literal(ws,r,c,_safe_text(value) if isinstance(value,str) else value)
        for attr,value_style in self.styles[key].items():setattr(cell,attr,value_style)
        cell.alignment = self.alignment
        cell.border = self.border
        if value == '×': cell.fill=self.yellow
        if percent: cell.number_format='0%'
        elif isinstance(value,date): cell.number_format='m"月"d"日"'
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 23)
        ws.column_dimensions[letter(c)].width = max(ws.column_dimensions[letter(c)].width or 0, 12)
        return cell

    def merge(self, ws, r1,c1,r2,c2,value,ref=(1,1),percent=False):
        for r in range(r1,r2+1):
            for c in range(c1,c2+1): self.cell(ws,r,c,ref=ref)
        ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
        return self.cell(ws,r1,c1,value,ref,percent)

    def note(self, ws, text):
        r=ws.max_row+2; last=ws.max_column
        cell=self.merge(ws,r,1,r,last,text)
        cell.font=Font(name='宋体',size=10,color='595959')
        cell.fill=PatternFill('solid',fgColor='FFFFFF')
        cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='left')
        ws.row_dimensions[r].height=44
        ws.print_area=f'A1:{letter(last)}{r}'


def _weeks(calendar):
    weeks={}
    for index,(day,_) in enumerate(calendar):
        weeks.setdefault(day-timedelta(days=day.weekday()),[]).append(index)
    return list(weeks.values())


def _rate(rows, indexes):
    values=[r[4+i] for r in rows for i in indexes]
    if not values or any(v is None or v=='待核对' for v in values):
        return None
    required=sum(v in ('√','×') for v in values)
    return values.count('√')/required if required else None


def _week_label(calendar, indexes, n):
    return f'第{n}周（{calendar[indexes[0]][0].day}–{calendar[indexes[-1]][0].day}）'


def _learning(layout, ws, calendar, rows, month):
    weeks=_weeks(calendar); left_end=len(weeks)+3; base=max(10,left_end+2)
    by_class=defaultdict(list)
    for row in rows: by_class[row[0]].append(row)
    if not by_class: by_class['学习班']=[]
    top=1
    for class_name,members in by_class.items():
        by_group=defaultdict(list)
        for member in members: by_group[member[1]].append(member)
        layout.merge(ws,top,1,top,left_end,f'{class_name}{month}月学习小组完成率')
        layout.cell(ws,top+1,1,'组别',(2,1)); layout.cell(ws,top+1,2,'统计项目',(2,2))
        for w,indexes in enumerate(weeks,1): layout.cell(ws,top+1,w+2,_week_label(calendar,indexes,w),(2,min(w+2,7)))
        layout.cell(ws,top+1,left_end,'小组月度',(2,8))
        grid=[]
        for indexes in weeks:
            grid.extend([('day',i) for i in indexes]); grid.extend([('个人',indexes),('小组',indexes)])
        end=base+3+len(grid)
        layout.merge(ws,top,base,top+2,base,'组别',(1,10))
        layout.cell(ws,top,base+1,'序',(1,11))
        layout.merge(ws,top,base+2,top,end,f'{class_name}{month}月读卡数据统计',(1,12))
        layout.cell(ws,top+1,base+1,'序号',(2,11)); layout.cell(ws,top+1,base+2,'姓名',(2,12))
        layout.cell(ws,top+1,base+3,'日期',(2,13))
        column=base+4; day_columns={}; group_columns=[]
        for w,indexes in enumerate(weeks,1):
            start_col=column
            for i in indexes:
                day=calendar[i][0]; day_columns[i]=column
                layout.cell(ws,top+1,column,day,(2,14))
                layout.cell(ws,top+2,column,'周'+'一二三四五六日'[day.weekday()],(3,14)); column+=1
            layout.merge(ws,top+1,column,top+1,column+1,_week_label(calendar,indexes,w),(2,19))
            layout.cell(ws,top+2,column,'个人',(3,19)); layout.cell(ws,top+2,column+1,'小组',(3,20))
            group_columns.append((column,column+1,start_col,indexes)); column+=2
        tail=column
        for c,label in enumerate(['应打卡数','未打卡数','实际打卡数','月度打卡率'],tail):
            layout.merge(ws,top,c,top+2,c,label,(1,48))
        ws.row_dimensions[top+1].height=36
        r=top+3; left_r=top+3; sequence=0
        for group_name,group in by_group.items():
            first=r
            for member in group:
                sequence+=1
                layout.cell(ws,r,base+1,sequence,(4,11)); layout.cell(ws,r,base+2,member[3],(4,12))
                layout.cell(ws,r,base+3,None,(4,13))
                for i,c in day_columns.items(): layout.cell(ws,r,c,member[4+i],(4,14))
                for personal,_,_,indexes in group_columns:
                    layout.cell(ws,r,personal,_rate([member],indexes),(4,19),True)
                for c,value in enumerate([member[-6],member[-4],member[-5],member[-3]],tail):
                    layout.cell(ws,r,c,value,(4,48),c==tail+3)
                r+=1
            layout.merge(ws,first,base,r-1,base,group_name,(4,10))
            layout.merge(ws,left_r,1,left_r+1,1,group_name,(4,1))
            for offset,label in enumerate(['读书打卡率','学习分享率']):
                layout.cell(ws,left_r+offset,2,label,(4+offset,2))
            for w,(personal,gcol,start_col,indexes) in enumerate(group_columns,1):
                rate=_rate(group,indexes)
                layout.merge(ws,first,gcol,r-1,gcol,rate,(4,20),True)
                if rate is not None:
                    required=sum(member[4+i] in ('√','×') for member in group for i in indexes)
                    ws.cell(first,gcol,f'=COUNTIF({letter(start_col)}{first}:{letter(personal-1)}{r-1},"√")/{required}')
                layout.cell(ws,left_r,w+2,rate,(4,3),True)
                layout.cell(ws,left_r+1,w+2,None,(5,3),True)
            layout.cell(ws,left_r,left_end,_rate(group,range(len(calendar))),(4,8),True)
            layout.cell(ws,left_r+1,left_end,None,(5,8),True); left_r+=2
        # Preserve the template's second matrix for ordinary/rotating shares.
        share_top=max(r,left_r)+2
        layout.merge(ws,share_top,base,share_top,tail+3,f'{class_name}轮值分享数据统计表',(123,10))
        for c,label in enumerate(['小组','轮值排班','姓名','日期'],base): layout.cell(ws,share_top+1,c,label,(124,min(c-base+10,13)))
        for i,c in day_columns.items():
            layout.cell(ws,share_top+1,c,calendar[i][0],(124,14))
            layout.cell(ws,share_top+2,c,'周'+'一二三四五六日'[calendar[i][0].weekday()],(125,14))
        for personal,gcol,_,_ in group_columns:
            layout.cell(ws,share_top+2,personal,'个人',(125,19)); layout.cell(ws,share_top+2,gcol,'小组',(125,20))
        for n,member in enumerate(members,share_top+3):
            for c,value in enumerate([member[1],None,member[3],None],base): layout.cell(ws,n,c,value,(126,min(c-base+10,13)))
            for c in range(base+4,tail+4): layout.cell(ws,n,c,None,(126,14))
        top=share_top+max(len(members),1)+6
    ws.column_dimensions[letter(base+2)].width=16; ws.column_dimensions['B'].width=18
    ws.freeze_panes=f'{letter(base+4)}4'; ws.print_title_rows='1:3'
    return weeks


def _analysis(layout, ws, center, calendar, rows, weeks, month):
    end=1+(len(weeks)+1)*2
    layout.merge(ws,1,1,1,end,f'{center}{month}月学习数据',(1,1))
    layout.merge(ws,2,1,3,1,'日期\n分中心 / 班级',(2,1))
    for n,indexes in enumerate([*weeks,list(range(len(calendar)))],1):
        col=n*2; label=_week_label(calendar,indexes,n) if n<=len(weeks) else f'{month}月'
        layout.merge(ws,2,col,2,col+1,label,(2,2))
        layout.cell(ws,3,col,'打卡率',(3,2)); layout.cell(ws,3,col+1,'分享率',(3,3))
    classes=defaultdict(list)
    for row in rows: classes[row[0]].append(row)
    for r,(label,members) in enumerate([(center,rows),*classes.items()],4):
        layout.cell(ws,r,1,label,(4,1)); ws.column_dimensions['A'].width=22
        for n,indexes in enumerate([*weeks,list(range(len(calendar)))],1):
            layout.cell(ws,r,n*2,_rate(members,indexes),(4,2),True)
            layout.cell(ws,r,n*2+1,None,(4,3),True)
    ws.freeze_panes='B4'; ws.print_title_rows='1:3'


def _abnormal(layout, ws, calendar, learning, abnormal, month):
    keys={(r[0],r[2]) for r in abnormal}; groups=defaultdict(list)
    for member in learning:
        if (member[0],member[2]) in keys: groups[member[0]].append(member)
    if not groups: groups['学习异常人员']=[]
    top=1
    for name,members in groups.items():
        end=len(calendar)+6
        layout.merge(ws,top,1,top,end,f'{name}{month}月读卡数据统计',(1,1))
        for c,label in enumerate(['姓名','日期',*[d for d,_ in calendar],'应打卡数','未打卡数','实际打卡数','月度打卡率'],1):
            layout.cell(ws,top+1,c,label,(2,min(c,26)))
        for i,(day,_) in enumerate(calendar,3): layout.cell(ws,top+2,i,'周'+'一二三四五六日'[day.weekday()],(3,3))
        for r,member in enumerate(members,top+3):
            values=[member[3],None,*member[4:4+len(calendar)],member[-6],member[-4],member[-5],member[-3]]
            for c,value in enumerate(values,1): layout.cell(ws,r,c,value,(4,min(c,26)),c==end)
        top+=max(1,len(members))+5
    ws.column_dimensions['A'].width=18; ws.freeze_panes='C4'; ws.print_title_rows='1:3'


def _activities(layout, ws, center, activities):
    top=1
    for event in activities or [[None,None,'活动数据',None,None,None,None,None,None,[]]]:
        people=event[9] or [(None,None)]
        end=top+len(people)+1
        layout.merge(ws,top,1,end,1,' '.join(str(v) for v in (event[1],event[2],event[4]) if v),(1,1))
        for c,label in enumerate(['姓名','分中心','签到','应到人数','实到人数','参会率','未参加次数'],2):
            layout.cell(ws,top,c,label,(1,c))
        for r,(name,status) in enumerate(people,top+1):
            layout.cell(ws,r,2,name,(2,2)); layout.cell(ws,r,3,center if name else None,(2,3))
            layout.cell(ws,r,4,status,(2,4)); layout.cell(ws,r,8,None,(2,8))
        for c,value in enumerate([event[5],event[6],event[7]],5): layout.merge(ws,top+1,c,end-1,c,value,(2,c),c==7)
        layout.merge(ws,end,3,end,7,'小计（按本场次去重；应到资料不足留空）',(2,3))
        top=end+2
    ws.column_dimensions['A'].width=25; ws.column_dimensions['B'].width=18


def _meetings(layout, ws, center, month, meetings):
    end=6+month
    layout.merge(ws,1,1,1,4,f'{center}小组学习会指标{month}月份完成情况统计',(1,1))
    layout.merge(ws,1,6,1,end,'小组学习会组员参会率',(1,6))
    layout.cell(ws,2,1,'核心指标',(2,1)); layout.merge(ws,2,2,2,4,'各班线下小组会议数量',(2,2))
    for c,label in enumerate(['分中心',*[f'{m}月' for m in range(1,month+1)]],6): layout.cell(ws,2,c,label,(2,min(c,15)))
    for c,label in enumerate(['班级',f'{month}月预定','截止当天实绩','完成率'],1): layout.cell(ws,3,c,label,(3,c))
    classes={r[0]:r for r in meetings}
    for r,(name,record) in enumerate(classes.items(),4):
        for c,value in enumerate([name,record[8],record[9],record[10]],1): layout.cell(ws,r,c,value,(5,c),c==4)
    layout.cell(ws,3,6,center,(3,6))
    for c in range(7,end+1): layout.cell(ws,3,c,None,(3,min(c,15)),True)
    completed=[r for r in meetings if r[2].startswith('已完成')]
    denominator=sum(r[4] for r in completed)
    layout.cell(ws,3,end,sum(r[5] or 0 for r in completed)/denominator if denominator else None,(3,15),True)
    top=max(15,len(classes)+7)
    layout.merge(ws,top,1,top,7,f'{center}小组会议统计明细',(15,1))
    for c,label in enumerate(['序号','参会人员','组别','会议时间','参会人数','小组参会率','所属分中心'],1): layout.cell(ws,top+1,c,label,(16,c))
    for n,record in enumerate(meetings,1):
        for c,value in enumerate([n,record[7],record[0]+record[1],record[3],record[5],record[6],center],1): layout.cell(ws,top+1+n,c,value,(17,c),c==6)
        ws.row_dimensions[top+1+n].height=max(30,((len(record[7] or '')+19)//20)*16)
    ws.column_dimensions['B'].width=40; ws.column_dimensions['C'].width=24


def _credits(layout, ws, center, month, learning, ledger):
    # Extend the template's monthly bands through the selected month.
    first_band=6; second=first_band+month; third=second+12; courses=third+month
    reference=layout.reference['学分']; course_labels=[reference.cell(3,c).value for c in range(36,77)]
    end=courses+len(course_labels)-1
    layout.merge(ws,1,1,1,end,f'{center}学员学分统计表',(1,1))
    for c,label in enumerate(['班级','序号','姓名','组别'],1): layout.merge(ws,2,c,4,c,label,(2,c))
    layout.cell(ws,2,5,'学习内容',(2,5)); layout.merge(ws,3,5,4,5,'本月已入账总分',(3,5))
    for start,count,label,ref in [(first_band,month,'每日读书＋读书优秀分享',(2,6)),
                                  (second,12,'班级学习日（上午＋下午＋空巴）',(2,15)),
                                  (third,month,'小组学习会',(2,27))]:
        layout.merge(ws,2,start,2,start+count-1,label,ref)
        for index in range(count):
            layout.cell(ws,3,start+index,f'{index+1}月',(3,ref[1]))
            layout.cell(ws,4,start+index,None,(4,ref[1]))
    for start,stop,label,ref in [(0,9,'线上课程',(2,36)),(10,20,'线下课程',(2,46)),
                               (21,28,'报告会',(2,57)),(29,40,'游学',(2,65))]:
        layout.merge(ws,2,courses+start,2,courses+stop,label,ref)
    for n,label in enumerate(course_labels):
        layout.cell(ws,3,courses+n,label,(3,36+n)); layout.cell(ws,4,courses+n,None,(4,36+n))
        ws.column_dimensions[letter(courses+n)].width=14
    people={}; points=defaultdict(lambda:defaultdict(float))
    for r in learning: people[r[0],r[2]]=(r[3],r[1])
    for r in ledger:
        key=(r[1],r[2]); people.setdefault(key,(r[3],None))
        points[key]['TOTAL']+=float(r[6]); points[key][r[5]]+=float(r[6])
    class_start={}
    for row,((class_name,code),(name,group)) in enumerate(people.items(),5):
        class_start.setdefault(class_name,[]).append(row)
        values={2:row-4,3:name,4:group,5:points[class_name,code].get('TOTAL'),
                first_band+month-1:sum(points[class_name,code][k] for k in ('读书打卡','优秀分享')) if any(k in points[class_name,code] for k in ('读书打卡','优秀分享')) else None,
                second+month-1:points[class_name,code].get('班会签到'),
                third+month-1:points[class_name,code].get('小组学习会')}
        for c in range(2,end+1): layout.cell(ws,row,c,values.get(c),(5,min(c,79)))
    for name,rows in class_start.items(): layout.merge(ws,rows[0],1,rows[-1],1,name,(5,1))
    ws.column_dimensions['A'].width=15; ws.column_dimensions['C'].width=16
    ws.row_dimensions[3].height=95; ws.freeze_panes='F5'; ws.print_title_rows='1:4'


def build_workbook(center, month, cutoff, calendar, learning, abnormal, activities, meetings, ledger):
    layout=Layout(); month_number=int(month[5:])
    ws=layout.sheet('学习数据'); weeks=_learning(layout,ws,calendar,learning,month_number)
    _analysis(layout,layout.sheet('学习分析'),center,calendar,learning,weeks,month_number)
    _abnormal(layout,layout.sheet('学习异常人员'),calendar,learning,abnormal,month_number)
    _activities(layout,layout.sheet('活动数据'),center,activities)
    _meetings(layout,layout.sheet('小组学习会数据'),center,month_number,meetings)
    _credits(layout,layout.sheet('学分明细','学分'),center,month_number,learning,ledger)
    for ws in layout.book:
        layout.note(ws,f'{center}　统计月份：{month}，截至{cutoff}；仅含当前授权班级。空白为资料不足或未提供，来源明确未打卡标×并标黄。轮值分享、历史月份学分和无法匹配模板课程的资料留空。')
    return layout.book
