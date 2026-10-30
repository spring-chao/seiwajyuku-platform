"""Scoped weekly reading matrix from saved sources, independent of credit posting."""
import io
import json
import re
import zipfile
from xml.etree import ElementTree as ET
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

from app.db import connect, execute
from app.services import credit_source_imports as sources
from app.services import learning_activity_credits as activity
from app.services import learning_credits as credits
from app.services.audit import write_audit
from app.services.credit_source_exports import _literal, _safe_text
from app.services.learning_org_catalog import class_catalog, selected_class_ids
from app.services.learning_requirement_policy import exemption_reason

PERMISSIONS = ('exports:normal', 'members:read', 'plans:credit_settlement_preview')
PAUSE_REASONS = {'COURSE': '参加课程', 'TEMPORARY': '临时特殊情况'}


def period(start, end):
    try:
        first, last = date.fromisoformat(start), date.fromisoformat(end)
    except (ValueError, TypeError) as exc:
        raise ValueError('请选择完整的统计开始和结束日期') from exc
    if first > last or last - first > timedelta(days=6) or first - timedelta(days=first.weekday()) != last - timedelta(days=last.weekday()):
        raise ValueError('每次请选择同一自然周内最多7天的日期')
    if (first.year, first.month) != (last.year, last.month):
        raise ValueError('跨月的一周请按自然月分段导出')
    if first < sources.START or last >= datetime.now(activity.SHANGHAI).date():
        raise ValueError('统计日期需在2026-10-01至昨天之间')
    return first, last


def _context(c, actor, class_id):
    user = credits.user_context(actor) or {}
    if not set(PERMISSIONS) <= set(user.get('permissions', [])):
        raise PermissionError('当前账号没有读书统计导出权限')
    if not class_id or not sources._has_scope(actor, class_id, PERMISSIONS):
        raise PermissionError('请选择授权范围内的班级')
    catalog = class_catalog(c, sources.effective_scope(actor, PERMISSIONS))
    ids = selected_class_ids(catalog, class_id)
    selected = next((r for r in catalog if class_id in r['related_class_ids']), None)
    if not selected or not ids:
        raise PermissionError('班级不在读书统计导出范围')
    return selected, ids


def _valid(row, day):
    text = day.isoformat()
    return (not row.get('valid_from') or str(row['valid_from'])[:10] <= text) and (not row.get('valid_until') or str(row['valid_until'])[:10] >= text)


def _choose(relations, day):
    valid = [r for r in relations if _valid(r, day)]
    if not valid:
        return None
    primary = max(r['is_primary'] for r in valid)
    best = {r['org_unit_id']: r for r in valid if r['is_primary'] == primary}
    return next(iter(best.values())) if len(best) == 1 else None


def _roster(c, class_ids, start, end):
    # Resolve formal relations at the requested date, not legacy text fields.
    rows = [dict(r) for r in execute(c,
        "SELECT m.id,m.member_code,m.name,m.status,r.org_unit_id,r.relation_type,r.is_primary,r.valid_from,r.valid_until,"
        "o.name AS org_name,o.parent_id FROM members m JOIN member_org_relations r ON r.member_id=m.id "
        "JOIN org_units o ON o.id=r.org_unit_id WHERE m.status IN ('ACTIVE','SUSPENDED') "
        "AND r.relation_type IN ('STUDY_CLASS','SPECIAL_COHORT','STUDY_GROUP') AND o.is_active=1 "
        "AND m.id IN (SELECT member_id FROM member_org_relations WHERE org_unit_id IN (" + ','.join('?' for _ in class_ids) + "))",
        tuple(sorted(class_ids))).fetchall()]
    by_member = defaultdict(list)
    for row in rows:
        by_member[int(row['id'])].append(row)
    result = []
    days = [start + timedelta(days=i) for i in range((end-start).days+1)]
    for mid, relations in by_member.items():
        classes = [r for r in relations if r['relation_type'] != 'STUDY_GROUP']
        assigned = {d: _choose(classes, d) for d in days}
        included = [d for d, r in assigned.items() if r and r['org_unit_id'] in class_ids]
        if not included:
            continue
        last_class = assigned[included[-1]]['org_unit_id']
        group = _choose([r for r in relations if r['relation_type'] == 'STUDY_GROUP' and r['parent_id'] == last_class], included[-1])
        result.append({'id': mid, 'name': relations[0]['name'], 'member_code': relations[0]['member_code'],
            'status': relations[0]['status'], 'group_id': group['org_unit_id'] if group else '',
            'group_name': group['org_name'] if group else '未分组', 'included_days': set(included)})
    if len(result) > 1000:
        raise ValueError('班级名册超过1000人，请核对组织范围')
    return sorted(result, key=lambda r: (r['group_name'] == '未分组', r['group_name'], r['group_id'], r['name'], r['id']))


def roster(actor, class_id, occurred_from, occurred_to):
    start, end = period(occurred_from, occurred_to)
    c = connect()
    try:
        _, ids = _context(c, actor, class_id)
        return {'members': [{k: r[k] for k in ('id','name','member_code','group_name')} for r in _roster(c, ids, start, end)]}
    finally:
        c.close()


def _pauses(values, members, start, end):
    paused = defaultdict(set)
    if len(values) > 200:
        raise ValueError('单次临时暂停设置最多200条')
    for value in values:
        mid = value['member_id']
        if mid not in members:
            raise ValueError('临时暂停学员不在所选班级统计名册')
        first, last = date.fromisoformat(value['occurred_from']), date.fromisoformat(value['occurred_to'])
        if first < start or last > end or first > last or value['reason'] not in PAUSE_REASONS:
            raise ValueError('临时暂停日期需在本次统计区间内，并选择有效原因')
        paused[mid].update(first + timedelta(days=i) for i in range((last-first).days+1))
    return paused


def _staff(c, ids, end):
    # Teacher names are presentation context; they are not added to the learner denominator.
    return [dict(r) for r in execute(c,
        "SELECT DISTINCT m.id,m.name,va.appointment_key FROM volunteer_appointments va JOIN members m ON m.id=va.member_id "
        "LEFT JOIN volunteer_service_units vsu ON vsu.id=va.volunteer_service_unit_id "
        "WHERE va.status='ACTIVE' AND m.status='ACTIVE' AND va.appointment_key IN ('volunteer_class_counselor','volunteer_deputy_class_teacher') "
        'AND COALESCE(vsu.service_target_org_unit_id,va.org_unit_id) IN ('+','.join('?' for _ in ids)+') '
        'AND (va.starts_at IS NULL OR va.starts_at<=?) AND (va.ends_at IS NULL OR va.ends_at>=?) ORDER BY va.appointment_key,m.name,m.id',
        (*sorted(ids), end.isoformat()+'T23:59:59', end.isoformat())).fetchall()]


def export_weekly(actor, class_id, occurred_from, occurred_to, pauses=()):
    start, end = period(occurred_from, occurred_to)
    c = connect()
    try:
        selected, ids = _context(c, actor, class_id)
        if not sources.available(c):
            raise ValueError('读书来源存储尚未准备，暂不能导出')
        days = []
        for index in range((end-start).days+1):
            day = start + timedelta(days=index)
            workday = activity.is_workday(c, day)
            if workday is None:
                raise ValueError(f'{day}的读书工作日日历未配置，请先准备日历')
            if workday:
                days.append(day)
        if not days:
            raise ValueError('所选区间没有读书工作日')
        members = _roster(c, ids, start, end)
        if not members:
            raise ValueError('所选班级在统计区间没有有效名册')
        paused = _pauses(pauses, {m['id'] for m in members}, start, end)
        raw = [dict(r) for r in execute(c,
            'SELECT r.member_id,r.occurred_on,r.record_json,i.id AS import_id,i.original_filename,i.file_sha256,i.source_snapshot_json '
            'FROM learning_credit_source_rows r JOIN learning_credit_source_imports i ON i.id=r.import_id '
            "WHERE i.source_kind='DAILY_READING' AND i.class_org_unit_id IN (" + ','.join('?' for _ in ids) + ') '
            'AND r.occurred_on>=? AND r.occurred_on<=? ORDER BY i.id,r.id',
            (*sorted(ids), start.isoformat(), end.isoformat())).fetchall()]
        records = defaultdict(set); imports = {}; ignored = {}
        for row in raw:
            record = json.loads(row['record_json'])
            state = record.get('已完成录音')
            records[(int(row['member_id']), str(row['occurred_on'])[:10])].add(state)
            imports[row['import_id']] = _safe_text(row['original_filename'])
            snapshot = json.loads(row['source_snapshot_json'])
            ignored[row['file_sha256']] = len(snapshot.get('ignored_rows', []))
        if not raw:
            raise ValueError('所选区间尚未保存读书打卡来源，请先同步或上传')
        data = {'class_name': selected['name'], 'center_name': selected.get('parent_name') or '',
            'start': start, 'end': end, 'days': days, 'groups': [], 'ignored': sum(ignored.values()),
            'sources': imports, 'staff': _staff(c, ids, end), 'total_completed': 0, 'total_required': 0, 'pending': False}
        groups = {}
        for member in members:
            exempt = exemption_reason(member['group_name'], selected['name'], metric='reading')
            values = []; completed = required = 0; pending = False
            for day in days:
                if day not in member['included_days']:
                    values.append('不在本班'); continue
                if day in paused[member['id']]:
                    values.append('暂停'); continue
                if member['status'] == 'SUSPENDED':
                    values.append('停学'); continue
                if exempt:
                    values.append('免要求'); continue
                required += 1
                states = records.get((member['id'], day.isoformat()), set())
                value = '√' if states == {'是'} else '×' if states == {'否'} else '待核对' if states else '待同步'
                values.append(value)
                completed += value == '√'
                pending |= value in ('待同步','待核对')
            item = {**member, 'values': values, 'completed': completed, 'required': required,
                'pending': pending, 'rate': None if pending else completed/required if required else 0}
            key = member['group_id']
            group = groups.setdefault(key, {'name': member['group_name'], 'members': [], 'completed': 0, 'required': 0, 'pending': False})
            group['members'].append(item); group['completed'] += completed; group['required'] += required; group['pending'] |= pending
        data['groups'] = list(groups.values())
        for group in data['groups']:
            data['total_completed'] += group['completed']; data['total_required'] += group['required']; data['pending'] |= group['pending']
        book = build_workbook(data)
        content = _workbook_bytes(book)
        write_audit(c, actor_user_id=actor, action='learning.reading_weekly.export', resource_type='reading_weekly_export',
            org_unit_id=class_id, after={'occurred_from': str(start), 'occurred_to': str(end), 'member_count': len(members),
            'source_count': len(imports), 'ignored_count': data['ignored'], 'temporary_pauses': list(pauses),
            'contact_fields': False, 'ledger_write': False})
        c.commit()
        return content
    finally:
        c.close()


def _workbook_bytes(book):
    """Keep formula results visible in mobile viewers that do not recalculate."""
    raw = io.BytesIO(); book.save(raw)
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    output = io.BytesIO()
    with zipfile.ZipFile(raw) as original, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as result:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(data)
                for node in root.iter(ns+'c'):
                    formula = node.find(ns+'f')
                    match = re.fullmatch(r'COUNTIF\(([A-Z]+\d+:[A-Z]+\d+),"√"\)/(\d+)', formula.text or '') if formula is not None else None
                    if match:
                        value = node.find(ns+'v')
                        if value is None: value = ET.SubElement(node, ns+'v')
                        value.text = str(sum(cell.value == '√' for row in book.active[match[1]] for cell in row)/int(match[2]))
                data = ET.tostring(root, encoding='utf-8')
            result.writestr(item, data)
    return output.getvalue()


def build_workbook(data):
    book = Workbook(); ws = book.active; ws.title = '每周读书打卡统计'
    ws.sheet_view.showGridLines = False
    days = data['days']; personal = len(days)+3; group_col = personal+1; last = get_column_letter(group_col)
    edge = Side(style='thin', color='000000')
    def cell(row, col, value=None, color='FFFFFF', bold=False, size=11):
        v = _literal(ws, row, col, value)
        v.font = Font(name='宋体', size=size, bold=bold)
        v.fill = PatternFill('solid', fgColor=color)
        v.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        v.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        return v
    def merged(r1, c1, r2, c2, value, color='FFFFFF', size=11):
        for r in range(r1,r2+1):
            for col in range(c1,c2+1): cell(r,col,color=color)
        ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
        cell(r1,c1,value,color,bold=r1==1,size=size)
    merged(1,1,3,1,'组别','FFC000')
    merged(1,2,1,group_col,f"{data['class_name']}{data['start'].month}月读卡数据统计",'FFC000',18)
    cell(2,2,'姓名'); cell(3,2,'')
    for col, day in enumerate(days,3):
        cell(2,col,day).number_format='m"月"d"日"'
        cell(3,col,'周'+'一二三四五六日'[day.weekday()],bold=True)
    week=(data['start'].day-1+data['start'].replace(day=1).weekday())//7+1
    merged(2,personal,2,group_col,f"第{week}周（{data['start'].day}–{data['end'].day}）")
    cell(3,personal,'个人'); cell(3,group_col,'小组')
    row=4; colors=['FFF2CC','D9E2F3','D6DCE4','FCE4D6']
    for key,label in [('volunteer_class_counselor','班主任'),('volunteer_deputy_class_teacher','副班主任')]:
        names=[_safe_text(r['name']) for r in data.get('staff',[]) if r['appointment_key']==key]
        if names:
            cell(row,1,label,'D6DCE4'); cell(row,2,'、'.join(names))
            merged(row,3,row,group_col,'师资信息（不额外计入学员统计）')
            ws.row_dimensions[row].height=24; row+=1
    for index, group in enumerate(data['groups']):
        first=row
        for member in group['members']:
            cell(row,2,_safe_text(member['name']))
            values=member['values']
            for col,value in enumerate(values,3):
                cell(row,col,value,'FFFF00' if value == '×' else 'FFFFFF')
            if len(set(values))==1 and values[0] in ('暂停','停学','免要求'):
                merged(row,3,row,personal-1,values[0])
            cell(row,personal,'待核对' if member['pending'] else member['rate']).number_format='0%'
            ws.row_dimensions[row].height=19
            row+=1
        merged(first,1,row-1,1,group['name'],colors[index%len(colors)])
        value='待核对' if group['pending'] else '—'
        merged(first,group_col,row-1,group_col,value)
        if not group['pending'] and group['required']:
            # Only trusted formula text is assigned as a formula. Names and source text stay literal.
            ws.cell(first,group_col,f'=COUNTIF(C{first}:{get_column_letter(personal-1)}{row-1},"√")/{group["required"]}')
        ws.cell(first,group_col).number_format='0%'
    merged(row,1,row,2,'汇总','C00000'); ws.cell(row,1).font=Font(name='宋体',size=12,bold=True,color='FFFFFF')
    for col in range(3,personal): cell(row,col,'')
    cell(row,personal,data['total_completed'],'E2EFDA',True)
    value='待核对' if data['pending'] else data['total_completed']/data['total_required'] if data['total_required'] else '—'
    cell(row,group_col,value,'E2EFDA',True).number_format='0%'
    notes = [f"{data['center_name']}　统计区间：{data['start']}至{data['end']}。汇总数量为已完成打卡人天。",
        '√＝已完成录音；×＝来源明确未完成；待同步＝尚无来源；待核对＝记录冲突或资料待核对。完成率＝已完成人天÷应打卡人天。',
        '小组完成率＝小组√总数÷（小组在册人数×应打卡天数）；临时暂停、停学、免要求及不在本班日期扣除对应人天。按工作日日历（含调休）统计，分组按区间内最后有效本班归属；临时暂停只用于本次导出。']
    if data['ignored']: notes.append(f"来源另有{data['ignored']}条未通过身份或内容检查的记录，仍待核对；仅本组缺失或冲突的打卡影响本组完成率，已齐全的小组正常计算。")
    notes.append('保存来源：'+'；'.join(f"#{key} {value}" for key,value in data['sources'].items()))
    for offset,note in enumerate(notes,2):
        ws.merge_cells(start_row=row+offset,start_column=1,end_row=row+offset,end_column=group_col)
        v=_literal(ws,row+offset,1,note); v.font=Font(name='宋体',size=10,color='595959'); v.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[row+offset].height=32 if len(note)>70 else 26
    ws.column_dimensions['A'].width=17; ws.column_dimensions['B'].width=18
    for col in range(3,personal): ws.column_dimensions[get_column_letter(col)].width=13
    ws.column_dimensions[get_column_letter(personal)].width=14; ws.column_dimensions[last].width=14
    ws.row_dimensions[1].height=34; ws.row_dimensions[2].height=22; ws.row_dimensions[3].height=20
    ws.freeze_panes='C4'; ws.print_title_rows='1:3'
    ws.conditional_formatting.add(f'C4:{get_column_letter(personal-1)}{row-1}',
        CellIsRule(operator='equal', formula=['"×"'], fill=PatternFill('solid',fgColor='FFFF00')))
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.print_area=f'A1:{last}{row+len(notes)+1}'
    return book
