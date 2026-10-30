"""Read-only monthly progress workbook for weekly group sharing."""
from __future__ import annotations

import io
import re
import unicodedata
from calendar import monthrange
from datetime import UTC, date, datetime, timedelta, timezone

from app.db import connect, execute
from app.services.iam import accessible_org_ids
from app.services.learning_org_catalog import class_catalog
from app.services.study_membership import CURRENT_CLASS_ID_SQL, CURRENT_GROUP_ID_SQL
from app.services.study_meetings import StudyMeetingError


def report_period(month: str | None, as_of: str | None) -> tuple[str, str, str]:
    today = datetime.now(UTC).astimezone(timezone(timedelta(hours=8))).date()
    month = month or today.strftime('%Y-%m')
    try:
        if not re.fullmatch(r'\d{4}-\d{2}', month):
            raise ValueError
        first = date.fromisoformat(month + '-01')
        last = first.replace(day=monthrange(first.year, first.month)[1])
        cutoff = date.fromisoformat(as_of) if as_of else min(today, last)
    except (ValueError, TypeError):
        raise StudyMeetingError('请选择有效统计月份和截止日期') from None
    if first > today or not first <= cutoff <= min(today, last):
        raise StudyMeetingError('截止日期须在统计月份内，且不能晚于今天')
    return month, first.isoformat(), cutoff.isoformat()


def _region(class_id, organizations):
    node, visited = organizations.get(class_id), set()
    while node and node['id'] not in visited:
        visited.add(node['id'])
        if node['unit_type'] in {'REGIONAL_CENTER', 'OPERATING_UNIT'}:
            return node['id'], node['name']
        node = organizations.get(node['parent_id'])
    # Direct classes stay independent, as in the shared operating report.
    return class_id, organizations[class_id]['name']


def _group_policies(connection, class_ids, month):
    """Normal catalog groups default to monthly required; explicit waivers remain authoritative."""
    from app.services.learning_cycle_calendar import clock, month_label
    bindings, cycles = {}, {}
    for start in range(0,len(class_ids),200):
        batch=class_ids[start:start+200]; marks=','.join('?' for _ in batch)
        for raw in execute(connection,
            'SELECT b.*,v.version_label FROM class_learning_bindings b '
            'JOIN learning_plan_versions v ON v.id=b.plan_version_id '
            f'WHERE b.class_org_unit_id IN ({marks})',tuple(batch)).fetchall():
            bindings[int(raw['id'])]=dict(raw)
    ids=list(bindings)
    for start in range(0,len(ids),200):
        batch=ids[start:start+200];marks=','.join('?' for _ in batch)
        for raw in execute(connection,f'SELECT * FROM class_learning_cycles WHERE binding_id IN ({marks})',tuple(batch)).fetchall():
            row=dict(raw);row['binding']=bindings[int(row['binding_id'])]
            cycles.setdefault(row['class_org_unit_id'],[]).append(row)
    result={}
    for class_id in class_ids:
        rows=cycles.get(class_id,[])
        if not rows:
            result[class_id]='REQUIRED';continue
        if rows and {r['group_meeting_policy'] for r in rows}=={'REQUIRED'}:
            result[class_id]='REQUIRED';continue
        dated=[]; invalid=False
        for row in rows:
            try:
                period=month_label(clock(connection,row['binding'],row)['month'])
            except (ValueError,TypeError,KeyError):
                invalid=True
                continue
            if period<=month:dated.append((period,row))
        if invalid:
            result[class_id]='UNKNOWN';continue
        latest=max((p for p,_ in dated),default=None)
        rows=[r for p,r in dated if p==latest]
        policies={r['group_meeting_policy'] for r in rows}
        result[class_id]=next(iter(policies)) if len(policies)==1 else 'UNKNOWN'
    waived=set()
    cycles_by_id={int(r['id']):r for rows in cycles.values() for r in rows}
    for start in range(0,len(ids),200):
        batch=ids[start:start+200];marks=','.join('?' for _ in batch)
        for task in execute(connection,
            'SELECT t.group_org_unit_id,t.class_learning_cycle_id FROM group_learning_cycle_tasks t '
            'JOIN class_learning_cycles lc ON lc.id=t.class_learning_cycle_id '
            f"WHERE lc.binding_id IN ({marks}) AND t.task_type='GROUP_MEETING' AND t.status='WAIVED'",tuple(batch)).fetchall():
            row=cycles_by_id[int(task['class_learning_cycle_id'])]
            try:
                period=month_label(clock(connection,row['binding'],row)['month'])
            except (ValueError,TypeError,KeyError):
                result[row['class_org_unit_id']]='UNKNOWN';continue
            if period==month:waived.add(task['group_org_unit_id'])
    return result,waived


def _targets(connection, actor, filters, month, *, roster=None, attribution=None, scope_permissions=()):
    allowed = None
    for permission in ('exports:normal', 'plans:credit_settlement_preview', *scope_permissions):
        scope = accessible_org_ids(actor, permission)
        if scope is not None:
            allowed = set(scope) if allowed is None else allowed & set(scope)
    organizations = {r['id']: dict(r) for r in execute(connection,
        'SELECT id,name,parent_id,unit_type,is_active FROM org_units').fetchall()}
    classes, centers = {}, {}
    for item in class_catalog(connection, allowed):
        if filters.get('class_org_unit_id') and filters['class_org_unit_id'] not in item['related_class_ids']:
            continue
        center_id, label = _region(item['id'], organizations)
        if filters.get('management_org_unit_id') and filters['management_org_unit_id'] != center_id:
            continue
        centers.setdefault(center_id, {'id': center_id, 'name': label, 'groups': {}})
        for class_id in item['related_class_ids']:
            if organizations[class_id]['is_active'] and (allowed is None or class_id in allowed):
                classes[class_id] = {'id': class_id, 'name': item['name'], 'center_id': center_id}
    policies,waived=_group_policies(connection,sorted(classes),month)
    from app.services.learning_requirement_policy import exemption_reason
    groups = {}
    catalog = list(organizations.values())
    if attribution and attribution.get('state') == 'FROZEN':
        catalog = [{**dict(r), 'id':r['group_org_unit_id'], 'parent_id':r['class_org_unit_id'], 'name':organizations.get(r['group_org_unit_id'],{}).get('name') or r['group_name'], 'unit_type':'GROUP', 'is_active':1} for r in execute(connection, 'SELECT * FROM study_meeting_roster_groups WHERE period=?',(month,)).fetchall()]
    for group in catalog:
        if group['unit_type'] != 'GROUP' or not group['is_active'] or group['parent_id'] not in classes:
            continue
        parent = classes[group['parent_id']]
        if exemption_reason(group['name'], parent['name'], metric='group_meeting'):
            continue
        row = {**group, 'class_name': parent['name'], 'center_id': parent['center_id'],
               'policy':'WAIVED' if group['id'] in waived else policies.get(group['parent_id'],'UNKNOWN'), 'roster': set()}
        groups[group['id']] = row
        centers[parent['center_id']]['groups'][group['id']] = row
    if roster is None:
        from app.services.study_meeting_monthly_roster import roster_for_period
        roster, attribution = roster_for_period(connection, month)
    for member in roster:
        group = groups.get(member['group_org_unit_id'])
        if group and group['parent_id'] == member['class_org_unit_id']:
            group['roster'].add(int(member['member_id']))
    for center in centers.values():
        center['attribution'] = attribution
    return centers, groups


def sharing_dataset(*, actor_user_id, report_month=None, report_as_of=None, scope_permissions=(), **filters):
    from app.services.study_meeting_operations import _export_records

    month, first, cutoff = report_period(report_month, report_as_of)
    selection = {k: filters.get(k) for k in ('management_org_unit_id', 'class_org_unit_id')}
    records = _export_records(actor_user_id, None, {**selection, 'status': 'SUBMITTED',
        'meeting_date_from': first, 'meeting_date_to': cutoff}, scope_permissions=scope_permissions)
    # Preserve legitimate different-cycle meetings. Historic duplicates of the
    # same group/cycle are represented once; completed groups are also unique.
    chosen = {}
    for record in records:
        key = record['study_group_org_unit_id'], record['learning_cycle_id']
        if key not in chosen or int(record['id']) > int(chosen[key]['id']):
            chosen[key] = record
    records = list(chosen.values())
    connection = connect()
    try:
        # One request-local snapshot; report generation never advances cycles,
        # repairs memberships, confirms a meeting or posts credits.
        import sqlite3
        if isinstance(connection, sqlite3.Connection):
            execute(connection, 'BEGIN')
        from app.services.study_meeting_monthly_roster import roster_for_period
        roster, attribution = roster_for_period(connection, month)
        centers, groups = _targets(connection, actor_user_id, selection, month, roster=roster, attribution=attribution,
                                  scope_permissions=scope_permissions)
        facts = {int(r['id']): [] for r in records}
        for start in range(0, len(records), 200):
            batch = records[start:start + 200]
            marks = ','.join('?' for _ in batch)
            source = {int(r['id']): r for r in batch}
            actual = execute(connection,
                f'SELECT id,class_org_unit_id,study_group_org_unit_id,meeting_date,status FROM study_meeting_sessions WHERE id IN ({marks})',
                tuple(source)).fetchall()
            if len(actual) != len(batch):
                raise StudyMeetingError('导出期间学习会记录已变化，请重新导出')
            for row in actual:
                before = source[int(row['id'])]
                if any(str(row[field])[:10] != str(before[field])[:10] if field == 'meeting_date'
                       else row[field] != before[field] for field in
                       ('class_org_unit_id', 'study_group_org_unit_id', 'meeting_date', 'status')):
                    raise StudyMeetingError('导出期间学习会记录已变化，请重新导出')
            for row in execute(connection,
                'SELECT a.study_meeting_session_id,a.member_id,m.name '
                'FROM study_meeting_attendances a JOIN members m ON m.id=a.member_id '
                f'WHERE a.study_meeting_session_id IN ({marks}) ORDER BY m.name,a.member_id', tuple(source)).fetchall():
                facts[int(row['study_meeting_session_id'])].append(dict(row))
    finally:
        connection.close()
    from app.services.learning_requirement_policy import exemption_reason
    by_member = {int(r['member_id']): r for r in roster}
    monthly_home = {group_id:{} for group_id in groups}
    local, exceptions, excluded_records = {}, [], 0
    for record in records:
        group = groups.get(record['study_group_org_unit_id'])
        if not group or group['parent_id'] != record['class_org_unit_id']:
            if exemption_reason(record.get('group_name'), record.get('class_name'), metric='group_meeting'):
                excluded_records += 1
            else:
                exceptions.append(record)
            continue
        attendees = {int(a['member_id']): a['name'] for a in facts[int(record['id'])]}
        local.setdefault(group['id'], []).append((record, attendees))
        for member_id, name in attendees.items():
            member = by_member.get(member_id)
            # Only same-class attendance contributes to one's own group.
            # Cross-class movers retain their month baseline and cannot gain
            # credit for a meeting in their new class under the old class.
            if not member or member['class_org_unit_id'] != record['class_org_unit_id']:
                continue
            home_id = member.get('group_org_unit_id')
            if home_id in monthly_home:
                monthly_home[home_id][member_id] = name
    for center in centers.values():
        meetings = []
        for group in sorted(center['groups'].values(), key=lambda r: (r['class_name'], r['name'], r['id'])):
            hosted = local.get(group['id'], [])
            actual = {mid:name for _, attendees in hosted for mid,name in attendees.items()}
            home = monthly_home[group['id']]
            incoming = {mid:name for mid,name in home.items() if mid not in actual}
            dates = sorted({str(record['meeting_date'])[:10] for record,_ in hosted})
            present = len(home) if hosted or home else None
            names = '、'.join(actual.values())
            if incoming:
                names += ('；' if names else '') + '同班跨组参会：' + '、'.join(incoming.values())
            meetings.append({'id':max((int(r['id']) for r,_ in hosted),default=None),
                'date':dates[0] if len(dates)==1 else '、'.join(dates) if dates else None,
                'dates':dates,'class_name':group['class_name'],'group_id':group['id'],'group_name':group['name'],
                'names':names,'other_names':'、'.join(name for mid,name in actual.items() if mid not in group['roster']),
                'home':set(home),'present':present,'roster':len(group['roster']),
                'same_class_elsewhere_names':'、'.join(incoming.values()),
                'participation_rate':len(home)/len(group['roster']) if group['roster'] and (hosted or home) else None,
                'status':('已完成' if hosted else '待完成') if group['policy']=='REQUIRED' else
                    {'WAIVED':'已完成（免开展）' if hosted else '免开展',
                     'SUSPENDED':'已完成（暂停）' if hosted else '暂停开展',
                     'UNKNOWN':'已完成（政策待核对）' if hosted else '政策待核对'}[group['policy']]})
        required={g['id'] for g in center['groups'].values() if g['policy']=='REQUIRED'}
        unknown={g['id'] for g in center['groups'].values() if g['policy']=='UNKNOWN'}
        completed={row['group_id'] for row in meetings if row['id'] is not None}
        planned=None if unknown else len(required)
        finished=len(completed & (required|unknown))
        # Preserve the established completed-group weighted center rate.
        # Each group's monthly same-class participants are distinct once.
        present=sum((r['present'] or 0) for r in meetings if r['id'] is not None)
        denominator=sum(r['roster'] for r in meetings if r['id'] is not None)
        center['planning_note'] = ('周期政策相互冲突或月份无法唯一确认：' + '、'.join(g['class_name']+'／'+g['name'] for g in center['groups'].values() if g['policy']=='UNKNOWN')) if unknown else '按有效组织目录统计应开小组，精进组及班委组免要求'
        center.update(meetings=meetings,planned=planned,completed=finished,pending_policy_count=len(unknown),
            present=present,denominator=denominator,completion_rate=finished/planned if planned else None,
            participation_rate=present/denominator if denominator else None)
    return {'month':month,'cutoff':cutoff,'basis':'completed_groups_monthly_distinct',
        'attribution':attribution,'centers':sorted(centers.values(),key=lambda r:(r['name'],r['id'])),
        'exceptions':exceptions,'excluded_record_count':excluded_records,'source_count':len(records),'selection':selection}


def _workbook(data):
    # Server-side exports deliberately use the platform's existing Excel
    # dependency and typed snapshot values, so mobile previews need no recalc.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    book = Workbook()
    book.remove(book.active)
    blue, pale = '4472C4', 'D9E2F3'
    table_number = 0
    period = f"{data['month']}，截至 {data['cutoff']}"
    basis = '总参会率＝已完成组月度本组参会人数合计÷对应本组应参人数合计（加权，同班跨组去重计入本人小组）'

    def sheet(name, title, headers, rows, widths, percent_cols, summary=None):
        nonlocal table_number
        clean = re.sub(r'[\\/*?:\[\]]', '-', name).strip("'")[:31] or '分中心'
        candidate, suffix = clean, 2
        while candidate in book.sheetnames:
            candidate = clean[:26] + f'-{suffix}'
            suffix += 1
        ws = book.create_sheet(candidate)
        ws.sheet_view.showGridLines = False
        count = len(headers)
        end = get_column_letter(count)
        ws.merge_cells(f'A1:{end}1'); ws['A1'] = title
        ws['A1'].data_type = 's'
        ws['A1'].font = Font(name='Microsoft YaHei', size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 32
        ws.merge_cells(f'A2:{end}2'); ws['A2'] = period
        ws['A2'].data_type = 's'
        ws['A2'].alignment = Alignment(horizontal='center')
        if summary:
            rate_text = f"{summary['completion_rate']:.0%}" if summary['completion_rate'] is not None else '—'
            ws.merge_cells('A3:E3')
            planned_text=str(summary['planned'])+' 组' if summary['planned'] is not None else '待配置'
            ws['A3'] = f"本月预定 {planned_text}，已完成 {summary['completed']} 组，完成率 {rate_text}"
            ws.merge_cells('F3:H3'); ws['F3'] = '分中心总参会率'
            ws.merge_cells('I3:J3')
            ws['I3'] = summary['participation_rate'] if summary['participation_rate'] is not None else '—'
            ws['I3'].number_format = '0%'
            for address in ('A3','F3','I3'):
                ws[address].font = Font(name='Microsoft YaHei', size=11, bold=True, color=blue)
                ws[address].alignment = Alignment(vertical='center',horizontal='center' if address=='I3' else 'left')
            ws.row_dimensions[3].height = 30
        ws.merge_cells(f'A4:{end}4'); ws['A4'] = basis
        ws['A4'].font = Font(name='Microsoft YaHei', size=10, color='595959')
        ws.row_dimensions[4].height = 24
        for row_number, values in enumerate([headers, *rows], 6):
            for col, value in enumerate(values, 1):
                cell = ws.cell(row_number, col, value)
                if isinstance(value, str): cell.data_type = 's'
                cell.font = Font(name='Microsoft YaHei', size=11, bold=row_number == 6,
                                 color='FFFFFF' if row_number == 6 else '202020')
                centered = row_number == 6 or isinstance(value,(int,float,date)) or value in ('—','已完成','未完成','待完成')
                cell.alignment = Alignment(vertical='center', horizontal='center' if centered else 'left', wrap_text=True)
                edge = Side(style='thin', color='C7D0E0')
                cell.border = Border(left=edge,right=edge,top=edge,bottom=edge)
                cell.fill = PatternFill('solid', fgColor=blue if row_number == 6 else pale if row_number % 2 == 0 else 'FFFFFF')
                if row_number > 6 and col in percent_cols:
                    cell.number_format = '0%'
                elif isinstance(value, date): cell.number_format = 'yyyy/mm/dd'
            lines = max((sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in str(value or ''))
                         // max(1, int(widths[col-1])-2) + 1 for col,value in enumerate(values,1)),default=1)
            ws.row_dimensions[row_number].height = max(34 if row_number == 6 else 30, lines*15+10)
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        last = max(6, 6 + len(rows))
        if rows:
            table_number += 1
            table = Table(displayName=f'MeetingShare{table_number}', ref=f'A6:{end}{last}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False)
            ws.add_table(table)
        ws.auto_filter.ref = f'A6:{end}{last}'
        ws.freeze_panes = 'A7'
        note = last + 2
        ws.merge_cells(start_row=note, start_column=1, end_row=note + 1, end_column=count)
        ws.cell(note, 1, '正常有效小组每月预定1次；精进组、班委组不纳入统计，免开展、暂停开展不计完成目标。未配置周期的正常小组按组织目录计入应开数，冲突政策单列待核对。本人参加同班其他小组会也计入本人小组月度参会率，同人同月不重复；自己小组未开会仍为待完成。30日0点锁定当月归组（2月月末），跨班调整保持当月原归属，原始记录不改。')
        ws.cell(note, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(note, 1).font = Font(name='Microsoft YaHei', size=10, color='595959')
        ws.row_dimensions[note].height = 30
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:6'; ws.print_area = f'A1:{end}{note + 1}'
        return ws

    count_label = '本组参会人次'
    denominator_label = '在册人次（已完成场次）'
    summary_rows = [[c['name'], c['planned'] if c['planned'] is not None else '待配置', c['completed'],
        c['planned']-c['completed'] if c['planned'] is not None else '待配置',
        c['completion_rate'] if c['completion_rate'] is not None else '—', c['present'], c['denominator'],
        c['participation_rate'] if c['participation_rate'] is not None else '—', c['planning_note']] for c in data['centers']]
    sheet('分中心完成情况', '各分中心小组学习会月度完成情况',
        ['分中心／直属班', '本月预定（组）', '截至已完成（组）', '未完成（组）', '完成率', count_label, denominator_label, '总参会率','统计说明'],
        summary_rows, [26,18,19,17,14,23,25,16,48], (5,8))
    for center in data['centers']:
        rows = [[i, r['class_name'], r['group_name'], r['names'], r['other_names'],
            date.fromisoformat(r['date']) if r['date'] and len(r['date'])==10 else r['date'], r['present'], r['roster'],
            r['participation_rate'] if r['participation_rate'] is not None else '—', r['status']]
            for i,r in enumerate(center['meetings'],1)]
        sheet(center['name'], center['name']+'小组学习会统计明细',
            ['序号','班级','小组','参会人员（含跨组）','其中跨组／其他人员','会议日期','本组参会人数','本组在册人数','小组参会率','完成情况'],
            rows, [8,18,20,54,28,16,17,17,16,14], (9,), center)
    if data['exceptions']:
        sheet('需核对记录', '组织归属变化记录（未计入当前小组目标）',
            ['学习会编号','会议日期','原班级','小组','说明'],
            [[r['session_code'],date.fromisoformat(str(r['meeting_date'])[:10]),r['class_name'],r['group_name'],
              '小组已停用、当前不在筛选范围或所属班级已变化，请核对后重新导出'] for r in data['exceptions']],
            [30,18,24,28,65], ())
    output = io.BytesIO(); book.save(output)
    return output.getvalue()


def export_sharing_workbook(*, actor_user_id, **filters):
    from app.services.study_meeting_operations import _audit_export
    data = sharing_dataset(actor_user_id=actor_user_id, **filters)
    content = _workbook(data)
    _audit_export(actor_user_id, None, {**data['selection'], 'report_month':data['month'],
        'report_as_of':data['cutoff'], 'attendance_basis':data['basis'],
        'target_basis':'required_groups_one_per_month'}, 'sharing', data['source_count'],
        sum(len(c['meetings']) for c in data['centers']))
    return content
