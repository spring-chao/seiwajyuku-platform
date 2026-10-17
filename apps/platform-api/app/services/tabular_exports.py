"""Small, literal-text Excel reports; untrusted strings never become formulas."""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def add_sheet(book, name, headers, rows, *, percent_columns=()):
    sheet = book.create_sheet(name)
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    edge = Side(style='thin', color='BCCFC6')
    for row_index, values in enumerate([headers, *rows], 1):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            if row_index == 1:
                cell.fill = PatternFill('solid', fgColor='275747')
                cell.font = Font(bold=True, color='FFFFFF')
            else:
                if row_index % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EFF6F2')
                if col in percent_columns:
                    cell.number_format = '0.0%'
                elif isinstance(value, (date, datetime)):
                    cell.number_format = 'yyyy-mm-dd'
        sheet.row_dimensions[row_index].height = 30 if row_index == 1 else min(150, max(30, max((len(str(v)) for v in values), default=0) // 32 * 16 + 22))
    for col, label in enumerate(headers, 1):
        width = max([len(label) * 2 + 2, 14] + [min(54, len(str(row[col - 1] or '')) * 1.5 + 2) for row in rows[:200]])
        sheet.column_dimensions[get_column_letter(col)].width = min(54, width)
    sheet.auto_filter.ref = sheet.dimensions
    if rows:
        table = Table(displayName=f'ExportTable{len(book.worksheets)}', ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
        sheet.add_table(table)
    return sheet


def workbook_bytes(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, headers, rows, percentages in sheets:
        add_sheet(book, name, headers, rows, percent_columns=percentages)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
