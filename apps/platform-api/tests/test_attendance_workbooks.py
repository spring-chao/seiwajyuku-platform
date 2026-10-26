import copy
import io
from datetime import UTC,datetime
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.db import execute,fetch_one,transaction
from app.main import app
from app.services import signin_engine
from test_v12_mvp import _seed_group_leader_fixture
from test_signin_management import ENV,_operator


def seeded_export():
    f=_seed_group_leader_fixture();now=datetime.now(UTC).isoformat();token=uuid4().hex
    with transaction() as c:
        gid=execute(c,"INSERT INTO attendance_event_groups(source_key,external_group_id,org_unit_id,study_org_unit_id,title,activity_type,event_date,status,created_at,updated_at) VALUES ('signin',?,?,?,'合成实时班会','CLASS_MEETING','2026-10-25','ACTIVE',?,?)",(token,f['class_id'],f['class_id'],now,now)).lastrowid
        execute(c,"UPDATE org_units SET name='合成班精进组' WHERE id=?",(f['other_group_id'],))
        execute(c,"UPDATE members SET name='合成实时新姓名' WHERE id=?",(f['member_id'],))
        sids={}
        for index,code in enumerate(('MORNING','AFTERNOON','KONPA')):
            source_id=token+'-'+code
            sid=execute(c,"INSERT INTO attendance_sessions(event_group_id,external_session_id,session_code,session_name,session_order,status,created_at,updated_at) VALUES (?,?,?,?,?,'ACTIVE',?,?)",(gid,source_id,code,code,index,now,now)).lastrowid
            sids[code]=(sid,source_id)
            execute(c,"INSERT INTO attendance_records(attendance_session_id,external_record_id,member_id,name_snapshot,participant_type,score_eligible,attendance_status,received_at,created_at,updated_at) VALUES (?,?,?,'合成过时同步姓名','MEMBER',0,'ABSENT',?,?,?)",(sid,source_id,f['member_id'],now,now,now))
    events={key:{'event_id':value[1],'event_group_id':token,'event_date':'2026-10-25','activity_type':'class_meeting','session_code':key,'session_name':{'MORNING':'上午','AFTERNOON':'下午','KONPA':'晚上空巴'}[key],'org_unit_id':f['class_id'],'class_org_unit_id':f['class_id'],'lifecycle_status':'CONFIRMED'} for key,value in sids.items()}
    def row(mid,rid,checked=True):
        member=fetch_one('SELECT id,member_code,name FROM members WHERE id=?',(mid,))
        return {'registration_id':rid,'name':'=SYNTHETIC()' if mid==f['member_id'] else member['name'],'actual_attendee_name':'=SYNTHETIC()' if mid==f['member_id'] else member['name'],'registered_name':'合成原报名姓名','platform_member_id':str(mid),'member_code':member['member_code'],'checked':checked,'checked_at':'2026-10-25T01:00:00Z' if checked else None,'phone':f['phone'],'attendance_role':'HOME_CLASS_MEMBER','class_name':'合成班','group_name':''}
    facts={'MORNING':[row(f['member_id'],'one'),row(f['other_member_id'],'waived')],'AFTERNOON':[row(f['member_id'],'duplicate')],'KONPA':[row(f['other_member_id'],'night-only')]}
    calls=[]
    def engine(path,envelope,**kwargs):
        operation=path.rsplit('/',1)[-1];source_id=envelope['payload']['event_id'];calls.append((operation,source_id,envelope['allowed_org_unit_ids']))
        event=next(e for e in events.values() if e['event_id']==source_id);rows=facts[event['session_code']]
        if operation=='stats':return {'event':copy.deepcopy(event),'home_class_total':2,'home_class_checked':len(rows),'cross_class_checked':0,'guest_checked':0,'onsite_total':len(rows),'rate':100}
        assert operation in {'export','event_detail'}
        return {'event':copy.deepcopy(event),'sessions':copy.deepcopy(list(events.values())),'rows':copy.deepcopy(rows)}
    return f,gid,sids,events,calls,engine


def test_all_session_activity_and_onsite_downloads_share_live_rows_and_requirement_summary():
    f,gid,sids,events,calls,engine=seeded_export()
    with patch.dict(__import__('os').environ,ENV),TestClient(app) as client,patch.object(signin_engine,'engine_request',side_effect=engine):
        headers,_=_operator({'members:read','attendance:export'},f['class_id'])
        all_response=client.get(f'/api/v1/attendance/event-groups/{gid}/records.xlsx',headers=headers)
        assert all_response.status_code==200,all_response.text[:500]
        book=load_workbook(io.BytesIO(all_response.content));names=book['签到名单']
        assert set(book.sheetnames)=={'签到名单','场次汇总','签到汇总','小组参会统计'}
        assert names.max_row==5 and book['场次汇总'].max_row==4
        assert names['C2'].value=='合成实时新姓名' and names['B2'].value=='=SYNTHETIC()' and names['B2'].data_type=='s'
        assert names['N2'].value=='2026-10-25 09:00:00'
        values=str([list(ws.values) for ws in book]);assert f['phone'] not in values and '合成过时同步姓名' not in values
        assert '精进组免学习要求' in values
        assert all(row[0].value!='合成班精进组' for row in book['小组参会统计'].iter_rows(min_row=2))
        summary={r[0].value:r[1].value for r in book['签到汇总'].iter_rows(min_row=2)}
        assert summary['班会应参人数（排除精进组）']==1 and summary['整体去重到场']==1 and summary['总参与率（含本班师资）']==1
        assert book['签到汇总']['B3'].number_format=='General'
        single=client.post('/api/v1/attendance/manage/export-workbook',headers=headers,json={'event_id':sids['MORNING'][1]})
        assert single.status_code==200,single.text[:500]
        onsite=load_workbook(io.BytesIO(single.content))
        assert list(onsite['小组参会统计'].values)==list(book['小组参会统计'].values)
        assert list(onsite['签到汇总'].values)==list(book['签到汇总'].values)
        selected=client.get(f'/api/v1/attendance/event-groups/{gid}/records.xlsx?session_id={sids["MORNING"][0]}',headers=headers)
        assert selected.status_code==200
        assert list(load_workbook(io.BytesIO(selected.content))['签到名单'].values)==list(onsite['签到名单'].values)
        assert all(set(scope)=={f['class_id'],f['group_id'],f['other_group_id']} for _,_,scope in calls)
        audit=fetch_one("SELECT after_json FROM audit_logs WHERE action='attendance.workbook_export' ORDER BY id DESC LIMIT 1")
        assert 'LIVE_SIGNIN' in audit['after_json'] and 'phone' not in audit['after_json']


def test_download_requires_export_scope_and_never_silently_exports_stale_snapshot():
    f,gid,sids,events,calls,engine=seeded_export()
    with patch.dict(__import__('os').environ,ENV),TestClient(app) as client,patch.object(signin_engine,'engine_request',side_effect=engine):
        url=f'/api/v1/attendance/event-groups/{gid}/records.xlsx'
        assert client.get(url).status_code==401
        viewer,_=_operator({'members:read'},f['class_id'])
        assert client.get(url,headers=viewer).status_code==403
        outside,_=_operator({'attendance:export','members:read'},f['other_class_id'])
        assert client.get(url,headers=outside).status_code==403
        headers,_=_operator({'members:read','attendance:export'},f['class_id'])
        with patch.object(signin_engine,'engine_request',side_effect=signin_engine.SigninEngineError('合成离线',503)):
            assert client.get(url,headers=headers).status_code==503
        events['AFTERNOON']['event_date']='2026-10-26'
        assert client.get(url,headers=headers).status_code==409
