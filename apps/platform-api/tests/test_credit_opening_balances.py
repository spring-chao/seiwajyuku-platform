"""The same workflow tests run on SQLite locally and isolated MySQL in CI."""
import io
from decimal import Decimal
from unittest.mock import patch
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import connect, execute, fetch_one, transaction
from app.main import app
from app.services import credit_opening_balances as service
from app.services.learning_credits import _insert_entry, credit_ledger_overview
from app.services.wechat_credit_summary import get_member_credit_summary, get_member_credit_entries
from credit_batch_test_support import create_credit_batch_reviewer
from test_learning_credit_ledger import _admin_id


@pytest.fixture
def source(monkeypatch):
    monkeypatch.setenv("LEARNING_CREDIT_OPENING_BALANCE_ENABLED", "true")
    actor, reviewer = _admin_id(), create_credit_batch_reviewer()
    org = "opening-class-" + uuid4().hex[:12]
    with transaction() as connection:
        now = service._db_timestamp(connection)
        execute(connection, "INSERT INTO org_units(id,unit_code,name,unit_type,is_active,created_at,updated_at) VALUES (?,?,?,'CLASS',1,?,?)", (org, org, "期初测试班", now, now))
        ids = []
        for i in range(2):
            code = "000-opening-" + uuid4().hex[:12]
            cursor = execute(connection, "INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE',?,?)", (code, f"测试学长{i}", org, now, now))
            ids.append((int(cursor.lastrowid), code, f"测试学长{i}"))
    perms = {service.OPENING_PERMISSION, "plans:historical_credit_import_manage", "members:read", "exports:normal", "plans:credit_settlement_approve", "plans:credit_settlement_post"}
    def user(uid):
        return {"id": uid, "roles": ["system_admin"], "permissions": sorted(perms)} if uid in {actor, reviewer} else None
    with patch.object(service, "user_context", side_effect=user), patch.object(service, "accessible_org_ids", return_value=None):
        yield {"actor": actor, "reviewer": reviewer, "org": org, "ids": ids, "permissions": perms}


def workbook(source, values=("12.25", "20"), *, transform=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "期初学分"
    sheet.append(service.HEADERS)
    for (_, code, name), value in zip(source["ids"], values):
        sheet.append([code, name, "期初测试班", value])
    if transform:
        transform(sheet)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def submitted(source):
    content = workbook(source)
    result = service.preview(content, "2001-08-31", source["actor"])
    assert result["error_count"] == 0 and result["total_points"] == "32.25"
    created = service.register(content, "2001-08-31", "已确认.xlsx", result["fingerprint"], source["actor"])
    return created["id"], result["fingerprint"], content


def test_opening_roundtrip_same_person_review_atomic_post_and_renamed_replay(source):
    import_id, fingerprint, _ = submitted(source)
    with pytest.raises(ValueError, match="尚未确认"):
        service.action(import_id, source["actor"], "post", fingerprint)
    assert service.action(import_id, source["actor"], "approve", fingerprint)["status"] == "APPROVED"
    assert service.action(import_id, source["actor"], "post", fingerprint)["status"] == "POSTED"
    assert service.action(import_id, source["actor"], "post", fingerprint)["idempotent"]
    replay = workbook(source, values=(Decimal("12.250"), "20.00"))
    assert service.register(replay, "2001-08-31", "换个名字.xlsx", fingerprint, source["actor"])["idempotent"]
    conn = connect()
    try:
        total = service.opening_summary(conn, allowed={source["org"]})
        assert total["entry_count"] == 2 and total["total_points"] == "32.25"
    finally:
        conn.close()
    member = source["ids"][0][0]
    summary = get_member_credit_summary(member)
    assert summary["total_points"] == "12.25" and summary["opening_balance_points"] == "12.25"
    assert summary["current_year_points"] == "0.00" and summary["standard_learning_points"] == "0.00"
    assert get_member_credit_entries(member)["opening_balance"]["cutoff_date"] == "2001-08-31"
    assert service.preview(replay, "2001-08-31", source["actor"])["error_count"] == 2
    with patch("app.services.learning_credits.user_context", return_value={"permissions": ["plans:credit_settlement_preview"]}), patch("app.services.learning_credits.accessible_org_ids", return_value={source["org"]}):
        overview = credit_ledger_overview(actor_user_id=source["actor"])
    assert overview["total_points"] == "32.25" and overview["entry_count"] == 2
    assert overview["categories"][0]["credit_category"] == "OPENING_BALANCE"


@pytest.mark.parametrize("values,error", [(("-1", "20"), "非负"), (("0.001", "20"), "两位"), (("NaN", "20"), "非负"), (("=1+1", "20"), "公式"), ((True, "20"), "布尔")])
def test_parser_reports_actionable_excel_row(source, values, error):
    result = service.preview(workbook(source, values), "2001-08-31", source["actor"])
    assert result["rows"][0]["excel_row"] == 2
    assert any(error in message for message in result["rows"][0]["errors"])
    assert result["fingerprint"] is None


def test_blank_is_skipped_zero_is_retained_and_duplicate_identity_rejected(source):
    result = service.preview(workbook(source, ("0", None)), "2001-08-31", source["actor"])
    assert result["row_count"] == 1 and result["skipped_count"] == 1 and result["total_points"] == "0.00"
    content = workbook(source, transform=lambda sheet: sheet.append([source["ids"][0][1], source["ids"][0][2], "期初测试班", 10]))
    assert "多行" in service.preview(content, "2001-08-31", source["actor"])["rows"][2]["errors"][0]


def test_scope_mismatch_and_missing_permission_fail_before_any_insert(source):
    content = workbook(source)
    with patch.object(service, "accessible_org_ids", return_value=set()):
        result = service.preview(content, "2001-08-31", source["actor"])
        assert result["error_count"] == 2 and "授权范围" in result["rows"][0]["errors"][0]
    source["permissions"].remove("plans:historical_credit_import_manage")
    source["permissions"].remove(service.OPENING_PERMISSION)
    with pytest.raises(PermissionError):
        service.preview(content, "2001-08-31", source["actor"])
    assert TestClient(app).get("/api/v1/learning-credits/opening-balances").status_code == 401


def test_frozen_totals_and_live_member_status_are_rechecked(source):
    import_id, fingerprint, _ = submitted(source)
    with transaction() as connection:
        execute(connection, "UPDATE learning_credit_opening_rows SET points=999 WHERE import_id=? AND member_id=?", (import_id, source["ids"][0][0]))
    with pytest.raises(ValueError, match="变化"):
        service.action(import_id, source["reviewer"], "approve", fingerprint)
    with transaction() as connection:
        execute(connection, "UPDATE learning_credit_opening_rows SET points=12.25 WHERE import_id=? AND member_id=?", (import_id, source["ids"][0][0]))
    service.action(import_id, source["reviewer"], "approve", fingerprint)
    with transaction() as connection:
        execute(connection, "UPDATE members SET status='INACTIVE' WHERE id=?", (source["ids"][0][0],))
    with pytest.raises(ValueError, match="变化"):
        service.action(import_id, source["actor"], "post", fingerprint)
    assert not fetch_one("SELECT member_id FROM learning_credit_opening_balances WHERE import_id=?", (import_id,))


def regular_entry(source, occurred_at="2001-08-31T12:00:00+08:00"):
    return {"member_id": source["ids"][0][0], "credit_category": "STANDARD_LEARNING", "credit_type": "COURSE_COMPLETION", "points": "2.00", "source_type": "TEST", "source_id": uuid4().hex, "class_org_unit_id": source["org"], "rule_key": "TEST", "rule_version": "TEST", "rule_snapshot": {}, "occurred_at": occurred_at, "idempotency_key": uuid4().hex}


def test_earlier_formal_entries_block_opening_and_cutoff_blocks_future_overlap(source):
    with transaction() as connection:
        _insert_entry(connection, regular_entry(source), status="POSTED", actor_user_id=source["actor"])
    assert "重复计算" in service.preview(workbook(source), "2001-08-31", source["actor"])["rows"][0]["errors"][0]


def test_regular_post_after_opening_only_accepts_later_events(source):
    import_id, fingerprint, _ = submitted(source)
    service.action(import_id, source["reviewer"], "approve", fingerprint)
    service.action(import_id, source["actor"], "post", fingerprint)
    with pytest.raises(ValueError, match="重复入账"):
        with transaction() as connection:
            _insert_entry(connection, regular_entry(source), status="POSTED", actor_user_id=source["actor"])
    with transaction() as connection:
        _insert_entry(connection, regular_entry(source, "2001-09-01T00:00:00+08:00"), status="POSTED", actor_user_id=source["actor"])
    assert get_member_credit_summary(source["ids"][0][0])["total_points"] == "14.25"


def test_atomic_post_rolls_back_all_rows_on_mid_import_failure(source):
    import_id, fingerprint, _ = submitted(source)
    service.action(import_id, source["reviewer"], "approve", fingerprint)
    original = service.execute
    count = 0
    def fail(connection, sql, params=()):
        nonlocal count
        if sql.startswith("INSERT INTO learning_credit_opening_balances"):
            count += 1
            if count == 2:
                raise RuntimeError("isolated fault")
        return original(connection, sql, params)
    with patch.object(service, "execute", side_effect=fail), pytest.raises(RuntimeError):
        service.action(import_id, source["actor"], "post", fingerprint)
    assert not fetch_one("SELECT member_id FROM learning_credit_opening_balances WHERE import_id=?", (import_id,))
    assert fetch_one("SELECT status FROM learning_credit_opening_imports WHERE id=?", (import_id,))["status"] == "APPROVED"


def test_downloaded_template_preserves_text_codes_and_has_no_phone(source):
    content = service.template_bytes(source["actor"])
    book = load_workbook(io.BytesIO(content))
    rows = list(book["期初学分"].values)
    row = next(r for r in rows if r[0] == source["ids"][0][1])
    assert row[3] is None and len(row) == 4 and rows[0] == service.HEADERS
    assert "填写说明" in book.sheetnames


def test_disabled_opening_does_not_open_other_credit_gates(source, monkeypatch):
    monkeypatch.setenv("LEARNING_CREDIT_OPENING_BALANCE_ENABLED", "false")
    content = workbook(source)
    result = service.preview(content, "2001-08-31", source["actor"])
    assert result["error_count"] == 0
    with pytest.raises(ValueError, match="暂未启用"):
        service.register(content, "2001-08-31", "source.xlsx", result["fingerprint"], source["actor"])


def test_cancel_keeps_source_and_audit_without_posting(source):
    import_id, fingerprint, content = submitted(source)
    assert service.action(import_id, source["actor"], "cancel", fingerprint)["status"] == "CANCELLED"
    assert service.action(import_id, source["actor"], "cancel", fingerprint)["idempotent"]
    with pytest.raises(ValueError, match="已取消"):
        service.action(import_id, source["reviewer"], "approve", fingerprint)
    with pytest.raises(ValueError, match="已取消"):
        service.register(content, "2001-08-31", "source.xlsx", fingerprint, source["actor"])
    assert fetch_one("SELECT COUNT(*) AS n FROM learning_credit_opening_rows WHERE import_id=?", (import_id,))["n"] == 2
    assert not fetch_one("SELECT member_id FROM learning_credit_opening_balances WHERE import_id=?", (import_id,))


def test_live_learning_staff_can_review_own_source_with_scoped_permission(source):
    from app.services.iam import create_user, user_context, accessible_org_ids
    uid = create_user(source["actor"], username="opening-staff-" + uuid4().hex[:12], display_name="测试学习专职", password="test-only-staff-password", roles=["employee_learning_management"], scopes=[{"scope_type": "UNIT", "org_unit_id": source["org"]}])
    with patch.object(service, "user_context", side_effect=user_context), patch.object(service, "accessible_org_ids", side_effect=accessible_org_ids):
        user = user_context(uid)
        assert service.OPENING_PERMISSION in user["permissions"]
        assert "plans:credit_settlement_post" not in user["permissions"]
        assert "exports:sensitive" not in user["permissions"]
        content = workbook(source)
        checked = service.preview(content, "2001-08-31", uid)
        receipt = service.register(content, "2001-08-31", "学习条线.xlsx", checked["fingerprint"], uid)
        service.action(receipt["id"], uid, "approve", checked["fingerprint"])
        service.action(receipt["id"], uid, "post", checked["fingerprint"])
        batch = fetch_one("SELECT * FROM learning_credit_opening_imports WHERE id=?", (receipt["id"],))
        assert batch["created_by"] == batch["approved_by"] == batch["posted_by"] == uid
        with transaction() as connection:
            execute(connection, "UPDATE data_scope_grants SET org_unit_id='org-suzhou' WHERE user_id=?", (uid,))
        assert service.preview(content, "2001-08-31", uid)["error_count"] == 2
        with pytest.raises(PermissionError, match="范围外"):
            service.action(receipt["id"], uid, "post", checked["fingerprint"])
        with transaction() as connection:
            execute(connection, "UPDATE user_roles SET valid_until='2000-01-01' WHERE user_id=?", (uid,))
        with pytest.raises(PermissionError):
            service.preview(content, "2001-08-31", uid)
        with transaction() as connection:
            execute(connection, "UPDATE app_users SET is_active=0 WHERE id=?", (uid,))
        assert user_context(uid) is None
