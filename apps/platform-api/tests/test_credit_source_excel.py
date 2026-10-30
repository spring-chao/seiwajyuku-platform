"""The same source/import/export acceptance runs on SQLite and isolated MySQL."""
import io
import json
from pathlib import Path
from datetime import date
from unittest.mock import patch
import pytest
from openpyxl import Workbook,load_workbook
from fastapi.testclient import TestClient
from app.main import app
from app.db import connect,execute,transaction,fetch_one
from app.services import credit_source_imports as sources
from app.services import credit_source_exports as exports
from app.services import learning_credits as credits
from app.services import credit_settlement_batches as batches
from app.services import attendance_scoring
from app.services import learning_activity_credits as activity
from app.services.learning_cycle_schedule import parse_utc_datetime
from credit_batch_test_support import create_credit_batch_reviewer
from test_learning_activity_credits import _fixture,_calendar,_admin_id

DAY='2026-10-09'
NOTE='学习条线已核对本次原始记录，按现有规则计分'

@pytest.fixture
def source_fixture(monkeypatch):
    for key in ('LEARNING_CREDIT_BATCH_DRY_RUN_ENABLED','LEARNING_CREDIT_BATCH_APPROVAL_ENABLED','LEARNING_CREDIT_BATCH_POST_ENABLED','LEARNING_CREDIT_SETTLEMENT_ENABLED'):monkeypatch.setenv(key,'true')
    monkeypatch.setenv('LEARNING_CYCLE_MONTHLY_REFRESH_ENABLED','false')
    fixture=_fixture();fixture['actor']=_admin_id();fixture['reviewer']=create_credit_batch_reviewer()
    _calendar(fixture,year=2026,day_types={DAY:'NORMAL_WORKDAY','2026-10-10':'ADJUSTED_WORKDAY','2026-10-11':'WEEKEND'})
    c=connect()
    try:member=dict(execute(c,'SELECT member_code,name FROM members WHERE id=?',(fixture['member_id'],)).fetchone());fixture.update(member)
    finally:c.close()
    with transaction() as c:
        now=credits._db_timestamp(c)
        for index in range(1,37):
            plan=execute(c,"INSERT INTO learning_plan_cycles(plan_version_id,cohort_month,cycle_index,year_index,cycle_label,created_at,updated_at) VALUES (?,1,?,?,?, ?,?)",(fixture['plan_id'],index,(index-1)//12+1,'测试学习日'+str(index),now,now)).lastrowid
            if index==1:
                fixture['cycle_id']=int(execute(c,"INSERT INTO class_learning_cycles(binding_id,class_org_unit_id,learning_cycle_index,plan_cycle_id,opened_at,class_meeting_status,group_meeting_policy,cycle_status,created_at,updated_at) VALUES (?,?,1,?,'2026-10-01','PLANNED','REQUIRED','OPEN',?,?)",(fixture['binding_id'],fixture['class_id'],plan,now,now)).lastrowid)
    user=credits.user_context(fixture['actor']);permissions=set(user['permissions'])|{'plans:period_write','attendance:adjudicate','plans:credit_activity_fact_manage','plans:credit_settlement_manage','plans:credit_settlement_preview','plans:credit_settlement_post','plans:credit_settlement_approve','exports:normal','members:read'}
    def principal(actor):return {'id':actor,'roles':['system_admin'],'permissions':sorted(permissions)}
    with patch.object(credits,'user_context',side_effect=principal),patch.object(credits,'accessible_org_ids',return_value=None):yield fixture


def workbook(f,kind,values=None):
    wb=Workbook();s=wb.active;s.title=sources.SHEETS[kind];s.append(sources.HEADERS[kind]);identity=[f['member_code'],f['name']]
    if kind=='CLASS_MEETING':s.append(identity+(values or ['迟到','正常到场','早退']))
    elif kind=='DAILY_READING':s.append(identity+(values or [date.fromisoformat(DAY),'是']))
    else:s.append(identity+(values or [date.fromisoformat(DAY),'share-test-1','测试优秀分享','是']))
    out=io.BytesIO();wb.save(out);return out.getvalue()


def test_312_reading_rows_save_in_bulk_and_retry_without_extra_facts(source_fixture,monkeypatch,record_property):
    import sys
    from app import db
    f=source_fixture;identities=[(f['member_code'],f['name'])]
    with transaction() as c:
        now=credits._db_timestamp(c)
        for i in range(25):
            code=f['member_code']+'-bulk-'+str(i);name='批量合成学员'+str(i)
            mid=execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE',?,?)",(code,name,f['class_id'],now,now)).lastrowid
            execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,'STUDY_CLASS',1,'2020-01-01',?,?)",(mid,f['class_id'],now,now))
            identities.append((code,name))
    days={f'2026-10-{day:02d}':'NORMAL_WORKDAY' for day in range(1,13)}
    _calendar(f,year=2026,day_types=days,label='bulk-'+f['suffix'])
    wb=Workbook();sheet=wb.active;sheet.title=sources.SHEETS['DAILY_READING'];sheet.append(sources.HEADERS['DAILY_READING'])
    for code,name in identities:
        for day in days:sheet.append([code,name,day,'是'])
    out=io.BytesIO();wb.save(out);data=out.getvalue()
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}','',f['actor'])
    assert preview['row_count']==312 and preview['error_count']==0
    calls=[];original=db.execute
    def counted(c,sql,params=()):
        calls.append(sql.split()[0])
        return original(c,sql,params)
    # Count imported aliases too, including repeated IAM reads. No production DB.
    with monkeypatch.context() as scoped:
        for module_name,module in list(sys.modules.items()):
            if module_name.startswith('app.') and getattr(module,'execute',None) is original:
                scoped.setattr(module,'execute',counted)
        saved=sources.register(data,'DAILY_READING',f['class_id'],'{}','','合成312行.xlsx',preview['fingerprint'],f['actor'])
    record_property('source_save_sql_calls',len(calls))
    assert len(calls)<1000, f'source save still executes {len(calls)} SQL calls'
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_source_rows WHERE import_id=?',(saved['id'],))['n']==312
    assert fetch_one("SELECT COUNT(*) AS n FROM learning_credit_activity_facts WHERE class_org_unit_id=?",(f['class_id'],))['n']==312
    assert fetch_one("SELECT COUNT(*) AS n FROM audit_logs WHERE action='learning.activity_fact.create' AND org_unit_id=?",(f['class_id'],))['n']==312
    assert sources.find_saved_source(preview['fingerprint'],f['class_id'],'DAILY_READING',f['actor'])=={'id':saved['id']}
    retry=sources.register(data,'DAILY_READING',f['class_id'],'{}','补充','改名.xlsx',preview['fingerprint'],f['actor'])
    assert retry=={'id':saved['id'],'idempotent':True}
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_source_rows WHERE import_id=?',(saved['id'],))['n']==312
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE member_id=?',(f['member_id'],))['n']==0


def test_bulk_source_audit_failure_rolls_back_entire_receipt_and_retry(source_fixture):
    f=source_fixture;data=workbook(f,'DAILY_READING')
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}','',f['actor'])
    with patch.object(activity,'write_audit',side_effect=RuntimeError('isolated audit failure')),pytest.raises(RuntimeError):
        sources.register(data,'DAILY_READING',f['class_id'],'{}','','test.xlsx',preview['fingerprint'],f['actor'])
    assert sources.find_saved_source(preview['fingerprint'],f['class_id'],'DAILY_READING',f['actor'])=={'id':None}
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_activity_facts WHERE class_org_unit_id=?',(f['class_id'],))['n']==0
    saved=sources.register(data,'DAILY_READING',f['class_id'],'{}','','test.xlsx',preview['fingerprint'],f['actor'])
    assert saved['id'] and not saved['idempotent']


def test_saved_source_lookup_and_bulk_writer_enforce_scope_and_permissions(source_fixture):
    f=source_fixture
    with patch.object(sources,'_has_scope',return_value=False),pytest.raises(PermissionError):
        sources.find_saved_source('a'*64,f['class_id'],'DAILY_READING',f['actor'])
    with patch.object(activity,'user_context',return_value={'permissions':[]}),transaction() as c,pytest.raises(PermissionError):
        activity.record_learning_activity_facts(actor_user_id=f['actor'],class_org_unit_id=f['class_id'],facts=[{}],_connection=c)
    with patch.object(activity,'_scope_allows',return_value=False),transaction() as c,pytest.raises(PermissionError):
        activity.record_learning_activity_facts(actor_user_id=f['actor'],class_org_unit_id=f['class_id'],facts=[{}],_connection=c)


@pytest.mark.parametrize('bad',[
    {'metadata':{'points':99}}, {'binding_id':-1}, {'source_type':'HQ_READING_EXPORT'},
    {'participation_status':'WRONG'}, {'member_id':-1},
])
def test_bulk_fact_validation_fails_before_any_insert(source_fixture,bad):
    f=source_fixture
    good=dict(activity_type='DAILY_READING',member_id=f['member_id'],occurred_on=DAY,source_type='BULK_TEST',source_id='good',binding_id=f['binding_id'])
    with transaction() as c,pytest.raises(credits.LearningCreditError):
        activity.record_learning_activity_facts(actor_user_id=f['actor'],class_org_unit_id=f['class_id'],facts=[good,{**good,'source_id':'bad',**bad}],_connection=c)
        pytest.fail('invalid batch accepted')
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_activity_facts WHERE class_org_unit_id=?',(f['class_id'],))['n']==0


def test_preview_context_is_reloaded_when_saving_and_lookup_is_authorized_http(source_fixture):
    from app.api.auth import current_user
    f=source_fixture;data=workbook(f,'DAILY_READING')
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}','',f['actor'])
    with transaction() as c:
        execute(c,"UPDATE member_org_relations SET valid_until='2026-10-01' WHERE member_id=? AND org_unit_id=?",(f['member_id'],f['class_id']))
    with pytest.raises(credits.LearningCreditError,match='已变化'):
        sources.register(data,'DAILY_READING',f['class_id'],'{}','','test.xlsx',preview['fingerprint'],f['actor'])
    app.dependency_overrides[current_user]=lambda:{'id':f['actor'],'permissions':credits.user_context(f['actor'])['permissions']}
    try:
        with TestClient(app) as client:
            params={'fingerprint':preview['fingerprint'],'class_id':f['class_id'],'kind':'DAILY_READING'}
            assert client.get('/api/v1/learning-credits/excel/saved-source',params=params).json()['data']=={'id':None}
            with patch.object(sources,'_has_scope',return_value=False):
                assert client.get('/api/v1/learning-credits/excel/saved-source',params=params).status_code==403
    finally:app.dependency_overrides.pop(current_user,None)


@pytest.mark.parametrize('kind',['CLASS_MEETING','DAILY_READING','EXCELLENT_SHARE'])
def test_excel_source_direct_confirmation_posts_once_without_fabricated_review(source_fixture,kind):
    f=source_fixture;iid,_,_=registered(f,kind)
    source=sources.create_batches(iid,f['actor'])
    assert source['member_count']==1 and source['member_results'][0]['member_name']==f['name']
    assert source['batches'][0]['period_start'] and source['batches'][0]['member_count']==1
    assert source['settlement_summary']['entry_count']>0
    posted=sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'])
    assert all(b['status']=='POSTED' for b in posted['batches'])
    row=fetch_one('SELECT approved_by,approved_at,approval_fingerprint FROM learning_credit_settlement_batches WHERE id=?',(source['batches'][0]['id'],))
    assert all(v is None for v in row.values())
    assert fetch_one("SELECT COUNT(*) AS n FROM audit_logs WHERE action='learning_credit.source.confirm' AND resource_id=?",(str(source['batches'][0]['id']),))['n']==1
    assert sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'])['settlement_summary']==posted['settlement_summary']


def test_direct_confirmation_requires_current_preview_and_post_scope(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING');source=sources.create_batches(iid,f['actor'])
    with pytest.raises(credits.LearningCreditError,match='预览已变化'):
        sources.process(iid,f['actor'],'confirm','0'*64)
    with patch.object(sources,'_has_scope',return_value=False),pytest.raises(PermissionError):
        sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'])

    with pytest.raises(credits.LearningCreditError):
        batches.post_daily_reading_batch(actor_user_id=f['actor'],batch_id=source['batches'][0]['id'])
    with transaction() as c:
        execute(c,"UPDATE learning_credit_activity_facts SET participation_status='CANCELLED' WHERE member_id=? AND source_type='CREDIT_EXCEL_DAILY_READING'",(f['member_id'],))
    with pytest.raises(credits.LearningCreditError):
        sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'])


def test_source_confirmation_failure_requires_explicit_resume_without_review(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING')
    source=sources.create_batches(iid,f['actor'])
    with patch.object(credits,'post_credit_entry',side_effect=RuntimeError('isolated failure')):
        with pytest.raises(credits.LearningCreditError,match='未全部完成'):
            sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'])
    failed=sources.get_import(iid,f['actor'])
    assert failed['batches'][0]['status']=='PARTIAL_FAILED'
    with pytest.raises(credits.LearningCreditError,match='明确请求续跑'):
        sources.process(iid,f['actor'],'confirm',failed['settlement_fingerprint'])
    completed=sources.process(iid,f['actor'],'confirm',failed['settlement_fingerprint'],True)
    assert completed['batches'][0]['status']=='POSTED'
    assert sources.create_batches(iid,f['actor'])['batches'][0]['id']==source['batches'][0]['id']


def test_source_register_waits_for_competing_operation_and_retry_is_idempotent(source_fixture):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from app.services import credit_opening_balances as opening
    f=source_fixture;data=workbook(f,'DAILY_READING')
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    locked=Event();release=Event()
    def competing():
        with transaction() as c,opening._operation(c):
            locked.set();release.wait(3)
    with ThreadPoolExecutor(max_workers=2) as pool:
        holder=pool.submit(competing);assert locked.wait(2)
        saved=pool.submit(sources.register,data,'DAILY_READING',f['class_id'],'{}',NOTE,'测试.xlsx',preview['fingerprint'],f['actor'])
        release.set();result=saved.result(timeout=10);holder.result()
    retry=sources.register(data,'DAILY_READING',f['class_id'],'{}',NOTE,'测试.xlsx',preview['fingerprint'],f['actor'])
    assert result['id']==retry['id'] and retry['idempotent']


def test_same_reading_file_with_changed_note_reuses_receipt_and_ledger(source_fixture):
    f=source_fixture;iid,_,data=registered(f,'DAILY_READING')
    note='重新核对原表，重试保存来源'
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}',note,f['actor'])
    assert preview['fingerprint']==sources.preview(data,'DAILY_READING',f['class_id'],'{}','',f['actor'])['fingerprint']
    retry=sources.register(data,'DAILY_READING',f['class_id'],'{}',note,'改名后.xlsx',preview['fingerprint'],f['actor'])
    assert retry=={'id':iid,'idempotent':True}
    source=sources.create_batches(iid,f['actor'])
    before=fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']
    done=sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    assert done['posting_progress']=={'state':'COMPLETE','total_entries':1,'posted_entries':1,'remaining_entries':0,'completed':True}
    again=sources.process(iid,f['actor'],'confirm',done['settlement_fingerprint'],incremental=True)
    assert again['posting_progress']==done['posting_progress']
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before+1


@pytest.mark.parametrize('note',[None,'','补'])
def test_optional_note_is_not_required_by_upload_api_or_confirmation(source_fixture,note):
    from app.api.auth import current_user
    f=source_fixture;data=workbook(f,'DAILY_READING')
    form={'kind':'DAILY_READING','class_id':f['class_id']}
    if note is not None:form['source_note']=note
    app.dependency_overrides[current_user]=lambda:{'id':f['actor'],'permissions':credits.user_context(f['actor'])['permissions']}
    try:
        with TestClient(app) as client:
            response=client.post('/api/v1/learning-credits/excel/preview',data=form,files={'workbook':('测试.xlsx',data)})
            assert response.status_code==200,response.text
            fingerprint=response.json()['data']['fingerprint']
            response=client.post('/api/v1/learning-credits/excel/imports',data={**form,'expected_fingerprint':fingerprint},files={'workbook':('测试.xlsx',data)})
            assert response.status_code==200,response.text
            iid=response.json()['data']['id']
            source=client.post(f'/api/v1/learning-credits/excel/imports/{iid}/batches').json()['data']
            response=client.post(f'/api/v1/learning-credits/excel/imports/{iid}/actions/confirm',json={'expected_fingerprint':source['settlement_fingerprint'],'incremental':True})
            assert response.status_code==200,response.text
            assert response.json()['data']['posting_progress']['completed']
    finally:app.dependency_overrides.pop(current_user,None)


def test_incremental_reading_commits_short_steps_and_resumes_without_duplicates(source_fixture):
    f=source_fixture
    identities=[(f['member_code'],f['name'])]
    with transaction() as c:
        now=credits._db_timestamp(c)
        for index in range(31):
            code=f['member_code']+'-step-'+str(index);name='分段测试'+str(index)
            mid=execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE',?,?)",(code,name,f['class_id'],now,now)).lastrowid
            execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,'STUDY_CLASS',1,'2020-01-01',?,?)",(mid,f['class_id'],now,now))
            identities.append((code,name))
    wb=Workbook();sheet=wb.active;sheet.title=sources.SHEETS['DAILY_READING'];sheet.append(sources.HEADERS['DAILY_READING'])
    for code,name in identities:sheet.append([code,name,DAY,'是'])
    out=io.BytesIO();wb.save(out)
    iid,_,_=registered(f,'DAILY_READING',out.getvalue());source=sources.create_batches(iid,f['actor'])
    counts=[];before=fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']
    while not source['posting_progress']['completed']:
        source=sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
        counts.append(source['posting_progress']['posted_entries'])
        assert source['batches'][0]['status']!='POSTING'
        assert counts[-1]-(counts[-2] if len(counts)>1 else 0)<=25
        assert len(counts)<=32
        if len(counts)==1:
            regenerated=sources.create_batches(iid,f['actor'])
            assert regenerated['batches'][0]['id']==source['batches'][0]['id']
            assert regenerated['posting_progress']['posted_entries']==counts[-1]
    assert counts[-1]==32 and len(counts)>=2
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before+32
    assert sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)['posting_progress']['posted_entries']==32
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before+32


def test_249_reading_entries_use_bounded_steps_without_per_item_batch_preview(source_fixture,monkeypatch,record_property):
    import sys,time
    from app import db
    f=source_fixture;identities=[(f['member_code'],f['name'])]
    with transaction() as c:
        now=credits._db_timestamp(c)
        for i in range(33):
            code=f['member_code']+'-post-'+str(i);name='入账合成学员'+str(i)
            mid=execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE',?,?)",(code,name,f['class_id'],now,now)).lastrowid
            execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,'STUDY_CLASS',1,'2020-01-01',?,?)",(mid,f['class_id'],now,now))
            identities.append((code,name))
    days=[f'2026-10-{day:02d}' for day in range(8,16)]
    _calendar(f,year=2026,day_types={d:'NORMAL_WORKDAY' for d in days},label='post-'+f['suffix'])
    wb=Workbook();sheet=wb.active;sheet.title=sources.SHEETS['DAILY_READING'];sheet.append(sources.HEADERS['DAILY_READING'])
    for i,(code,name) in enumerate(identities):
        for day in days:sheet.append([code,name,day,'否' if day==days[0] and i<23 else '是'])
    out=io.BytesIO();wb.save(out)
    iid,_,_=registered(f,'DAILY_READING',out.getvalue());source=sources.create_batches(iid,f['actor'])
    assert source['settlement_summary']['proposed_points']=='249.00'
    calls=[];original=db.execute;requests=0;started=time.monotonic()
    ledger_inserts=[]
    def counted(c,sql,params=()):
        calls.append(sql.split()[0])
        if sql.startswith('INSERT INTO learning_credit_entries '):ledger_inserts.append(sql)
        return original(c,sql,params)
    with monkeypatch.context() as scoped,patch.object(batches,'_assert_current_preview',wraps=batches._assert_current_preview) as live:
        for module_name,module in list(sys.modules.items()):
            if module_name.startswith('app.') and getattr(module,'execute',None) is original:scoped.setattr(module,'execute',counted)
        while not source['posting_progress']['completed']:
            previous=source['posting_progress']['posted_entries']
            source=sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True);requests+=1
            assert 0<source['posting_progress']['posted_entries']-previous<=25
        record_property('posting_sql_calls',len(calls));record_property('posting_requests',requests)
        record_property('posting_full_preview_calls',live.call_count);record_property('posting_seconds',round(time.monotonic()-started,3))
        assert live.call_count==requests*2
    assert requests<=16 and len(calls)<5000
    assert len(ledger_inserts)==requests  # one multi-row ledger INSERT per step
    assert source['settlement_summary']['posted_points']=='249.00'
    assert fetch_one("SELECT COUNT(*) AS n FROM learning_credit_entries WHERE class_org_unit_id=?",(f['class_id'],))['n']==249
    assert sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)['settlement_summary']['posted_points']=='249.00'


def test_reading_step_rolls_back_all_items_on_mid_step_failure(source_fixture):
    f=source_fixture;identities=[(f['member_code'],f['name'])]
    with transaction() as c:
        now=credits._db_timestamp(c)
        for i in range(2):
            code=f['member_code']+'-rollback-'+str(i);name='回滚合成'+str(i)
            mid=execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE',?,?)",(code,name,f['class_id'],now,now)).lastrowid
            execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,'STUDY_CLASS',1,'2020-01-01',?,?)",(mid,f['class_id'],now,now))
            identities.append((code,name))
    wb=Workbook();s=wb.active;s.title=sources.SHEETS['DAILY_READING'];s.append(sources.HEADERS['DAILY_READING'])
    for code,name in identities:s.append([code,name,DAY,'是'])
    out=io.BytesIO();wb.save(out);iid,_,_=registered(f,'DAILY_READING',out.getvalue());source=sources.create_batches(iid,f['actor'])
    original=credits.execute
    def fail_after_insert(connection, sql, params=()):
        result=original(connection, sql, params)
        if sql.startswith('INSERT INTO learning_credit_entries '):
            raise RuntimeError('isolated failure after multi-row insert')
        return result
    with patch.object(credits,'execute',side_effect=fail_after_insert),pytest.raises(credits.LearningCreditError,match='未全部完成'):
        sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    failed=sources.get_import(iid,f['actor']);assert failed['posting_progress']['posted_entries']==0
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE class_org_unit_id=?',(f['class_id'],))['n']==0
    resumed=sources.process(iid,f['actor'],'confirm',failed['settlement_fingerprint'],True,True)
    assert resumed['posting_progress']['completed'] and resumed['posting_progress']['posted_entries']==3


def test_mysql_reading_step_concurrent_confirmations_never_duplicate(source_fixture):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    import os
    if not os.environ.get('DATABASE_URL','').startswith('mysql+pymysql://'):
        pytest.skip('concurrent row locking requires isolated MySQL')
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING');source=sources.create_batches(iid,f['actor'])
    barrier=Barrier(2)
    def confirm():
        barrier.wait(timeout=10)
        try:return sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
        except credits.LearningCreditError:return None
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures=[pool.submit(confirm) for _ in range(2)]
        results=[future.result(timeout=30) for future in futures]
    assert any(result and result['posting_progress']['completed'] for result in results)
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE idempotency_key=?',(f"DAILY_READING:{f['member_id']}:{DAY}",))['n']==1
    assert sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)['posting_progress']['posted_entries']==1


def test_abandoned_source_posting_shows_live_progress_and_uses_existing_recovery_gate(source_fixture):
    from datetime import UTC,datetime,timedelta
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING');source=sources.create_batches(iid,f['actor'])
    real_insert=credits.post_reading_entries
    class Interrupted(BaseException):pass
    def post_then_interrupt(**kwargs):
        real_insert(**kwargs)
        raise Interrupted()
    with patch.object(credits,'post_reading_entries',side_effect=post_then_interrupt),pytest.raises(Interrupted):
        sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    # The item and ledger write share one atomic transaction; even a process
    # interruption cannot leave an unlinked ledger fact to be posted twice.
    stopped=sources.get_import(iid,f['actor'])
    assert stopped['posting_progress']['state']=='POSTING'
    assert stopped['posting_progress']['posted_entries']==0
    with pytest.raises(credits.LearningCreditError,match='30分钟'):
        sources.reconcile(iid,f['actor'])
    with transaction() as c:
        execute(c,'UPDATE learning_credit_settlement_batches SET updated_at=? WHERE id=?',(credits._db_timestamp(c,datetime.now(UTC)-timedelta(minutes=31)),source['batches'][0]['id']))
    permissions=credits.user_context(f['actor'])['permissions']
    with patch.object(credits,'user_context',return_value={'permissions':[p for p in permissions if p!='plans:credit_settlement_reconcile']}),pytest.raises(PermissionError):
        sources.reconcile(iid,f['actor'])
    resumed=sources.reconcile(iid,f['actor'])
    assert resumed['posting_progress']['state']=='PARTIAL_FAILED'
    completed=sources.process(iid,f['actor'],'confirm',resumed['settlement_fingerprint'],True,True)
    assert completed['posting_progress']['completed']


def test_changed_workbook_with_overlapping_days_posts_only_new_member_day(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING');source=sources.create_batches(iid,f['actor'])
    before=fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']
    sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    wb=Workbook();sheet=wb.active;sheet.title=sources.SHEETS['DAILY_READING'];sheet.append(sources.HEADERS['DAILY_READING'])
    for day in [DAY,'2026-10-10']:sheet.append([f['member_code'],f['name'],day,'是'])
    out=io.BytesIO();wb.save(out)
    second,_,_=registered(f,'DAILY_READING',out.getvalue());source=sources.create_batches(second,f['actor'])
    while not source['posting_progress']['completed']:
        source=sources.process(second,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before+2
    assert fetch_one("SELECT COUNT(*) AS n FROM learning_credit_entries WHERE idempotency_key=?",(f"DAILY_READING:{f['member_id']}:{DAY}",))['n']==1


def test_legacy_identical_receipts_share_one_list_row_but_keep_audit_records(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING');sources.create_batches(iid,f['actor'])
    with transaction() as c:
        original=dict(execute(c,'SELECT * FROM learning_credit_source_imports WHERE id=?',(iid,)).fetchone())
        snapshot=json.loads(original['source_snapshot_json']);snapshot.pop('fingerprint_version',None)
        execute(c,'UPDATE learning_credit_source_imports SET fingerprint=?,source_snapshot_json=? WHERE id=?',(sources._hash(snapshot),sources._json(snapshot),iid))
        snapshot['note']='原件已再次核对确认'
        original.update(source_note=snapshot['note'],source_snapshot_json=sources._json(snapshot),fingerprint=sources._hash(snapshot))
        columns=[key for key in original if key!='id']
        duplicate=execute(c,'INSERT INTO learning_credit_source_imports('+','.join(columns)+') VALUES ('+','.join('?' for _ in columns)+')',tuple(original[key] for key in columns)).lastrowid
        execute(c,'INSERT INTO learning_credit_source_rows(import_id,member_id,excel_row,occurred_on,source_reference,record_json,fact_id) SELECT ?,member_id,excel_row,occurred_on,source_reference,record_json,fact_id FROM learning_credit_source_rows WHERE import_id=?',(duplicate,iid))
        execute(c,'INSERT INTO learning_credit_source_batches(import_id,batch_id) SELECT ?,batch_id FROM learning_credit_source_batches WHERE import_id=?',(duplicate,iid))
    rows=[r for r in sources.workbench(f['actor'])['imports'] if r['class_org_unit_id']==f['class_id']]
    assert len(rows)==1 and rows[0]['duplicate_upload_count']==2
    assert sources.get_import(iid,f['actor'])['row_count']==1
    assert sources.get_import(duplicate,f['actor'])['row_count']==1


def test_receipt_validates_facts_in_one_bulk_read(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING')
    statements=[];real_execute=sources.execute
    def record(c,sql,params=()):statements.append(sql);return real_execute(c,sql,params)
    with patch.object(sources,'execute',side_effect=record):sources.get_import(iid,f['actor'])
    assert len([q for q in statements if 'SELECT * FROM learning_credit_activity_facts' in q])==1


def test_class_choices_share_management_canonical_ids_and_keep_scoped_historical_aliases(source_fixture):
    from app.services.learning_org_catalog import class_catalog, selected_class_ids
    from app.services.organization_management import list_learning_org_units
    from app.services.study_meetings import study_meeting_filter_options
    f=source_fixture
    original=fetch_one('SELECT * FROM org_units WHERE id=?',(f['class_id'],))
    duplicate=f['class_id']+'-duplicate';inactive=f['class_id']+'-inactive'
    foreign=f['class_id']+'-other-center'
    with transaction() as c:
        for oid,parent,active in [(duplicate,original['parent_id'],1),(inactive,original['parent_id'],0),(foreign,'org-suzhou',1)]:
            # Ensure the same name in a distinct parent is never merged.
            if oid==foreign and parent==original['parent_id']:
                parent='org-wuxi'
            execute(c,'INSERT INTO org_units(id,unit_code,name,unit_type,parent_id,is_active,created_at,updated_at) VALUES (?,?,?,\'CLASS\',?,?,?,?)',
                (oid,oid,original['name'],parent,active,'2099-01-01','2099-01-01'))
        catalog=class_catalog(c)
        chosen=next(r for r in catalog if f['class_id'] in r['related_class_ids'])
        assert chosen['id']==f['class_id']
        assert selected_class_ids(catalog,f['class_id'])=={f['class_id'],duplicate,inactive}
        limited=class_catalog(c,{duplicate})
        assert limited[0]['id']==duplicate and limited[0]['related_class_ids']==[duplicate]
    for options in [exports.export_options(f['actor']),sources.workbench(f['actor']),
                    study_meeting_filter_options(actor_user_id=f['actor']),list_learning_org_units(f['actor'])]:
        ids={r['id'] for r in options['classes']}
        assert f['class_id'] in ids and duplicate not in ids and inactive not in ids


def options(f,kind):return json.dumps({'cycle_id':f['cycle_id'],'actual_start':DAY+'T09:00:00+08:00'}) if kind=='CLASS_MEETING' else '{}'


def registered(f,kind,data=None):
    data=data or workbook(f,kind)
    p=sources.preview(data,kind,f['class_id'],options(f,kind),NOTE,f['actor'])
    assert p['error_count']==0,p['rows']
    receipt=sources.register(data,kind,f['class_id'],options(f,kind),NOTE,'测试原文件.xlsx',p['fingerprint'],f['actor'])
    return receipt['id'],p,data


@pytest.mark.parametrize('kind',['CLASS_MEETING','DAILY_READING','EXCELLENT_SHARE'])
def test_sources_use_existing_rules_post_once_and_trace_file_rows(source_fixture,kind):
    f=source_fixture;iid,p,data=registered(f,kind)
    before=fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']
    state=sources.create_batches(iid,f['actor'])
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before
    assert len(state['batches'])==1 and DecimalText(state['batches'][0]['proposed_points']) in ('16.00','1.00','2.00')
    sources.process(iid,f['actor'],'submit')
    with pytest.raises(PermissionError,match='审批'):sources.process(iid,f['actor'],'approve')
    sources.process(iid,f['reviewer'],'approve');posted=sources.process(iid,f['actor'],'post')
    assert posted['batches'][0]['status']=='POSTED'
    assert sources.process(iid,f['actor'],'post')['batches'][0]['status']=='POSTED'
    assert sources.register(data,kind,f['class_id'],options(f,kind),NOTE,'重命名.xlsx',p['fingerprint'],f['actor'])['idempotent']
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries')['n']==before+1
    entries=credits.list_credit_entries(actor_user_id=f['actor'],member_id=f['member_id'])
    assert entries[0]['import_sources'][0]['original_filename']=='测试原文件.xlsx'
    assert entries[0]['import_sources'][0]['excel_row']==2
    if kind=='CLASS_MEETING':
        cycle=fetch_one('SELECT class_meeting_status,source_event_group_id FROM class_learning_cycles WHERE id=?',(f['cycle_id'],))
        assert cycle['class_meeting_status']=='HELD' and cycle['source_event_group_id']==posted['event_group_id']
        c=connect();records=execute(c,'SELECT r.* FROM attendance_records r JOIN attendance_sessions s ON s.id=r.attendance_session_id WHERE s.event_group_id=?',(posted['event_group_id'],)).fetchall();c.close()
        assert all(r['checked_at'] is None for r in records)
        attendance_scoring.recalculate_event_group(posted['event_group_id'])
        scores=attendance_scoring.member_scores(f['member_id'])
        assert sum(float(r['final_points']) for r in scores)==16
    result=exports.export_credits(f['actor'],f['class_id'],DAY,DAY)
    book=load_workbook(io.BytesIO(result),data_only=False)
    assert book.sheetnames==['学分汇总','新增学分明细','已导入历史年度']
    assert book['新增学分明细']['K2'].data_type=='s' and '测试原文件.xlsx' in book['新增学分明细']['K2'].value
    assert book['新增学分明细']['G2'].data_type=='n'
    assert book['新增学分明细']['F2'].is_date


def DecimalText(value):
    from decimal import Decimal
    return format(Decimal(str(value)),'.2f')


def test_bulk_export_sources_match_single_lookup_without_per_entry_queries(source_fixture):
    f=source_fixture
    iid,_,_=registered(f,'DAILY_READING')
    sources.create_batches(iid,f['actor']); sources.process(iid,f['actor'],'submit')
    sources.process(iid,f['reviewer'],'approve'); sources.process(iid,f['actor'],'post')
    c=connect()
    try:
        original=dict(execute(c,'SELECT * FROM learning_credit_entries WHERE member_id=?',(f['member_id'],)).fetchone())
        expected=exports.entry_sources(c,original)
        assert expected
        entries=[{**original,'id':original['id']+10000+i} for i in range(2000)]
        with patch.object(exports,'execute',wraps=execute) as query, patch.object(sources,'available',wraps=sources.available) as availability:
            result=exports.entry_sources_batch(c,entries)
            assert query.call_count < 10 and availability.call_count == 1
        assert len(result)==2000 and all(value==expected for value in result.values())
        bases=exports.opening.opening_totals_by_member(c,allowed={f['class_id']})
        assert bases.get(f['member_id'],0)==__import__('decimal').Decimal(exports.opening.opening_summary(c,member_id=f['member_id'],allowed={f['class_id']})['total_points'])
    finally:c.close()


def test_export_can_target_one_member_without_other_roster_or_ledger(source_fixture):
    f = source_fixture
    iid, _, _ = registered(f, 'DAILY_READING')
    sources.create_batches(iid, f['actor'])
    sources.process(iid, f['actor'], 'submit')
    sources.process(iid, f['reviewer'], 'approve')
    sources.process(iid, f['actor'], 'post')
    with transaction() as c:
        now = credits._db_timestamp(c)
        other = execute(c, "INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?, '另一位合成学员', ?, 'ACTIVE', ?, ?)",
                        ('OTHER-' + f['suffix'], f['class_id'], now, now)).lastrowid
        execute(c, "INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,created_at,updated_at) VALUES (?, ?, 'STUDY_CLASS', 1, ?, ?)",
                (other, f['class_id'], now, now))
    book = load_workbook(io.BytesIO(exports.export_credits(f['actor'], f['class_id'], DAY, DAY, member_id=f['member_id'])))
    assert book['学分汇总'].max_row == 2
    assert book['学分汇总']['A2'].value == f['member_code']
    assert book['新增学分明细'].max_row == 2
    assert book['新增学分明细']['B2'].value == f['member_code']
    assert '另一位合成学员' not in str([list(sheet.values) for sheet in book])
    with pytest.raises(PermissionError):
        exports.export_credits(f['actor'], f['class_id'], DAY, DAY, member_id=999999999)
    with TestClient(app) as client:
        assert client.get('/api/v1/learning-credits/excel/export', params={'member_id': f['member_id']}).status_code == 401


def test_offline_cycle_advances_only_after_successful_settlement(source_fixture,monkeypatch):
    f=source_fixture
    monkeypatch.setenv('LEARNING_CYCLE_MONTHLY_REFRESH_ENABLED','true')
    iid,_,_=registered(f,'CLASS_MEETING')
    def current():
        return fetch_one("SELECT * FROM class_learning_cycles WHERE binding_id=? AND cycle_status='OPEN'",(f['binding_id'],))
    before=current()
    assert before['id']==f['cycle_id'] and before['actual_class_meeting_at'] is None
    batch=sources.create_batches(iid,f['actor'])['batches'][0]
    assert current()['id']==before['id'] and current()['actual_class_meeting_at'] is None
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve')
    assert current()['id']==before['id']
    # The normal batch API must have the same side effect as the Excel UI.
    assert batches.post_class_meeting_batch(actor_user_id=f['actor'],batch_id=batch['id'])['status']=='POSTED'
    after=current()
    assert after['learning_cycle_index']==2
    assert parse_utc_datetime(after['opened_at']).isoformat().startswith('2026-10-09T01:00:00')
    old=fetch_one('SELECT * FROM class_learning_cycles WHERE id=?',(before['id'],))
    assert old['cycle_status']=='CLOSED' and old['class_meeting_status']=='HELD'
    assert old['source_event_group_id']==sources.get_import(iid,f['actor'])['event_group_id']
    entry=fetch_one("SELECT learning_cycle_id FROM learning_credit_entries WHERE member_id=? AND credit_type='CLASS_MEETING_SCORE'",(f['member_id'],))
    assert entry['learning_cycle_id']==before['id']
    assert batches.post_class_meeting_batch(actor_user_id=f['actor'],batch_id=batch['id'])['idempotent']
    assert current()==after


def test_failed_settlement_preserves_cycle_then_explicit_retry_advances_once(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'CLASS_MEETING')
    batch=sources.create_batches(iid,f['actor'])['batches'][0]
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve')
    with patch.object(credits,'post_credit_entry',side_effect=RuntimeError('test failure')):
        result=batches.post_class_meeting_batch(actor_user_id=f['actor'],batch_id=batch['id'])
    assert result['status']=='PARTIAL_FAILED'
    old=fetch_one('SELECT * FROM class_learning_cycles WHERE id=?',(f['cycle_id'],))
    assert old['cycle_status']=='OPEN' and old['actual_class_meeting_at'] is None
    result=batches.post_class_meeting_batch(actor_user_id=f['actor'],batch_id=batch['id'],resume_partial_failure=True)
    assert result['status']=='POSTED'
    assert fetch_one("SELECT learning_cycle_index FROM class_learning_cycles WHERE binding_id=? AND cycle_status='OPEN'",(f['binding_id'],))['learning_cycle_index']==2


def test_last_cycle_settlement_completes_round_without_cycle_37(source_fixture):
    f=source_fixture
    with transaction() as c:
        plan=execute(c,'SELECT id FROM learning_plan_cycles WHERE plan_version_id=? AND cycle_index=36',(f['plan_id'],)).fetchone()
        execute(c,'UPDATE class_learning_cycles SET learning_cycle_index=36,plan_cycle_id=? WHERE id=?',(plan['id'],f['cycle_id']))
    iid,_,_=registered(f,'CLASS_MEETING');sources.create_batches(iid,f['actor'])
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    assert fetch_one('SELECT status FROM class_learning_bindings WHERE id=?',(f['binding_id'],))['status']=='COMPLETED'
    assert not fetch_one("SELECT id FROM class_learning_cycles WHERE binding_id=? AND cycle_status='OPEN'",(f['binding_id'],))


def test_cycle_failure_rolls_back_batch_terminal_state_and_can_reconcile(source_fixture):
    from datetime import UTC,datetime,timedelta
    from app.services import learning_cycles
    f=source_fixture;iid,_,_=registered(f,'CLASS_MEETING')
    batch=sources.create_batches(iid,f['actor'])['batches'][0]
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve')
    with patch.object(learning_cycles,'_plan_cycle_for_track',return_value=None),pytest.raises(credits.LearningCreditError,match='缺少'):
        batches.post_class_meeting_batch(actor_user_id=f['actor'],batch_id=batch['id'])
    assert fetch_one('SELECT status FROM learning_credit_settlement_batches WHERE id=?',(batch['id'],))['status']=='POSTING'
    assert fetch_one('SELECT cycle_status,actual_class_meeting_at FROM class_learning_cycles WHERE id=?',(f['cycle_id'],))=={'cycle_status':'OPEN','actual_class_meeting_at':None}
    with transaction() as c:
        execute(c,'UPDATE learning_credit_settlement_batches SET updated_at=? WHERE id=?',(credits._db_timestamp(c,datetime.now(UTC)-timedelta(minutes=31)),batch['id']))
    result=batches.reconcile_stale_posting_batch(actor_user_id=f['actor'],batch_id=batch['id'])
    assert result['status']=='POSTED'
    assert fetch_one("SELECT learning_cycle_index FROM class_learning_cycles WHERE binding_id=? AND cycle_status='OPEN'",(f['binding_id'],))['learning_cycle_index']==2


def test_reading_weekend_incomplete_and_invalid_rows_have_actionable_results(source_fixture):
    f=source_fixture
    p=sources.preview(workbook(f,'DAILY_READING',[date(2026,10,11),'是']),'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    assert p['error_count']==0 and p['proposed_points']=='0.00' and p['rows'][0]['credit_status']=='NO_CREDIT'
    p=sources.preview(workbook(f,'DAILY_READING',[date.fromisoformat(DAY),'否']),'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    assert p['proposed_points']=='0.00'
    p=sources.preview(workbook(f,'EXCELLENT_SHARE',[date.fromisoformat(DAY),'ref','普通分享','否']),'EXCELLENT_SHARE',f['class_id'],'{}',NOTE,f['actor'])
    assert p['error_count']==1 and p['fingerprint'] is None
    bad=workbook(f,'CLASS_MEETING',['正常到场',None,'缺席'])
    assert sources.preview(bad,'CLASS_MEETING',f['class_id'],options(f,'CLASS_MEETING'),NOTE,f['actor'])['error_count']==1
    assert TestClient(app).get('/api/v1/learning-credits/excel').status_code==401


def test_source_scope_permission_and_frozen_row_tampering_fail_before_credits(source_fixture):
    f=source_fixture;iid,p,data=registered(f,'DAILY_READING')
    with patch.object(credits,'_scope_allows',return_value=False):
        with pytest.raises(PermissionError):sources.get_import(iid,f['actor'])
        with pytest.raises(PermissionError):exports.export_credits(f['actor'],f['class_id'],DAY,DAY)
    with transaction() as c:execute(c,"UPDATE learning_credit_source_rows SET source_reference='tampered' WHERE import_id=?",(iid,))
    with pytest.raises(ValueError,match='来源'):sources.create_batches(iid,f['actor'])
    assert not fetch_one('SELECT batch_id FROM learning_credit_source_batches WHERE import_id=?',(iid,))


def test_monthly_share_cap_and_stable_reference_duplicates_are_preserved(source_fixture):
    f=source_fixture;wb=Workbook();s=wb.active;s.title=sources.SHEETS['EXCELLENT_SHARE'];s.append(sources.HEADERS['EXCELLENT_SHARE'])
    for index in range(7):s.append([f['member_code'],f['name'],date.fromisoformat(DAY),'reference-'+str(index),'测试分享','是'])
    raw=io.BytesIO();wb.save(raw)
    iid,p,data=registered(f,'EXCELLENT_SHARE',raw.getvalue())
    assert sum(r['credit_status']=='READY' for r in p['rows'])==5
    state=sources.create_batches(iid,f['actor'])
    assert state['batches'][0]['proposed_entry_count']==5
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    assert fetch_one("SELECT SUM(points) AS points FROM learning_credit_entries WHERE member_id=? AND credit_type='EXCELLENT_SHARE'",(f['member_id'],))['points']==5


def test_templates_prefill_exact_identity_and_preserve_validation(source_fixture):
    f=source_fixture
    for kind in sources.KINDS:
        book=load_workbook(io.BytesIO(sources.template(kind,f['class_id'],f['actor'])))
        s=book[sources.SHEETS[kind]]
        assert s['A2'].value==f['member_code'] and s['B2'].value==f['name']
        assert len(s.data_validations.dataValidation)>0 and s.freeze_panes=='C2'


def test_permission_specific_scope_cannot_be_widened_by_workbench_permission(source_fixture):
    f=source_fixture;raw=workbook(f,'DAILY_READING')
    p=sources.preview(raw,'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    def scopes(actor,permission=None):return [] if permission in {'plans:credit_activity_fact_manage','exports:normal'} else None
    with patch.object(credits,'accessible_org_ids',side_effect=scopes):
        with pytest.raises(PermissionError):sources.register(raw,'DAILY_READING',f['class_id'],'{}',NOTE,'scope.xlsx',p['fingerprint'],f['actor'])
        assert exports.export_options(f['actor'])['classes']==[]
        with pytest.raises(PermissionError):exports.export_credits(f['actor'],f['class_id'],DAY,DAY)
    assert not fetch_one('SELECT id FROM learning_credit_source_imports WHERE fingerprint=?',(p['fingerprint'],))


def test_filtered_reading_batches_exclude_unrelated_facts_and_reject_swapped_link(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING')
    unrelated=activity.record_learning_activity_fact(actor_user_id=f['actor'],activity_type='DAILY_READING',member_id=f['member_id'],class_org_unit_id=f['class_id'],occurred_on=DAY,source_type='OTHER_TEST',source_id='unrelated-source',participation_status='CONFIRMED',binding_id=f['binding_id'])
    result=sources.create_batches(iid,f['actor']);bid=result['batches'][0]['id']
    snapshot=json.loads(fetch_one('SELECT source_snapshot_json FROM learning_credit_settlement_batches WHERE id=?',(bid,))['source_snapshot_json'])
    assert unrelated['id'] not in snapshot['fact_ids_filter']
    broader=batches.dry_run_daily_reading_batch(actor_user_id=f['actor'],class_org_unit_id=f['class_id'],occurred_from=DAY,occurred_to=DAY)
    with transaction() as c:
        execute(c,'DELETE FROM learning_credit_source_batches WHERE import_id=?',(iid,))
        execute(c,'INSERT INTO learning_credit_source_batches(import_id,batch_id) VALUES (?,?)',(iid,broader['id']))
    with pytest.raises(ValueError,match='原文件'):sources.process(iid,f['actor'],'submit')


def test_source_date_and_formula_input_are_rejected_without_writing(source_fixture):
    f=source_fixture
    for values in ([date(2026,9,30),'是'],[date(2099,10,9),'是'],[date.fromisoformat(DAY),'=1+1']):
        p=sources.preview(workbook(f,'DAILY_READING',values),'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
        assert p['error_count']==1 and p['fingerprint'] is None


def test_posted_meeting_reupload_skips_and_refuses_mutating_absent_slot(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'CLASS_MEETING',workbook(f,'CLASS_MEETING',['正常到场','缺席','缺席']))
    sources.create_batches(iid,f['actor']);sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    p=sources.preview(workbook(f,'CLASS_MEETING',['正常到场','缺席','缺席']),'CLASS_MEETING',f['class_id'],options(f,'CLASS_MEETING'),NOTE,f['actor'])
    assert p['error_count']==0 and p['proposed_points']=='0.00'
    p=sources.preview(workbook(f,'CLASS_MEETING',['正常到场','正常到场','缺席']),'CLASS_MEETING',f['class_id'],options(f,'CLASS_MEETING'),NOTE,f['actor'])
    assert p['error_count']==1 and '已经入账' in ' '.join(p['rows'][0]['errors'])


def test_closed_month_class_confirmation_preserves_current_month(source_fixture,monkeypatch):
    f=source_fixture
    monkeypatch.setenv('LEARNING_CYCLE_MONTHLY_REFRESH_ENABLED','true')
    from app.services import learning_cycles,learning_cycle_monthly
    with transaction() as c:
        execute(c,"UPDATE class_learning_cycles SET cycle_status='CLOSED',closed_at='2026-10-10' WHERE id=?",(f['cycle_id'],))
    iid,_,_=registered(f,'CLASS_MEETING')
    sources.create_batches(iid,f['actor'])
    assert fetch_one('SELECT class_meeting_status FROM class_learning_cycles WHERE id=?',(f['cycle_id'],))['class_meeting_status']=='PLANNED'
    sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    cycle=fetch_one('SELECT cycle_status,closed_at,class_meeting_status FROM class_learning_cycles WHERE id=?',(f['cycle_id'],))
    assert cycle['cycle_status']=='CLOSED' and str(cycle['closed_at'])[:10]=='2026-10-10' and cycle['class_meeting_status']=='HELD'


def test_offline_early_leave_cancellation_keeps_source_and_updates_proposal(source_fixture):
    from app.api.attendance import create_adjudication,AdjudicationPayload
    f=source_fixture;iid,_,_=registered(f,'CLASS_MEETING')
    receipt=sources.get_import(iid,f['actor']);slot=next(r for r in receipt['rows'] if r['source_reference']=='KONPA')
    create_adjudication(slot['attendance_record_id'],AdjudicationPayload(adjudication_type='CANCEL_EARLY_LEAVE',reason='专职人员复核原表后撤销早退'),{'id':f['actor']})
    assert json.loads(sources.get_import(iid,f['actor'])['rows'][-1]['record_json'])['state']=='早退'
    score=fetch_one('SELECT is_early_leave,final_points FROM attendance_score_records WHERE attendance_record_id=?',(slot['attendance_record_id'],))
    assert not score['is_early_leave'] and DecimalText(score['final_points'])=='4.00'
    assert DecimalText(sources.create_batches(iid,f['actor'])['batches'][0]['proposed_points'])=='17.00'


def test_calendar_preparation_preserves_existing_published_policy(source_fixture):
    f=source_fixture;before=activity.get_business_calendar(calendar_year=2026)
    result=sources.prepare_reading_calendar(f['actor'])
    assert result['idempotent'] and result['reading_calendar']['status']=='PUBLISHED'
    assert activity.get_business_calendar(calendar_year=2026)==before
    canonical=Path(__file__).resolve().parents[3]/'data/learning-calendars/china-mainland-2026.json'
    asset=Path(sources.__file__).resolve().parents[1]/'assets/credit-reading-calendar-2026.json'
    assert canonical.read_bytes()==asset.read_bytes()


def test_source_lists_resolve_scope_once_and_fetch_rows_only_for_detail(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING')
    with patch.object(credits,'accessible_org_ids',return_value=None) as resolver:
        state=sources.workbench(f['actor'])
        assert resolver.call_count==4
    header=next(r for r in state['imports'] if r['id']==iid)
    assert header['row_count']==1 and header['rows']==[] and 'source_snapshot_json' not in header
    assert len(sources.get_import(iid,f['actor'])['rows'])==1
    with patch.object(credits,'accessible_org_ids',return_value=None) as resolver:
        assert exports.export_options(f['actor'])['classes']
        assert resolver.call_count==4


def test_list_scope_is_permission_intersection_and_refreshed_per_request(source_fixture):
    f=source_fixture;iid,_,_=registered(f,'DAILY_READING')
    def scopes(actor,permission=None):return set() if permission=='members:read' else None
    with patch.object(credits,'accessible_org_ids',side_effect=scopes):
        assert sources.workbench(f['actor'])['classes']==[]
        assert sources.workbench(f['actor'])['imports']==[]
        assert exports.export_options(f['actor'])['classes']==[]
    # Removing the request-local restriction is observed immediately, without
    # a TTL/global role cache retaining a previous request's scopes.
    assert any(r['id']==iid for r in sources.workbench(f['actor'])['imports'])


def test_full_export_keeps_authorized_opening_balance_without_class_assignment(source_fixture):
    f=source_fixture
    with transaction() as c:
        now=credits._db_timestamp(c)
        execute(c,"DELETE FROM member_org_relations WHERE member_id=? AND relation_type='STUDY_CLASS'",(f['member_id'],))
        execute(c,'UPDATE members SET org_unit_id=? WHERE id=?',(f['center_id'],f['member_id']))
        fp=sources._hash(['unassigned-opening',f['suffix']])
        iid=execute(c,"INSERT INTO learning_credit_opening_imports(content_fingerprint,file_sha256,original_filename,cutoff_date,status,created_by,posted_by,created_at,posted_at) VALUES (?,?,'隔离期初.xlsx','2026-09-30','POSTED',?,?,?,?)",(fp,fp,f['actor'],f['actor'],now,now)).lastrowid
        execute(c,"INSERT INTO learning_credit_opening_balances(member_id,import_id,org_unit_id,points,cutoff_date,posted_by,posted_at) VALUES (?,?,?,12.75,'2026-09-30',?,?)",(f['member_id'],iid,f['center_id'],f['actor'],now))
    book=load_workbook(io.BytesIO(exports.export_credits(f['actor'])))
    rows=list(book['学分汇总'].iter_rows(min_row=2,values_only=True))
    row=next((r for r in rows if r[0]==f['member_code']),None)
    assert row is not None and row[4]==12.75 and row[6]==12.75
    assert '余额归属' in row[2]
    with patch.object(credits,'accessible_org_ids',return_value={f['class_id']}):
        restricted=load_workbook(io.BytesIO(exports.export_credits(f['actor'])))
        assert all(r[0].value!=f['member_code'] for r in restricted['学分汇总'].iter_rows(min_row=2))


def test_incomplete_recording_archives_source_then_can_be_confirmed_later(source_fixture):
    f=source_fixture
    first,_,_=registered(f,'DAILY_READING',workbook(f,'DAILY_READING',[date.fromisoformat(DAY),'否']))
    old=sources.get_import(first,f['actor'])
    assert old['source_only'] and old['rows'][0]['fact_id'] is None
    assert not sources.create_batches(first,f['actor'])['batches']
    second,_,_=registered(f,'DAILY_READING')
    assert second!=first and not sources.get_import(second,f['actor'])['source_only']
    sources.create_batches(second,f['actor']);sources.process(second,f['actor'],'submit');sources.process(second,f['reviewer'],'approve');sources.process(second,f['actor'],'post')
    assert fetch_one("SELECT SUM(points) AS p FROM learning_credit_entries WHERE member_id=? AND credit_type='DAILY_READING'",(f['member_id'],))['p']==1
    assert json.loads(sources.get_import(first,f['actor'])['rows'][0]['record_json'])['已完成录音']=='否'


@pytest.mark.parametrize('wide',[False,True])
def test_original_reading_export_trusts_supplied_dates_and_keeps_source(source_fixture,wide):
    f=source_fixture
    with transaction() as c:
        execute(c,"UPDATE members SET phone_last4='2986',phone_masked='199****2986' WHERE id=?",(f['member_id'],))
    headers=sources.READING_REPORT_HEADERS[int(wide)]
    book=Workbook();sheet=book.active;sheet.title='导出数据';sheet.append(headers)
    def record(day,status):
        values={'姓名':f['name'],'账户':'199****2986','小组':'测试组','日期':day,'内容标题':'经营之心','录音':status}
        return [values.get(label,'') for label in headers]
    sheet.append(record('2026-10-09','未提交'))
    sheet.append(record('2026-10-10','按时完成'))
    sheet.append(record('2026-10-11','按时完成'))
    raw=io.BytesIO();book.save(raw);data=raw.getvalue()
    preview=sources.preview(data,'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    assert preview['error_count']==0 and preview['proposed_points']=='2.00'
    assert [r['credit_status'] for r in preview['rows']]==['NO_CREDIT','READY','READY']
    import_id=sources.register(data,'DAILY_READING',f['class_id'],'{}',NOTE,'读书会打卡统计列表.xlsx',preview['fingerprint'],f['actor'])['id']
    saved=sources.get_import(import_id,f['actor'])
    assert [r['fact_id'] is not None for r in saved['rows']]==[False,True,True]
    assert json.loads(saved['rows'][1]['record_json'])['录音']=='按时完成'
    sources.create_batches(import_id,f['actor']);sources.process(import_id,f['actor'],'submit')
    sources.process(import_id,f['reviewer'],'approve');sources.process(import_id,f['actor'],'post')
    assert fetch_one("SELECT SUM(points) AS p FROM learning_credit_entries WHERE member_id=? AND credit_type='DAILY_READING'",(f['member_id'],))['p']==2
    exported=load_workbook(io.BytesIO(exports.export_credits(f['actor'],f['class_id'],'2026-10-01','2026-10-13')))
    assert '读书会打卡统计列表.xlsx' in exported['新增学分明细']['K2'].value


def test_original_reading_export_rejects_unmatched_masked_account(source_fixture):
    f=source_fixture;book=Workbook();sheet=book.active;sheet.title='导出数据';sheet.append(sources.READING_REPORT_HEADERS[0])
    sheet.append([f['name'],'199****9999','测试组',DAY,'经营之心','过期完成'])
    raw=io.BytesIO();book.save(raw)
    preview=sources.preview(raw.getvalue(),'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    assert preview['error_count']==1 and preview['fingerprint'] is None
    assert '账户尾号' in preview['rows'][0]['errors'][0]


def test_reading_partial_import_skips_invalid_and_duplicate_rows_then_posts_once(source_fixture):
    f=source_fixture
    book=Workbook();sheet=book.active;sheet.title=sources.SHEETS['DAILY_READING'];sheet.append(sources.HEADERS['DAILY_READING'])
    sheet.append(['UNKNOWN','无匹配学员',date.fromisoformat(DAY),'是'])
    sheet.append([f['member_code'],f['name'],date.fromisoformat(DAY),'=1+1'])
    sheet.append([f['member_code'],f['name'],date.fromisoformat(DAY),'是'])
    sheet.append([f['member_code'],f['name'],date.fromisoformat(DAY),'是'])
    output=io.BytesIO();book.save(output);raw=output.getvalue()
    p=sources.preview(raw,'DAILY_READING',f['class_id'],'{}',NOTE,f['actor'])
    assert p['error_count']==3 and p['ignored_count']==3 and p['accepted_count']==1
    assert p['fingerprint'] and p['proposed_points']=='1.00'
    iid=sources.register(raw,'DAILY_READING',f['class_id'],'{}',NOTE,'混合读书记录.xlsx',p['fingerprint'],f['actor'])['id']
    receipt=sources.get_import(iid,f['actor'])
    assert [r['excel_row'] for r in receipt['rows']]==[4]
    assert [r['excel_row'] for r in receipt['ignored_rows']]==[2,3,5]
    frozen=fetch_one('SELECT source_snapshot_json FROM learning_credit_source_imports WHERE id=?',(iid,))['source_snapshot_json']
    assert '=1+1' not in frozen and 'UNKNOWN' not in frozen
    sources.create_batches(iid,f['actor']);sources.process(iid,f['actor'],'submit')
    sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    assert sources.register(raw,'DAILY_READING',f['class_id'],'{}',NOTE,'重传.xlsx',p['fingerprint'],f['actor'])['idempotent']
    assert fetch_one("SELECT SUM(points) AS p FROM learning_credit_entries WHERE member_id=? AND credit_type='DAILY_READING'",(f['member_id'],))['p']==1


def test_credit_export_uses_workdays_plus_fixed_basic_points(source_fixture):
    from decimal import Decimal
    from app.services.credit_monthly_targets import monthly_target
    f=source_fixture
    _calendar(f,year=2026,label='full-month-'+f['suffix'],day_types={f'2026-10-{d:02}':'NORMAL_WORKDAY' for d in range(1,32)})
    with transaction() as c:
        now=credits._db_timestamp(c)
        execute(c,"UPDATE class_learning_cycles SET planned_class_meeting_at='2026-10-15T01:00:00+00:00' WHERE id=?",(f['cycle_id'],))
        plan=execute(c,'SELECT plan_cycle_id FROM class_learning_cycles WHERE id=?',(f['cycle_id'],)).fetchone()['plan_cycle_id']
        for task in ('GROUP_MEETING','GROUP_MEETING'):
            execute(c,'INSERT INTO learning_plan_tasks(plan_cycle_id,task_type,title,is_required,sort_order,created_at,updated_at) VALUES (?,?,?,1,1,?,?)',(plan,task,'小组学习',now,now))
        rule_id=execute(c,'SELECT credit_rule_version_id FROM class_learning_bindings WHERE id=?',(f['binding_id'],)).fetchone()['credit_rule_version_id']
        from app.services.credit_rule_mapping import bound_generic_rule
        expected=Decimal(31+5+18+4)
        assert monthly_target(c,f['class_id'],date(2026,10,9))['points']==expected
        group_id='export-group-'+f['suffix']
        execute(c,"INSERT INTO org_units(id,unit_code,name,unit_type,parent_id,is_active,created_at,updated_at) VALUES (?,?,'合成学习小组','GROUP',?,1,?,?)",(group_id,group_id,f['class_id'],now,now))
        execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,created_at,updated_at) VALUES (?,?,'STUDY_GROUP',1,?,?)",(f['member_id'],group_id,now,now))
    iid,_,_=registered(f,'DAILY_READING')
    sources.create_batches(iid,f['actor']);sources.process(iid,f['actor'],'submit');sources.process(iid,f['reviewer'],'approve');sources.process(iid,f['actor'],'post')
    # Export starts after the earning day: the monthly numerator still starts
    # at month-day 1, independently of the detail date filter.
    book=load_workbook(io.BytesIO(exports.export_credits(f['actor'],f['class_id'],'2026-10-10','2026-10-15')))
    row=dict(zip([c.value for c in book['学分汇总'][1]],[c.value for c in book['学分汇总'][2]]))
    assert row['小组']=='合成学习小组' and row['当月基础满分']==float(expected)
    assert row['学分率']==pytest.approx(1/float(expected))
    assert book['学分汇总']['J2'].number_format=='0.0%'
    with transaction() as c:execute(c,"UPDATE learning_plan_versions SET status='DRAFT' WHERE id=?",(f['plan_id'],))
    book=load_workbook(io.BytesIO(exports.export_credits(f['actor'],f['class_id'],DAY,DAY)))
    assert book['学分汇总']['I2'].value==float(expected) and book['学分汇总']['J2'].value==pytest.approx(1/float(expected))


def test_month_target_uses_latest_complete_calendar_independent_of_plan(source_fixture):
    from decimal import Decimal
    from app.services.credit_monthly_targets import monthly_target
    f=source_fixture
    types={f'2026-10-{d:02}':'NORMAL_WORKDAY' for d in range(1,32)}
    types['2026-10-01']='HOLIDAY';types['2026-10-02']='WEEKEND';types['2026-10-03']='ADJUSTED_WORKDAY'
    _calendar(f,year=2026,label='basic-target-'+f['suffix'],day_types=types)
    with transaction() as c:
        execute(c,'UPDATE class_learning_cycles SET planned_class_meeting_at=NULL WHERE id=?',(f['cycle_id'],))
        execute(c,"UPDATE learning_plan_versions SET status='DRAFT' WHERE id=?",(f['plan_id'],))
        result=monthly_target(c,f['class_id'],date(2026,10,15))
        assert result['points']==Decimal(29+27)
        assert '29个工作日' in result['reason']
        version=execute(c,"SELECT id FROM learning_business_calendar_versions WHERE calendar_key='CHINA_MAINLAND' AND calendar_year=2026 AND status='PUBLISHED' ORDER BY id DESC LIMIT 1").fetchone()['id']
        execute(c,"DELETE FROM learning_business_calendar_days WHERE calendar_version_id=? AND business_date='2026-10-31'",(version,))
        assert monthly_target(c,f['class_id'],date(2026,10,15))['points'] is None
        assert '2026-10-31' in monthly_target(c,None,date(2026,10,15))['reason']


def test_monthly_basic_rate_excludes_courses_and_keeps_reversal_net(source_fixture):
    from decimal import Decimal
    f=source_fixture
    _calendar(f,year=2026,label='basic-earned-'+f['suffix'],day_types={f'2026-10-{d:02}':'NORMAL_WORKDAY' for d in range(1,32)})
    iid,_,_=registered(f,'DAILY_READING');source=sources.create_batches(iid,f['actor'])
    sources.process(iid,f['actor'],'confirm',source['settlement_fingerprint'],incremental=True)
    with transaction() as c:
        template=dict(execute(c,"SELECT * FROM learning_credit_entries WHERE member_id=? AND credit_type='DAILY_READING'",(f['member_id'],)).fetchone())
        base={**template,'rule_snapshot':json.loads(template['rule_snapshot_json'])}
        for kind,points in [('EXCELLENT_SHARE',3),('CLASS_MEETING_SCORE',18),('GROUP_MEETING_ATTENDANCE',4),('COURSE_COMPLETION',40)]:
            row=credits._insert_entry(c,{**base,'credit_type':kind,'points':str(points),'idempotency_key':'basic-export-'+f['suffix']+kind},status='POSTED',actor_user_id=f['actor'])
            if kind=='EXCELLENT_SHARE':
                credits._insert_entry(c,{**base,'credit_type':kind,'source_type':'REVERSAL','points':'-1','reversal_of_entry_id':row['id'],'idempotency_key':'basic-reversal-'+f['suffix']},status='POSTED',actor_user_id=f['actor'])
    book=load_workbook(io.BytesIO(exports.export_credits(f['actor'],f['class_id'],DAY,'2026-10-15')))
    sheet=book['学分汇总'];row=dict(zip([c.value for c in sheet[1]],[c.value for c in sheet[2]]))
    assert row['当月基础满分']==58
    assert row['2026-10-01至2026-10-15当月基础获得学分']==25
    assert row['学分率']==pytest.approx(25/58)
    assert row['2026-10-09至2026-10-15新增学分']==65
