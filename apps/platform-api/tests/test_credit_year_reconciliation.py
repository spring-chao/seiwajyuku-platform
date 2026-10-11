"""Deficits, identity overrides, stale facts and atomic year rollover on both DBs."""
import io
import json
from decimal import Decimal
from unittest.mock import patch
import pytest
from openpyxl import load_workbook
from app.db import connect,execute,fetch_one,transaction
from app.services import credit_year_allocations as year
from app.services import credit_year_reconciliation as rec
from app.services import credit_opening_balances as opening
from app.services.wechat_credit_summary import get_member_credit_summary
from app.migrations import MIGRATION_ROOT,_split_mysql
import sqlite3
from test_credit_year_allocations import annual,content
from test_credit_opening_balances import source

NOTE='第二学年已经完成，年度原表为准，只补缺额，保留原余额'

def preview(a,data=None,overrides='{}'):
    return year.preview(data or content(a,('100','200')),a['binding_id'],2,'2001-08-31',NOTE,a['actor'],rec.MODE,overrides)

def submit(a,data=None,overrides='{}'):
    data=data or content(a,('100','200'));p=preview(a,data,overrides)
    assert p['error_count']==0
    r=year.register(data,a['binding_id'],2,'2001-08-31',NOTE,'年度.xlsx',p['fingerprint'],a['actor'],rec.MODE,overrides)
    return r['id'],p['fingerprint'],p

def test_deficit_preserves_original_and_missing_member_starts_year_three(annual):
    # Add a third missing member, leaving the two original source rows immutable.
    with transaction() as c:
        now=opening._db_timestamp(c)
        mid=int(execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES ('synthetic-annual-missing','缺失余额测试',?,'ACTIVE',?,?)",(annual['org'],now,now)).lastrowid)
    data=content(annual,('100','200'))
    book=load_workbook(io.BytesIO(data));s=book.active;s['A7']=3;s['B7']='缺失余额测试';s['D7']='50';out=io.BytesIO();book.save(out);data=out.getvalue()
    iid,fp,p=submit(annual,data)
    assert p['total_points']=='350.00' and p['supplement_points']=='317.75'
    with pytest.raises(ValueError,match='先复核'):year.action(iid,annual['actor'],'post',fp)
    header,rows=year.workbench(annual['actor'])['imports'][0],None
    c=connect();batch,rows=year._batch(c,annual['actor'],iid);c.close()
    oid=rows[0]['opening_import_id']
    ofp=fetch_one('SELECT content_fingerprint FROM learning_credit_opening_imports WHERE id=?',(oid,))['content_fingerprint']
    with pytest.raises(ValueError,match='年度'):opening.action(oid,annual['actor'],'approve',ofp)
    assert year.action(iid,annual['actor'],'approve',fp)['status']=='APPROVED'
    assert year.action(iid,annual['actor'],'post',fp)['status']=='POSTED'
    assert year.action(iid,annual['actor'],'post',fp)['idempotent']
    assert year.register(data,annual['binding_id'],2,'2001-08-31',NOTE,'重命名.xlsx',fp,annual['actor'],rec.MODE,'{}')['idempotent']
    assert fetch_one('SELECT points,import_id FROM learning_credit_opening_balances WHERE member_id=?',(annual['ids'][0][0],))=={'points':Decimal('12.25'),'import_id':annual['opening_import_id']}
    c=connect();summary=opening.opening_summary(c,allowed={annual['org']});c.close()
    assert summary['entry_count']==3 and summary['total_points']=='350.00'
    for member,score in [(annual['ids'][0][0],'100.00'),(annual['ids'][1][0],'200.00'),(mid,'50.00')]:
        result=get_member_credit_summary(member)
        assert result['total_points']==score and result['current_learning_year']==3 and result['current_learning_year_points']=='0.00'
        assert {r['year_index']:r['points'] for r in result['learning_years']}=={1:None,2:score,3:'0.00'}
        assert not result['has_unallocated_learning_year_credits']
        assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE member_id=?',(member,))['n']==0
    assert preview(annual)['error_count']==2

def test_alias_mapping_is_explicit_scoped_and_retained_without_renaming(annual):
    data=content(annual,('100','200'))
    book=load_workbook(io.BytesIO(data));book.active['B5']='原表别名';out=io.BytesIO();book.save(out);data=out.getvalue()
    assert preview(annual,data)['error_count']==1
    codes=json.dumps({'原表别名':annual['ids'][0][1]})
    iid,fp,_=submit(annual,data,codes)
    year.action(iid,annual['actor'],'approve',fp);year.action(iid,annual['actor'],'post',fp)
    assert fetch_one('SELECT name FROM members WHERE id=?',(annual['ids'][0][0],))['name']==annual['ids'][0][2]
    assert fetch_one('SELECT source_name,member_code FROM learning_credit_year_reconciliation_rows WHERE import_id=? AND member_id=?',(iid,annual['ids'][0][0]))=={'source_name':'原表别名','member_code':annual['ids'][0][1]}
    with pytest.raises(ValueError,match='不存在'):preview(annual,data,json.dumps({'不存在':annual['ids'][0][1]}))
    assert preview(annual,data,json.dumps({'原表别名':'outside-scope'}))['error_count']>=1
    with patch.object(opening,'accessible_org_ids',return_value=set()),pytest.raises(PermissionError):preview(annual,data,codes)

@pytest.mark.parametrize('change',['base','delta','member','opening_row'])
def test_stale_or_tampered_deficit_never_partially_posts(annual,change):
    iid,fp,_=submit(annual)
    year.action(iid,annual['actor'],'approve',fp)
    with transaction() as c:
        if change=='base':execute(c,'UPDATE learning_credit_opening_balances SET points=13 WHERE member_id=?',(annual['ids'][0][0],))
        elif change=='delta':execute(c,'UPDATE learning_credit_year_reconciliation_rows SET supplement_points=0 WHERE import_id=?',(iid,))
        elif change=='member':execute(c,"UPDATE members SET status='INACTIVE' WHERE id=?",(annual['ids'][0][0],))
        else:execute(c,'UPDATE learning_credit_opening_rows SET points=0 WHERE import_id=(SELECT opening_import_id FROM learning_credit_year_rows WHERE import_id=? LIMIT 1)',(iid,))
    with pytest.raises(ValueError):year.action(iid,annual['actor'],'post',fp)
    assert fetch_one('SELECT status FROM learning_credit_year_imports WHERE id=?',(iid,))['status']=='APPROVED'
    assert not fetch_one('SELECT member_id FROM learning_credit_year_allocations WHERE import_id=?',(iid,))
    assert not fetch_one('SELECT id FROM learning_credit_opening_adjustments WHERE year_import_id=?',(iid,))
    assert not fetch_one('SELECT binding_id FROM learning_credit_class_progress WHERE binding_id=?',(annual['binding_id'],))

def test_lower_source_is_rejected_and_same_total_posts_zero_deficit(annual):
    bad=preview(annual,content(annual,('10','20')))
    assert bad['error_count']==1 and bad['fingerprint'] is None
    iid,fp,p=submit(annual,content(annual))
    assert p['supplement_points']=='0.00'
    year.action(iid,annual['actor'],'approve',fp);year.action(iid,annual['actor'],'post',fp)
    assert not fetch_one('SELECT id FROM learning_credit_opening_adjustments WHERE year_import_id=?',(iid,))
    assert get_member_credit_summary(annual['ids'][0][0])['learning_years'][1]['points']=='12.25'


def test_retained_deficit_source_refuses_schema_reversal(annual):
    iid,fp,_=submit(annual)
    c=connect()
    try:
        dialect='sqlite' if isinstance(c,sqlite3.Connection) else 'mysql'
        with pytest.raises(Exception,match='constraint|CHECK'):
            for sql in _split_mysql((MIGRATION_ROOT/'rollback'/dialect/'0071_learning_credit_year_reconciliation.down.sql').read_text()):execute(c,sql)
        assert rec.available(c)
        assert execute(c,'SELECT COUNT(*) AS n FROM learning_credit_year_reconciliation_rows WHERE import_id=?',(iid,)).fetchone()['n']==2
        assert opening.storage_available(c) and year.available(c)
    finally:c.close()


def test_ledger_fact_added_after_preview_prevents_registration(annual):
    from app.services.learning_credits import _insert_entry
    from uuid import uuid4
    data=content(annual,('100','200'));p=preview(annual,data)
    # Inject a conflicting retained fact in the disposable DB to test revalidation.
    with transaction() as c, patch.object(opening,'guard_regular_post'):
        _insert_entry(c,{'member_id':annual['ids'][0][0],'credit_category':'STANDARD_LEARNING','credit_type':'CLASS_MEETING_SCORE','points':'2','source_type':'TEST','source_id':uuid4().hex,'class_org_unit_id':annual['org'],'rule_key':'TEST','rule_version':'TEST','rule_snapshot':{},'occurred_at':'2001-08-10','idempotency_key':uuid4().hex},status='POSTED',actor_user_id=None)
    with pytest.raises(ValueError,match='已变化'):year.register(data,annual['binding_id'],2,'2001-08-31',NOTE,'年度.xlsx',p['fingerprint'],annual['actor'],rec.MODE,'{}')
    assert not fetch_one('SELECT id FROM learning_credit_year_imports WHERE content_fingerprint=?',(p['fingerprint'],))
