import io
import json
from datetime import date, datetime
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.auth import current_user
from app.db import execute, fetch_one, transaction
from app.main import app
from app.services import monthly_core_export as monthly
from app.services import learning_credits as credits
from app.services import credit_source_imports as sources
from app.services import reading_weekly_export as weekly
from test_reading_weekly_export import reading_fixture, save
from test_credit_source_excel import source_fixture


def workbook(f, month='2026-10'):
    return load_workbook(io.BytesIO(monthly.export_monthly(f['actor'], f['center_id'], month)))


def test_monthly_saved_source_missing_days_yellow_and_audit_without_posting(reading_fixture):
    f = reading_fixture
    save(f, [('2026-10-05','是'),('2026-10-06','否')])
    save(f, [('2026-10-05','是')], 'overlap.xlsx')
    book = workbook(f)
    assert book.sheetnames == ['学习数据','学习分析','学习异常人员','活动数据','小组学习会数据','学分明细']
    ws = book['学习数据']
    columns = {c.value.date().isoformat():c.column for c in ws[2] if isinstance(c.value,datetime)}
    row = next(c.row for r in ws for c in r if c.value == f['name'])
    assert ws.cell(row,columns['2026-10-05']).value == '√'
    missing = ws.cell(row,columns['2026-10-07'])
    assert missing.value is None
    failed = ws.cell(row,columns['2026-10-06'])
    assert failed.value == '×' and failed.fill.fgColor.rgb == '00FFFF00'
    assert ws.cell(row,ws.max_column).value is None
    assert ws.cell(row,ws.max_column-1).value == 1
    assert book['学习分析']['C4'].value is None
    assert ws['B5'].value == '学习分享率'
    assert any('轮值分享数据统计表' in str(c.value) for r in ws for c in r)
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE member_id=?',(f['member_id'],))['n'] == 0
    audit = fetch_one("SELECT after_json FROM audit_logs WHERE action='learning.monthly_core.export' AND org_unit_id=? ORDER BY id DESC LIMIT 1",(f['center_id'],))
    assert json.loads(audit['after_json'])['ledger_write'] is False


def test_historical_month_without_sources_exports_blanks_and_literal_names(reading_fixture):
    f = reading_fixture
    with transaction() as c:
        execute(c, 'UPDATE members SET name=? WHERE id=?', ('=HYPERLINK("https://example.invalid")',f['member_id']))
    book = workbook(f,'2026-09')
    ws = book['学习数据']
    name = next(c for r in ws for c in r if str(c.value).startswith('=HYPERLINK'))
    assert name.data_type == 's'
    assert ws.cell(name.row,14).value is None
    assert not any(c.data_type == 'f' for s in book for r in s for c in r)


def test_monthly_scope_intersects_all_permissions_and_rejects_foreign_center(reading_fixture):
    f = reading_fixture
    with patch.object(sources, 'effective_scope', return_value={f['class_id']}):
        options = monthly.export_options(f['actor'])
        assert options['centers'] == [{'id':f['center_id'],'name':'C7测试中心'}]
        book = workbook(f)
        assert book['学习数据']['A1'].value.startswith('C7测试班')
        with pytest.raises(PermissionError): monthly.export_monthly(f['actor'],'org-suzhou','2026-10')
    with patch.object(sources, 'effective_scope', return_value=set()):
        with pytest.raises(PermissionError): workbook(f)
    with patch.object(credits, 'user_context', return_value={'permissions':['exports:normal']}):
        with pytest.raises(PermissionError): workbook(f)


def test_period_on_first_day_allows_previous_full_month_and_current_empty_month():
    with patch.object(monthly, 'datetime') as clock:
        clock.now.return_value = datetime(2026,11,1,8)
        assert monthly.period('2026-10') == (date(2026,10,1), date(2026,10,31))
        assert monthly.period('2026-11') == (date(2026,11,1), date(2026,10,31))
        with pytest.raises(ValueError): monthly.period('2026-12')
        with pytest.raises(ValueError): monthly.period('2026-13')


def test_monthly_api_auth_scope_and_download(reading_fixture):
    f = reading_fixture; client = TestClient(app)
    params = {'center_id':f['center_id'],'month':'2026-10'}
    assert client.get('/api/v1/learning-credits/excel/monthly-core-export',params=params).status_code == 401
    app.dependency_overrides[current_user] = lambda: {'id':f['actor'],'permissions':credits.user_context(f['actor'])['permissions']}
    try:
        result = client.get('/api/v1/learning-credits/excel/monthly-core-options')
        assert result.status_code == 200 and 'phone' not in result.text
        result = client.get('/api/v1/learning-credits/excel/monthly-core-export',params=params)
        assert result.status_code == 200 and result.content[:2] == b'PK'
        with patch.object(sources,'effective_scope',return_value=set()):
            assert client.get('/api/v1/learning-credits/excel/monthly-core-export',params=params).status_code == 403
    finally:
        app.dependency_overrides.pop(current_user,None)


def test_group_formula_uses_seven_members_times_five_days():
    days = [date(2026,10,d) for d in range(19,24)]
    members = [{'name':f'合成学员{i}', 'values':['√']*5, 'pending':False, 'rate':1} for i in range(7)]
    members[0]['values'][1] = '×'; members[0]['rate'] = .8
    book = weekly.build_workbook({'class_name':'测试班','center_name':'测试中心','start':days[0],'end':days[-1],
        'days':days,'groups':[{'name':'一组','members':members,'required':35,'completed':34,'pending':False}],
        'total_completed':34,'total_required':35,'pending':False,'ignored':0,'sources':{}})
    assert book.active['I4'].value == '=COUNTIF(C4:G10,"√")/35'
    assert book.active['D4'].fill.fgColor.rgb == '00FFFF00'
    raw = weekly._workbook_bytes(book)
    assert load_workbook(io.BytesIO(raw),data_only=True).active['I4'].value == pytest.approx(34/35)


def test_template_asset_contains_only_styles_and_course_headings():
    from app.services.monthly_core_layout import ASSET
    reference = load_workbook(ASSET)
    for sheet in reference:
        for row in sheet:
            for cell in row:
                assert cell.value is None or sheet.title == '学分' and cell.row == 3 and 36 <= cell.column <= 76
                assert not cell.hyperlink and not cell.comment and cell.data_type != 'f'


def test_template_credit_bands_and_net_credits_are_typed_and_missing_history_stays_blank():
    from app.services.monthly_core_layout import build_workbook
    days=[(date(2026,10,19),True)]
    people=[['合成班','第一组','m1','合成学员','√',1,1,0,1,None,'']]
    entries=[['1','合成班','m1','合成学员',date(2026,10,19),'读书打卡',1,'source','1','POSTED',None],
             ['2','合成班','m1','合成学员',date(2026,10,19),'读书打卡',-1,'source','2','POSTED','1']]
    book=build_workbook('合成中心','2026-10',date(2026,10,29),days,people,[],[],[],entries)
    ws=book['学分明细']
    assert ws['E5'].value == 0 and ws['E5'].data_type == 'n'
    assert ws['O5'].value == 0 and ws['F5'].value is None  # October is filled; January is unknown.
    assert ws['F2'].value == '每日读书＋读书优秀分享'
    assert ws['P2'].value == '班级学习日（上午＋下午＋空巴）'
    assert ws['AB2'].value == '小组学习会'
    assert book['学习数据']['A1'].fill.fgColor == load_workbook(__import__('app.services.monthly_core_layout',fromlist=['ASSET']).ASSET)['学习数据']['A1'].fill.fgColor
