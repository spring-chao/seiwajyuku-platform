"""Exercise the production export against the same isolated SQLite/MySQL fixtures."""
import io
import json
from datetime import date
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.auth import current_user
from app.db import execute, fetch_one, transaction
from app.main import app
from app.services import credit_source_imports as sources
from app.services import learning_credits as credits
from app.services import learning_activity_credits as activity
from app.services import reading_weekly_export as weekly
from test_credit_source_excel import source_fixture
from test_learning_activity_credits import _calendar


@pytest.fixture
def reading_fixture(source_fixture):
    f=source_fixture
    _calendar(f, year=2026, day_types={**{f'2026-10-{d:02}':'NORMAL_WORKDAY' for d in range(5,10)},
        '2026-10-10':'ADJUSTED_WORKDAY','2026-10-11':'WEEKEND'}, label='weekly-'+f['suffix'])
    with transaction() as c:
        now=credits._db_timestamp(c); gid='weekly-group-'+f['suffix']
        execute(c,"INSERT INTO org_units(id,unit_code,name,unit_type,parent_id,is_active,created_at,updated_at) VALUES (?,?,'第一组','GROUP',?,1,?,?)",(gid,gid,f['class_id'],now,now))
        execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,'STUDY_GROUP',1,'2020-01-01',?,?)",(f['member_id'],gid,now,now))
        mid=int(execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,'临时暂停学员',?,'ACTIVE',?,?)",('weekly-'+f['suffix'],f['class_id'],now,now)).lastrowid)
        for oid,kind in [(f['class_id'],'STUDY_CLASS'),(gid,'STUDY_GROUP')]:
            execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,valid_from,created_at,updated_at) VALUES (?,?,?,1,'2020-01-01',?,?)",(mid,oid,kind,now,now))
        f.update(paused_id=mid,group_id=gid)
    return f


def save(f, days, label='weekly.xlsx'):
    book=Workbook(); sheet=book.active; sheet.title='读书打卡'; sheet.append(sources.HEADERS['DAILY_READING'])
    for day,state in days: sheet.append([f['member_code'],f['name'],day,state])
    data=io.BytesIO(); book.save(data)
    p=sources.preview(data.getvalue(),'DAILY_READING',f['class_id'],'{}','',f['actor'])
    assert p['error_count']==0
    return sources.register(data.getvalue(),'DAILY_READING',f['class_id'],'{}','',label,p['fingerprint'],f['actor'])['id']


def export(f, pauses=(), start='2026-10-05', end='2026-10-09'):
    return load_workbook(io.BytesIO(weekly.export_weekly(f['actor'],f['class_id'],start,end,pauses)))


def pause(mid, start='2026-10-05', end='2026-10-09'):
    return {'member_id':mid,'occurred_from':start,'occurred_to':end,'reason':'TEMPORARY'}


def test_partial_and_full_temporary_pause_exclude_person_days_and_keep_original_facts(reading_fixture):
    f=reading_fixture; save(f, [('2026-10-05','是'),('2026-10-06','是'),('2026-10-07','否')])
    original=fetch_one('SELECT status FROM members WHERE id=?',(f['paused_id'],))
    book=export(f,[pause(f['member_id'],'2026-10-07'),pause(f['paused_id'])]); ws=book.active
    member=next(r for r in ws.iter_rows() if r[1].value==f['name']); full=next(r for r in ws.iter_rows() if r[1].value=='临时暂停学员')
    assert [v.value for v in member[2:7]]==['√','√','暂停','暂停','暂停']
    assert member[7].value==1 and full[2].value=='暂停' and full[7].value==0
    assert ws['I4'].value=='=COUNTIF(C4:G5,"√")/2'
    summary=next(r for r in ws.iter_rows() if r[0].value=='汇总')
    assert summary[7].value==2 and summary[8].value==1
    assert fetch_one('SELECT status FROM members WHERE id=?',(f['paused_id'],))==original
    assert fetch_one('SELECT COUNT(*) AS n FROM learning_credit_entries WHERE member_id=?',(f['member_id'],))['n']==0
    audit=fetch_one("SELECT after_json FROM audit_logs WHERE action='learning.reading_weekly.export' AND org_unit_id=? ORDER BY id DESC LIMIT 1",(f['class_id'],))
    assert json.loads(audit['after_json'])['ledger_write'] is False


def test_missing_days_are_pending_explicit_incomplete_is_zero_and_overlap_deduplicates(reading_fixture):
    f=reading_fixture; save(f,[('2026-10-05','是'),('2026-10-06','否')]); save(f,[('2026-10-05','是')],'overlap.xlsx')
    ws=export(f,[pause(f['paused_id'])]).active
    member=next(r for r in ws.iter_rows() if r[1].value==f['name'])
    assert [c.value for c in member[2:7]]==['√','×','待同步','待同步','待同步']
    assert member[3].fill.fgColor.rgb == '00FFFF00'
    assert member[7].value=='待核对'
    summary=next(r for r in ws.iter_rows() if r[0].value=='汇总')
    assert summary[7].value==1 and summary[8].value=='待核对'
    ws=export(f,[pause(f['paused_id'],end='2026-10-06')],end='2026-10-06').active
    assert ws['H4'].value is None  # Only two daily columns, percentages move left.
    member=next(r for r in ws.iter_rows() if r[1].value==f['name'])
    assert member[4].value==0.5


def test_adjusted_workday_is_included_and_no_fake_weekend_target(reading_fixture):
    f=reading_fixture; save(f,[(f'2026-10-{d:02}','是') for d in range(5,11)])
    ws=export(f,[pause(f['paused_id'],end='2026-10-11')],end='2026-10-11').active
    assert ws['H2'].value.date()==date(2026,10,10) and ws['H3'].value=='周六'
    assert ws['I4'].value==1
    summary=next(r for r in ws.iter_rows() if r[0].value=='汇总')
    assert summary[8].value==6 and summary[9].value==1


def test_pauses_roster_and_export_revalidate_scope_and_permissions(reading_fixture):
    f=reading_fixture; save(f,[('2026-10-05','是')])
    with patch.object(sources,'_has_scope',return_value=False):
        with pytest.raises(PermissionError): export(f)
        with pytest.raises(PermissionError): weekly.roster(f['actor'],f['class_id'],'2026-10-05','2026-10-09')
    with patch.object(credits,'user_context',return_value={'permissions':['exports:normal']}),pytest.raises(PermissionError): export(f)
    with pytest.raises(ValueError,match='不在'): export(f,[pause(99999999)])
    with pytest.raises(ValueError,match='区间'): export(f,[pause(f['member_id'],'2026-10-04')])
    assert TestClient(app).post('/api/v1/learning-credits/excel/reading-weekly-export',json={}).status_code==401
    app.dependency_overrides[current_user]=lambda:{'id':f['actor'],'permissions':credits.user_context(f['actor'])['permissions']}
    try:
        params={'class_id':f['class_id'],'occurred_from':'2026-10-05','occurred_to':'2026-10-09'}
        client=TestClient(app)
        result=client.get('/api/v1/learning-credits/excel/reading-weekly-roster',params=params)
        assert result.status_code==200 and len(result.json()['data']['members'])==2
        assert 'phone' not in result.text
        response=client.post('/api/v1/learning-credits/excel/reading-weekly-export',json={**params,'pauses':[pause(f['paused_id'])]})
        assert response.status_code==200 and response.content[:2]==b'PK'
        with patch.object(sources,'_has_scope',return_value=False):
            assert client.post('/api/v1/learning-credits/excel/reading-weekly-export',json=params).status_code==403
    finally: app.dependency_overrides.pop(current_user,None)


def test_calendar_missing_cross_month_and_no_source_fail_explicitly(reading_fixture):
    f=reading_fixture
    with pytest.raises(ValueError,match='尚未保存'): export(f)
    with pytest.raises(ValueError,match='跨月'): weekly.period('2026-10-26','2026-11-01')
    with pytest.raises(ValueError,match='自然周'): weekly.period('2026-10-05','2026-10-12')
    with patch.object(activity,'is_workday',return_value=None),pytest.raises(ValueError,match='日历'): export(f)


def test_unmatched_sources_do_not_block_complete_groups_and_names_stay_literal(reading_fixture):
    f=reading_fixture; iid=save(f,[(f'2026-10-{d:02}','是') for d in range(5,10)])
    with transaction() as c:
        row=execute(c,'SELECT source_snapshot_json FROM learning_credit_source_imports WHERE id=?',(iid,)).fetchone()
        snapshot=json.loads(row['source_snapshot_json']); snapshot['ignored_rows']=[{'excel_row':99,'reasons':['身份待核对']}]
        execute(c,'UPDATE learning_credit_source_imports SET source_snapshot_json=? WHERE id=?',(json.dumps(snapshot),iid))
        execute(c,"UPDATE members SET name='=HYPERLINK(\"https://example.invalid\")' WHERE id=?",(f['member_id'],))
    ws=export(f,[pause(f['paused_id'])]).active
    member=next(r for r in ws.iter_rows() if isinstance(r[1].value,str) and r[1].value.startswith('=HYPERLINK'))
    assert member[1].data_type=='s' and member[7].value==1 and ws['I4'].value=='=COUNTIF(C4:G5,"√")/5'
    assert all(c.value.startswith('=COUNTIF(') for row in ws for c in row if c.data_type=='f')
    assert '1条' in '\n'.join(str(c.value or '') for row in ws for c in row)


def test_reading_export_does_not_transfer_frozen_workbook_per_daily_row(reading_fixture):
    f=reading_fixture;iid=save(f,[('2026-10-05','是')])
    with transaction() as c:
        raw=execute(c,'SELECT source_snapshot_json FROM learning_credit_source_imports WHERE id=?',(iid,)).fetchone()
        snapshot=json.loads(raw['source_snapshot_json']);snapshot['ignored_rows']=[{'reason':'synthetic'}];snapshot['padding']='synthetic'*100000
        execute(c,'UPDATE learning_credit_source_imports SET source_snapshot_json=? WHERE id=?',(json.dumps(snapshot),iid))
        metadata=weekly.source_metadata(c,[iid]*2000)
        assert len(metadata)==1 and metadata[0]['ignored_count']==1 and 'source_snapshot_json' not in metadata[0]
    from app.db import execute as real_execute
    with patch.object(weekly,'execute',wraps=real_execute) as query:
        export(f,[pause(f['paused_id'],end='2026-10-05')],end='2026-10-05')
    joined=[call.args[1] for call in query.call_args_list if 'JOIN learning_credit_source_imports' in call.args[1]]
    assert joined and all('source_snapshot_json' not in sql for sql in joined)
