import csv
import io
import json
from unittest.mock import patch

import pytest

from app.db import execute, fetch_one, transaction
from app.services.study_meeting_operations import meeting_credit_preview, export_meeting_records
from app.services.study_meetings import confirm_study_meeting_course_completion, StudyMeetingError
from app.services.study_meetings import correct_meeting_courses, get_study_meeting_record_for_operations
from app.services.study_meeting_attendees import correct_attendees
from app.services.study_meeting_review import confirm_meeting_review
from app.services.study_meetings import list_study_meeting_records, StudyMeetingPermissionError
from app.services.learning_credits import _insert_entry, dry_run_study_meeting_settlement
from app.services.credit_settlement_batches import dry_run_study_meeting_batch, submit_study_meeting_batch_for_approval
from test_learning_credit_ledger import _admin_id, _use_credit_plan, _submit_without_evidence
from test_study_meeting_evidence import create
from test_v12_mvp import _seed_group_leader_fixture


@pytest.fixture(autouse=True)
def flags(monkeypatch):
    monkeypatch.setenv("STUDY_MEETING_SUBMISSION_ENABLED", "true")
    monkeypatch.setenv("STUDY_MEETING_COURSE_EDIT_ENABLED", "true")
    monkeypatch.setenv("STUDY_MEETING_ATTENDEE_EDIT_ENABLED", "true")
    monkeypatch.setenv("LEARNING_CREDIT_SETTLEMENT_ENABLED", "false")
    monkeypatch.setenv("STUDY_MEETING_PRE_SETTLEMENT_REVIEW_ENABLED", "false")


def prepared(course=False):
    f = _seed_group_leader_fixture()
    _use_credit_plan(f, "OPERATOR_VIDEO" if course else None)
    if course:
        with transaction() as connection:
            execute(connection, "UPDATE learning_plan_credit_rule_versions SET status='PUBLISHED' WHERE id=(SELECT course_credit_rule_version_id FROM class_learning_bindings WHERE class_org_unit_id=?)", (f["class_id"],))
    session = create(f, ["OPERATOR_VIDEO"] if course else None)
    _submit_without_evidence(session)
    return f, session


def _resubmit(f,monkeypatch,tmp_path,*,home=None,cross=None,courses=None):
    from app.services.study_meetings import create_study_meeting,submit_study_meeting
    from app.services.study_meeting_evidence import upload_evidence
    from test_study_meeting_evidence import photo
    monkeypatch.setenv('STUDY_MEETING_EVIDENCE_ENABLED','true')
    monkeypatch.setenv('STUDY_EVIDENCE_STORAGE_BACKEND','local')
    monkeypatch.setenv('STUDY_EVIDENCE_LOCAL_ROOT',str(tmp_path))
    keys=courses or []
    draft=create_study_meeting(member_id=f['member_id'],group_org_unit_id=f['group_id'],meeting_date=None,
        member_ids=[f['member_id']] if home is None else home,
        cross_group_member_ids=[f['other_member_id']] if cross is None else cross,
        has_course=bool(keys),course_keys=keys)
    upload_evidence(member_id=f['member_id'],session_id=draft['id'],content=photo(),content_type='image/jpeg')
    return submit_study_meeting(member_id=f['member_id'],session_id=draft['id'])


def _review_post(session):
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    detail=get_study_meeting_record_for_operations(**args)
    return confirm_and_settle_meeting_review(**args,expected_fingerprint=detail['review_fingerprint'])


@pytest.mark.parametrize('reviewed',[False,True])
def test_resubmission_before_or_after_review_replaces_only_after_success(monkeypatch,tmp_path,reviewed):
    from app.services.study_meetings import create_study_meeting,submit_study_meeting
    f,first=prepared()
    if reviewed:
        d=get_study_meeting_record_for_operations(actor_user_id=_admin_id(),session_id=first['id'])
        confirm_meeting_review(actor_user_id=_admin_id(),session_id=first['id'],expected_fingerprint=d['review_fingerprint'])
    draft=create(f)
    assert draft['id']!=first['id'] and draft['status']=='DRAFT'
    assert [r['id'] for r in list_study_meeting_records(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'])]==[first['id']]
    with pytest.raises(StudyMeetingError,match='合影'):
        submit_study_meeting(member_id=f['member_id'],session_id=draft['id'])
    assert fetch_one('SELECT status FROM study_meeting_sessions WHERE id=?',(first['id'],))['status']=='SUBMITTED'
    second=_resubmit(f,monkeypatch,tmp_path)
    assert second['id']==draft['id'] and second['meeting_date']==first['meeting_date']
    assert fetch_one('SELECT status FROM study_meeting_sessions WHERE id=?',(first['id'],))['status']=='CANCELLED'
    assert [r['id'] for r in list_study_meeting_records(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'])]==[second['id']]
    assert get_study_meeting_record_for_operations(actor_user_id=_admin_id(),session_id=second['id'])['review_status']=='UNCONFIRMED'
    assert submit_study_meeting(member_id=f['member_id'],session_id=second['id'])['id']==second['id']


def test_posted_resubmission_unchanged_does_not_duplicate_credit(monkeypatch,tmp_path):
    f,first=prepared();_automatic_flags(monkeypatch)
    assert _review_post(first)['credit_settlement']['status']=='POSTED'
    before=fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n']
    second=_resubmit(f,monkeypatch,tmp_path)
    posted=_review_post(second)
    assert posted['credit_settlement']['status']=='POSTED',posted['credit_settlement']
    assert fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n']==before
    assert _review_post(second)['credit_settlement']['status']=='POSTED'


def test_posted_resubmission_remove_readd_and_repeat_apply_only_differences(monkeypatch,tmp_path):
    f,first=prepared();_automatic_flags(monkeypatch)
    assert _review_post(first)['credit_settlement']['status']=='POSTED'
    original=fetch_one("SELECT * FROM learning_credit_entries WHERE idempotency_key=?",
                       (f"GROUP_MEETING:{f['other_member_id']}:{f['learning_cycle_id']}",))
    second=_resubmit(f,monkeypatch,tmp_path,cross=[])
    preview=meeting_credit_preview(actor_user_id=_admin_id(),session_id=second['id'])
    assert preview['adjustments'][0]['member_id']==f['other_member_id']
    assert preview['adjustments'][0]['points']<0
    assert _review_post(second)['credit_settlement']['status']=='POSTED'
    assert _review_post(second)['credit_settlement']['status']=='POSTED'
    total=fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['other_member_id'],))['n']
    assert float(total)==0
    third=_resubmit(f,monkeypatch,tmp_path)
    assert _review_post(third)['credit_settlement']['status']=='POSTED'
    assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['other_member_id'],))['n'])==float(original['points'])
    assert fetch_one('SELECT * FROM learning_credit_entries WHERE id=?',(original['id'],))==original


def test_posted_resubmission_course_removed_keeps_attendance_and_readds_once(monkeypatch,tmp_path):
    f,first=prepared(course=True);_automatic_flags(monkeypatch)
    assert _review_post(first)['credit_settlement']['status']=='POSTED'
    second=_resubmit(f,monkeypatch,tmp_path,courses=[])
    d=_review_post(second)
    assert d['credit_settlement']['status']=='POSTED',d['credit_settlement']
    assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['member_id'],))['n'])==4
    third=_resubmit(f,monkeypatch,tmp_path,courses=['OPERATOR_VIDEO'])
    d=_review_post(third)
    assert d['credit_settlement']['status']=='POSTED',d['credit_settlement']
    assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['member_id'],))['n'])==44


def test_resubmission_preserves_other_group_credit_until_last_participation_removed(monkeypatch,tmp_path):
    f,first=prepared();_automatic_flags(monkeypatch)
    assert _review_post(first)['credit_settlement']['status']=='POSTED'
    with transaction() as c:
        execute(c,"INSERT INTO volunteer_appointments(person_id,member_id,appointment_key,org_unit_id,scope_type,starts_at,ends_at,status,source_reference,created_at,updated_at) "
            "SELECT person_id,member_id,appointment_key,?,scope_type,starts_at,ends_at,status,'resubmit-test',created_at,updated_at "
            "FROM volunteer_appointments WHERE member_id=? AND org_unit_id=?",
            (f['other_group_id'],f['member_id'],f['group_id']))
    other={**f,'group_id':f['other_group_id']}
    other_meeting=_resubmit(other,monkeypatch,tmp_path,home=[f['other_member_id']],cross=[f['member_id']])
    assert _review_post(other_meeting)['credit_settlement']['status']=='POSTED'
    replacement=_resubmit(f,monkeypatch,tmp_path,cross=[])
    assert _review_post(replacement)['credit_settlement']['status']=='POSTED'
    assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['other_member_id'],))['n'])==4
    removed=_resubmit(other,monkeypatch,tmp_path,home=[],cross=[f['member_id']])
    done=_review_post(removed)
    assert done['credit_settlement']['status']=='POSTED',done['credit_settlement']
    assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(f['other_member_id'],))['n'])==0


def test_replaced_meeting_cannot_post_stale_approved_preview(monkeypatch,tmp_path):
    from app.services import credit_settlement_batches as batches,learning_credits as credits
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    f,first=prepared();_automatic_flags(monkeypatch)
    args={'actor_user_id':_admin_id(),'session_id':first['id']}
    d=get_study_meeting_record_for_operations(**args)
    confirm_meeting_review(**args,expected_fingerprint=d['review_fingerprint'])
    batch=batches.dry_run_study_meeting_batch(**args)
    with transaction() as c:
        row=execute(c,'SELECT * FROM learning_credit_settlement_batches WHERE id=?',(batch['id'],)).fetchone()
        items=execute(c,"SELECT * FROM learning_credit_settlement_batch_items WHERE batch_id=? AND status<>'BLOCKED' ORDER BY id",(batch['id'],)).fetchall()
        fingerprint=batches._approval_fingerprint(row,items)
    newer=_resubmit(f,monkeypatch,tmp_path,cross=[])
    with pytest.raises(credits.LearningCreditError):
        batches._post_batch(actor_user_id=_admin_id(),batch_id=batch['id'],source_type='STUDY_MEETING',
            meeting_confirmation={'session_id':first['id'],'review_fingerprint':d['review_fingerprint'],'fingerprint':fingerprint})
    assert _review_post(newer)['credit_settlement']['status']=='POSTED'
    assert fetch_one('SELECT COUNT(*) n FROM learning_credit_entries WHERE member_id=?',(f['other_member_id'],))['n']==0


def test_partial_difference_posting_retries_without_repeating_deductions(monkeypatch,tmp_path):
    from app.services import learning_credits as credits
    f,first=prepared(course=True);_automatic_flags(monkeypatch)
    assert _review_post(first)['credit_settlement']['status']=='POSTED'
    second=_resubmit(f,monkeypatch,tmp_path,courses=[])
    original=credits._insert_entry
    attempts=0
    def interrupted(connection,item,**kwargs):
        nonlocal attempts
        if item['idempotency_key'].startswith('SMR:'):
            attempts+=1
            if attempts==2:raise RuntimeError('synthetic interrupted adjustment')
        return original(connection,item,**kwargs)
    with patch.object(credits,'_insert_entry',side_effect=interrupted):
        partial=_review_post(second)
    assert partial['credit_settlement']['status']=='PARTIAL_FAILED'
    assert _review_post(second)['credit_settlement']['status']=='POSTED'
    assert _review_post(second)['credit_settlement']['status']=='POSTED'
    for mid in [f['member_id'],f['other_member_id']]:
        assert float(fetch_one('SELECT SUM(points) n FROM learning_credit_entries WHERE member_id=?',(mid,))['n'])==4


def test_unsubmitted_drafts_are_hidden_from_formal_records_and_exports():
    from app.services.study_meeting_operations import export_meeting_workbook
    from openpyxl import load_workbook
    f=_seed_group_leader_fixture()
    draft=create(f)
    assert list_study_meeting_records(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'])==[]
    assert list_study_meeting_records(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'],status='DRAFT')==[]
    book=load_workbook(io.BytesIO(export_meeting_workbook(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'])))
    assert book['小组汇总'].max_row==1
    assert fetch_one('SELECT status FROM study_meeting_sessions WHERE id=?',(draft['id'],))['status']=='DRAFT'
    _submit_without_evidence(draft)
    assert [r['id'] for r in list_study_meeting_records(actor_user_id=_admin_id(),class_org_unit_id=f['class_id'])]==[draft['id']]


def test_region_class_filters_apply_to_records_export_and_authorized_options():
    from app.services.study_meetings import study_meeting_filter_options
    from app.services.study_meeting_operations import export_meeting_workbook
    from openpyxl import load_workbook
    f,session=prepared();other,other_session=prepared()
    actor=_admin_id()
    with patch('app.services.study_meetings.accessible_org_ids',return_value={f['class_id']}):
        options=study_meeting_filter_options(actor_user_id=actor)
        assert [r['id'] for r in options['classes']]==[f['class_id']]
        region=options['classes'][0]['management_org_unit_id']
        assert region and len(options['regions'])==1
        assert list_study_meeting_records(actor_user_id=actor,class_org_unit_id=other['class_id'])==[]
        assert list_study_meeting_records(actor_user_id=actor,management_org_unit_id='unknown')==[]
        selected=list_study_meeting_records(actor_user_id=actor,management_org_unit_id=region,class_org_unit_id=f['class_id'])
        assert [r['id'] for r in selected]==[session['id']]
    report=load_workbook(io.BytesIO(export_meeting_workbook(actor_user_id=actor,class_org_unit_id=f['class_id'],management_org_unit_id=region)))
    assert report['小组汇总'].max_row==2
    assert report['小组汇总']['A2'].value==session['session_code']
    # The export permission intersection still applies after selecting a class.
    with patch('app.services.study_meetings.accessible_org_ids',side_effect=lambda _actor,permission=None: set() if permission=='exports:normal' else None):
        assert list_study_meeting_records(actor_user_id=actor,class_org_unit_id=f['class_id'],scope_permissions=('exports:normal',))==[]


def test_preview_confirmation_removal_export_share_real_calculation():
    f, session = prepared(course=True)
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    before = fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"]
    preview = meeting_credit_preview(**args)
    assert [item["expected_points"] for item in preview["members"]] == [44, 44]
    confirm_study_meeting_course_completion(**args, course_key="OPERATOR_VIDEO", member_id=f["member_id"], completion_status="NOT_COMPLETED", expected_completion_status="CONFIRMED")
    preview = meeting_credit_preview(**args)
    member = next(item for item in preview["members"] if item["member_id"] == f["member_id"])
    assert member["expected_points"] == 4 and member["course_points"] == 0
    other = next(item for item in preview["members"] if item["member_id"] == f["other_member_id"])
    assert other["course_points"] == 40 and other["expected_points"] == 44
    assert preview["totals"]["proposed_points"] == 48
    with pytest.raises(StudyMeetingError, match="其他人修改"):
        confirm_study_meeting_course_completion(**args, course_key="OPERATOR_VIDEO", member_id=f["member_id"], completion_status="CONFIRMED", expected_completion_status="CONFIRMED")
    correct_attendees(**args, member_ids=[f["member_id"]], expected_member_ids=[f["member_id"], f["other_member_id"]])
    preview = meeting_credit_preview(**args)
    assert len(preview["members"]) == 1 and preview["totals"]["proposed_points"] == 4
    rows = list(csv.DictReader(io.StringIO(export_meeting_records(**args).lstrip("\ufeff"))))
    assert len(rows) == 1 and float(rows[0]["预计新增学分"]) == 4
    assert "未观看（人工确认）" in rows[0]["观看确认"]
    assert "phone" not in rows[0] and f["phone"] not in str(rows)
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"] == before
    assert fetch_one("SELECT id FROM audit_logs WHERE action='exports.study_meetings.normal' AND resource_id=?", (str(session["id"]),))


def test_preview_deduplicates_posted_credit(monkeypatch):
    f, session = prepared()
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    proposals = dry_run_study_meeting_settlement(**args)
    # Seed a prior approved settlement in the isolated ledger; the retired
    # direct settlement endpoint deliberately remains closed.
    with transaction() as connection:
        for entry in proposals["entries"]:
            _insert_entry(connection, {**entry, "occurred_at": session["meeting_date"]}, status="POSTED", actor_user_id=_admin_id())
    preview = meeting_credit_preview(**args)
    assert all(item["expected_points"] == 0 for item in preview["members"])
    assert all(item["reasons"] for item in preview["members"])


def test_export_scopes_permissions_and_csv_formula_safety():
    f, session = prepared()
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    with patch("app.services.study_meeting_operations.user_context", return_value={"permissions": ["exports:normal"]}), pytest.raises(PermissionError):
        export_meeting_records(**args)
    with patch("app.services.study_meeting_operations.accessible_org_ids", return_value={f["other_class_id"]}), pytest.raises(PermissionError):
        export_meeting_records(**args)
    with patch("app.services.learning_credits.accessible_org_ids", return_value={f["other_class_id"]}), pytest.raises(PermissionError):
        meeting_credit_preview(**args)
    with transaction() as connection:
        execute(connection, "UPDATE members SET name='=FORMULA()' WHERE id=?", (f["member_id"],))
    rows = list(csv.DictReader(io.StringIO(export_meeting_records(**args).lstrip("\ufeff"))))
    assert next(item for item in rows if int(item["学员编号"]) == f["member_id"])["姓名"] == "'=FORMULA()"


def test_pre_settlement_review_has_an_independent_gate_and_keeps_ledger_untouched(monkeypatch):
    f, session = prepared(course=True)
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    monkeypatch.setenv("LEARNING_CREDIT_SETTLEMENT_ENABLED", "true")
    monkeypatch.setenv("STUDY_MEETING_COURSE_EDIT_ENABLED", "false")
    monkeypatch.setenv("STUDY_MEETING_ATTENDEE_EDIT_ENABLED", "false")
    assert not get_study_meeting_record_for_operations(**args)["can_edit_courses"]
    monkeypatch.setenv("STUDY_MEETING_PRE_SETTLEMENT_REVIEW_ENABLED", "true")
    assert get_study_meeting_record_for_operations(**args)["can_edit_courses"]
    assert get_study_meeting_record_for_operations(**args)["can_edit_attendees"]
    before = fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"]
    confirm_study_meeting_course_completion(**args, course_key="OPERATOR_VIDEO", member_id=f["member_id"], completion_status="NOT_COMPLETED")
    correct_attendees(**args, member_ids=[f["member_id"]], expected_member_ids=[f["member_id"], f["other_member_id"]])
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"] == before


@pytest.mark.parametrize("protection", ["ledger", "approval", "posting", "course_ledger"])
def test_posted_or_frozen_sources_reject_all_manual_corrections(protection, monkeypatch):
    f, session = prepared(course=protection == "course_ledger")
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    monkeypatch.setenv("LEARNING_CREDIT_SETTLEMENT_ENABLED", "true")
    monkeypatch.setenv("STUDY_MEETING_PRE_SETTLEMENT_REVIEW_ENABLED", "true")
    if protection in {"ledger", "course_ledger"}:
        preview = dry_run_study_meeting_settlement(**args)
        entry = next(item for item in preview["entries"] if item["entry_kind"] == ("COURSE_COMPLETION" if protection == "course_ledger" else "GROUP_MEETING_ATTENDANCE"))
        with transaction() as connection:
            _insert_entry(connection, {**entry, "occurred_at": session["meeting_date"]}, status="POSTED", actor_user_id=_admin_id())
    else:
        monkeypatch.setenv("LEARNING_CREDIT_BATCH_DRY_RUN_ENABLED", "true")
        batch = dry_run_study_meeting_batch(**args)
        submit_study_meeting_batch_for_approval(actor_user_id=_admin_id(), batch_id=batch["id"])
        if protection == "posting":
            with transaction() as connection:
                execute(connection, "UPDATE learning_credit_settlement_batches SET status='POSTING' WHERE id=?", (batch["id"],))
    detail = get_study_meeting_record_for_operations(**args)
    assert not detail["can_edit_courses"] and not detail["can_edit_attendees"]
    assert not detail["can_confirm_review"]
    with pytest.raises(StudyMeetingError):
        confirm_meeting_review(**args, expected_fingerprint=detail["review_fingerprint"])
    assert detail["correction_lock_reason"]
    before = fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"]
    with pytest.raises(StudyMeetingError):
        correct_attendees(**args, member_ids=[f["member_id"]], expected_member_ids=[f["member_id"], f["other_member_id"]])
    with pytest.raises(StudyMeetingError):
        correct_meeting_courses(**args, course_keys=[], expected_course_keys=["OPERATOR_VIDEO"] if protection == "course_ledger" else [])
    if protection == "course_ledger":
        with pytest.raises(StudyMeetingError):
            confirm_study_meeting_course_completion(**args, course_key="OPERATOR_VIDEO", member_id=f["member_id"], completion_status="NOT_COMPLETED")
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"] == before


def test_home_rate_excludes_cross_group_and_counts_active_roster():
    f, session = prepared()
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    detail = get_study_meeting_record_for_operations(**args)
    assert detail["total_count"] == 2 and detail["home_count"] == 1
    assert detail["home_group_member_count"] == 1 and detail["home_group_attendance_rate"] == 100
    with transaction() as connection:
        # A second current home member did not attend this meeting.
        execute(connection, "UPDATE member_org_relations SET org_unit_id=? WHERE member_id=? AND relation_type='STUDY_CLASS'", (f['class_id'], f['foreign_member_id']))
        execute(connection, "INSERT INTO member_org_relations(member_id, org_unit_id, relation_type, is_primary, created_at, updated_at) "
            "VALUES (?, ?, 'STUDY_GROUP', 0, '2020-01-01', '2020-01-01')", (f["foreign_member_id"], f["group_id"]))
    detail = get_study_meeting_record_for_operations(**args)
    assert detail["home_group_member_count"] == 2 and detail["home_group_attendance_rate"] == 50
    record = next(row for row in list_study_meeting_records(actor_user_id=_admin_id()) if row["id"] == session["id"])
    assert record["home_group_attendance_rate"] == detail["home_group_attendance_rate"]
    with transaction() as connection:
        execute(connection, "UPDATE members SET status='INACTIVE' WHERE id=?", (f["foreign_member_id"],))
    assert get_study_meeting_record_for_operations(**args)["home_group_member_count"] == 1
    with transaction() as connection:
        execute(connection, "UPDATE member_org_relations SET valid_until='2000-01-01' WHERE org_unit_id=?", (f["group_id"],))
    detail = get_study_meeting_record_for_operations(**args)
    assert detail["home_group_member_count"] == 0 and detail["home_group_attendance_rate"] is None


def test_whole_meeting_review_is_audited_idempotent_and_invalidated_by_edits():
    f, session = prepared(course=True)
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    detail = get_study_meeting_record_for_operations(**args)
    assert detail["review_status"] == "UNCONFIRMED" and detail["can_confirm_review"]
    count = fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"]
    confirmed = confirm_meeting_review(**args, expected_fingerprint=detail["review_fingerprint"])
    assert confirmed["review_status"] == "CONFIRMED" and confirmed["reviewed_at"] and confirmed["reviewer_name"]
    confirm_meeting_review(**args, expected_fingerprint=detail["review_fingerprint"])
    audit = fetch_one("SELECT COUNT(*) n FROM audit_logs WHERE action='study_meeting.review_confirm' AND resource_id=?", (str(session["id"]),))
    assert audit["n"] == 1
    record = next(row for row in list_study_meeting_records(actor_user_id=_admin_id()) if row["id"] == session["id"])
    assert record["review_status"] == "CONFIRMED"
    updated = confirm_study_meeting_course_completion(**args, course_key="OPERATOR_VIDEO", member_id=f["member_id"], completion_status="NOT_COMPLETED")
    assert updated["review_status"] == "UNCONFIRMED" and updated["reviewed_at"] is None
    with pytest.raises(StudyMeetingError, match="内容已变化"):
        confirm_meeting_review(**args, expected_fingerprint=detail["review_fingerprint"])
    updated = confirm_meeting_review(**args, expected_fingerprint=updated["review_fingerprint"])
    assert updated["review_status"] == "CONFIRMED"
    updated = correct_attendees(**args, member_ids=[f["member_id"]], expected_member_ids=[f["member_id"], f["other_member_id"]])
    assert updated["review_status"] == "UNCONFIRMED"
    confirm_meeting_review(**args, expected_fingerprint=updated["review_fingerprint"])
    updated = correct_meeting_courses(**args, course_keys=[], expected_course_keys=["OPERATOR_VIDEO"])
    assert updated["review_status"] == "UNCONFIRMED"
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries")["n"] == count
    payload = fetch_one("SELECT after_json FROM audit_logs WHERE action='study_meeting.review_confirm' AND resource_id=? ORDER BY id DESC LIMIT 1", (str(session["id"]),))
    assert set(json.loads(payload["after_json"])) == {"fingerprint", "review_status"}


def test_whole_meeting_review_enforces_scope_permission_write_gate_and_posted_lock(monkeypatch):
    f, session = prepared()
    args = {"actor_user_id": _admin_id(), "session_id": session["id"]}
    fingerprint = get_study_meeting_record_for_operations(**args)["review_fingerprint"]
    with patch("app.services.study_meeting_review.accessible_org_ids", return_value={f["other_class_id"]}), pytest.raises(StudyMeetingPermissionError):
        confirm_meeting_review(**args, expected_fingerprint=fingerprint)
    with patch("app.services.study_meeting_review.can_edit_attendees", return_value=False), pytest.raises(StudyMeetingPermissionError):
        confirm_meeting_review(**args, expected_fingerprint=fingerprint)
    monkeypatch.setenv("DEPLOYMENT_READ_ONLY", "true")
    with pytest.raises(StudyMeetingPermissionError):
        confirm_meeting_review(**args, expected_fingerprint=fingerprint)
    monkeypatch.setenv("DEPLOYMENT_READ_ONLY", "false")
    preview = dry_run_study_meeting_settlement(**args)
    with transaction() as connection:
        _insert_entry(connection, {**preview["entries"][0], "occurred_at": session["meeting_date"]}, status="POSTED", actor_user_id=_admin_id())
    assert not get_study_meeting_record_for_operations(**args)["can_confirm_review"]
    with pytest.raises(StudyMeetingError, match="更正流程"):
        confirm_meeting_review(**args, expected_fingerprint=fingerprint)


def test_review_endpoint_requires_auth_and_rejects_stale_facts():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.services.iam import authenticate
    f, session = prepared()
    with TestClient(app) as client:
        url = f"/api/v1/study-meetings/records/{session['id']}/review-confirmation"
        assert client.post(url, json={"expected_fingerprint": "0" * 64}).status_code == 401
        token = authenticate("admin", "test-admin-password")["access_token"]
        headers = {"Authorization": "Bearer " + token}
        assert client.post(url, json={"expected_fingerprint": "0" * 64}, headers=headers).status_code == 400
        detail = get_study_meeting_record_for_operations(actor_user_id=_admin_id(), session_id=session["id"])
        response = client.post(url, json={"expected_fingerprint": detail["review_fingerprint"]}, headers=headers)
        assert response.status_code == 200, response.text
        assert response.json()["data"]["review_status"] == "CONFIRMED"


def _automatic_flags(monkeypatch):
    for key in ('LEARNING_CREDIT_SETTLEMENT_ENABLED','LEARNING_CREDIT_BATCH_DRY_RUN_ENABLED',
                'LEARNING_CREDIT_BATCH_POST_ENABLED','STUDY_MEETING_PRE_SETTLEMENT_REVIEW_ENABLED'):
        monkeypatch.setenv(key,'true')


def test_review_automatically_posts_once_without_fabricating_independent_approval(monkeypatch):
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    f,session=prepared();_automatic_flags(monkeypatch)
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    detail=get_study_meeting_record_for_operations(**args)
    fingerprint=detail['review_fingerprint']
    result=confirm_and_settle_meeting_review(**args,expected_fingerprint=fingerprint)
    assert result['review_status']=='CONFIRMED'
    assert result['credit_settlement']['status']=='POSTED',result['credit_settlement']
    assert float(result['credit_settlement']['posted_points'])==8
    again=confirm_and_settle_meeting_review(**args,expected_fingerprint=fingerprint)
    assert again['credit_settlement']['status']=='POSTED'
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries WHERE source_type='STUDY_MEETING_ATTENDANCE' AND source_id=?",(str(session['id']),))['n']==2
    row=fetch_one("SELECT approved_by,approval_fingerprint FROM learning_credit_settlement_batches WHERE batch_no LIKE ?",(f"LC-STUDY-{session['id']}-%",))
    assert row['approved_by'] is None and row['approval_fingerprint'] is None


def test_review_posting_failure_stays_retryable_in_meeting_without_duplicates(monkeypatch):
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    from app.services import learning_credits as credits
    f,session=prepared();_automatic_flags(monkeypatch)
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    fp=get_study_meeting_record_for_operations(**args)['review_fingerprint']
    original=credits.post_credit_entry;calls=0
    def fail_second(**kwargs):
        nonlocal calls
        calls+=1
        if calls==2:raise RuntimeError('isolated posting interruption')
        return original(**kwargs)
    with patch.object(credits,'post_credit_entry',side_effect=fail_second):
        partial=confirm_and_settle_meeting_review(**args,expected_fingerprint=fp)
    assert partial['review_status']=='CONFIRMED'
    assert partial['credit_settlement']['status']=='PARTIAL_FAILED'
    assert partial['can_confirm_review'] and partial['credit_settlement']['message']
    done=confirm_and_settle_meeting_review(**args,expected_fingerprint=fp)
    assert done['credit_settlement']['status']=='POSTED'
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries WHERE source_type='STUDY_MEETING_ATTENDANCE' AND source_id=?",(str(session['id']),))['n']==2


def test_automatic_meeting_post_requires_current_review_scope_and_post_gate(monkeypatch):
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    from app.services import credit_settlement_batches as batches
    from app.services import learning_credits as credits
    f,session=prepared();_automatic_flags(monkeypatch)
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    detail=get_study_meeting_record_for_operations(**args)
    batch=batches.dry_run_study_meeting_batch(**args)
    with transaction() as c:
        row=execute(c,'SELECT * FROM learning_credit_settlement_batches WHERE id=?',(batch['id'],)).fetchone()
        items=execute(c,"SELECT * FROM learning_credit_settlement_batch_items WHERE batch_id=? AND status<>'BLOCKED' ORDER BY id",(batch['id'],)).fetchall()
        fp=batches._approval_fingerprint(row,items)
    with pytest.raises(credits.LearningCreditError,match='尚未核实'):
        batches._post_batch(actor_user_id=args['actor_user_id'],batch_id=batch['id'],source_type='STUDY_MEETING',
            meeting_confirmation={'session_id':session['id'],'review_fingerprint':detail['review_fingerprint'],'fingerprint':fp})
    monkeypatch.setenv('LEARNING_CREDIT_BATCH_POST_ENABLED','false')
    result=confirm_and_settle_meeting_review(**args,expected_fingerprint=detail['review_fingerprint'])
    assert result['review_status']=='CONFIRMED' and result['credit_settlement']['status']=='DRY_RUN'
    assert '尚未开启' in result['credit_settlement']['message']
    monkeypatch.setenv('LEARNING_CREDIT_BATCH_POST_ENABLED','true')
    with patch('app.services.study_meeting_review.accessible_org_ids',side_effect=lambda actor,permission=None: set() if permission=='plans:credit_settlement_post' else None):
        result=confirm_and_settle_meeting_review(**args,expected_fingerprint=detail['review_fingerprint'])
        assert '授权范围' in result['credit_settlement']['message']
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries WHERE source_type='STUDY_MEETING_ATTENDANCE' AND source_id=?",(str(session['id']),))['n']==0


def test_mysql_concurrent_meeting_confirmations_do_not_duplicate_credits(monkeypatch):
    import os
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from app.services.study_meeting_review import confirm_and_settle_meeting_review
    if not os.environ.get('DATABASE_URL','').startswith('mysql+pymysql://'):
        pytest.skip('concurrent automatic settlement requires isolated MySQL')
    f,session=prepared();_automatic_flags(monkeypatch)
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    fp=get_study_meeting_record_for_operations(**args)['review_fingerprint']
    barrier=Barrier(2)
    def confirm():
        barrier.wait(timeout=10)
        return confirm_and_settle_meeting_review(**args,expected_fingerprint=fp)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures=[pool.submit(confirm) for _ in range(2)]
        results=[future.result(timeout=45) for future in futures]
    # A competing attempt may observe an active posting or deadlock and ask
    # for an explicit retry. It must never duplicate formal daily/cycle keys.
    assert any(r['review_status']=='CONFIRMED' for r in results)
    final=confirm_and_settle_meeting_review(**args,expected_fingerprint=fp)
    assert final['credit_settlement']['status']=='POSTED',final['credit_settlement']
    assert fetch_one("SELECT COUNT(*) n FROM learning_credit_entries WHERE source_type='STUDY_MEETING_ATTENDANCE' AND source_id=?",(str(session['id']),))['n']==2
