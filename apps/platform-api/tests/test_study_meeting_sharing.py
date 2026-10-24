import io
import json
from datetime import date
from unittest.mock import patch
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import execute, fetch_one, transaction
from app.main import app
from app.services.iam import authenticate
from app.services.study_meeting_operations import export_meeting_workbook
from app.services.study_meeting_sharing import sharing_dataset, report_period
from app.services.study_meetings import StudyMeetingError
from test_study_meeting_operations import prepared, flags
from test_learning_credit_ledger import _admin_id


def source():
    f,s=prepared()
    with transaction() as c:
        execute(c,"UPDATE study_meeting_sessions SET meeting_date='2020-02-10' WHERE id=?",(s['id'],))
    return f,s


def report(f,**kwargs):
    return sharing_dataset(actor_user_id=_admin_id(),report_month='2020-02',report_as_of='2020-02-20',
                           class_org_unit_id=f['class_id'],**kwargs)


def test_progress_includes_uncompleted_groups_and_rate_uses_only_completed_meetings():
    f,s=source()
    before=fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n']
    data=report(f)
    center=data['centers'][0]
    assert (center['planned'],center['completed'],center['completion_rate'])==(2,1,.5)
    assert (center['present'],center['denominator'],center['participation_rate'])==(1,1,1)
    assert [r['status'] for r in center['meetings']]==['已完成','未完成']
    assert center['meetings'][0]['other_names'] and center['meetings'][0]['present']==1
    assert fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n']==before


def copy_meeting(c,s,**changes):
    values=dict(execute(c,'SELECT * FROM study_meeting_sessions WHERE id=?',(s['id'],)).fetchone())
    values.pop('id'); values.update(session_code='SHARE-'+uuid4().hex,**changes)
    sid=execute(c,'INSERT INTO study_meeting_sessions ('+','.join(values)+') VALUES ('+','.join('?' for _ in values)+')',tuple(values.values())).lastrowid
    for fact in execute(c,'SELECT * FROM study_meeting_attendances WHERE study_meeting_session_id=?',(s['id'],)).fetchall():
        row=dict(fact);row.pop('id');row['study_meeting_session_id']=sid
        execute(c,'INSERT INTO study_meeting_attendances ('+','.join(row)+') VALUES ('+','.join('?' for _ in row)+')',tuple(row.values()))
    return sid


def test_latest_same_group_cycle_deduplicates_and_drafts_cancelled_future_are_not_completion():
    f,s=source()
    with transaction() as c:
        duplicate=copy_meeting(c,s,meeting_date='2020-02-11')
        copy_meeting(c,s,status='DRAFT',study_group_org_unit_id=f['other_group_id'])
        copy_meeting(c,s,status='CANCELLED',study_group_org_unit_id=f['other_group_id'])
        copy_meeting(c,s,meeting_date='2020-02-25',study_group_org_unit_id=f['other_group_id'])
    center=report(f)['centers'][0]
    assert center['completed']==1 and center['present']==1
    assert center['meetings'][0]['id']==duplicate


def test_center_rate_is_weighted_and_current_canonical_membership_drives_counts():
    f,s=source()
    with transaction() as c:
        copy_meeting(c,s,study_group_org_unit_id=f['other_group_id'])
        for i in range(2):
            mid=execute(c,"INSERT INTO members(member_code,name,org_unit_id,status,created_at,updated_at) VALUES (?,?,?,'ACTIVE','2020-01-01','2020-01-01')",
                (uuid4().hex,'合成未参会学员',f['class_id'])).lastrowid
            for kind,target in [('STUDY_CLASS',f['class_id']),('STUDY_GROUP',f['other_group_id'])]:
                execute(c,"INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,created_at,updated_at) VALUES (?,?,?,1,'2020-01-01','2020-01-01')",(mid,target,kind))
    center=report(f)['centers'][0]
    assert center['completed']==2 and center['completion_rate']==1
    assert center['present']==2 and center['denominator']==4 and center['participation_rate']==.5
    assert center['participation_rate'] != pytest.approx((1+1/3)/2)


def test_zero_completed_center_remains_visible_and_scope_intersection_covers_targets_and_facts():
    f,s=source();other,_=source()
    with transaction() as c:
        execute(c,"UPDATE study_meeting_sessions SET status='DRAFT' WHERE id=?",(s['id'],))
    center=report(f)['centers'][0]
    assert center['planned']==2 and center['completed']==0 and center['participation_rate'] is None
    with patch('app.services.study_meeting_sharing.accessible_org_ids',side_effect=lambda actor,permission: {f['class_id']} if permission=='exports:normal' else {other['class_id']}):
        assert report(f)['centers']==[]
    with patch('app.services.study_meeting_operations.user_context',return_value={'permissions':['exports:normal']}):
        with pytest.raises(PermissionError):report(f)


def test_renamed_same_named_centers_do_not_merge_and_direct_classes_have_own_rows():
    f,s=source();other,_=source()
    # Fixtures intentionally have equal names but distinct organization IDs.
    data=sharing_dataset(actor_user_id=_admin_id(),report_month='2020-02',report_as_of='2020-02-20')
    ids={f['class_id'],other['class_id']}
    found=[c for c in data['centers'] if any(g['parent_id'] in ids for g in c['groups'].values())]
    assert len(found)==2 and found[0]['id']!=found[1]['id']
    with transaction() as c:
        execute(c,"UPDATE org_units SET parent_id='org-suzhou' WHERE id=?",(f['class_id'],))
    assert report(f)['centers'][0]['name']=='V1.2测试班'


def test_group_move_does_not_silently_count_fact_under_different_class():
    f,s=source()
    with transaction() as c:
        execute(c,'UPDATE org_units SET parent_id=? WHERE id=?',(f['other_class_id'],f['group_id']))
    data=report(f)
    assert data['centers'][0]['completed']==0 and len(data['exceptions'])==1


@pytest.mark.parametrize('month,cutoff',[('2020-13','2020-01-01'),('2020-02','2020-03-01'),('2999-01','2999-01-01')])
def test_report_period_rejects_invalid_cross_month_or_future_dates(month,cutoff):
    with pytest.raises(StudyMeetingError):report_period(month,cutoff)


def test_workbook_has_ready_to_share_sheets_typed_rates_dates_and_no_credits_or_phone():
    f,s=source()
    raw=export_meeting_workbook(actor_user_id=_admin_id(),layout='sharing',report_month='2020-02',
        report_as_of='2020-02-20',class_org_unit_id=f['class_id'])
    book=load_workbook(io.BytesIO(raw))
    assert book.sheetnames[0]=='分中心完成情况' and len(book.sheetnames)==2
    summary=book.worksheets[0]; detail=book.worksheets[1]
    assert [summary.cell(7,col).value for col in (2,3,4,5,6,7,8)]==[2,1,1,.5,1,1,1]
    assert summary['E7'].number_format=='0%'
    assert detail['F7'].is_date and detail['I7'].value==1
    assert detail['J8'].value=='未完成'
    assert detail['A6'].fill.fgColor.rgb.endswith('4472C4')
    assert detail.freeze_panes=='A7' and detail.page_setup.orientation=='landscape'
    assert f['phone'] not in str([list(ws.values) for ws in book])
    assert '学分' not in str([list(ws.values) for ws in book])
    audit=fetch_one("SELECT after_json FROM audit_logs WHERE action='exports.study_meetings.normal' ORDER BY id DESC LIMIT 1")
    assert json.loads(audit['after_json'])['layout']=='sharing'


def test_sharing_export_http_auth_validation_and_literal_names():
    f,s=source()
    with transaction() as c:
        execute(c,"UPDATE members SET name='=FORMULA()' WHERE id=?",(f['member_id'],))
    with TestClient(app) as client:
        url='/api/v1/study-meetings/records-export'
        assert client.get(url,params={'layout':'sharing'}).status_code==401
        token=authenticate('admin','test-admin-password')['access_token']
        headers={'Authorization':'Bearer '+token}
        response=client.get(url,headers=headers,params={'layout':'sharing','report_month':'2020-02',
            'report_as_of':'2020-02-20','class_org_unit_id':f['class_id']})
        assert response.status_code==200,response.text
        book=load_workbook(io.BytesIO(response.content))
        assert book.worksheets[1]['D7'].data_type=='s' and '=FORMULA()' in book.worksheets[1]['D7'].value
        assert client.get(url,headers=headers,params={'layout':'sharing','report_month':'bad'}).status_code==400
        assert client.get(url,headers=headers,params={'layout':'unknown'}).status_code==422


@pytest.mark.parametrize('policy,label',[('WAIVED','免开展'),('SUSPENDED','暂停开展')])
def test_waived_or_suspended_groups_are_not_uncompleted_targets(policy,label):
    f,s=source()
    with transaction() as c:
        execute(c,'UPDATE class_learning_cycles SET group_meeting_policy=? WHERE id=?',(policy,f['learning_cycle_id']))
    center=report(f)['centers'][0]
    assert center['planned']==0 and center['completed']==0 and center['completion_rate'] is None
    assert center['meetings'][-1]['status']==label
    # A real voluntary meeting remains visible and contributes its own rate.
    assert center['participation_rate']==1


def test_missing_cycle_policy_is_pending_and_not_a_fake_zero_or_required_target():
    f,s=source()
    with transaction() as c:
        execute(c,'UPDATE org_units SET parent_id=? WHERE id=?',(f['other_class_id'],f['other_group_id']))
    data=sharing_dataset(actor_user_id=_admin_id(),report_month='2020-02',report_as_of='2020-02-20',
        class_org_unit_id=f['other_class_id'])
    center=data['centers'][0]
    assert center['planned'] is None and center['completion_rate'] is None
    assert center['meetings'][0]['status']=='预定待配置'


def test_later_open_cycle_policy_does_not_replace_selected_month_and_group_waiver_is_respected():
    f,s=source()
    with transaction() as c:
        old=dict(execute(c,'SELECT * FROM class_learning_cycles WHERE id=?',(f['learning_cycle_id'],)).fetchone())
        execute(c,"UPDATE class_learning_cycles SET class_meeting_status='HELD',cycle_status='CLOSED',actual_class_meeting_at='2020-02-21T01:00:00+00:00' WHERE id=?",(old['id'],))
        old.pop('id');old.update(learning_cycle_index=2,group_meeting_policy='WAIVED',cycle_status='OPEN',
            opened_at='2020-03-01T00:00:00+00:00',actual_class_meeting_at=None)
        execute(c,'INSERT INTO class_learning_cycles ('+','.join(old)+') VALUES ('+','.join('?' for _ in old)+')',tuple(old.values()))
    assert report(f)['centers'][0]['planned']==2
    march=sharing_dataset(actor_user_id=_admin_id(),report_month='2020-03',report_as_of='2020-03-20',class_org_unit_id=f['class_id'])
    assert march['centers'][0]['planned']==0
    with transaction() as c:
        task=execute(c,"SELECT id FROM group_learning_cycle_tasks WHERE class_learning_cycle_id=? AND group_org_unit_id=? AND task_type='GROUP_MEETING'",(f['learning_cycle_id'],f['other_group_id'])).fetchone()
        if task:
            execute(c,"UPDATE group_learning_cycle_tasks SET status='WAIVED' WHERE id=?",(task['id'],))
        else:
            execute(c,"INSERT INTO group_learning_cycle_tasks(class_learning_cycle_id,group_org_unit_id,task_type,status,created_at,updated_at) VALUES (?,?,'GROUP_MEETING','WAIVED','2020-01-01','2020-01-01')",(f['learning_cycle_id'],f['other_group_id']))
    center=report(f)['centers'][0]
    assert center['planned']==1 and center['completed']==1 and center['completion_rate']==1
    assert center['meetings'][-1]['status']=='免开展'


def test_sharing_does_not_silently_truncate_over_500_records():
    f,s=source()
    with patch('app.services.study_meeting_operations.list_study_meeting_records',return_value=[s]*501):
        with pytest.raises(StudyMeetingError,match='500'):report(f)


def test_changed_source_fails_closed_and_formula_like_organization_title_stays_text():
    f,s=source()
    from app.services.study_meeting_operations import _export_records
    snapshot=_export_records(_admin_id(),None,{'class_org_unit_id':f['class_id'],'status':'SUBMITTED'})
    with transaction() as c:
        execute(c,"UPDATE study_meeting_sessions SET status='CANCELLED' WHERE id=?",(s['id'],))
    with patch('app.services.study_meeting_operations._export_records',return_value=snapshot):
        with pytest.raises(StudyMeetingError,match='已变化'):report(f)
    with transaction() as c:
        execute(c,"UPDATE study_meeting_sessions SET status='SUBMITTED' WHERE id=?",(s['id'],))
        parent=execute(c,'SELECT parent_id FROM org_units WHERE id=?',(f['class_id'],)).fetchone()
        execute(c,"UPDATE org_units SET name='=UNTRUSTED()' WHERE id=?",(parent['parent_id'],))
    raw=export_meeting_workbook(actor_user_id=_admin_id(),layout='sharing',report_month='2020-02',
        report_as_of='2020-02-20',class_org_unit_id=f['class_id'])
    book=load_workbook(io.BytesIO(raw))
    assert book.worksheets[1]['A1'].value.startswith('=UNTRUSTED()')
    assert book.worksheets[1]['A1'].data_type=='s'
