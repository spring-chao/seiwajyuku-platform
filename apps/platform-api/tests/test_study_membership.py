import io
import json
from datetime import UTC, datetime, timedelta
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import connect, execute, fetch_one, transaction
from app.main import app
from app.services.iam import authenticate
from app.services.members import list_members_page, update_member, resolve_member_scope
from app.services.study_membership import group_roster
from app.services.study_membership_audit import audit_study_membership
from app.services.study_meetings import _active_group_members, get_study_meeting_record_for_operations
from app.services.study_meeting_attendees import _roster, correct_attendees
from app.services.study_meeting_operations import export_meeting_workbook
from app.services.volunteer_positions import get_current_member_study_orgs
from test_study_meeting_operations import prepared as prepare_meeting, flags
from test_learning_credit_ledger import _admin_id


def prepared():
    f, session = prepare_meeting()
    with transaction() as c:
        execute(c, 'UPDATE members SET org_unit_id=(SELECT parent_id FROM org_units WHERE id=?) '
                   'WHERE id IN (?,?)', (f['class_id'], f['member_id'], f['other_member_id']))
    return f, session


def secondary(f, *, primary=0):
    with transaction() as c:
        execute(c, 'INSERT INTO member_org_relations(member_id,org_unit_id,relation_type,is_primary,created_at,updated_at) '
            "VALUES (?,?,'STUDY_GROUP',?,'2020-01-01','2020-01-01')", (f['member_id'],f['other_group_id'],primary))


def test_display_filter_roster_and_permission_all_follow_one_formal_group():
    f, session = prepared(); secondary(f)
    actor = _admin_id()
    own = list_members_page(actor, group_org_unit_id=f['group_id'], page_size=100)['items']
    other = list_members_page(actor, group_org_unit_id=f['other_group_id'], page_size=100)['items']
    assert any(row['id'] == f['member_id'] and row['group_org_unit_id'] == f['group_id'] for row in own)
    assert all(row['id'] != f['member_id'] for row in other)
    c = connect()
    try:
        assert f['member_id'] in {r['member_id'] for r in _active_group_members(c,f['group_id'])}
        assert f['member_id'] not in {r['member_id'] for r in _active_group_members(c,f['other_group_id'])}
        assert next(r for r in _roster(c,session) if r['member_id']==f['member_id'])['attendance_type']=='HOME_GROUP'
        assert get_current_member_study_orgs(c,f['member_id'])['group']['org_unit_id']==f['group_id']
    finally: c.close()
    with pytest.raises(PermissionError): resolve_member_scope(f['member_id'],f['class_id'],{f['other_group_id']})


def test_equal_primary_groups_fail_closed_instead_of_using_largest_relation_id():
    f, _ = prepared(); secondary(f, primary=1)
    c = connect()
    try:
        assert get_current_member_study_orgs(c,f['member_id'])['group'] is None
        assert group_roster(c,member_ids=[f['member_id']])==[]
    finally: c.close()
    assert all(row['id']!=f['member_id'] for row in list_members_page(_admin_id(),group_org_unit_id=f['group_id'],page_size=100)['items'])


def test_future_and_date_end_boundaries_are_consistent_across_readers():
    f,_=prepared(); secondary(f,primary=1)
    today=datetime.now(UTC).date(); tomorrow=today+timedelta(days=1)
    with transaction() as c:
        execute(c,"UPDATE member_org_relations SET valid_from=? WHERE member_id=? AND org_unit_id=? AND relation_type='STUDY_GROUP'",(tomorrow.isoformat(),f['member_id'],f['other_group_id']))
        execute(c,"UPDATE member_org_relations SET valid_until=? WHERE member_id=? AND org_unit_id=? AND relation_type='STUDY_GROUP'",(today.isoformat(),f['member_id'],f['group_id']))
    c=connect()
    try:
        assert get_current_member_study_orgs(c,f['member_id'])['group']['org_unit_id']==f['group_id']
        assert group_roster(c,group_id=f['other_group_id'],member_ids=[f['member_id']])==[]
    finally:c.close()


def test_member_edit_retires_competing_rows_and_reuses_existing_target_without_unique_collision():
    f,session=prepared(); secondary(f)
    before=fetch_one('SELECT COUNT(*) n FROM study_meeting_attendances')['n']
    update_member(_admin_id(),f['member_id'],{'group_org_unit_id':f['other_group_id']})
    c=connect()
    try:
        assert get_current_member_study_orgs(c,f['member_id'])['group']['org_unit_id']==f['other_group_id']
        assert group_roster(c,group_id=f['group_id'],member_ids=[f['member_id']])==[]
    finally:c.close()
    assert fetch_one("SELECT COUNT(*) n FROM member_org_relations WHERE member_id=? AND relation_type='STUDY_GROUP'",(f['member_id'],))['n']==2
    assert fetch_one('SELECT COUNT(*) n FROM study_meeting_attendances')['n']==before
    raw=fetch_one('SELECT attendance_type,home_study_group_org_unit_id FROM study_meeting_attendances WHERE study_meeting_session_id=? AND member_id=?',(session['id'],f['member_id']))
    assert raw['attendance_type']=='HOME_GROUP' and raw['home_study_group_org_unit_id']==f['group_id']


def test_current_cross_group_is_shown_beside_preserved_snapshot_and_can_be_corrected_before_posting():
    f,session=prepared(); secondary(f)
    update_member(_admin_id(),f['member_id'],{'group_org_unit_id':f['other_group_id']})
    args={'actor_user_id':_admin_id(),'session_id':session['id']}
    detail=get_study_meeting_record_for_operations(**args)
    person=next(a for a in detail['attendees'] if a['member_id']==f['member_id'])
    assert person['attendance_type']=='HOME_GROUP' and person['current_attendance_type']=='CROSS_GROUP'
    assert person['membership_review_required']
    assert detail['current_home_count']==0
    sheet=load_workbook(io.BytesIO(export_meeting_workbook(**args,layout='member'))).active
    headers=[c.value for c in sheet[1]]
    row=next(dict(zip(headers,values)) for values in sheet.iter_rows(min_row=2,values_only=True) if values[6]==f['member_id'])
    assert row['参加类型']=='本组' and row['按当前归组参加类型']=='跨组'
    ids=[f['member_id'],f['other_member_id']]
    corrected=correct_attendees(**args,member_ids=ids,expected_member_ids=ids,note='合成数据归组分类复核')
    assert next(a for a in corrected['attendees'] if a['member_id']==f['member_id'])['attendance_type']=='CROSS_GROUP'
    assert corrected['membership_review_count']==0


def test_audit_is_scoped_readonly_and_never_exposes_contacts():
    f,session=prepared(); secondary(f)
    actor=_admin_id()
    before=fetch_one('SELECT COUNT(*) n FROM audit_logs')['n']
    result=audit_study_membership(actor_user_id=actor)
    assert any(row['member_id']==f['member_id'] for row in result['membership_issues'])
    assert 'phone' not in json.dumps(result) and 'ciphertext' not in json.dumps(result)
    assert fetch_one('SELECT COUNT(*) n FROM audit_logs')['n']==before
    with patch('app.services.study_membership_audit.accessible_org_ids',return_value=set()):
        assert audit_study_membership(actor_user_id=actor)['checked_members']==0
    with patch('app.services.study_membership_audit.user_context',return_value={'permissions':['org:manage']}),pytest.raises(PermissionError):
        audit_study_membership(actor_user_id=actor)
    with TestClient(app) as client:
        endpoint='/api/v1/iam/org-units/study-membership-audit'
        assert client.get(endpoint).status_code==401
        token=authenticate('admin','test-admin-password')['access_token']
        assert client.get(endpoint,headers={'Authorization':'Bearer '+token}).status_code==200


def test_same_group_batch_cleans_residual_relation_and_keeps_deactivation_guard():
    from app.services.member_group_reassignment import preview_group_reassignment, apply_group_reassignment
    from app.services.organization_management import _reference_counts
    f, _ = prepared(); secondary(f)
    c = connect()
    try:
        unit = {'id': f['other_group_id'], 'unit_type': 'GROUP'}
        before = _reference_counts(c, unit)
        assert before['unresolved_member_relations'] == 1
    finally:
        c.close()
    args = {'actor_user_id': _admin_id(), 'class_org_unit_id': f['class_id'], 'assignments': [
        {'member_id': f['member_id'], 'group_org_unit_id': f['group_id']} ]}
    preview = preview_group_reassignment(**args)
    assert preview['changed_count'] == 1 and preview['rows'][0]['needs_relation_cleanup']
    apply_group_reassignment(**args, preview_fingerprint=preview['preview_fingerprint'], reason='核对并清理旧小组关系')
    assert preview_group_reassignment(**args)['changed_count'] == 0


def test_explicit_target_can_resolve_group_conflict_without_contact_or_history_changes():
    from app.services.member_group_reassignment import preview_group_reassignment, apply_group_reassignment, group_reassignment_roster
    f, session = prepared(); secondary(f, primary=1)
    actor = _admin_id()
    personal = dict(fetch_one('SELECT phone_ciphertext,enterprise_financial_ciphertext FROM members WHERE id=?', (f['member_id'],)))
    facts = dict(fetch_one('SELECT * FROM study_meeting_attendances WHERE study_meeting_session_id=? AND member_id=?', (session['id'],f['member_id'])))
    roster = group_reassignment_roster(actor_user_id=actor, class_org_unit_id=f['class_id'])
    member = next(row for row in roster['members'] if row['id'] == f['member_id'])
    assert not member.get('unavailable') and member['group_org_unit_id'] is None
    assert member['group_name'].startswith('待核对：')
    args = {'actor_user_id': actor, 'class_org_unit_id': f['class_id'], 'assignments': [
        {'member_id': f['member_id'], 'group_org_unit_id': f['group_id']} ]}
    preview = preview_group_reassignment(**args)
    apply_group_reassignment(**args, preview_fingerprint=preview['preview_fingerprint'], reason='依据已确认主档消除归组冲突')
    assert dict(fetch_one('SELECT phone_ciphertext,enterprise_financial_ciphertext FROM members WHERE id=?', (f['member_id'],))) == personal
    assert dict(fetch_one('SELECT * FROM study_meeting_attendances WHERE study_meeting_session_id=? AND member_id=?', (session['id'],f['member_id']))) == facts
