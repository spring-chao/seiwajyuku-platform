import io
import cProfile
import pstats
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import execute, fetch_one, transaction
from app.main import app
from app.services.iam import authenticate
from app.services.member_exports import export_members_xlsx, member_export_fields
from app.services.study_meeting_operations import export_meeting_workbook
from app.services.study_meeting_operations import meeting_credit_preview
from app.services.study_meetings import create_study_meeting, StudyMeetingError
from app.services.study_meetings import get_study_meeting_record_for_operations
from test_study_meeting_operations import prepared, flags
from test_learning_credit_ledger import _admin_id
from test_study_meeting_evidence import create


def book(raw):
    return load_workbook(io.BytesIO(raw), data_only=False)


def test_group_workbook_has_one_row_names_together_and_numeric_rates():
    f, session = prepared()
    report = book(export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id']))
    assert report.sheetnames == ['小组汇总', '参会率汇总']
    sheet = report.worksheets[0]
    assert sheet.max_row == 2
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row['其中跨组姓名'] in row['姓名'] and '、' in row['姓名']
    assert row['小组月度参会率'] == 1
    assert 0 < row['班级参会率（筛选范围）'] <= 1
    assert sheet['L2'].number_format == '0.0%'
    assert sheet['B2'].is_date
    assert sheet['R2'].data_type == 'n'
    assert sheet['L2'].border.bottom.style == 'thin'
    assert sheet['K1'].border.right.style == 'thin'
    assert sheet['A2'].fill.fgColor.rgb.endswith('EFF6F2')
    assert next(iter(sheet.tables.values())).ref == sheet.dimensions
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == sheet.dimensions
    assert f['phone'] not in str(list(sheet.values))
    stats = report['参会率汇总']
    assert stats['A3'].border.top.style == 'thin'
    assert next(iter(stats.tables)) != next(iter(sheet.tables))
    assert {stats.cell(i, 1).value for i in range(2, stats.max_row + 1)} == {'班级', '分中心'}


def test_member_workbook_retains_individual_points_and_confirmation():
    f, session = prepared(course=True)
    sheet = book(export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'], layout='member')).active
    assert sheet.max_row == 3
    assert sheet['H1'].value == '姓名'
    assert sheet['O2'].value == 44
    assert sheet['O2'].data_type == 'n'


def test_bulk_export_avoids_per_meeting_iam_connections_and_keeps_every_member():
    # Count actual DB connections rather than asserting wall-clock time or
    # mocking away authorization. Previously even five meetings opened >700.
    fixtures = [prepared(course=True) for _ in range(5)]
    profile = cProfile.Profile()
    with profile:
        report = book(export_meeting_workbook(actor_user_id=_admin_id(), layout='member'))
    calls = sum(value[1] for (filename, _, name), value in pstats.Stats(profile).stats.items()
                if filename.endswith('db.py') and name == 'connect')
    assert calls < 100
    rows = list(report.active.values)[1:]
    for _, session in fixtures:
        own = [row for row in rows if row[0] == session['session_code']]
        assert len(own) == 2 and all(row[14] == 44 for row in own)


def test_export_reuses_canonical_preview_after_posting_without_writing_credits():
    from app.services.learning_credits import _insert_entry, dry_run_study_meeting_settlement
    f, session = prepared(course=True)
    args = {'actor_user_id': _admin_id(), 'session_id': session['id']}
    entry = dry_run_study_meeting_settlement(**args)['entries'][0]
    with transaction() as connection:
        _insert_entry(connection, {**entry, 'occurred_at': session['meeting_date']}, status='POSTED', actor_user_id=_admin_id())
    before = fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n']
    expected = {member['member_id']: member for member in meeting_credit_preview(**args)['members']}
    rows = list(book(export_meeting_workbook(**args, layout='member')).active.values)[1:]
    for row in rows:
        member = expected[row[6]]
        assert row[12:15] == (member['group_points'], member['course_points'], member['expected_points'])
    assert fetch_one('SELECT COUNT(*) n FROM learning_credit_entries')['n'] == before
    audit = fetch_one("SELECT after_json FROM audit_logs WHERE action='exports.study_meetings.normal' ORDER BY id DESC LIMIT 1")
    import json
    assert json.loads(audit['after_json'])['row_count'] == 2


def test_export_rejects_changed_ownership_before_building_the_file():
    from app.services import study_meeting_operations as operations
    f, session = prepared()
    original = operations._session_payload
    def changed(connection, sid):
        payload = original(connection, sid)
        payload['class_org_unit_id'] = f['other_class_id']
        return payload
    with patch.object(operations, '_session_payload', side_effect=changed), pytest.raises(StudyMeetingError, match='归属已变化'):
        export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'])


def test_export_preview_and_export_scopes_are_intersected_and_limit_is_preserved():
    from app.services import study_meeting_operations as operations
    f, session = prepared()
    with patch('app.services.study_meetings.accessible_org_ids', return_value={f['other_class_id']}), pytest.raises(PermissionError):
        export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'])
    with patch.object(operations, 'user_context', return_value={'permissions': ['exports:normal']}), pytest.raises(PermissionError):
        export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'])
    with patch.object(operations, 'list_study_meeting_records', return_value=[{'class_org_unit_id': f['class_id']}] * 501), pytest.raises(StudyMeetingError, match='500场'):
        export_meeting_workbook(actor_user_id=_admin_id())


@pytest.mark.parametrize('layout,expected', [('group', '小组汇总'), ('member', '人员明细')])
def test_no_matching_meetings_still_produces_a_valid_header_only_workbook(layout, expected):
    sheet = book(export_meeting_workbook(actor_user_id=_admin_id(), layout=layout, meeting_date_to='1900-01-01')).active
    assert sheet.title == expected and sheet.max_row == 1


def test_historic_duplicate_group_rows_use_latest_and_do_not_double_count_rates():
    f, session = prepared()
    with transaction() as c:
        original = dict(execute(c, 'SELECT * FROM study_meeting_sessions WHERE id=?', (session['id'],)).fetchone())
        original.pop('id')
        original['session_code'] += '-legacy-copy'
        columns = ','.join(original)
        placeholders = ','.join('?' for _ in original)
        duplicate_id = execute(c, f'INSERT INTO study_meeting_sessions ({columns}) VALUES ({placeholders})', tuple(original.values())).lastrowid
        for attendee in execute(c, 'SELECT * FROM study_meeting_attendances WHERE study_meeting_session_id=?', (session['id'],)).fetchall():
            values = dict(attendee); values.pop('id'); values['study_meeting_session_id'] = duplicate_id
            execute(c, 'INSERT INTO study_meeting_attendances (' + ','.join(values) + ') VALUES (' + ','.join('?' for _ in values) + ')', tuple(values.values()))
    records = [get_study_meeting_record_for_operations(actor_user_id=_admin_id(), session_id=sid) for sid in [session['id'], duplicate_id]]
    with patch('app.services.study_meeting_operations._export_records', return_value=records):
        report = book(export_meeting_workbook(actor_user_id=_admin_id()))
    assert report.active.max_row == 2
    assert report.active['A2'].value.endswith('-legacy-copy')
    assert '2条历史记录' in report.active['S2'].value
    assert report['参会率汇总']['E2'].value == 2
    assert report['参会率汇总']['G2'].value == 1


def test_group_workbook_retains_same_named_people_and_rejects_scope():
    f, session = prepared()
    with patch('app.services.study_meeting_operations.accessible_org_ids', return_value={f['other_class_id']}):
        with pytest.raises(PermissionError):
            export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'])
    with transaction() as c:
        execute(c, "UPDATE members SET name='=FORMULA()' WHERE id IN (?,?)", (f['member_id'], f['other_member_id']))
    sheet = book(export_meeting_workbook(actor_user_id=_admin_id(), session_id=session['id'])).active
    assert sheet['H2'].value == '=FORMULA()、=FORMULA()' and sheet['H2'].data_type == 's'


def test_member_export_fields_filters_groups_and_permission_intersection():
    f, session = prepared()
    fields = ['name', 'class_name', 'group_name']
    sheet = book(export_members_xlsx(_admin_id(), fields=fields, class_org_unit_id=f['class_id'], layout='group')).active
    assert sheet['A1'].value == '姓名' and sheet['D1'].value == '人数'
    assert sheet.max_row >= 2
    assert f['phone'] not in str(list(sheet.values))
    with pytest.raises(ValueError):
        export_members_xlsx(_admin_id(), fields=['phone'])
    with pytest.raises(ValueError):
        export_members_xlsx(_admin_id(), fields=['phone_masked'], layout='group')
    with patch('app.services.members.accessible_org_ids', side_effect=lambda actor, permission=None: set() if permission == 'exports:normal' else None):
        empty = book(export_members_xlsx(_admin_id(), fields=fields, class_org_unit_id=f['class_id'])).active
        assert empty.max_row == 1
    with patch('app.services.member_exports.user_context', return_value={'permissions': ['members:read']}):
        with pytest.raises(PermissionError):
            export_members_xlsx(_admin_id(), fields=fields)


def test_member_export_downloads_all_filtered_pages_not_just_twenty():
    items = [{'id': i, 'name': f'合成学员{i}', 'status': 'ACTIVE'} for i in range(125)]
    def page(actor, **kwargs):
        start = (kwargs['page'] - 1) * 100
        return {'items': items[start:start + 100], 'pagination': {'total': 125, 'total_pages': 2}}
    with patch('app.services.member_exports.list_members_page', side_effect=page):
        sheet = book(export_members_xlsx(_admin_id(), fields=['name'])).active
        assert sheet.max_row == 126


def test_exports_http_auth_and_default_page_size():
    f, session = prepared()
    with TestClient(app) as client:
        url = '/api/v1/exports/members/workbook'
        assert client.get(url, params={'fields': 'name'}).status_code == 401
        token = authenticate('admin', 'test-admin-password')['access_token']
        headers = {'Authorization': 'Bearer ' + token}
        catalog_url = '/api/v1/exports/members/fields'
        assert client.get(catalog_url).status_code == 401
        catalog = client.get(catalog_url, headers=headers).json()['data']
        assert len(catalog) > 30
        assert next(field for field in catalog if field['key'] == 'phone')['available'] is False
        response = client.get(url, params={'fields': 'name', 'class_org_unit_id': f['class_id']}, headers=headers)
        assert response.status_code == 200, response.text
        assert response.headers['cache-control'] == 'private, no-store'
        assert book(response.content).active['A1'].value == '姓名'
        assert client.get('/api/v1/members', headers=headers).json()['data']['pagination']['page_size'] == 20
        response = client.get('/api/v1/study-meetings/records-export', params={'session_id': session['id']}, headers=headers)
        assert response.status_code == 200, response.text
        assert book(response.content).active.title == '小组汇总'


def test_full_member_field_catalog_exports_typed_values_and_preserves_zero():
    f, session = prepared()
    with transaction() as c:
        execute(c, 'UPDATE members SET gender=?, birthday=?, company_name=?, position=?, '
                'employee_count=0, membership_years=0, membership_years_overridden=1, '
                'company_products=?, notes=? WHERE id=?',
                ('FEMALE', '2000-01-02', '=Synthetic Company', '合成职务',
                 '合成产品', 'PROTECTED_NOTE_FIXTURE', f['member_id']))
    catalog = member_export_fields(_admin_id())
    fields = [field['key'] for field in catalog if field['available']]
    assert {'birthday', 'company_name', 'position', 'employee_count', 'volunteer_positions'} <= set(fields)
    with TestClient(app) as client:
        token = authenticate('admin', 'test-admin-password')['access_token']
        response = client.get('/api/v1/exports/members/workbook', params={
            'fields': ','.join(fields), 'purpose': '合成资料测试', 'class_org_unit_id': f['class_id']},
            headers={'Authorization': 'Bearer ' + token})
    assert response.status_code == 200, response.text
    sheet = book(response.content).active
    data = [dict(zip(fields, row)) for row in sheet.iter_rows(min_row=2)]
    row = next(row for row in data if row['company_name'].value == '=Synthetic Company')
    assert row['company_name'].data_type == 's'
    assert row['gender'].value == '女'
    assert row['birthday'].is_date and row['birthday'].value.year == 2000
    assert row['employee_count'].value == 0 and row['employee_count'].data_type == 'n'
    assert row['membership_years'].value == 0
    assert 'PROTECTED_NOTE_FIXTURE' not in str(list(sheet.values))
    assert f['phone'] not in str(list(sheet.values))
    audit = fetch_one("SELECT purpose FROM audit_logs WHERE action='exports.members.normal' ORDER BY id DESC LIMIT 1")
    assert audit['purpose'] == '合成资料测试'


def test_expanded_member_fields_keep_profile_permissions_scopes_and_purpose():
    f, session = prepared()
    basic = {'permissions': ['members:read', 'exports:normal']}
    with patch('app.services.member_exports.user_context', return_value=basic):
        catalog = {field['key']: field for field in member_export_fields(_admin_id())}
        assert catalog['company_name']['available']
        assert not catalog['birthday']['available'] and not catalog['company_address']['available']
        with pytest.raises(PermissionError):
            export_members_xlsx(_admin_id(), fields=['birthday'])
    for protected in ['phone', 'annual_sales', 'profit_margin', 'company_size', 'notes']:
        with pytest.raises(ValueError):
            export_members_xlsx(_admin_id(), fields=[protected])
    with pytest.raises(ValueError, match='用途'):
        export_members_xlsx(_admin_id(), fields=['company_address'])
    with patch('app.services.members.accessible_org_ids', side_effect=lambda actor, permission=None: set() if permission == 'members:detail_view' else None):
        assert book(export_members_xlsx(_admin_id(), fields=['birthday'], class_org_unit_id=f['class_id'])).active.max_row == 1
    with patch('app.services.members.accessible_org_ids', side_effect=lambda actor, permission=None: set() if permission == 'members:enterprise_view' else None):
        assert book(export_members_xlsx(_admin_id(), fields=['company_address'], purpose='合成范围测试', class_org_unit_id=f['class_id'])).active.max_row == 1
    with patch('app.services.member_exports.get_settings') as settings:
        settings.return_value.identity_authorization_enabled = False
        assert not next(field for field in member_export_fields(_admin_id()) if field['key'] == 'volunteer_positions')['available']
        with pytest.raises(PermissionError):
            export_members_xlsx(_admin_id(), fields=['volunteer_positions'])


def test_group_cycle_rejects_duplicate_submitted_but_retries_identical_draft():
    f, session = prepared()
    count = fetch_one('SELECT COUNT(*) AS n FROM study_meeting_sessions')['n']
    with pytest.raises(StudyMeetingError, match='不能重复上传'):
        create(f)
    assert fetch_one('SELECT COUNT(*) AS n FROM study_meeting_sessions')['n'] == count
    with transaction() as c:
        execute(c, "UPDATE study_meeting_sessions SET status='CANCELLED' WHERE id=?", (session['id'],))
    new = create(f)
    assert new['id'] != session['id']
    assert create(f)['id'] == new['id']
    with transaction() as c:
        execute(c, "UPDATE study_meeting_sessions SET status='CANCELLED' WHERE id=?", (new['id'],))
    # Concurrent first creations serialize on the group, not on a missing child.
    count = fetch_one('SELECT COUNT(*) AS n FROM study_meeting_sessions')['n']
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: create(f), range(2)))
    assert len({r['id'] for r in results}) == 1
    assert fetch_one('SELECT COUNT(*) AS n FROM study_meeting_sessions')['n'] == count + 1
